#!/usr/bin/env python3
"""Thư viện nguồn của project: khung thư mục references/ và chỉ mục references/_INDEX.xlsx.
Source library of a project: the references/ folder layout and the references/_INDEX.xlsx index.

Dùng / Usage (chạy từ thư mục gốc của project / run from the project root):
  python tools/library.py init                 tạo khung thư mục và chỉ mục nếu chưa có (không ghi đè)
                                               create the folders and the index if missing (never overwrites)
  python tools/library.py check [--days 180]   kiểm chỉ mục: thiếu file, file chưa vào chỉ mục, trùng số hiệu,
                                               tình trạng mâu thuẫn, ngày kiểm quá cũ / check the index
  python tools/library.py check --content-audit  thêm dấu hiệu nguồn chỉ đọc, PDF engine tùy chọn;
                                               add read-only source indicators; optional PDF engine, no verification
  python tools/library.py find <từ khóa> [--group "Pháp luật"] [--all]
                                               tìm theo số hiệu, trích yếu, lĩnh vực, điều (không phân biệt dấu)
                                               search number, title, domain, articles (accent-insensitive)
  python tools/library.py stats                đếm theo nhóm và tình trạng / counts by group and status
  python tools/library.py sync [--apply] [--rename ID=đường_dẫn#sha256 ...]
                                               xem trước / ghi metadata của file nguồn; không xác minh nội dung.
                                               Thư mục nguồn cũ chỉ được quét khi người dùng ghi vào PROJECT.md (dòng
                                               "Ngoại lệ thư mục nguồn cũ"). Đổi tên/di chuyển chỉ áp dụng khi có
                                               --rename lấy từ bản xem trước; trùng hash hoặc tên không tự gộp.
                                               preview / record source metadata; does not verify content.
                                               Old source folders are scanned only when the user lists them in PROJECT.md
                                               ("Existing source-folder exceptions"). A move or rename applies only with
                                               --rename copied from the preview; equal hashes or names never merge rows.
  python tools/library.py content preview --request FILE | apply --request FILE --expect-plan TOKEN
  python tools/library.py content find [từ khóa] [--source-id ID] [--entry-id ID]
  python tools/library.py source-event preview --request FILE | apply --request FILE --expect-plan TOKEN
                                               ghi/tra nội dung đã đọc theo phiên bản nguồn (sheet "LAW content") và sự kiện
                                               ngữ nghĩa explicit; find chỉ đọc sổ, không đọc nguồn, không xác minh
                                               record/look up read content per source version (sheet "LAW content") and explicit
                                               semantic source-events; find reads the workbook only, never the source, verifies nothing
  python tools/library.py intake [--file BASENAME --code CODE] [--apply --expect-plan TOKEN] [--json]
                                               nhập create-only từ staging project đã khai; mặc định chỉ xem trước,
                                               không xác minh danh tính/toàn văn/hiệu lực; giữ nguyên staging.
                                               create-only from configured staging; preview by default, staging retained.
  Thêm --json để lấy kết quả dạng JSON (bảng điều khiển dùng) / add --json for JSON output.

find/check/stats đọc sheet "Danh mục"; sync --apply chỉ ghi sổ này và sheet lịch sử LAW, không sửa file nguồn.
find/check/stats read "Danh mục"; sync --apply writes only this index and LAW history, never source files.
Cần Python 3.8+ và openpyxl (pip install openpyxl). Không dùng mạng. / Needs Python 3.8+ and openpyxl. No network.
"""
import datetime, hashlib, io, json, os, re, shutil, stat, sys, traceback, unicodedata
from result_schema import emit, make

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REF = os.path.join(ROOT, 'references')
INDEX = os.path.join(REF, '_INDEX.xlsx')
TEMPLATE = os.path.join(ROOT, 'skills', 'local', 'shared', 'source-library', 'assets', 'source-index-template.xlsx')
SHEET = 'Danh mục'
HEADER_ROW = 4
FILE_EXT = ('.pdf', '.doc', '.docx', '.xls', '.xlsx', '.xlsm', '.odt', '.ods', '.rtf', '.txt', '.md', '.html', '.htm',
            '.png', '.jpg', '.jpeg', '.tif', '.tiff', '.ppt', '.pptx', '.zip')
GROUP_DIRS = {  # nhóm -> thư mục / group -> folder
    'legal': ('Pháp luật', 'Văn bản quy phạm pháp luật và văn bản của cơ quan nhà nước, chia theo lĩnh vực: legal/<linh-vuc>/.',
              'Laws, decrees, circulars and other state documents, one sub-folder per domain: legal/<domain>/.'),
    'policies': ('Nội bộ', 'Quy chế, nội quy, quy định, quyết định nội bộ của đơn vị, chia theo mảng: policies/<mang>/.',
                 'The organisation\'s own rules, regulations and internal decisions, one sub-folder per area: policies/<area>/.'),
    'templates': ('Mẫu biểu', 'Mẫu văn bản, biểu mẫu, mẫu báo cáo đã duyệt: templates/<mang>/.',
                  'Approved document, form and report templates: templates/<area>/.'),
    'source-documents': ('Tài liệu nguồn', 'Tài liệu nghiệp vụ khác dùng làm nguồn: hồ sơ, số liệu, tài liệu tham khảo.',
                         'Other business documents used as sources: case files, data, reference material.'),
}
# khóa cột -> tiền tố tiêu đề (phần tiếng Việt) / column key -> header prefix
KEYS = {
    'id': 'Mã', 'group': 'Nhóm', 'domain': 'Lĩnh vực', 'issuer': 'Cơ quan', 'number': 'Số, ký hiệu', 'type': 'Loại văn bản',
    'issued': 'Ngày ban hành', 'effective': 'Ngày có hiệu lực', 'title': 'Trích yếu', 'replaces': 'Thay thế',
    'replaced_by': 'Bị thay thế', 'status': 'Tình trạng', 'checked': 'Ngày kiểm', 'verify': 'Mức độ xác minh',
    'fulltext': 'Có toàn văn', 'file': 'Đường dẫn file', 'articles': 'Điều', 'sensitivity': 'Mức nhạy cảm',
    'used_by': 'Gói', 'url': 'Nguồn', 'notes': 'Ghi chú',
    'law_revision': 'LAW revision', 'law_sha': 'LAW SHA-256', 'law_observed': 'LAW observed', 'law_state': 'LAW file state',
}
# cột cần có để đọc/duyệt sổ (find, stats, bảng điều khiển); metadata LAW không bắt buộc để xem sổ cũ
# columns needed to read/browse the index; the LAW metadata columns are not required to view an older index
REQUIRED_VIEW = ('group', 'number', 'title', 'status', 'fulltext', 'file')
BROWSE_LIMIT = 500
LANG = 'vi'


def L(vi, en):
    return en if LANG == 'en' else vi


RESULT_REQUESTED = '--result-json' in sys.argv  # cờ ẩn: phát một kết quả có cấu trúc / hidden flag: emit one structured result
LAST_RESULT = None
CREATED = []  # khung đã tạo trong lượt này / folders and files created in this run, for partial results


def set_result(status, code, message, **kw):
    global LAST_RESULT
    LAST_RESULT = make(status, code=code, message=message, **kw)


def stop(msg, status='blocked', **kw):
    """Như sys.exit(msg) (mã 1), nhưng ghi kết quả có cấu trúc trước / Like sys.exit(msg) (code 1), recording the structured result first."""
    set_result(status, 1, msg, changes=False, **kw)
    sys.exit(msg)


def fold(s):
    s = unicodedata.normalize('NFD', str(s or '')).replace('đ', 'd').replace('Đ', 'D')
    return ''.join(c for c in s if unicodedata.category(c) != 'Mn').lower().strip()


def need_openpyxl():
    try:
        import openpyxl  # noqa: F401
        return openpyxl
    except ImportError:
        msg = L('CHƯA KIỂM ĐƯỢC: thư viện nguồn cần openpyxl. Chạy "python -m pip install openpyxl" rồi thử lại. Các công cụ văn bản vẫn dùng được.',
                'NOT CHECKED: the source library requires openpyxl. Run "python -m pip install openpyxl" and retry. Text tools still work.')
        print(msg, file=sys.stderr)
        set_result('not_checked', 2, msg, changes=False, unchecked=[{'tool': 'library.py', 'reason': 'openpyxl_missing'}],
                   next_actions=[L('Chạy: python -m pip install openpyxl', 'Run: python -m pip install openpyxl')])
        sys.exit(2)


def as_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    m = re.search(r'(\d{1,2})/(\d{1,2})/(\d{4})', str(v or ''))
    if m:
        try:
            return datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    return None


def fmt(v):
    d = as_date(v) if isinstance(v, (datetime.date, datetime.datetime)) else None
    return d.strftime('%d/%m/%Y') if d else ('' if v is None else str(v).strip())


def index_path(root=None):
    return os.path.join(root or ROOT, 'references', '_INDEX.xlsx')


BOOK_ERRORS = {  # loại lỗi sổ -> (tiếng Việt, English); khác sổ hợp lệ 0 dòng / index problem kind -> message; never a valid empty index
    'no_index': ('Chưa có references/_INDEX.xlsx. Chạy: python tools/library.py init',
                 'references/_INDEX.xlsx not found. Run: python tools/library.py init'),
    'denied': ('Không có quyền đọc references/_INDEX.xlsx; chưa biết sổ có hợp lệ hay không.',
               'No permission to read references/_INDEX.xlsx; its validity is unknown.'),
    'unreadable': ('Không đọc được references/_INDEX.xlsx (file hỏng, đang bị khóa hoặc không phải sổ Excel hợp lệ).',
                   'references/_INDEX.xlsx cannot be read (corrupt, locked or not a valid Excel workbook).'),
    'missing_sheet': ('Chỉ mục thiếu sheet "%s".' % SHEET, 'The index has no "%s" sheet.' % SHEET),
    'no_header': ('Sheet "%s" không có dòng tiêu đề (dòng %d).' % (SHEET, HEADER_ROW),
                  'The "%s" sheet has no header row (row %d).' % (SHEET, HEADER_ROW)),
    'missing_columns': ('Chỉ mục thiếu cột bắt buộc: %s', 'The index is missing required columns: %s'),
}


def book_message(res):
    vi, en = BOOK_ERRORS[res['kind']]
    text = L(vi, en)
    return text % ', '.join(res.get('missing') or []) if res['kind'] == 'missing_columns' else text


# Bước xử lý cụ thể cho từng loại lỗi sổ. Không bước nào dựng đè hay xóa sổ: dữ liệu trong sổ là của bạn.
# A concrete next step per kind of index problem. None rebuilds over or deletes the index: its data is yours.
BOOK_STEPS = {
    'denied': ('Kiểm quyền đọc file và đóng ứng dụng đang khóa nó, rồi chạy lại.',
               'Check read permission and close the application that locks it, then run again.'),
    'unreadable': ('Đóng Excel nếu đang mở file. Nếu vẫn lỗi, giữ nguyên file hỏng và khôi phục references/_INDEX.xlsx từ bản sao lưu gần nhất của chính bạn (bản sao thư mục dự án, hướng dẫn nâng cấp mục 1). Đừng chép mẫu rỗng hay chạy lại init để dựng đè.',
                   'Close Excel if it is open. If it still fails, keep the broken file as it is and restore references/_INDEX.xlsx from your own most recent backup (a copy of the project folder, upgrade guide section 1). Do not copy an empty template over it or run init again to rebuild.'),
    'missing_sheet': ('Mở sổ bằng Excel: nếu sheet đã bị đổi tên thì đặt lại đúng tên "%s"; nếu dữ liệu đã mất thì khôi phục từ bản sao lưu. Đừng chép mẫu rỗng đè lên sổ.' % SHEET,
                      'Open the index in Excel: if the sheet was renamed, give it back the exact name "%s"; if the data is gone, restore it from a backup. Do not copy an empty template over the index.' % SHEET),
    'no_header': ('Khôi phục dòng tiêu đề (dòng %d) theo mẫu skills/local/shared/source-library/assets/source-index-template.xlsx hoặc từ bản sao lưu; chỉ sửa tiêu đề, giữ nguyên các dòng dữ liệu.' % HEADER_ROW,
                  'Restore the header row (row %d) from skills/local/shared/source-library/assets/source-index-template.xlsx or from a backup; fix only the headings and keep the data rows.' % HEADER_ROW),
    'missing_columns': ('Thêm lại các cột còn thiếu vào sheet với đúng tên tiêu đề như mẫu source-index-template.xlsx; không xóa cột hay dữ liệu đang có.',
                        'Add the missing columns back to the sheet with the exact headings of the source-index-template.xlsx template; do not delete existing columns or data.'),
}


def book_steps(kind):
    """Câu hướng dẫn xử lý cho một loại lỗi sổ, hoặc '' nếu không có. / The next-step sentence for an index problem kind, or ''."""
    vi_en = BOOK_STEPS.get(kind)
    return L(*vi_en) if vi_en else ''


def read_book(root=None, require_columns=True):
    """Đọc sổ một lần, không thoát chương trình, không ghi gì. Trả {'ok': True, cols, data, snapshot} hoặc {'ok': False, kind}.
    Loại lỗi: openpyxl_missing | no_index | denied | unreadable | missing_sheet | no_header | missing_columns. Sổ hợp lệ không có dòng
    dữ liệu là ok với data rỗng, khác mọi loại lỗi. snapshot = SHA-256 của đúng các byte đã đọc.
    Read the index once, never exit, never write. A valid index with no data rows is ok with empty data, distinct from every error kind.
    snapshot is the SHA-256 of exactly the bytes that were read."""
    try:
        import openpyxl as opx
    except ImportError:
        return {'ok': False, 'kind': 'openpyxl_missing'}
    path = index_path(root)
    try:
        with open(path, 'rb') as f:
            raw = f.read()
    except FileNotFoundError:
        return {'ok': False, 'kind': 'no_index'}
    except PermissionError:
        return {'ok': False, 'kind': 'denied'}
    except OSError:
        return {'ok': False, 'kind': 'unreadable'}
    return parse_book(raw, require_columns)


def parse_book(raw, require_columns=True, data_only=True):
    """Phân tích đúng các byte đã đọc (dùng chung với content_find để catalog và nội dung cùng một snapshot).
    Parse exactly the bytes already read (shared with content_find so the catalog and the content come from one snapshot)."""
    import openpyxl as opx
    snapshot = hashlib.sha256(raw).hexdigest()
    try:
        wb = opx.load_workbook(io.BytesIO(raw), data_only=data_only, read_only=True)
    except Exception:  # not a zip, corrupt part, unsupported file: all "cannot be read", never "empty"
        return {'ok': False, 'kind': 'unreadable'}
    try:
        if SHEET not in wb.sheetnames:
            return {'ok': False, 'kind': 'missing_sheet'}
        rows = list(wb[SHEET].iter_rows(values_only=True))
    except Exception:
        return {'ok': False, 'kind': 'unreadable'}
    finally:
        wb.close()
    if len(rows) < HEADER_ROW:
        return {'ok': False, 'kind': 'no_header'}
    header = [str(h or '') for h in rows[HEADER_ROW - 1]]
    cols = {}
    for k, pre in KEYS.items():
        for i, h in enumerate(header):
            if h.startswith(pre) and i not in cols.values():
                cols[k] = i
                break
    missing = [KEYS[k] for k in REQUIRED_VIEW if k not in cols]
    if not cols:
        return {'ok': False, 'kind': 'no_header'}
    if missing and require_columns:
        return {'ok': False, 'kind': 'missing_columns', 'missing': missing}
    data = []
    for n, r in enumerate(rows[HEADER_ROW:], HEADER_ROW + 1):
        if not any(c not in (None, '') for c in r):
            continue
        rec = {k: (r[i] if i < len(r) else None) for k, i in cols.items()}
        rec['_row'] = n
        data.append(rec)
    return {'ok': True, 'cols': cols, 'data': data, 'snapshot': snapshot, 'missing': missing}


def load(require_columns=True):
    res = read_book(None, require_columns)
    if res['ok']:
        return res['cols'], res['data']
    if res['kind'] == 'openpyxl_missing':
        need_openpyxl()  # in thông báo, ghi kết quả not_checked và thoát mã 2 / prints, records not_checked, exits 2
    step = book_steps(res['kind'])
    stop(book_message(res) + (' ' + step if step else ''), status='error' if res['kind'] in ('denied', 'unreadable') else 'blocked', reason='index_' + res['kind'],
         next_actions=[step] if step else None)


def library_files():
    out = []
    if not os.path.isdir(REF):
        return out
    shipped = package_readme_filter()
    for d, dirs, files in os.walk(REF):
        dirs[:] = [x for x in dirs if not x.startswith(('.', '_'))]
        for f in files:
            if f.startswith(('~$', '.', '_INDEX')):
                continue
            if f.lower().endswith(FILE_EXT):
                full = os.path.join(d, f)
                if f == 'README.md' and shipped(full, d):
                    continue  # README do package giao và còn nguyên; README bạn viết hoặc sửa vẫn là nguồn / a shipped, unchanged README; one you wrote or edited stays a source
                out.append(os.path.relpath(full, REF).replace(os.sep, '/'))
    return sorted(out)


def package_readme_filter():
    """Hàm (đường dẫn, thư mục) -> True nếu README đó do package giao và chưa sửa, theo ĐÚNG điều kiện ownership của library_sync (hash trong MANIFEST
    hoặc đúng README khung của nhóm), không theo tên file. Không nạp được library_sync thì không loại README nào.
    (path, folder) -> True when that README is shipped by the package and unchanged, by library_sync's own ownership predicate, never by basename."""
    try:
        import library_sync
        if not library_sync.link_free(os.path.join(ROOT, 'MANIFEST.sha256'), ROOT):
            return lambda path, folder: False
        hashes = library_sync.manifest_hashes(ROOT)
    except Exception:  # không xác định được quyền sở hữu thì không bỏ qua README nào / ownership unknown: skip no README
        return lambda path, folder: False
    groups = dict((os.path.normcase(os.path.join(REF, name)), name) for name in GROUP_DIRS)

    def shipped(path, folder):
        # Ownership must not introduce a content read through an out-of-scope link.
        if not library_sync.link_free(path, ROOT) or not library_sync.contained(path, REF):
            return False
        try:
            return library_sync.package_owned(ROOT, hashes, path, library_sync.digest(path), folder, groups.get(os.path.normcase(folder)))
        except OSError:
            return False
    return shipped


def norm_path(p):
    s = str(p or '').strip().replace('\\', '/')
    s = re.sub(r'^\./', '', s)
    return re.sub(r'^references/', '', s)


def group_readme(d):
    """README that `init` writes into a group folder; library_sync uses it to tell it from a README the user wrote."""
    grp, vi, en = GROUP_DIRS[d]
    return ('# references/%s/\n\nNhóm "%s" trong chỉ mục `references/_INDEX.xlsx`. %s\n\n%s\n\n'
            'Mỗi lĩnh vực một thư mục con, tên không dấu, nối bằng gạch ngang. Thư mục con có thể có README ngắn: phạm vi, người cập nhật, '
            'nguồn chính thống. Thêm file thì thêm dòng vào chỉ mục; xem skill `source-library`.\n' % (d, grp, vi, en))


# ---------------------------------------------------------------- lệnh / commands
def cmd_init(args):
    made = CREATED
    os.makedirs(REF, exist_ok=True)
    for d, (grp, vi, en) in GROUP_DIRS.items():
        p = os.path.join(REF, d)
        if not os.path.isdir(p):
            os.makedirs(p)
            made.append('references/%s/' % d)
        rd = os.path.join(p, 'README.md')
        if not os.path.isfile(rd):
            io.open(rd, 'w', encoding='utf-8', newline='\n').write(group_readme(d))
            made.append('references/%s/README.md' % d)
    if not os.path.isfile(INDEX):
        if not os.path.isfile(TEMPLATE):
            stop(L('Thiếu mẫu chỉ mục: %s', 'Index template missing: %s') % os.path.relpath(TEMPLATE, ROOT), status='error')
        shutil.copy2(TEMPLATE, INDEX)
        made.append('references/_INDEX.xlsx')
    out = {'created': made}
    set_result('success', 0, L('Đã tạo %d mục trong khung thư viện.', 'Created %d items in the library frame.') % len(made) if made
               else L('Khung thư viện đã có đủ; không tạo gì.', 'The library is already set up; nothing created.'),
               changes=bool(made), files=[{'path': x, 'action': 'create'} for x in made],
               next_actions=[L('Đặt file vào thư mục nhóm/lĩnh vực, thêm dòng vào sheet "Danh mục", rồi chạy: python tools/library.py check',
                               'Put files in group/domain folders, add rows to the "Danh mục" sheet, then run: python tools/library.py check')])
    if args.get('json'):
        print(json.dumps(out, ensure_ascii=False))
        return 0
    if made:
        print(L('Đã tạo:', 'Created:'))
        for m in made:
            print('  - ' + m)
    else:
        print(L('Khung thư viện đã có đủ; không tạo gì.', 'The library is already set up; nothing created.'))
    print(L('Tiếp theo: đặt file vào thư mục nhóm/lĩnh vực, thêm dòng vào sheet "Danh mục", rồi chạy python tools/library.py check.',
            'Next: put files in group/domain folders, add rows to the "Danh mục" sheet, then run python tools/library.py check.'))
    return 0


def cmd_check(args):
    cols, data = load(require_columns=False)  # thiếu cột được báo thành lỗi bên dưới / missing columns are reported as findings below
    days = int(args.get('days') or 180)
    today = datetime.date.today()
    errs, warns = [], []
    missing_cols = [KEYS[k] for k in ('group', 'number', 'title', 'status', 'fulltext', 'file') if k not in cols]
    if missing_cols:
        errs.append((0, L('Chỉ mục thiếu cột: %s', 'Index is missing columns: %s') % ', '.join(missing_cols)))
    seen = {}
    ids = {}  # Mã -> các dòng; cùng quy tắc với library_sync (cắt khoảng trắng, phân biệt hoa thường) / Mã -> rows; same rule as library_sync (stripped, case-sensitive)
    indexed = set()
    for r in data:
        n = r['_row']
        ident = str(r.get('id') or '').strip() if 'id' in cols else ''
        if ident:
            ids.setdefault(ident, []).append(n)
        num = str(r.get('number') or '').strip()
        if not (num or str(r.get('title') or '').strip()):
            errs.append((n, L('thiếu cả Số, ký hiệu lẫn Trích yếu', 'both Number and Title are empty')))
        if not r.get('group'):
            warns.append((n, L('chưa ghi Nhóm', 'Group is empty')))
        key = fold(num)
        if key and not key.startswith('['):
            if key in seen:
                errs.append((n, L('trùng Số, ký hiệu "%s" với dòng %d', 'duplicate Number "%s" (row %d)') % (num, seen[key])))
            else:
                seen[key] = n
        path = norm_path(r.get('file'))
        ft = fold(r.get('fulltext'))
        if path:
            indexed.add(path)
            if '[' in path:
                warns.append((n, L('đường dẫn file còn chỗ trống: %s', 'file path has a placeholder: %s') % path))
            elif not os.path.isfile(os.path.join(REF, *path.split('/'))):
                (errs if ft == 'co' else warns).append((n, L('không thấy file references/%s', 'file not found: references/%s') % path))
        if fold(r.get('verify')) == 'da doi chieu toan van' and not (path and os.path.isfile(os.path.join(REF, *path.split('/')))):
            warns.append((n, L('ghi "Đã đối chiếu toàn văn" nhưng không có file trong thư viện', '"checked against full text" but no file in the library')))
        if ft == 'co' and not path:
            errs.append((n, L('Có toàn văn = Có nhưng không có đường dẫn file', 'Full text = Có but no file path')))
        if path and ft != 'co' and os.path.isfile(os.path.join(REF, *path.split('/'))):
            warns.append((n, L('có file nhưng cột Có toàn văn chưa ghi "Có"', 'file exists but Full text is not "Có"')))
        st = fold(r.get('status'))
        rb = str(r.get('replaced_by') or '').strip()
        if st == 'con hieu luc' and rb and rb not in ('—', '-'):
            d = as_date(rb)
            if d is None or d <= today:
                warns.append((n, L('Tình trạng "Còn hiệu lực" nhưng có ghi bị thay thế: %s', 'status "in force" but replaced by: %s') % rb))
        if not st:
            warns.append((n, L('chưa ghi Tình trạng', 'Status is empty')))
        chk = as_date(r.get('checked'))
        if r.get('group') and fold(r.get('group')) == 'phap luat':
            if chk is None:
                warns.append((n, L('văn bản pháp luật chưa ghi Ngày kiểm tình trạng', 'legal document without a status check date')))
            elif (today - chk).days > days:
                warns.append((n, L('kiểm tình trạng lần cuối %s, quá %d ngày', 'status last checked %s, over %d days ago') % (fmt(chk), days)))
        if fold(r.get('verify')) == 'can doi chieu':
            warns.append((n, L('mức độ xác minh: Cần đối chiếu', 'verification: needs checking')))
        if r.get('law_state') == 'NEEDS_REVIEW':
            warns.append((n, L('file đã đổi; cần đối chiếu revision mới', 'file changed; review the new revision')))
        if fold(r.get('sensitivity')) == 'mat':
            warns.append((n, L('văn bản Mật: không đưa lên dịch vụ đám mây', 'confidential: do not upload to cloud services')))
    for ident, rows in ids.items():
        if len(rows) > 1:
            errs.append((rows[1], L('trùng Mã "%s" ở các dòng %s. Mỗi nguồn cần một Mã riêng; library.py sync cũng từ chối sổ này.',
                                    'duplicate ID "%s" in rows %s. Each source needs its own ID; library.py sync also refuses this index.') % (ident, ', '.join(str(x) for x in rows))))
    unindexed = [f for f in library_files() if f not in indexed]
    for f in unindexed:
        warns.append((0, L('file chưa có trong chỉ mục: references/%s', 'file not in the index: references/%s') % f))
    if args.get('content_audit'):
        import library_audit
        warns.extend(library_audit.audit(ROOT, data, sys.modules[__name__]))
    out = {'rows': len(data), 'errors': [{'row': n, 'msg': m} for n, m in errs], 'warnings': [{'row': n, 'msg': m} for n, m in warns],
           'unindexed': unindexed}
    findings = [{'row': n or None, 'severity': 'error', 'message': m} for n, m in errs] + \
        [{'row': n or None, 'severity': 'warning', 'message': m} for n, m in warns]
    msg = L('Chỉ mục: %d dòng; %d lỗi, %d cảnh báo.', 'Index: %d rows; %d errors, %d warnings.') % (len(data), len(errs), len(warns))
    dup_step = L('Mã trùng: giữ một bản sao của references/_INDEX.xlsx, xác định dòng nào là dòng thật (đúng file, đã được kết quả dẫn chiếu), rồi sửa hoặc bỏ dòng còn lại trực tiếp trong Excel và chạy lại check. Đừng dựng lại sổ và đừng đổi Mã của dòng đã được dẫn chiếu.',
                 'Duplicate ID: keep a copy of references/_INDEX.xlsx, work out which row is the real one (the right file, already cited by results), then fix or drop the others directly in Excel and run check again. Do not rebuild the index and do not change the ID of a row that is already cited.')
    if errs:
        set_result('issues', 1, msg, changes=False, rows=len(data), errors=len(errs), warnings=len(warns), findings=findings,
                   next_actions=[dup_step] if any(len(v) > 1 for v in ids.values()) else None)
    else:
        set_result('success', 0, msg, changes=False, rows=len(data), errors=len(errs), warnings=len(warns), findings=findings)
    if args.get('json'):
        print(json.dumps(out, ensure_ascii=False))
    else:
        print(L('Chỉ mục: %d dòng (sheet "%s").', 'Index: %d rows ("%s" sheet).') % (len(data), SHEET))
        for tag, lst in ((L('LỖI', 'ERROR'), errs), (L('CẢNH BÁO', 'WARNING'), warns)):
            for n, m in lst:
                print('%s%s: %s' % (tag, (L(' dòng %d', ' row %d') % n) if n else '', m))
        print(L('\n%s: %d lỗi, %d cảnh báo.', '\n%s: %d errors, %d warnings.') % (L('ĐẠT', 'PASS') if not errs else L('KHÔNG ĐẠT', 'FAIL'), len(errs), len(warns)))
        if any(len(v) > 1 for v in ids.values()):
            print('  > ' + dup_step)
    return 1 if errs else 0


def show(r):
    path = norm_path(r.get('file'))
    has = fold(r.get('fulltext')) == 'co' and path and os.path.isfile(os.path.join(REF, *path.split('/')))
    return {'row': r['_row'], 'id': fmt(r.get('id')), 'group': fmt(r.get('group')), 'number': fmt(r.get('number')),
            'type': fmt(r.get('type')), 'title': fmt(r.get('title')), 'status': fmt(r.get('status')),
            'checked': fmt(r.get('checked')), 'verify': fmt(r.get('verify')), 'effective': fmt(r.get('effective')),
            'replaced_by': fmt(r.get('replaced_by')), 'articles': fmt(r.get('articles')),
            'file': (path[3:] if path.startswith('../') else 'references/' + path) if path else '', 'fulltext': bool(has),
            'revision': r.get('law_revision'), 'file_state': fmt(r.get('law_state'))}


def matches(r, words, grp):
    """Một bộ lọc cho CLI find và bảng điều khiển / one filter for the CLI find and the panel."""
    if grp and fold(r.get('group')) != grp:
        return False
    hay = fold(' '.join(str(r.get(k) or '') for k in ('id', 'number', 'title', 'domain', 'issuer', 'type', 'articles', 'notes', 'used_by')))
    return all(w in hay for w in words)


def revision_of(v):
    """Revision đã ghi: số nguyên, None khi trống, nguyên văn khi không hợp lệ / the recorded revision: int, None if empty, the text if invalid."""
    if v in (None, ''):
        return None
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v == int(v):
        return int(v)
    return str(v).strip()


def display_file(path):
    return (path[3:] if path.startswith('../') else 'references/' + path) if path else ''


def record(r, id_counts=None):
    """Siêu dữ liệu đã ghi trong sổ của một dòng, không đụng tới file nguồn (không stat, không hash).
    Metadata recorded in the index for one row. Never touches the source file (no stat, no hash)."""
    path = norm_path(r.get('file'))
    ident, revision = fmt(r.get('id')), revision_of(r.get('law_revision'))
    if not ident:
        identity = 'no_id'
    elif (id_counts or {}).get(ident, 1) > 1:
        identity = 'dup_id'
    elif not (isinstance(revision, int) and revision >= 1):
        identity = 'no_revision'
    else:
        identity = 'ok'
    return {'row': r['_row'], 'id': ident, 'revision': revision, 'identity': identity,
            'group': fmt(r.get('group')), 'domain': fmt(r.get('domain')), 'issuer': fmt(r.get('issuer')),
            'number': fmt(r.get('number')), 'type': fmt(r.get('type')), 'title': fmt(r.get('title')),
            'issued': fmt(r.get('issued')), 'effective': fmt(r.get('effective')),
            'replaces': fmt(r.get('replaces')), 'replaced_by': fmt(r.get('replaced_by')),
            'status': fmt(r.get('status')), 'checked': fmt(r.get('checked')), 'verify': fmt(r.get('verify')),
            'articles': fmt(r.get('articles')), 'sensitivity': fmt(r.get('sensitivity')), 'used_by': fmt(r.get('used_by')),
            'url': fmt(r.get('url')), 'notes': fmt(r.get('notes')),
            'fulltext_declared': fold(r.get('fulltext')) == 'co', 'locator': path, 'file': display_file(path),
            'file_state': fmt(r.get('law_state')), 'observed': fmt(r.get('law_observed')),
            'sha12': str(r.get('law_sha') or '').strip()[:12]}


def id_counts(data):
    counts = {}
    for r in data:
        ident = fmt(r.get('id'))
        if ident:
            counts[ident] = counts.get(ident, 0) + 1
    return counts


def browse(root=None, query='', group='', limit=BROWSE_LIMIT):
    """Duyệt/tìm/lọc nhóm trên sổ đã lưu. Chỉ đọc sổ một lần: không quét thư mục, không hash, không đọc nguồn, không ghi gì.
    Browse/search/group filter on the saved index. Reads the index once: no folder scan, no hash, no source read, no write."""
    book = read_book(root)
    if not book['ok']:
        return dict(book, message='' if book['kind'] == 'openpyxl_missing' else book_message(book))
    data = book['data']
    counts = id_counts(data)
    words, grp = fold(query).split(), fold(group)
    picked = [r for r in data if matches(r, words, grp)]
    seen = []
    for r in data:
        g = fmt(r.get('group'))
        if g and g not in seen:
            seen.append(g)
    return {'ok': True, 'snapshot': book['snapshot'], 'rows': [record(r, counts) for r in picked[:limit]],
            'matched': len(picked), 'truncated': len(picked) > limit, 'total': len(data),
            'groups': seen or [v[0] for v in GROUP_DIRS.values()],
            'summary': {'rows': len(data),
                        'declared_fulltext': sum(1 for r in data if fold(r.get('fulltext')) == 'co'),
                        'not_declared': sum(1 for r in data if fold(r.get('fulltext')) != 'co'),
                        'needs_check': sum(1 for r in data if fold(r.get('verify')) == 'can doi chieu'),
                        'identified': sum(1 for r in data if record(r, counts)['identity'] == 'ok')},
            'missing_columns': book['missing']}


def rebind(root, sel, book=None):
    """Đọc lại sổ và gắn lựa chọn của giao diện vào đúng dòng đã thấy. Trả (record, None) hoặc (None, lý do).
    Lý do: bad_selection | index:<loại lỗi sổ> | stale (sổ đã đổi, tải lại) | changed (dòng khác với lúc chọn).
    Re-read the index and bind a panel selection to the exact row that was seen: (record, None) or (None, reason)."""
    if not isinstance(sel, dict):
        return None, 'bad_selection'
    row, snap, locator, rev = sel.get('row'), sel.get('snapshot'), sel.get('locator'), sel.get('revision')
    if (type(row) is not int or not isinstance(snap, str) or not isinstance(locator, str) or not isinstance(sel.get('id'), str)
            or isinstance(rev, bool) or not (rev is None or isinstance(rev, (int, str)))):
        return None, 'bad_selection'
    book = book or read_book(root)
    if not book['ok']:
        return None, 'index:' + book['kind']
    if book['snapshot'] != snap:
        return None, 'stale'
    found = next((r for r in book['data'] if r['_row'] == row), None)
    if found is None:
        return None, 'stale'
    rec = record(found, id_counts(book['data']))
    if sel['id'] != rec['id'] or rev != rec['revision'] or locator != rec['locator']:
        return None, 'changed'
    return rec, None


def locate_source(root, locator):
    """Kiểm tiếp cận khi mở một nguồn: (đường dẫn đầy đủ, None) hoặc (None, lý do).
    Lý do: no_locator | bad_locator | out_of_scope | unsupported_type | link | missing | denied | not_regular | unreadable. Chỉ trong references/ hoặc thư mục
    ngoại lệ PROJECT.md hiện còn hiệu lực; từ chối tuyệt đối, '..', liên kết. Đây là kiểm tại thời điểm mở, không phải atomic no-follow.
    Access check when opening one source: only inside references/ or a current PROJECT.md exception folder; never atomic no-follow."""
    import library_sync
    # record() already normalised the workbook path. A second prefix removal
    # would open a different file when the locator itself starts with references/.
    loc = locator
    if not loc:
        return None, 'no_locator'
    parts = loc.split('/')
    old = parts[0] == '..'
    if ('[' in loc or ']' in loc or ':' in loc or loc.startswith('/') or any(p in ('', '.') for p in parts)
            or any(p == '..' for p in parts[1:]) or (old and len(parts) < 3)):
        return None, 'bad_locator'
    if any(part.startswith(('.', '_', '~$')) for part in (parts[2:] if old else parts)):
        return None, 'out_of_scope'  # sổ _INDEX.xlsx, file ẩn, file tạm của Office: không phải nguồn do sync quản lý / the index, hidden and Office temp files are not managed sources
    if os.path.splitext(loc)[1].lower() not in FILE_EXT:
        return None, 'unsupported_type'  # reuse the managed source formats; exclude unsupported program/script/shortcut suffixes
    ref = os.path.join(root, 'references')
    full = os.path.normpath(os.path.join(ref, *parts))
    if old:
        try:
            zones = library_sync.scopes(root)
        except library_sync.SyncError:
            return None, 'out_of_scope'  # phạm vi không rõ thì không mở / an unclear scope never opens
        if not library_sync.in_scope(root, loc, zones):
            return None, 'out_of_scope'
        base = next((z['real'] for z in zones if parts[1] == z['path'].split('/')[0] and loc[3:].startswith(z['path'] + '/')), None)
        if base is None:
            return None, 'out_of_scope'
    else:
        base = os.path.realpath(ref)
    if not library_sync.link_free(full, root) or not library_sync.contained(full, base):
        return None, 'link'
    try:
        st = os.lstat(full)
    except FileNotFoundError:
        return None, 'missing'
    except PermissionError:
        return None, 'denied'
    except OSError:
        return None, 'unreadable'
    if stat.S_ISLNK(st.st_mode) or (getattr(st, 'st_file_attributes', 0) or 0) & 0x400:
        return None, 'link'
    if not stat.S_ISREG(st.st_mode):
        return None, 'not_regular'
    return full, None


def select_text(recs, snapshot, lang='vi'):
    """Văn bản sao chép để giao việc: danh tính + locator + trạng thái đã ghi, nhắc đối chiếu. Không phải duyệt hay gửi.
    Text to paste into a task: identity, locator and recorded status, with a reminder to check. Not an approval, not a dispatch."""
    en = lang == 'en'
    none = 'not recorded' if en else 'chưa ghi nhận'
    lines = ['Task sources (from references/_INDEX.xlsx, index SHA-256 %s):' % snapshot[:12] if en
             else 'Nguồn giao việc (từ references/_INDEX.xlsx, SHA-256 sổ %s):' % snapshot[:12]]
    for r in recs:
        lines.append('- %s | revision %s | %s | %s %s | %s %s | %s %s' % (
            r['id'], r['revision'], r['file'] or r['locator'],
            'Status:' if en else 'Tình trạng:', r['status'] or none,
            'Verification:' if en else 'Xác minh:', r['verify'] or none,
            'Recorded file state:' if en else 'Trạng thái file đã ghi:', r['file_state'] or none))
    lines.append('Check the ID, revision and path against the current index before use; this list verifies nothing about the content.' if en
                 else 'Trước khi dùng, đối chiếu ID, revision và đường dẫn với sổ hiện hành; danh sách này không xác minh nội dung.')
    return '\n'.join(lines)


# ---------------------------------------------------------------- nội dung đã ghi theo phiên bản nguồn (B-01) / recorded content per source version (B-01)
CONTENT_SHEET = 'LAW content'
CONTENT_MARKER = 'law.source-content/1'
CONTENT_HEADER_ROW = 4
CONTENT_HEADERS = ('entry_id', 'source_id', 'source_revision', 'source_sha256', 'source_path', 'locator', 'read_scope', 'kind',
                   'text', 'limits', 'recorded_by', 'recorded_at', 'check_level', 'check_scope', 'check_evidence_ref',
                   'checked_by', 'checked_at', 'user_decision', 'user_confirmation_ref', 'supersedes', 'basis_status',
                   'basis_replaces', 'basis_replaced_by', 'basis_effective')
# Cột do người gọi khai; phần còn lại (source_path, recorded_at, basis_*) do công cụ lấy từ sổ lúc ghi.
# Columns the caller declares; the rest (source_path, recorded_at, basis_*) is taken from the index by the tool at write time.
CONTENT_REQUEST_KEYS = ('entry_id', 'source_id', 'source_revision', 'source_sha256', 'locator', 'read_scope', 'kind', 'text',
                        'limits', 'recorded_by', 'check_level', 'check_scope', 'check_evidence_ref', 'checked_by', 'checked_at',
                        'user_decision', 'user_confirmation_ref', 'supersedes')
CONTENT_REQUIRED = ('entry_id', 'source_id', 'source_revision', 'source_sha256', 'locator', 'read_scope', 'kind', 'text',
                    'recorded_by')
CONTENT_KINDS = ('EXTRACT', 'FACT', 'CALCULATION', 'INFERENCE')
CHECK_LEVELS = ('NOT_CHECKED', 'DECLARED_CHECK')
SEMANTIC_FIELDS = (('status', 'basis_status'), ('replaces', 'basis_replaces'),
                   ('replaced_by', 'basis_replaced_by'), ('effective', 'basis_effective'))
CELL_LIMIT = 32767
SHA256_RE = re.compile(r'[0-9a-f]{64}')
BAD_CHARS = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]')


class ContentError(Exception):
    """Từ chối có mã của đường nội dung/sự kiện nguồn; chưa ghi gì. / A coded refusal of the content or source-event path; nothing was written."""

    def __init__(self, code, vi, en):
        Exception.__init__(self, '[%s] %s' % (code, en if LANG == 'en' else vi))
        self.code = code


def _text(cell):
    value = cell.value
    return '' if value is None else (value if isinstance(value, str) else str(value))


def _malformed(row, vi, en):
    return ContentError('MALFORMED_CONTENT_SHEET', 'Sheet "%s" không hợp lệ (dòng %s): %s' % (CONTENT_SHEET, row, vi),
                        'Sheet "%s" is malformed (row %s): %s' % (CONTENT_SHEET, row, en))


def parse_content(ws):
    """Đọc sheet nội dung, chỉ đọc. Marker/cột/công thức/khóa/liên kết nội bộ sai thì ContentError, không coi là rỗng.
    Read the content sheet without writing. A bad marker, column, formula, key or internal link raises; it is never treated as empty."""
    rows = list(ws.iter_rows(min_row=1))
    width = len(CONTENT_HEADERS)
    if len(rows) < CONTENT_HEADER_ROW or not rows[0] or _text(rows[0][0]) != CONTENT_MARKER:
        raise _malformed(1, 'thiếu hoặc sai marker A1 (%s)' % CONTENT_MARKER, 'A1 marker missing or wrong (%s)' % CONTENT_MARKER)
    if getattr(rows[0][0], 'data_type', 's') == 'f':
        raise _malformed(1, 'A1 là công thức', 'A1 is a formula')
    header = rows[CONTENT_HEADER_ROW - 1]
    if tuple(_text(c) for c in header[:width]) != CONTENT_HEADERS or any(_text(c) for c in header[width:]) or \
            any(getattr(c, 'data_type', 's') == 'f' for c in header):
        raise _malformed(CONTENT_HEADER_ROW, 'tiêu đề cột khác phiên bản đã tài liệu hóa', 'headers differ from the documented version')
    entries, by_id, targets = [], {}, {}
    for number, cells in enumerate(rows[CONTENT_HEADER_ROW:], CONTENT_HEADER_ROW + 1):
        if not any(c.value not in (None, '') for c in cells):
            continue
        if any(getattr(c, 'data_type', 's') == 'f' for c in cells):
            raise _malformed(number, 'ô được quản lý là công thức', 'a managed cell is a formula')
        if any(_text(c) for c in cells[width:]):
            raise _malformed(number, 'có dữ liệu ngoài vùng cột được quản lý', 'data outside the managed columns')
        entry = {key: (_text(cells[i]) if i < len(cells) else '') for i, key in enumerate(CONTENT_HEADERS)}
        raw_revision = cells[2].value if len(cells) > 2 else None
        if isinstance(raw_revision, bool) or not (isinstance(raw_revision, int) or
                                                  (isinstance(raw_revision, float) and raw_revision.is_integer())) or raw_revision < 1:
            raise _malformed(number, 'source_revision phải là số nguyên dương', 'source_revision must be a positive integer')
        entry['source_revision'] = int(raw_revision)
        entry['_row'] = number
        if not entry['entry_id'] or not entry['source_id']:
            raise _malformed(number, 'thiếu entry_id hoặc source_id', 'entry_id or source_id is empty')
        if any(not entry[key].strip() for key in CONTENT_REQUIRED if key not in ('source_revision', 'source_sha256')) or \
                not entry['source_path'].strip() or not entry['recorded_at'].strip():
            raise _malformed(number, 'thiếu trường nội dung hoặc xuất xứ bắt buộc', 'a required content or provenance field is empty')
        if entry['entry_id'] in by_id:
            raise _malformed(number, 'trùng entry_id ' + entry['entry_id'], 'duplicate entry_id ' + entry['entry_id'])
        if not SHA256_RE.fullmatch(entry['source_sha256']):
            raise _malformed(number, 'source_sha256 không phải 64 ký tự hex', 'source_sha256 is not 64 hex characters')
        if entry['kind'] not in CONTENT_KINDS or entry['check_level'] not in CHECK_LEVELS:
            raise _malformed(number, 'kind hoặc check_level ngoài danh sách', 'kind or check_level is outside the allowed list')
        if entry['check_level'] == 'DECLARED_CHECK' and not all(entry[k].strip() for k in ('check_scope', 'check_evidence_ref', 'checked_by', 'checked_at')):
            raise _malformed(number, 'DECLARED_CHECK thiếu phạm vi/chứng cứ/người/thời điểm', 'DECLARED_CHECK lacks scope, evidence, actor or time')
        if entry['check_level'] == 'NOT_CHECKED' and any(entry[k] for k in ('check_scope', 'check_evidence_ref', 'checked_by', 'checked_at')):
            raise _malformed(number, 'NOT_CHECKED có thông tin kiểm', 'NOT_CHECKED carries check details')
        if bool(entry['user_decision'].strip()) != bool(entry['user_confirmation_ref'].strip()):
            raise _malformed(number, 'user_decision và user_confirmation_ref phải đi cùng nhau', 'user_decision and user_confirmation_ref must come together')
        target = entry['supersedes']
        if target:
            old = by_id.get(target)
            if old is None or old['source_id'] != entry['source_id']:
                raise _malformed(number, 'supersedes phải trỏ entry đứng trước cùng SourceID', 'supersedes must point to an earlier entry of the same SourceID')
            if target in targets:
                raise _malformed(number, 'entry %s bị thay bởi hai dòng (nhánh mơ hồ)' % target, 'entry %s is superseded twice (ambiguous branch)' % target)
            targets[target] = entry['entry_id']
        by_id[entry['entry_id']] = entry
        entries.append(entry)
    return entries


def basis_state(entry, catalog):
    """(trạng thái, lý do) của căn cứ một entry so với sổ hiện tại; chỉ so snapshot đã ghi, không đọc nguồn.
    CURRENT_CATALOG_BASIS chỉ nghĩa là khớp sổ, không chứng minh bytes hiện tại, hiệu lực hay nội dung đúng.
    (state, reason) of an entry's basis against the current index; recorded snapshots only, no source read."""
    recs = catalog.get(entry['source_id']) or []
    if not recs:
        return 'MISSING_SOURCE', 'source_not_in_catalog'
    if len(recs) > 1:
        return 'MISSING_BASIS', 'duplicate_source_id'
    rec = recs[0]
    revision, sha = revision_of(rec.get('law_revision')), str(rec.get('law_sha') or '').strip().lower()
    if not (isinstance(revision, int) and revision >= 1 and SHA256_RE.fullmatch(sha)):
        return 'MISSING_BASIS', 'no_valid_revision_or_sha'
    if revision > entry['source_revision']:
        return 'NEEDS_REVIEW', 'catalog_revision_newer'
    if revision < entry['source_revision']:
        return 'BASIS_CONFLICT', 'catalog_revision_older'
    if sha != entry['source_sha256']:
        return 'BASIS_CONFLICT', 'sha_differs_same_revision'
    if any(fmt(rec.get(key)) != entry[snap] for key, snap in SEMANTIC_FIELDS):
        return 'BASIS_CONFLICT', 'semantic_differs_same_revision'
    return 'CURRENT_CATALOG_BASIS', 'matches_catalog'


def content_find(root=None, query='', source_id='', entry_id=''):
    """Tra nội dung đã ghi, chỉ đọc sổ: không quét/hash nguồn, không OCR, không model, không mạng. Trạng thái theo các chiều độc lập:
    basis_state, superseded_by, recorded_file_state, bytes_observation (luôn UNOBSERVED).
    Look up recorded content from the workbook only. State on independent dimensions; bytes are never observed here."""
    try:
        import openpyxl as opx
    except ImportError:
        return {'ok': False, 'kind': 'openpyxl_missing', 'message': ''}
    try:
        with open(index_path(root), 'rb') as f:
            raw = f.read()
    except FileNotFoundError:
        return {'ok': False, 'kind': 'no_index', 'message': book_message({'kind': 'no_index'})}
    except OSError:
        return {'ok': False, 'kind': 'unreadable', 'message': book_message({'kind': 'unreadable'})}
    book = parse_book(raw, require_columns=False, data_only=False)
    if not book['ok']:
        return dict(book, message=book_message(book))
    try:
        wb = opx.load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    except Exception:
        return {'ok': False, 'kind': 'unreadable', 'message': book_message({'kind': 'unreadable'})}
    try:
        present = CONTENT_SHEET in wb.sheetnames
        entries = parse_content(wb[CONTENT_SHEET]) if present else []
    except ContentError as e:
        return {'ok': False, 'kind': 'content_malformed', 'code': e.code, 'message': str(e)}
    finally:
        wb.close()
    catalog = {}
    for rec in book['data']:
        ident = fmt(rec.get('id'))
        if ident:
            catalog.setdefault(ident, []).append(rec)
    replaced = {e['supersedes']: e['entry_id'] for e in entries if e['supersedes']}
    words = fold(query).split()
    out = []
    for entry in entries:
        if (source_id and entry['source_id'] != source_id) or (entry_id and entry['entry_id'] != entry_id):
            continue
        hay = fold(' '.join(entry[k] for k in ('entry_id', 'source_id', 'locator', 'read_scope', 'kind', 'text', 'limits',
                                               'recorded_by', 'check_scope', 'user_decision', 'source_path')))
        if not all(w in hay for w in words):
            continue
        state, reason = basis_state(entry, catalog)
        recs = catalog.get(entry['source_id']) or []
        one = recs[0] if len(recs) == 1 else None
        view = {k: entry[k] for k in CONTENT_HEADERS}
        view.update(row=entry['_row'], basis_state=state, basis_reason=reason, superseded_by=replaced.get(entry['entry_id'], ''),
                    recorded_file_state=fmt(one.get('law_state')) if one else '', bytes_observation='UNOBSERVED',
                    current_source_path=norm_path(one.get('file')) if one else '',
                    check_note=L('DECLARED_CHECK: lời khai kiểm theo phạm vi đã ghi, không phải công cụ đã xác minh',
                                 'DECLARED_CHECK: a declared check within the recorded scope, not tool-verified')
                    if entry['check_level'] == 'DECLARED_CHECK' else '')
        out.append(view)
    return {'ok': True, 'snapshot': book['snapshot'], 'content_sheet': 'present' if present else 'absent',
            'total': len(entries), 'matched': len(out), 'entries': out, 'bytes_observation': 'UNOBSERVED'}


def cmd_find(args):
    q = fold(args.get('query') or '')
    if not q:
        stop(L('Cần từ khóa. Ví dụ: python tools/library.py find "30/2020"', 'A keyword is required. Example: python tools/library.py find "30/2020"'))
    _, data = load()
    grp = fold(args.get('group') or '')
    words = q.split()
    hits = [show(r) for r in data if matches(r, words, grp)]
    set_result('success', 0, L('%d văn bản khớp "%s".', '%d documents match "%s".') % (len(hits), args.get('query')), changes=False,
               query=args.get('query'), count=len(hits), rows=[h['row'] for h in hits])
    limit = None if args.get('all') else 30
    if args.get('json'):
        print(json.dumps({'query': args.get('query'), 'count': len(hits), 'hits': hits[:limit]}, ensure_ascii=False))
        return 0
    if not hits:
        print(L('Không thấy văn bản khớp "%s" trong chỉ mục. Chưa có trong thư viện thì không trích nguyên văn; tra nguồn chính thống và thêm vào chỉ mục.',
                'No document matches "%s". Not in the library means no verbatim quotes; look it up officially and add it to the index.') % args.get('query'))
        return 0
    print(L('%d văn bản khớp "%s":', '%d documents match "%s":') % (len(hits), args.get('query')))
    for h in hits[:limit]:
        print('- [%s %d] %s - %s' % (L('dòng', 'row'), h['row'], ' '.join(x for x in (h['type'], h['number']) if x), h['title']))
        print('    %s: %s%s | %s: %s | %s: %s' % (
            L('Tình trạng', 'Status'), h['status'] or '?', (' (%s %s)' % (L('kiểm', 'checked'), h['checked'])) if h['checked'] else '',
            L('Xác minh', 'Verified'), h['verify'] or '?',
            L('Toàn văn', 'Full text'), h['file'] if h['fulltext'] else L('không có, không trích nguyên văn', 'none, no verbatim quotes')))
        if h['replaced_by']:
            print('    %s: %s' % (L('Bị thay thế bởi', 'Replaced by'), h['replaced_by']))
        if h['articles']:
            print('    %s: %s' % (L('Điều hay dùng', 'Key articles'), h['articles']))
    if limit and len(hits) > limit:
        print(L('... còn %d dòng; thêm --all để xem hết.', '... %d more; add --all to see all.') % (len(hits) - limit))
    return 0


def cmd_stats(args):
    _, data = load()
    groups = {}
    for r in data:
        g = fmt(r.get('group')) or L('(chưa ghi nhóm)', '(no group)')
        s = fmt(r.get('status')) or L('(chưa ghi)', '(empty)')
        groups.setdefault(g, {}).setdefault(s, 0)
        groups[g][s] += 1
    no_ft = sum(1 for r in data if fold(r.get('fulltext')) != 'co')
    verify = sum(1 for r in data if fold(r.get('verify')) == 'can doi chieu')
    out = {'rows': len(data), 'groups': groups, 'no_fulltext': no_ft, 'needs_check': verify,
           'files': len(library_files()), 'index': os.path.isfile(INDEX)}
    set_result('success', 0, L('Chỉ mục: %d văn bản; %d file trong references/.', 'Index: %d documents; %d files in references/.') % (len(data), out['files']),
               changes=False, rows=len(data), groups=groups, no_fulltext=no_ft, needs_check=verify,
               reference_files=out['files'], index=out['index'])
    if args.get('json'):
        print(json.dumps(out, ensure_ascii=False))
        return 0
    print(L('Chỉ mục: %d văn bản; %d file trong references/; %d chưa có toàn văn; %d cần đối chiếu.',
            'Index: %d documents; %d files in references/; %d without full text; %d need checking.') % (len(data), out['files'], no_ft, verify))
    for g, st in groups.items():
        print('- %s: %s' % (g, ', '.join('%s %d' % (k, v) for k, v in st.items())))
    return 0


def cmd_sync(args):
    import library_sync
    if args.get('query'):
        stop(L('sync không nhận đường dẫn riêng; dùng sync [--apply] [--rename ID=đường_dẫn#sha256] [--json].',
               'sync takes no path argument; use sync [--apply] [--rename ID=path#sha256] [--json].'))
    renames = args.get('rename') or []
    try:
        result = (library_sync.apply(ROOT, renames) if args.get('apply') else library_sync.plan(ROOT, renames))
    except library_sync.SyncError as e:
        stop(str(e), next_actions=[L('Kiểm nguồn/sổ, giải quyết xung đột rồi chạy lại sync.',
                                    'Check the source/index, resolve the conflict, then rerun sync.')])
    count = len(result['changes'])
    changed = result['bytes'] is not None
    label = (L('Đã cập nhật sổ nguồn.', 'Source index updated.') if args.get('apply') and changed else
             L('Xem trước cập nhật sổ nguồn.', 'Source index preview.') if changed else
             L('Sổ nguồn không đổi.', 'Source index unchanged.'))
    set_result('success' if args.get('apply') or not changed else 'preview', 0, label,
               changes=bool(args.get('apply') and changed), snapshot=result['snapshot'],
               planned=count, rows=result['changes'], skipped=result['skipped'],
               candidates=result['candidates'], scopes=result['scopes'],
               files=[{'path': 'references/_INDEX.xlsx', 'action': 'update'}] if changed else [],
               next_actions=[L('Xem các thay đổi, rồi chạy: python tools/library.py sync --apply',
                               'Review the changes, then run: python tools/library.py sync --apply')]
               if changed and not args.get('apply') else [])
    out = {'planned': count, 'changes': result['changes'], 'skipped': result['skipped'],
           'candidates': result['candidates'], 'scopes': result['scopes'],
           'index_changed': changed, 'applied': bool(args.get('apply') and changed), 'snapshot': result['snapshot'],
           'notices': result.get('notices') or [], 'protection': result.get('protection')}
    if args.get('json'):
        print(json.dumps(out, ensure_ascii=False))
    else:
        print('\n'.join(sync_lines(result, label, bool(args.get('apply')))))
    return 0


def sync_lines(result, label, apply, lang=None):
    """Chữ người đọc của sync, dùng chung cho CLI và bảng điều khiển / the human text of sync, shared by the CLI and the panel."""
    def T(vi, en):
        return en if (lang or LANG) == 'en' else vi
    changed = result['bytes'] is not None
    notices = result.get('notices') or []
    lines = [label]
    for scope in result['scopes']:
        lines.append('  %s: %s' % (T('Phạm vi nguồn cũ (PROJECT.md)', 'Old-source scope (PROJECT.md)'), scope))
    for item in result['changes']:
        line = '  %s: %s' % (item['action'], item['path'])
        if item.get('from'):
            line += ' <- %s (%s)' % (item['from'], item['id'])
        lines.append(line)
    for path, reason in result['skipped']:
        lines.append('  %s: %s (%s)' % (T('Bỏ qua', 'Skipped'), path, reason))
    for item in result['candidates']:
        lines.append('  %s: %s -> %s%s' % (T('Có thể là đổi tên, chưa áp dụng', 'Possible rename, not applied'),
                                          item['from'], item['to'],
                                          T(' [MƠ HỒ]', ' [AMBIGUOUS]') if item['ambiguous'] else ''))
        lines.append('    --rename "%s"' % item['token'])
    if changed and 'x14_data_validation' in notices:
        lines.append(T('Lưu ý: sổ có danh sách thả xuống do Excel lưu ở dạng mở rộng (x14); ghi lại sổ sẽ làm mất chúng nên `sync --apply` sẽ từ chối, chưa ghi gì. Công cụ không tự dựng lại.',
                       'Note: the index has drop-down lists that Excel saved in an extended form (x14); rewriting it would lose them, so `sync --apply` will refuse and write nothing. The tool does not recreate them.'))
    if changed and 'x14_scan_incomplete' in notices:
        lines.append(T('Lưu ý: chưa kiểm xong sổ có danh sách thả xuống x14 hay không; `sync --apply` sẽ từ chối, chưa ghi gì.',
                       'Note: checking the index for x14 drop-down lists did not finish; `sync --apply` will refuse and write nothing.'))
    if result.get('protection') == 'recheck_only':
        lines.append(T('Mức bảo vệ khi ghi: chỉ kiểm lại sát trước khi ghi (máy này không giữ được handle thư mục); còn một cửa sổ rất ngắn chưa đóng.',
                       'Write protection: re-check right before the write only (this machine cannot hold folder handles); a very short window stays open.'))
    if changed and not apply:
        lines.append(T('Thêm --apply sau khi xem trước.', 'Add --apply after reviewing the preview.'))
    return lines


def option(rest, name, default=''):
    """Lấy --name VALUE ra khỏi rest (thiếu giá trị thì trả chuỗi rỗng). / Take --name VALUE out of rest."""
    if name not in rest:
        return default
    i = rest.index(name)
    value = rest[i + 1] if i + 1 < len(rest) else ''
    del rest[i:i + 2]
    return value


def content_summary(res):
    """Phần JSON an toàn của một kết quả plan/apply: không chứa bytes workbook. / The JSON-safe part of a plan/apply result: no workbook bytes."""
    keep = ('action', 'entry_id', 'event_id', 'source_id', 'mode', 'plan_token', 'snapshot', 'sheet_created', 'revision_before',
            'revision_after', 'observed_before', 'current_basis_validated', 'notices', 'protection')
    out = {k: res[k] for k in keep if k in res}
    out['index_changed'] = res.get('bytes') is not None
    return out


def print_content_entries(res):
    print(L('%d/%d entry nội dung khớp (chỉ đọc sổ; bytes nguồn: UNOBSERVED, chưa kiểm file).',
            '%d/%d content entries match (workbook only; source bytes: UNOBSERVED, no file checked).') % (res['matched'], res['total']))
    if res['content_sheet'] == 'absent':
        print(L('Sổ chưa có sheet "%s": chưa ghi nội dung nào.', 'The workbook has no "%s" sheet: nothing recorded yet.') % CONTENT_SHEET)
    for e in res['entries']:
        print('- %s | %s rev %s | %s | %s' % (e['entry_id'], e['source_id'], e['source_revision'], e['kind'], e['locator']))
        print('    %s: %s (%s) | %s: %s | %s: %s' % (
            L('Căn cứ', 'Basis'), e['basis_state'], e['basis_reason'], L('Bị thay bởi', 'Superseded by'),
            e['superseded_by'] or '-', L('File đã ghi', 'Recorded file state'), e['recorded_file_state'] or '-'))
        print('    %s: %s | %s: %s' % (L('Kiểm', 'Check'), e['check_level'], L('Quyết định user', 'User decision'), e['user_decision'] or '-'))
        if e['check_note']:
            print('    ' + e['check_note'])
        print('    %s: %s' % (L('Giới hạn', 'Limits'), e['limits'] or '-'))
        print('    %s: %s' % (L('Nội dung (đã khai, chưa xác minh)', 'Text (declared, unverified)'), e['text']))


def cmd_content_find(rest, as_json):
    source_id, entry_id = option(rest, '--source-id'), option(rest, '--entry-id')
    res = content_find(ROOT, ' '.join(rest), source_id, entry_id)
    if not res['ok']:
        if res['kind'] == 'openpyxl_missing':
            need_openpyxl()
        stop(res['message'], status='error' if res['kind'] in ('denied', 'unreadable') else 'blocked', reason=res['kind'])
    set_result('success', 0, L('%d entry nội dung khớp.', '%d content entries match.') % res['matched'], changes=False,
               snapshot=res['snapshot'], matched=res['matched'], total=res['total'], bytes_observation='UNOBSERVED')
    if as_json:
        print(json.dumps(res, ensure_ascii=False))
    else:
        print_content_entries(res)
    return 0


def cmd_content(cmd, rest):
    """content preview|apply|find và source-event preview|apply. Cùng writer library_sync; kết quả không phải PASS/VERIFIED/quyền.
    content preview|apply|find and source-event preview|apply, through the same library_sync writer; a result is not PASS, VERIFIED or a grant."""
    import library_sync
    as_json = '--json' in rest
    rest = [a for a in rest if a != '--json']
    sub = rest.pop(0) if rest else ''
    if cmd == 'content' and sub == 'find':
        return cmd_content_find(rest, as_json)
    if sub not in ('preview', 'apply'):
        stop(L('Cú pháp: %s preview|apply --request FILE [--expect-plan TOKEN]', 'Usage: %s preview|apply --request FILE [--expect-plan TOKEN]') % cmd)
    path, token = option(rest, '--request'), option(rest, '--expect-plan')
    if not path or rest:
        stop(L('Cần --request FILE và không nhận tham số thừa.', 'Needs --request FILE and takes no extra arguments.'))
    if sub == 'apply' and not token:
        stop(L('apply cần --expect-plan TOKEN lấy từ bản xem trước.', 'apply needs --expect-plan TOKEN from the preview.'))
    plan_fn, apply_fn = ((library_sync.content_plan, library_sync.content_apply) if cmd == 'content'
                         else (library_sync.event_plan, library_sync.event_apply))
    try:
        request = library_sync.load_request(path)
        res = apply_fn(ROOT, request, token) if sub == 'apply' else plan_fn(ROOT, request)
    except (library_sync.SyncError, library_sync.ContentError, ContentError) as e:
        stop(str(e), reason=getattr(e, 'code', None) or getattr(e, 'reason', 'refused'),
             next_actions=[L('Sửa yêu cầu hoặc đồng bộ nguồn, rồi xem trước lại.', 'Fix the request or sync the source, then preview again.')])
    changed = res.get('bytes') is not None
    done = sub == 'apply' and changed
    labels = {'record': (L('Đã ghi entry nội dung.', 'Content entry recorded.'), L('Xem trước: sẽ ghi entry nội dung.', 'Preview: the content entry would be recorded.')),
              'event': (L('Đã ghi source-event.', 'Source-event recorded.'), L('Xem trước: sẽ ghi source-event.', 'Preview: the source-event would be recorded.'))}
    same = {'already_recorded': L('Đã ghi từ trước, không ghi lại (không xác nhận căn cứ hiện tại).',
                                  'Already recorded; nothing written (this does not confirm the current basis).'),
            'no_change': L('Không có thay đổi ngữ nghĩa; không tăng revision.', 'No semantic change; revision not increased.')}
    label = same[res['action']] if res['action'] in same else labels[res['action']][0 if sub == 'apply' else 1]
    set_result('success' if sub == 'apply' or not changed else 'preview', 0, label, changes=done,
               files=[{'path': 'references/_INDEX.xlsx', 'action': 'update'}] if changed else [],
               plan_token=res.get('plan_token'), snapshot=res.get('snapshot'), action=res['action'],
               next_actions=[L('Chạy lại với apply --expect-plan TOKEN.', 'Run again with apply --expect-plan TOKEN.')]
               if changed and sub == 'preview' else [])
    out = dict(content_summary(res), applied=done)
    if as_json:
        print(json.dumps(out, ensure_ascii=False))
        return 0
    print(label)
    for key in ('entry_id', 'event_id', 'source_id', 'revision_before', 'revision_after', 'plan_token'):
        if key in out:
            print('  %s: %s' % (key, out[key]))
    for notice in out.get('notices') or []:
        print('  ' + L('Lưu ý', 'Notice') + ': ' + notice)
    if sub == 'preview' and changed:
        print(L('Thêm apply --expect-plan TOKEN sau khi xem trước.', 'Add apply --expect-plan TOKEN after reviewing the preview.'))
    return 0


def main(argv):
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8', errors='replace')
    global LANG
    try:
        info = json.load(io.open(os.path.join(ROOT, '.agent', 'PACKAGE_INFO.json'), encoding='utf-8'))
        LANG = 'en' if info.get('language') == 'en' else 'vi'
    except (OSError, ValueError):
        pass
    if '--lang' in argv:
        i = argv.index('--lang')
        LANG = argv[i + 1] if i + 1 < len(argv) else LANG
        del argv[i:i + 2]
    if not argv or argv[0] in ('-h', '--help'):
        print(__doc__)
        return 0
    cmd, rest = argv[0], argv[1:]
    if cmd == 'intake':
        import library_intake
        return library_intake.command(rest, sys.modules[__name__])
    if cmd in ('content', 'source-event'):
        return cmd_content(cmd, rest)
    args = {'json': '--json' in rest, 'all': '--all' in rest, 'apply': '--apply' in rest}
    if cmd == 'check' and '--content-audit' in rest:
        args['content_audit'] = True
        rest.remove('--content-audit')
    rest = [a for a in rest if a not in ('--json', '--all', '--apply')]
    for opt in ('--days', '--group'):
        if opt in rest:
            i = rest.index(opt)
            args[opt[2:]] = rest[i + 1] if i + 1 < len(rest) else ''
            del rest[i:i + 2]
    while '--rename' in rest:  # có thể lặp / may repeat
        i = rest.index('--rename')
        args.setdefault('rename', []).append(rest[i + 1] if i + 1 < len(rest) else '')
        del rest[i:i + 2]
    args['query'] = ' '.join(rest)
    fn = {'init': cmd_init, 'check': cmd_check, 'find': cmd_find, 'stats': cmd_stats, 'sync': cmd_sync}.get(cmd)
    if not fn:
        msg = L('Lệnh không có: %s. Dùng: init | check | find | stats | sync | content | source-event', 'Unknown command: %s. Use: init | check | find | stats | sync | content | source-event') % cmd
        print(msg)
        set_result('blocked', 2, msg, changes=False)
        return 2
    return fn(args)


def finish(code, msg, fallback):
    """Phát đúng một kết quả; mã trong kết quả luôn bằng mã thoát / Emit exactly one result; its code always equals the exit code."""
    res = LAST_RESULT
    if res is None:
        res = make(fallback, code=code, message=msg)
    elif res['code'] != code:
        res = make('error', code=code, message=L('Mã kết quả không khớp mã thoát.', 'Result code does not match the exit code.'))
    emit(res)
    return code


def run_with_result(argv):
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    code, msg, fallback = 0, '', 'success'
    try:
        code = main(argv) or 0
    except SystemExit as e:
        code = e.code if isinstance(e.code, int) else (0 if e.code is None else 1)
        fallback = 'success' if code == 0 else 'blocked'
        if isinstance(e.code, str):
            msg = e.code
            print(msg, file=sys.stderr)  # như sys.exit(msg) / same as sys.exit(msg)
        elif code:  # argparse đã in lỗi cú pháp / argparse has printed the usage error
            msg = L('Tham số không hợp lệ; xem trợ giúp của lệnh.', 'Invalid arguments; see the command help.')
    except OSError as e:
        code, msg, fallback = 1, str(e), 'error'
        if CREATED:
            set_result('partial', 1, msg, changes=True, files=[{'path': x, 'action': 'create'} for x in CREATED],
                       next_actions=[L('Chạy lại python tools/library.py init: chỉ tạo phần còn thiếu, không ghi đè.',
                                       'Run python tools/library.py init again: it only creates what is still missing, never overwrites.')])
            print('CHƯA HOÀN TẤT / INCOMPLETE: %s. Có thể đã tạo một số mục; chưa tự hoàn tác. / Some items may have been created; no automatic rollback.' % e, file=sys.stderr)
        else:
            print('Lỗi / Error: %s' % e, file=sys.stderr)
    except Exception as e:  # lỗi bất ngờ vẫn phát kết quả / unexpected errors still emit a result
        traceback.print_exc()
        code, msg, fallback = 1, repr(e), 'error'
    return finish(code, msg, fallback)


if __name__ == '__main__':
    if RESULT_REQUESTED:
        sys.exit(run_with_result([x for x in sys.argv[1:] if x != '--result-json']))
    sys.exit(main(sys.argv[1:]))
