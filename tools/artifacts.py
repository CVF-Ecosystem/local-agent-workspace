#!/usr/bin/env python3
"""Phiên bản kết quả (artifact) và quan hệ nguồn–phụ lục–kết quả (B-02). / Artifact versions and source–appendix–result relations (B-02).

Dùng / Usage (chạy từ thư mục gốc project; --project DIR chỉ để trỏ project khác / run from the project root):
  python tools/artifacts.py find [từ khóa] [--artifact-id ID] [--source-id ID] [--revision N] [--json]
  python tools/artifacts.py register preview --request FILE | apply --request FILE --expect-plan TOKEN
  python tools/artifacts.py relate   preview --request FILE | apply --request FILE --expect-plan TOKEN
  python tools/artifacts.py prepare --request FILE [--json]
  python tools/artifacts.py open --artifact-id ID --revision N [--expect-snapshot HEX]
  python tools/artifacts.py copy --artifact-id ID --revision N [--expect-snapshot HEX] [--lang vi|en]

Hai sheet nằm trong references/_INDEX.xlsx, chỉ tạo khi ghi lần đầu: "LAW artifacts" (phiên bản kết quả, chỉ thêm dòng) và "LAW relations"
(RELATED nguồn↔nguồn, DERIVED_FROM từ một phiên bản kết quả tới nguồn hoặc phiên bản kết quả khác). Ghi qua đúng library_sync.locked_commit
(khóa, nhật ký strict, x14, Transaction expected_before) kèm plan token: sổ, yêu cầu, file và nguồn đã khai đổi sau bản xem trước thì từ chối.
Two sheets in references/_INDEX.xlsx, created on the first write. Writes go through library_sync.locked_commit with a plan token.

Đây là khai báo và dấu vết, KHÔNG xác minh nội dung, hiệu lực hay nghiệm thu; không chọn Nháp hiện tại, không chạy AI, không quét thư viện.
Declarations and provenance only: no content check, no validity, no approval, no current-draft pointer, no AI, no library scan.
Windows + NTFS cục bộ, hợp tác/ADVISORY: quan sát tại thời điểm, không chống đua hay sửa ngoài. / Cooperative; observations are point-in-time.
TEST ONLY: biến LOCAL_AGENT_WORKSPACE_TEST_OPEN_LOG (xem current_draft.py) ghi đường dẫn thay vì mở ứng dụng.
Mã thoát / exit: 0 ok; 1 từ chối hoặc lỗi / refused or error; 2 lỗi cách dùng / usage.
"""
import datetime
import io
import json
import os
import re
import stat
import sys

import library
import library_sync
import transaction

ContentError = library.ContentError
SyncError = library_sync.SyncError

ARTIFACT_SHEET = 'LAW artifacts'
RELATION_SHEET = 'LAW relations'
ARTIFACT_MARKER = 'law.artifacts/1'
RELATION_MARKER = 'law.relations/1'
HEADER_ROW = 4
ARTIFACT_HEADERS = ('artifact_id', 'revision', 'predecessor', 'path', 'sha256', 'size', 'title', 'kind', 'limits', 'origin',
                    'publish_rid', 'publish_identity', 'recorded_by', 'recorded_at')
RELATION_HEADERS = ('relation_key', 'type', 'from_kind', 'from_id', 'from_revision', 'to_kind', 'to_id', 'to_revision', 'to_sha256',
                    'basis_status', 'basis_replaces', 'basis_replaced_by', 'basis_effective', 'recorded_by', 'recorded_at')
# Loại do người gọi khai; trùng nhãn nội dung B-01 để INFERENCE đi cùng kết quả. / Declared by the caller; shares the B-01 labels.
KINDS = ('REPORT', 'APPENDIX', 'OTHER') + tuple(library.CONTENT_KINDS)
ORIGINS = ('P_CN1', 'USER_FILE')
MAX_DEPS = 50       # phụ thuộc trực tiếp mỗi lần đăng ký / direct dependencies per registration
MAX_NODES = 200     # số phiên bản trong một closure / artifact versions in one closure
MAX_DEPTH = 32      # độ sâu closure / closure depth
MAX_LIST = 500      # số dòng trả về của find / rows returned by find
WORKBOOK_PATH = 'references/_INDEX.xlsx'
ID_RE = re.compile(r'^[A-Za-z0-9][A-Za-z0-9._:-]{0,127}$')
RID_RE = re.compile(r'^[0-9a-f]{32}$')
IDENT_RE = re.compile(r'^[0-9a-f]{8}:[0-9a-f]{16}$')
REGISTER_KEYS = ('artifact_id', 'revision', 'predecessor', 'path', 'sha256', 'title', 'kind', 'limits', 'origin', 'rid', 'recorded_by',
                 'dependencies')
DEP_KEYS = ('kind', 'id', 'revision', 'sha256')
CURRENT = ('CURRENT_CATALOG_BASIS', 'CURRENT_ARTIFACT_VERSION')
RANK = {'CURRENT_CATALOG_BASIS': 0, 'CURRENT_ARTIFACT_VERSION': 0, 'NEEDS_REVIEW': 1, 'MISSING_BASIS': 2, 'MISSING_SOURCE': 3,
        'BASIS_CONFLICT': 4}
STATE_CODE = {'NEEDS_REVIEW': 'STALE_SOURCE_BASIS', 'BASIS_CONFLICT': 'STALE_SOURCE_BASIS', 'MISSING_SOURCE': 'UNKNOWN_SOURCE',
              'MISSING_BASIS': 'SOURCE_BASIS_MISSING'}
FILE_WHY = {'bad_path': 'PATH_UNSUPPORTED', 'link': 'LINK', 'missing': 'MISSING', 'not_regular': 'NOT_REGULAR', 'hardlink': 'HARDLINK',
            'unreadable': 'UNREADABLE', 'changed': 'CHANGED_WHILE_READING'}


def L(vi, en):
    return en if library.LANG == 'en' else vi


def err(code, vi, en):
    return ContentError(code, vi, en)


def key_text(key):
    return '%s@%d' % key


def now():
    return datetime.datetime.now().replace(microsecond=0).isoformat()


# ---------------------------------------------------------------- đường dẫn và file / paths and files
def path_problem(rel):
    """Mã lỗi hoặc None. Chỉ working/ hoặc output/, dạng đường dẫn của owner P-CN1, không tên file quyền. / Error code or None."""
    try:
        import publish_new
    except ImportError as e:  # the P-CN1 owner is part of the package: without it paths cannot be judged, so nothing is accepted
        raise err('OWNER_UNAVAILABLE', 'Không nạp được publish_new để kiểm đường dẫn: %s' % e, 'publish_new cannot be loaded to judge the path: %s' % e)
    if not isinstance(rel, str) or publish_new.path_form_problem(rel):
        return 'PATH_UNSUPPORTED_FORM'
    if publish_new.target_zone(rel) is None:
        return 'ZONE_NOT_LIGHT'
    if publish_new._authority_name(rel):
        return 'AUTHORITY_NAME'
    return None


def observe_file(root, rel):
    """(quan sát, None) hoặc (None, lý do). Không theo liên kết; từ chối reparse, file nhiều tên, không phải file thường.
    (observation, None) or (None, reason). Never follows a link; refuses reparse points, multi-name files and non-regular files."""
    if path_problem(rel):
        return None, 'bad_path'
    full = os.path.join(root, *rel.split('/'))
    if not transaction.link_free(full, root):
        return None, 'link'
    try:
        st = os.lstat(full)
    except FileNotFoundError:
        return None, 'missing'
    except OSError:
        return None, 'unreadable'
    if stat.S_ISLNK(st.st_mode) or (getattr(st, 'st_file_attributes', 0) or 0) & library_sync.REPARSE_POINT:
        return None, 'link'
    if not stat.S_ISREG(st.st_mode):
        return None, 'not_regular'
    if getattr(st, 'st_nlink', 1) != 1:
        return None, 'hardlink'
    try:
        sha = library_sync.digest(full)
        after = os.lstat(full)
    except OSError:
        return None, 'unreadable'
    if (st.st_size, st.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
        return None, 'changed'
    return {'sha256': sha, 'size': after.st_size, 'identity': '%x:%x' % (getattr(st, 'st_dev', 0), getattr(st, 'st_ino', 0))}, None


def file_error(rel, why, prefix='ARTIFACT_FILE'):
    return err('%s_%s' % (prefix, FILE_WHY.get(why, 'UNREADABLE')),
               'Không dùng được file kết quả (%s): %s' % (why, rel), 'The result file cannot be used (%s): %s' % (why, rel))


# ---------------------------------------------------------------- đọc và kiểm cấu trúc hai sheet / reading and validating the sheets
def _text(cell):
    value = getattr(cell, 'value', None)
    return '' if value is None else (value if isinstance(value, str) else str(value))


def _formula(cell):
    return getattr(cell, 'data_type', 's') == 'f'


def _number(cell):
    value = getattr(cell, 'value', None)
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return None


def _table(ws, marker, headers, code, label):
    def bad(row, vi, en):
        return err(code, 'Sheet "%s" không hợp lệ (dòng %s): %s' % (label, row, vi), 'Sheet "%s" is malformed (row %s): %s' % (label, row, en))
    rows = list(ws.iter_rows(min_row=1))
    width = len(headers)
    if len(rows) < HEADER_ROW or not rows[0] or _text(rows[0][0]) != marker:
        raise bad(1, 'thiếu hoặc sai marker A1 (%s)' % marker, 'A1 marker missing or wrong (%s)' % marker)
    if _formula(rows[0][0]):
        raise bad(1, 'A1 là công thức', 'A1 is a formula')
    header = rows[HEADER_ROW - 1]
    if tuple(_text(c) for c in header[:width]) != headers or any(_text(c) for c in header[width:]) or any(_formula(c) for c in header):
        raise bad(HEADER_ROW, 'tiêu đề cột khác phiên bản đã tài liệu hóa', 'headers differ from the documented version')
    out = []
    for number, cells in enumerate(rows[HEADER_ROW:], HEADER_ROW + 1):
        if not any(getattr(c, 'value', None) not in (None, '') for c in cells):
            continue
        if any(_formula(c) for c in cells):
            raise bad(number, 'ô được quản lý là công thức', 'a managed cell is a formula')
        if any(_text(c) for c in cells[width:]):
            raise bad(number, 'có dữ liệu ngoài vùng cột được quản lý', 'data outside the managed columns')
        out.append((number, [cells[i] if i < len(cells) else None for i in range(width)]))
    return out, bad


def parse_artifacts(ws):
    rows, bad = _table(ws, ARTIFACT_MARKER, ARTIFACT_HEADERS, 'MALFORMED_ARTIFACT_SHEET', ARTIFACT_SHEET)
    found, paths = {}, {}
    for number, cells in rows:
        row = {name: _text(cells[i]) for i, name in enumerate(ARTIFACT_HEADERS)}
        revision, size, predecessor = _number(cells[1]), _number(cells[5]), _number(cells[2])
        if revision is None or revision < 1:
            raise bad(number, 'revision phải là số nguyên dương', 'revision must be a positive integer')
        if size is None or size < 0:
            raise bad(number, 'size phải là số nguyên không âm', 'size must be a non-negative integer')
        if row['predecessor'].strip() and predecessor is None:
            raise bad(number, 'predecessor phải là số nguyên', 'predecessor must be an integer')
        row.update(revision=revision, size=size, predecessor=predecessor, _row=number)
        if not ID_RE.fullmatch(row['artifact_id']):
            raise bad(number, 'artifact_id sai dạng', 'artifact_id has a wrong form')
        if not library.SHA256_RE.fullmatch(row['sha256']):
            raise bad(number, 'sha256 không phải 64 ký tự hex', 'sha256 is not 64 hex characters')
        if any(not row[k].strip() for k in ('title', 'limits', 'recorded_by', 'recorded_at')):
            raise bad(number, 'thiếu title, limits hoặc xuất xứ', 'title, limits or provenance is empty')
        if row['kind'] not in KINDS or row['origin'] not in ORIGINS:
            raise bad(number, 'kind hoặc origin ngoài danh sách', 'kind or origin is outside the allowed list')
        if path_problem(row['path']):
            raise bad(number, 'path không hợp lệ hoặc ngoài working/output', 'path is invalid or outside working/output')
        if row['origin'] == 'P_CN1':
            if not (RID_RE.fullmatch(row['publish_rid']) and IDENT_RE.fullmatch(row['publish_identity'])):
                raise bad(number, 'P_CN1 cần publish_rid và publish_identity hợp lệ', 'P_CN1 needs a valid publish_rid and publish_identity')
        elif row['publish_rid'] or row['publish_identity']:
            raise bad(number, 'USER_FILE không mang rid hay identity', 'USER_FILE carries no rid or identity')
        key = (row['artifact_id'], revision)
        if key in found:
            raise bad(number, 'trùng %s' % key_text(key), 'duplicate %s' % key_text(key))
        folded = row['path'].casefold()
        if folded in paths:
            raise bad(number, 'path đã thuộc phiên bản khác (%s)' % key_text(paths[folded]), 'path already belongs to another version (%s)' % key_text(paths[folded]))
        paths[folded] = key
        found[key] = row
    latest = {}
    for (aid, revision) in found:
        latest[aid] = max(latest.get(aid, 0), revision)
    for (aid, revision), row in found.items():
        if revision > 1 and ((aid, revision - 1) not in found or row['predecessor'] != revision - 1):
            raise bad(row['_row'], 'revision %d cần predecessor %d của cùng ID' % (revision, revision - 1),
                      'revision %d needs predecessor %d of the same ID' % (revision, revision - 1))
        if revision == 1 and row['predecessor'] is not None:
            raise bad(row['_row'], 'revision 1 không có predecessor', 'revision 1 has no predecessor')
    return sorted(found.values(), key=lambda r: r['_row']), found, latest


def relation_key(rtype, from_kind, from_id, from_revision, to_kind, to_id):
    return library_sync.canon([rtype, from_kind, from_id, from_revision, to_kind, to_id])


def parse_relations(ws, found):
    rows, bad = _table(ws, RELATION_MARKER, RELATION_HEADERS, 'MALFORMED_RELATION_SHEET', RELATION_SHEET)
    edges, keys = [], set()
    for number, cells in rows:
        row = {name: _text(cells[i]) for i, name in enumerate(RELATION_HEADERS)}
        row['from_revision'], row['to_revision'] = (_number(cells[4]), _number(cells[7]))
        row['_row'] = number
        for name, number_cell in (('from_revision', cells[4]), ('to_revision', cells[7])):
            if row[name] is None and _text(number_cell).strip():
                raise bad(number, '%s phải là số nguyên' % name, '%s must be an integer' % name)
        if row['type'] not in ('RELATED', 'DERIVED_FROM') or not row['recorded_by'].strip() or not row['recorded_at'].strip():
            raise bad(number, 'type hoặc xuất xứ sai', 'type or provenance is wrong')
        if not row['from_id'] or not row['to_id'] or row['from_kind'] not in ('SOURCE', 'ARTIFACT') or row['to_kind'] not in ('SOURCE', 'ARTIFACT'):
            raise bad(number, 'thiếu hoặc sai tham chiếu', 'a reference is missing or wrong')
        if row['type'] == 'RELATED':
            if (row['from_kind'], row['to_kind']) != ('SOURCE', 'SOURCE') or row['from_id'] >= row['to_id']:
                raise bad(number, 'RELATED chỉ nối hai SourceID khác nhau, ghi theo thứ tự tăng dần', 'RELATED joins two different SourceIDs in ascending order')
            if row['from_revision'] is not None or row['to_revision'] is not None or row['to_sha256'] or \
                    any(row[k] for k in ('basis_status', 'basis_replaces', 'basis_replaced_by', 'basis_effective')):
                raise bad(number, 'RELATED không mang revision, hash hay căn cứ', 'RELATED carries no revision, hash or basis')
        else:
            if row['from_kind'] != 'ARTIFACT' or row['from_revision'] is None or row['from_revision'] < 1 or \
                    row['to_revision'] is None or row['to_revision'] < 1 or not library.SHA256_RE.fullmatch(row['to_sha256']):
                raise bad(number, 'DERIVED_FROM cần phiên bản kết quả nguồn và revision/hash đích đầy đủ', 'DERIVED_FROM needs an artifact version origin and a full target revision/hash')
            if (row['from_id'], row['from_revision']) not in found:
                raise err('ARTIFACT_DANGLING', 'Quan hệ dòng %s trỏ phiên bản kết quả không có: %s@%s' % (number, row['from_id'], row['from_revision']),
                          'Relation row %s points to a missing artifact version: %s@%s' % (number, row['from_id'], row['from_revision']))
            if row['to_kind'] == 'ARTIFACT':
                target = found.get((row['to_id'], row['to_revision']))
                if target is None:
                    raise err('ARTIFACT_DANGLING', 'Quan hệ dòng %s trỏ phiên bản kết quả không có: %s@%s' % (number, row['to_id'], row['to_revision']),
                              'Relation row %s points to a missing artifact version: %s@%s' % (number, row['to_id'], row['to_revision']))
                if target['sha256'] != row['to_sha256'] or row['from_id'] == row['to_id'] or \
                        any(row[k] for k in ('basis_status', 'basis_replaces', 'basis_replaced_by', 'basis_effective')):
                    raise bad(number, 'hash đích khác phiên bản đã đăng ký, tự tham chiếu hoặc mang căn cứ nguồn', 'target hash differs from the registered version, self reference or carries source basis')
        if row['relation_key'] != relation_key(row['type'], row['from_kind'], row['from_id'], row['from_revision'], row['to_kind'], row['to_id']):
            raise bad(number, 'relation_key không khớp nội dung dòng', 'relation_key does not match the row')
        if row['relation_key'] in keys:
            raise bad(number, 'trùng quan hệ', 'duplicate relation')
        keys.add(row['relation_key'])
        edges.append(row)
    return edges


def find_cycle(found, edges):
    graph = {}
    for edge in edges:
        if edge['type'] == 'DERIVED_FROM' and edge['to_kind'] == 'ARTIFACT':
            graph.setdefault((edge['from_id'], edge['from_revision']), []).append((edge['to_id'], edge['to_revision']))
    state = {}
    for start in graph:
        if start in state:
            continue
        stack = [(start, iter(graph.get(start, ())))]
        state[start] = 1
        while stack:
            node, it = stack[-1]
            for nxt in it:
                if state.get(nxt) == 1:
                    return nxt
                if nxt not in state:
                    state[nxt] = 1
                    stack.append((nxt, iter(graph.get(nxt, ()))))
                    break
            else:
                state[node] = 2
                stack.pop()
    return None


def parse_registry(wb):
    """Đọc cả hai sheet, chỉ đọc. Thiếu cả hai: legacy hợp lệ, rỗng. Có một: lỗi cấu trúc. Dòng nguồn lịch sử không khớp sổ KHÔNG là lỗi ở đây.
    Read both sheets without writing. Neither: valid legacy, empty. Exactly one: a structure error. A stale historical SourceID is not an error here."""
    names = wb.sheetnames
    has_a, has_r = ARTIFACT_SHEET in names, RELATION_SHEET in names
    reg = {'present': has_a and has_r, 'artifacts': [], 'by_key': {}, 'latest': {}, 'edges': [], 'deps': {}, 'related': []}
    if not has_a and not has_r:
        return reg
    if has_a != has_r:
        raise err('ARTIFACT_SHEETS_INCOMPLETE', 'Chỉ có một trong hai sheet "%s"/"%s"; không coi là rỗng' % (ARTIFACT_SHEET, RELATION_SHEET),
                  'Only one of the sheets "%s"/"%s" exists; it is not treated as empty' % (ARTIFACT_SHEET, RELATION_SHEET))
    reg['artifacts'], reg['by_key'], reg['latest'] = parse_artifacts(wb[ARTIFACT_SHEET])
    reg['edges'] = parse_relations(wb[RELATION_SHEET], reg['by_key'])
    cycle = find_cycle(reg['by_key'], reg['edges'])
    if cycle is not None:
        raise err('ARTIFACT_CYCLE', 'Quan hệ DERIVED_FROM có vòng tại %s' % key_text(cycle), 'DERIVED_FROM relations have a cycle at %s' % key_text(cycle))
    for edge in reg['edges']:
        if edge['type'] == 'RELATED':
            reg['related'].append(edge)
        else:
            reg['deps'].setdefault((edge['from_id'], edge['from_revision']), []).append(edge)
    return reg


def closure(reg, key):
    """(các phiên bản trong closure, các cạnh) từ một phiên bản, hữu hạn; vượt giới hạn thì báo lỗi, không cắt.
    (versions, edges) reachable from one version; finite; a limit breach is an error, never a truncation."""
    depth, order, edges, queue = {key: 0}, [key], [], [key]
    while queue:
        cur = queue.pop(0)
        for edge in reg['deps'].get(cur, []):
            edges.append((cur, edge))
            if edge['to_kind'] != 'ARTIFACT':
                continue
            nxt = (edge['to_id'], edge['to_revision'])
            if nxt not in depth:
                if depth[cur] + 1 > MAX_DEPTH or len(depth) >= MAX_NODES:
                    raise err('CLOSURE_LIMIT', 'Closure của %s vượt giới hạn (%d phiên bản, sâu %d); không cắt bớt' % (key_text(key), MAX_NODES, MAX_DEPTH),
                              'The closure of %s exceeds the limit (%d versions, depth %d); it is not truncated' % (key_text(key), MAX_NODES, MAX_DEPTH))
                depth[nxt] = depth[cur] + 1
                order.append(nxt)
                queue.append(nxt)
    return order, edges


# ---------------------------------------------------------------- sổ và nguồn / the workbook and sources
def build_catalog(book):
    catalog = {}
    for rec in book['data']:
        ident = library.fmt(rec.get('id'))
        if ident:
            catalog.setdefault(ident, []).append(rec)
    return catalog


def load_book(root=None):
    """Đọc sổ một lần: catalog + registry cùng một snapshot. Không ghi, không quét, không đọc nguồn.
    {'ok': True, snapshot, book, catalog, registry} hoặc {'ok': False, kind, message[, code]}."""
    root = root or library.ROOT
    try:
        import openpyxl as opx
    except ImportError:
        return {'ok': False, 'kind': 'openpyxl_missing', 'message': ''}
    index = library.index_path(root)
    if not transaction.link_free(index, root):
        return {'ok': False, 'kind': 'unsafe_index',
                'message': L('Đường dẫn sổ qua liên kết/reparse; chưa đọc sổ.', 'The index path crosses a link/reparse point; the workbook was not read.')}
    try:
        with open(index, 'rb') as f:
            raw = f.read()
    except FileNotFoundError:
        return {'ok': False, 'kind': 'no_index', 'message': library.book_message({'kind': 'no_index'})}
    except OSError:
        return {'ok': False, 'kind': 'unreadable', 'message': library.book_message({'kind': 'unreadable'})}
    book = library.parse_book(raw, require_columns=False, data_only=False)
    if not book['ok']:
        return dict(book, message=library.book_message(book))
    try:
        wb = opx.load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    except Exception:
        return {'ok': False, 'kind': 'unreadable', 'message': library.book_message({'kind': 'unreadable'})}
    try:
        registry = parse_registry(wb)
    except ContentError as e:
        return {'ok': False, 'kind': 'registry_malformed', 'code': e.code, 'message': str(e)}
    finally:
        wb.close()
    return {'ok': True, 'snapshot': book['snapshot'], 'book': book, 'catalog': build_catalog(book), 'registry': registry}


def need_book(root):
    loaded = load_book(root)
    if not loaded['ok']:
        if loaded['kind'] == 'openpyxl_missing':
            library.need_openpyxl()
        raise err('INDEX_UNUSABLE', loaded['message'], loaded['message'])
    return loaded


def observe_source(root, catalog, source_id, revision=None, sha=None):
    """Quan sát một nguồn theo sổ đã đọc: đúng một dòng, revision/SHA khớp (nếu được khai) và bytes file khớp SHA trong sổ.
    Observe one source against the workbook already read: one row, expected revision/SHA, and file bytes equal to the recorded SHA."""
    recs = catalog.get(source_id) or []
    if not recs:
        raise err('UNKNOWN_SOURCE', 'SourceID không có trong sổ: ' + source_id, 'SourceID is not in the index: ' + source_id)
    if len(recs) > 1:
        raise err('DUPLICATE_SOURCE', 'SourceID trùng trong sổ: ' + source_id, 'Duplicate SourceID in the index: ' + source_id)
    rec = recs[0]
    revision_now = library.revision_of(rec.get('law_revision'))
    sha_now = str(rec.get('law_sha') or '').strip().lower()
    if not (isinstance(revision_now, int) and revision_now >= 1 and library.SHA256_RE.fullmatch(sha_now)):
        raise err('SOURCE_BASIS_MISSING', 'Nguồn %s chưa có revision/SHA hợp lệ trong sổ; chạy sync trước' % source_id,
                  'Source %s has no valid revision/SHA in the index; run sync first' % source_id)
    if (revision is not None and revision_now != revision) or (sha is not None and sha_now != sha):
        raise err('STALE_SOURCE_BASIS', 'Căn cứ nguồn %s đã đổi (sổ: revision %s, SHA %s…)' % (source_id, revision_now, sha_now[:12]),
                  'The basis of source %s changed (index: revision %s, SHA %s...)' % (source_id, revision_now, sha_now[:12]))
    locator = library.norm_path(rec.get('file'))
    full, why = library.locate_source(root, locator)
    if full is None:
        raise err('SOURCE_UNAVAILABLE', 'Không mở được file nguồn (%s): %s' % (why, locator), 'The source file cannot be opened (%s): %s' % (why, locator))
    try:
        before = os.stat(full)
        actual = library_sync.digest(full)
        after = os.stat(full)
    except OSError as e:
        raise err('SOURCE_UNAVAILABLE', 'Không đọc được file nguồn: %s' % e, 'The source file cannot be read: %s' % e)
    if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
        raise err('SOURCE_UNAVAILABLE', 'Nguồn đổi khi đọc: ' + locator, 'The source changed while it was read: ' + locator)
    if actual != sha_now:
        raise err('STALE_SOURCE_BASIS', 'File nguồn %s đã đổi so với sổ (SHA thực %s…); chạy sync rồi lập lại' % (source_id, actual[:12]),
                  'Source %s differs from the index (actual SHA %s...); run sync and try again' % (source_id, actual[:12]))
    semantic = {key: library.fmt(rec.get(key)) for key, _ in library.SEMANTIC_FIELDS}
    return {'source_id': source_id, 'revision': revision_now, 'sha256': sha_now, 'locator': locator, 'file_sha256': actual,
            'semantic': semantic, 'path': library.display_file(locator)}


def edge_state(edge, catalog, reg):
    """(trạng thái, lý do) của một cạnh DERIVED_FROM so với sổ hiện tại. Chỉ so snapshot đã ghi; không đọc bytes.
    State of one edge against the current workbook; recorded snapshots only. A referenced version stays the version; it is never redirected to a newer one."""
    if edge['to_kind'] == 'SOURCE':
        entry = {'source_id': edge['to_id'], 'source_revision': edge['to_revision'], 'source_sha256': edge['to_sha256'],
                 'basis_status': edge['basis_status'], 'basis_replaces': edge['basis_replaces'],
                 'basis_replaced_by': edge['basis_replaced_by'], 'basis_effective': edge['basis_effective']}
        return library.basis_state(entry, catalog)
    if reg['latest'].get(edge['to_id'], 0) != edge['to_revision']:
        return 'NEEDS_REVIEW', 'upstream_superseded'
    return 'CURRENT_ARTIFACT_VERSION', 'latest_version'


def describe(reg, catalog, key):
    """Hồ sơ một phiên bản: trạng thái theo từng chiều độc lập, tổ tiên với kind/limits, unknown, INFERENCE. Chỉ đọc metadata.
    Profile of one version on independent dimensions, ancestors with their kind/limits, unknown basis, INFERENCE. Metadata only."""
    row = reg['by_key'][key]
    order, edges = closure(reg, key)
    dependencies, worst = [], None
    for owner, edge in edges:
        state, reason = edge_state(edge, catalog, reg)
        dependencies.append({'from': key_text(owner), 'direct': owner == key, 'kind': edge['to_kind'], 'id': edge['to_id'],
                             'revision': edge['to_revision'], 'sha256': edge['to_sha256'], 'sha12': edge['to_sha256'][:12],
                             'basis': {snap: edge[snap] for _, snap in library.SEMANTIC_FIELDS} if edge['to_kind'] == 'SOURCE' else {},
                             'state': state, 'reason': reason})
        if worst is None or RANK[state] > RANK[worst]:
            worst = state
    lineage = [{'key': key_text(k), 'title': reg['by_key'][k]['title'], 'kind': reg['by_key'][k]['kind'], 'origin': reg['by_key'][k]['origin'],
                'limits': reg['by_key'][k]['limits'], 'declared_basis': bool(reg['deps'].get(k)),
                'superseded': reg['latest'][k[0]] != k[1]} for k in order]
    # Căn cứ chưa khai của một tổ tiên không trở thành CURRENT nhờ cạnh nguồn khác; stale đã biết vẫn ưu tiên.
    # An ancestor's unknown basis cannot become CURRENT through another source edge; known stale states still take precedence.
    if not edges:
        state = 'NO_DECLARED_BASIS'
    elif RANK[worst] > 0:
        state = worst
    else:
        state = 'CURRENT_CATALOG_BASIS' if all(x['declared_basis'] for x in lineage) else 'NO_DECLARED_BASIS'
    latest = reg['latest'][key[0]]
    out = {k: row[k] for k in ARTIFACT_HEADERS}
    out.update(key=key_text(key), row=row['_row'], sha12=row['sha256'][:12], latest_revision=latest,
               superseded_by=key[1] + 1 if latest > key[1] else '', dependency_state=state, dependencies=dependencies,
               diagnostics=[d for d in dependencies if d['state'] not in CURRENT], lineage=lineage,
               unknown_basis=[x['key'] for x in lineage if not x['declared_basis']],
               inference=[x['key'] for x in lineage if x['kind'] == 'INFERENCE'],
               openable=row['path'].lower().endswith(_open_ext()), bytes_observation='UNOBSERVED')
    return out


def _open_ext():
    import current_draft
    return current_draft.ALLOWED_EXT


def pick(item, snapshot):
    """Lựa chọn gắn với đúng sổ + phiên bản + file đã thấy; giao diện gửi lại khi mở hoặc sao chép. / The selection a client sends back."""
    return {'snapshot': snapshot, 'id': item['artifact_id'], 'revision': item['revision'], 'path': item['path'], 'sha256': item['sha256']}


def find(root=None, query='', artifact_id='', source_id='', revision=None):
    """Tra phiên bản kết quả, chỉ đọc sổ: không quét/hash nguồn hay file, không model, không mạng, không ghi.
    Look up artifact versions from the workbook only; bytes are never observed (always UNOBSERVED)."""
    loaded = load_book(root)
    if not loaded['ok']:
        return loaded
    reg, catalog = loaded['registry'], loaded['catalog']
    words = library.fold(query).split()
    items = []
    try:
        for row in reg['artifacts']:
            key = (row['artifact_id'], row['revision'])
            if (artifact_id and row['artifact_id'] != artifact_id) or (revision is not None and row['revision'] != revision):
                continue
            item = describe(reg, catalog, key)
            if source_id and not any(d['kind'] == 'SOURCE' and d['id'] == source_id for d in item['dependencies']):
                continue
            hay = library.fold(' '.join(str(item[k]) for k in ('artifact_id', 'title', 'kind', 'limits', 'path', 'origin', 'recorded_by')))
            if all(w in hay for w in words):
                item['pick'] = pick(item, loaded['snapshot'])
                items.append(item)
    except ContentError as e:
        return {'ok': False, 'kind': 'closure_limit', 'code': e.code, 'message': str(e)}
    return {'ok': True, 'snapshot': loaded['snapshot'], 'registry': 'present' if reg['present'] else 'absent', 'total': len(reg['artifacts']),
            'matched': len(items), 'truncated': len(items) > MAX_LIST, 'items': items[:MAX_LIST], 'bytes_observation': 'UNOBSERVED',
            'related': [{'a': e['from_id'], 'b': e['to_id'], 'recorded_by': e['recorded_by'], 'recorded_at': e['recorded_at']}
                        for e in reg['related'] if not source_id or source_id in (e['from_id'], e['to_id'])]}


def links(root, selection):
    """Phụ lục/nguồn liên quan và kết quả liên quan của một dòng nguồn trên giao diện: gắn lại dòng, chỉ đọc sổ, không side effect.
    Related sources/appendices and results of one source row: re-binds the row, reads the workbook only."""
    loaded = load_book(root)
    if not loaded['ok']:
        return {'ok': False, 'reason': 'index:' + loaded['kind'], 'message': loaded.get('message', '')}
    rec, why = library.rebind(root, selection, loaded['book'])
    if why:
        return {'ok': False, 'reason': why}
    if not rec['id']:
        return {'ok': True, 'snapshot': loaded['snapshot'], 'source_id': '', 'related': [], 'results': []}
    reg, catalog = loaded['registry'], loaded['catalog']
    related, seen = [], set()
    for edge in reg['related']:
        if rec['id'] in (edge['from_id'], edge['to_id']):
            other = edge['to_id'] if edge['from_id'] == rec['id'] else edge['from_id']
            if other in seen:
                continue
            seen.add(other)
            recs = catalog.get(other) or []
            info = library.record(recs[0]) if len(recs) == 1 else None
            related.append({'id': other, 'title': info['title'] if info else '', 'revision': info['revision'] if info else None,
                            'file': info['file'] if info else '', 'in_catalog': info is not None, 'recorded_by': edge['recorded_by']})
    results = []
    try:
        for row in reg['artifacts']:
            item = describe(reg, catalog, (row['artifact_id'], row['revision']))
            hits = [d for d in item['dependencies'] if d['kind'] == 'SOURCE' and d['id'] == rec['id']]
            if hits:
                results.append(dict(item, via_direct=any(d['direct'] for d in hits), pick=pick(item, loaded['snapshot'])))
    except ContentError as e:
        return {'ok': False, 'reason': e.code, 'message': str(e)}
    keep = ('artifact_id', 'revision', 'title', 'kind', 'origin', 'path', 'sha256', 'size', 'dependency_state', 'superseded_by',
            'via_direct', 'unknown_basis', 'inference', 'openable', 'pick', 'recorded_at')
    return {'ok': True, 'snapshot': loaded['snapshot'], 'source_id': rec['id'], 'related': related[:MAX_LIST],
            'results': [{k: r[k] for k in keep} for r in results[:MAX_LIST]], 'truncated': len(results) > MAX_LIST or len(related) > MAX_LIST}


# ---------------------------------------------------------------- closure đọc lại bytes / re-reading a closure's bytes
def source_input(obs):
    return {'path': obs['path'], 'digest': 'sha256:' + obs['file_sha256'], 'source_id': obs['source_id'], 'revision': str(obs['revision'])}


def artifact_input(row):
    return {'path': row['path'], 'digest': 'sha256:' + row['sha256'], 'revision': str(row['revision']),
            'scope': 'law-artifact:%s' % key_text((row['artifact_id'], row['revision']))}


def inspect_closure(root, reg, catalog, keys):
    """Đọc/hash lại finite closure của các phiên bản được dùng. Phiên bản đã bị thay, căn cứ nguồn cũ/thiếu/xung đột, file đổi/mất thì từ chối.
    Re-read and hash the finite closure of the versions in use. A superseded version, a stale/missing/conflicting basis or a changed file refuses."""
    nodes, edges, seen = [], [], set()
    for key in keys:
        order, found = closure(reg, key)
        edges += found
        for node in order:
            if node not in seen:
                seen.add(node)
                nodes.append(node)
    for node in nodes:
        if reg['latest'][node[0]] != node[1]:
            raise err('UPSTREAM_SUPERSEDED', 'Phiên bản %s đã có bản mới hơn (revision %d); không dùng làm đầu vào' % (key_text(node), reg['latest'][node[0]]),
                      'Version %s has a newer revision (%d); it is not used as an input' % (key_text(node), reg['latest'][node[0]]))
    for owner, edge in edges:
        state, reason = edge_state(edge, catalog, reg)
        if state not in CURRENT:
            raise err(STATE_CODE[state], 'Căn cứ của %s đã không còn khớp: %s %s (%s)' % (key_text(owner), edge['to_id'], state, reason),
                      'The basis of %s no longer matches: %s %s (%s)' % (key_text(owner), edge['to_id'], state, reason))
    inputs, files, sources = {}, {}, {}
    for node in nodes:
        row = reg['by_key'][node]
        obs, why = observe_file(root, row['path'])
        if obs is None:
            raise file_error(row['path'], why)
        if obs['sha256'] != row['sha256'] or obs['size'] != row['size']:
            raise err('ARTIFACT_FILE_CHANGED', 'File của %s đã khác bản đã đăng ký: %s' % (key_text(node), row['path']),
                      'The file of %s differs from the registered version: %s' % (key_text(node), row['path']))
        files[row['path']] = obs['sha256']
        inputs[row['path']] = artifact_input(row)
    for owner, edge in edges:
        if edge['to_kind'] == 'SOURCE':
            obs = observe_source(root, catalog, edge['to_id'], edge['to_revision'], edge['to_sha256'])
            sources[obs['locator']] = obs['file_sha256']
            inputs[obs['path']] = source_input(obs)
    return {'nodes': nodes, 'inputs': inputs, 'files': files, 'sources': sources}


def prepare(root, request):
    """Tạo consumer mới từ phiên bản/nguồn đã chọn: đọc lại sổ + hash lại closure ngay lúc dùng. Không ghi, không receipt, không chọn Nháp.
    Trả snapshot sổ và danh sách inputs (path/digest) để recipe A-11/P-CN1 bind vào publish task. Not a grant, not a verification of content."""
    if not isinstance(request, dict) or set(request) - {'inputs', 'expected_snapshot'} or not isinstance(request.get('inputs'), list) or not request['inputs']:
        raise library_sync.bad_request('cần {"inputs": [...]} không rỗng', 'needs a non-empty {"inputs": [...]}')
    if len(request['inputs']) > MAX_DEPS:
        raise library_sync.bad_request('quá %d đầu vào' % MAX_DEPS, 'more than %d inputs' % MAX_DEPS)
    loaded = need_book(root)
    if 'expected_snapshot' in request and request['expected_snapshot'] != loaded['snapshot']:
        raise err('STALE_SELECTION', 'Sổ đã đổi từ lúc chọn; chọn lại', 'The workbook changed since the selection; select again')
    reg, catalog = loaded['registry'], loaded['catalog']
    keys, source_ids = [], []
    for item in request['inputs']:
        if not isinstance(item, dict):
            raise library_sync.bad_request('mỗi đầu vào là một đối tượng', 'each input is an object')
        if item.get('kind') == 'ARTIFACT' and set(item) <= {'kind', 'id', 'revision', 'sha256'}:
            revision = library_sync.clean_revision('revision', item.get('revision'))
            row = reg['by_key'].get((library_sync.clean_text('id', item.get('id'), True), revision))
            if row is None:
                raise err('UNKNOWN_ARTIFACT', 'Không có phiên bản %s@%s' % (item.get('id'), revision), 'No such version %s@%s' % (item.get('id'), revision))
            if 'sha256' in item and library_sync.clean_sha('sha256', item['sha256']) != row['sha256']:
                raise err('STALE_ARTIFACT_BASIS', 'Hash đã chọn khác phiên bản đã đăng ký', 'The selected hash differs from the registered version')
            keys.append((row['artifact_id'], revision))
        elif item.get('kind') == 'SOURCE' and set(item) <= {'kind', 'id'}:
            source_ids.append(library_sync.clean_text('id', item.get('id'), True))
        else:
            raise library_sync.bad_request('đầu vào cần kind ARTIFACT (id, revision) hoặc SOURCE (id)', 'an input needs kind ARTIFACT (id, revision) or SOURCE (id)')
    got = inspect_closure(root, reg, catalog, keys)
    inputs = dict(got['inputs'])
    for sid in source_ids:
        obs = observe_source(root, catalog, sid)
        inputs[obs['path']] = source_input(obs)
    inputs[WORKBOOK_PATH] = {'path': WORKBOOK_PATH, 'digest': 'sha256:' + loaded['snapshot']}
    chosen = [describe(reg, catalog, k) for k in keys]
    return {'snapshot': loaded['snapshot'], 'inputs': sorted(inputs.values(), key=lambda i: i['path']), 'selected': chosen,
            'closure': [key_text(n) for n in got['nodes']], 'sources': source_ids,
            'unknown_basis': sorted({u for c in chosen for u in c['unknown_basis']}),
            'inference': sorted({u for c in chosen for u in c['inference']}),
            'limits': [{'key': x['key'], 'limits': x['limits']} for c in chosen for x in c['lineage']],
            'bytes_observation': 'OBSERVED_AT_PREPARE', 'verifies_content': False}


# ---------------------------------------------------------------- yêu cầu đăng ký / requests
def _dependency(item):
    if not isinstance(item, dict) or set(item) != set(DEP_KEYS) or item['kind'] not in ('SOURCE', 'ARTIFACT'):
        raise library_sync.bad_request('mỗi phụ thuộc cần đúng khóa kind (SOURCE|ARTIFACT), id, revision, sha256',
                                       'each dependency needs exactly kind (SOURCE|ARTIFACT), id, revision, sha256')
    return {'kind': item['kind'], 'id': library_sync.clean_text('id', item['id'], True), 'revision': library_sync.clean_revision('revision', item['revision']),
            'sha256': library_sync.clean_sha('sha256', item['sha256'])}


def register_request(request):
    """Chuẩn hóa yêu cầu đăng ký một phiên bản; khóa lạ hoặc sai kiểu bị từ chối. / Normalise a registration request."""
    if not isinstance(request, dict):
        raise library_sync.bad_request('phải là một đối tượng', 'it must be an object')
    unknown = sorted(set(request) - set(REGISTER_KEYS))
    if unknown:
        raise library_sync.bad_request('khóa không được biết: ' + ', '.join(map(str, unknown)), 'unknown keys: ' + ', '.join(map(str, unknown)))
    missing = [k for k in ('artifact_id', 'revision', 'path', 'title', 'kind', 'limits', 'origin', 'recorded_by') if k not in request]
    if missing:
        raise library_sync.bad_request('thiếu: ' + ', '.join(missing), 'missing: ' + ', '.join(missing))
    out = {'artifact_id': library_sync.clean_text('artifact_id', request['artifact_id'], True),
           'revision': library_sync.clean_revision('revision', request['revision']),
           'path': library_sync.clean_text('path', request['path'], True), 'title': library_sync.clean_text('title', request['title'], True),
           'kind': library_sync.clean_text('kind', request['kind'], True), 'limits': library_sync.clean_text('limits', request['limits'], True),
           'origin': library_sync.clean_text('origin', request['origin'], True),
           'recorded_by': library_sync.clean_text('recorded_by', request['recorded_by'], True), 'rid': '', 'sha256': '', 'predecessor': None}
    if not ID_RE.fullmatch(out['artifact_id']):
        raise library_sync.bad_request('artifact_id chỉ gồm chữ, số, . _ : - (tối đa 128)', 'artifact_id may use letters, digits and . _ : - (128 max)')
    if out['kind'] not in KINDS or out['origin'] not in ORIGINS:
        raise library_sync.bad_request('kind thuộc %s; origin thuộc %s' % ('/'.join(KINDS), '/'.join(ORIGINS)),
                                       'kind is one of %s; origin is one of %s' % ('/'.join(KINDS), '/'.join(ORIGINS)))
    problem = path_problem(out['path'])
    if problem:
        raise err('PATH_REFUSED', 'Đường dẫn không được đăng ký (%s): chỉ working/ hoặc output/' % problem, 'The path cannot be registered (%s): only working/ or output/' % problem)
    if 'rid' in request:
        out['rid'] = library_sync.clean_text('rid', request['rid'], True)
    if (out['origin'] == 'P_CN1') != bool(out['rid']) or (out['rid'] and not RID_RE.fullmatch(out['rid'])):
        raise library_sync.bad_request('P_CN1 cần rid 32 hex; USER_FILE không có rid', 'P_CN1 needs a 32-hex rid; USER_FILE has none')
    if 'sha256' in request:
        out['sha256'] = library_sync.clean_sha('sha256', request['sha256'])
    if request.get('predecessor') is not None:
        out['predecessor'] = library_sync.clean_revision('predecessor', request['predecessor'])
    deps = request.get('dependencies', [])
    if not isinstance(deps, list) or len(deps) > MAX_DEPS:
        raise library_sync.bad_request('dependencies là danh sách tối đa %d' % MAX_DEPS, 'dependencies is a list of at most %d' % MAX_DEPS)
    out['dependencies'] = sorted((_dependency(d) for d in deps), key=lambda d: (d['kind'], d['id']))
    if len({(d['kind'], d['id']) for d in out['dependencies']}) != len(out['dependencies']):
        raise library_sync.bad_request('một phụ thuộc bị khai hai lần', 'a dependency is listed twice')
    if any(d['kind'] == 'ARTIFACT' and d['id'] == out['artifact_id'] for d in out['dependencies']):
        raise library_sync.bad_request('một kết quả không phụ thuộc chính ID của nó (dùng predecessor)', 'a result cannot depend on its own ID (use predecessor)')
    if len(library_sync.canon(out)) > library.CELL_LIMIT:
        raise library_sync.bad_request('yêu cầu dài hơn giới hạn ô; không cắt bớt', 'the request is longer than the cell limit; it is not truncated')
    return out


def relate_request(request):
    if not isinstance(request, dict) or set(request) != {'source_a', 'source_b', 'recorded_by'}:
        raise library_sync.bad_request('cần đúng source_a, source_b, recorded_by', 'needs exactly source_a, source_b, recorded_by')
    a, b = (library_sync.clean_text(k, request[k], True).strip() for k in ('source_a', 'source_b'))
    if a == b:
        raise library_sync.bad_request('hai SourceID phải khác nhau', 'the two SourceIDs must differ')
    low, high = sorted((a, b))
    return {'source_a': low, 'source_b': high, 'recorded_by': library_sync.clean_text('recorded_by', request['recorded_by'], True)}


# ---------------------------------------------------------------- P-CN1: evidence của owner / the owner's evidence
def pcn1_target(root, req):
    """Bằng chứng xuất bản của owner P-CN1 đúng rid: đường dẫn, digest, size, identity; bytes hiện tại vẫn khớp. Không tạo receipt mới.
    The pinned P-CN1 owner's publish evidence for the rid, and the current bytes still equal it. No new receipt."""
    import current_draft as cd
    try:
        env = cd.Env(root)
        cand, why = cd._candidate(env, req['rid'])
        if cand is None:
            raise err('EVIDENCE_UNUSABLE', 'rid không có bằng chứng xuất bản hợp lệ (%s)' % why, 'The rid has no valid publish evidence (%s)' % why)
        if cand['path'] != req['path']:
            raise err('RID_PATH_MISMATCH', 'rid xuất bản %s, không phải %s' % (cand['path'], req['path']), 'The rid published %s, not %s' % (cand['path'], req['path']))
        digest = cand['digest'][len('sha256:'):]
        if req['sha256'] and req['sha256'] != digest:
            raise err('RID_DIGEST_MISMATCH', 'sha256 khai khác bằng chứng xuất bản', 'The declared sha256 differs from the publish evidence')
        state, reason = cd._judge(cand, cd._facts(env, cand['path']))
        if state != 'matched':
            raise err('TARGET_NOT_MATCHING', 'File không còn khớp bản đã xuất bản (%s)' % (reason or state), 'The file no longer matches what was published (%s)' % (reason or state))
        evidence = env.observer.load_evidence(root, req['rid'])
        inputs = evidence['prepared']['task']['inputs']
    except cd.Unavailable as e:
        raise err('OWNER_UNAVAILABLE', 'Owner P-CN1 không dùng được (%s); chưa đăng ký' % e.reason, 'The P-CN1 owner is unusable (%s); nothing registered' % e.reason)
    except (KeyError, TypeError):
        raise err('EVIDENCE_UNUSABLE', 'Bằng chứng prepared thiếu inputs', 'The prepared evidence has no inputs')
    identity = '%s:%s' % (cand['identity']['volume_serial'], cand['identity']['file_index'])
    return {'sha256': digest, 'size': cand['size'], 'identity': identity, 'inputs': inputs, 'work_id': cand['work_id']}


def check_prepared(required, prepared):
    """Mọi input của closure (và sổ nếu có phụ thuộc) phải nằm trong prepared.inputs với đúng digest/source_id/revision.
    Every closure input (and the workbook when there are dependencies) must be in prepared.inputs with the same digest/source_id/revision."""
    by_path = {item.get('path'): item for item in prepared if isinstance(item, dict)}
    for item in required:
        found = by_path.get(item['path'])
        if found is None:
            code = 'WORKBOOK_NOT_IN_PREPARED' if item['path'] == WORKBOOK_PATH else 'INPUT_NOT_IN_PREPARED'
            raise err(code, 'Input %s không nằm trong prepared.inputs của lần xuất bản' % item['path'], 'Input %s is not in the publish prepared.inputs' % item['path'])
        if found.get('digest') != item['digest'] or any(found.get(k) != item[k] for k in ('source_id', 'revision') if k in item):
            code = 'REGISTRY_CHANGED_SINCE_PREPARE' if item['path'] == WORKBOOK_PATH else 'PREPARED_INPUT_MISMATCH'
            raise err(code, 'Input %s đã khác lúc prepare/xuất bản' % item['path'], 'Input %s differs from what was prepared/published' % item['path'])


def unregistered(error):
    body = str(error)
    body = body.split('] ', 1)[1] if body.startswith('[') and '] ' in body else body
    return err(error.code, body + ' Đã tạo file và biên nhận nhưng chưa đăng ký; không rebase, không dựng lịch sử, không publish lại.',
               body + ' The file and receipt exist but are not registered; no rebase, no reconstructed history, no republish.')


# ---------------------------------------------------------------- ghi sổ / writing
def history_ok(wb):
    """Các sheet quản lý khác phải hợp lệ trước khi ghi, như mọi writer B-01. / The other managed sheets must be valid before a write."""
    library_sync.history_state(wb['LAW history'] if 'LAW history' in wb.sheetnames else None)
    if library.CONTENT_SHEET in wb.sheetnames:
        library.parse_content(wb[library.CONTENT_SHEET])


def open_writable(root):
    original, snapshot, wb = library_sync.open_index(root)
    book = library.parse_book(original, require_columns=False, data_only=False)
    if not book['ok']:
        wb.close()
        raise SyncError(library.book_message(book))
    return original, snapshot, wb, build_catalog(book)


def last_row(rows):
    return max([r['_row'] for r in rows] + [HEADER_ROW])


def append_rows(wb, reg, artifact_rows, relation_rows):
    """Thêm dòng vào hai sheet (tạo cả hai nếu chưa có). Trả giới hạn dòng cũ để kiểm các ô ngoài vùng thêm không đổi."""
    limits = {ARTIFACT_SHEET: last_row(reg['artifacts']), RELATION_SHEET: last_row(reg['edges'])}
    if not reg['present']:
        limits = {ARTIFACT_SHEET: 0, RELATION_SHEET: 0}
        for name, marker, headers in ((ARTIFACT_SHEET, ARTIFACT_MARKER, ARTIFACT_HEADERS), (RELATION_SHEET, RELATION_MARKER, RELATION_HEADERS)):
            sheet = wb.create_sheet(name)
            library_sync.put(sheet, 1, 1, marker)
            for col, head in enumerate(headers, 1):
                library_sync.put(sheet, HEADER_ROW, col, head)
    for name, headers, rows, at in ((ARTIFACT_SHEET, ARTIFACT_HEADERS, artifact_rows, last_row(reg['artifacts'])),
                                    (RELATION_SHEET, RELATION_HEADERS, relation_rows, last_row(reg['edges']))):
        for n, values in enumerate(rows, 1):
            for col, head in enumerate(headers, 1):
                if values.get(head) not in (None, ''):
                    library_sync.put(wb[name], at + n, col, values[head])
    return limits


def stage(opx, original, wb, reg, artifact_rows, relation_rows, check):
    """Dựng bản sẽ ghi, đọc lại và so với sổ gốc: chỉ dòng thêm vào hai sheet được đổi. / Build the bytes to write, reload and compare with the original."""
    limits = append_rows(wb, reg, artifact_rows, relation_rows)
    staged = library_sync.save_workbook(wb)
    back = opx.load_workbook(io.BytesIO(staged), data_only=False)
    try:
        again = parse_registry(back)
        check(again)
    finally:
        back.close()
    library_sync.verify_staged(opx, original, staged, lambda name, key: name in limits and key[0] > limits[name])
    return staged


def artifact_row(req, observed, stamp):
    return {'artifact_id': req['artifact_id'], 'revision': req['revision'], 'predecessor': req['predecessor'], 'path': req['path'],
            'sha256': observed['sha256'], 'size': observed['size'], 'title': req['title'], 'kind': req['kind'], 'limits': req['limits'],
            'origin': req['origin'], 'publish_rid': req['rid'], 'publish_identity': observed['identity'] if req['origin'] == 'P_CN1' else '',
            'recorded_by': req['recorded_by'], 'recorded_at': stamp}


def register_plan(root, request):
    """Xem trước đăng ký một phiên bản kết quả và các phụ thuộc trực tiếp. Chỉ đọc sổ, file kết quả và closure đã khai; không ghi, không tạo sheet.
    Preview registering one artifact version and its direct dependencies. Reads the workbook, the result file and the declared closure only."""
    req = register_request(request)
    original, snapshot, wb, catalog = open_writable(root)
    opx = library.need_openpyxl()
    try:
        history_ok(wb)
        reg = parse_registry(wb)
        key = (req['artifact_id'], req['revision'])
        extra = {'artifact_key': key_text(key), 'sheets_created': not reg['present']}
        stored = reg['by_key'].get(key)
        if stored is not None:
            edges = {(e['to_kind'], e['to_id'], e['to_revision'], e['to_sha256']) for e in reg['deps'].get(key, [])}
            asked = {(d['kind'], d['id'], d['revision'], d['sha256']) for d in req['dependencies']}
            if (all(stored[k] == req[k] for k in ('path', 'title', 'kind', 'limits', 'origin', 'recorded_by')) and stored['publish_rid'] == req['rid'] and
                    edges == asked and (not req['sha256'] or req['sha256'] == stored['sha256']) and
                    req['predecessor'] in (None, stored['predecessor'])):
                return dict(extra, action='already_recorded', snapshot=snapshot, bytes=None, notices=[], sources={}, files={}, observation=None,
                            current_basis_validated=False, plan_token=library_sync.plan_token('artifact', req, snapshot, None), sheets_created=False)
            raise err('ARTIFACT_CONFLICT', 'Phiên bản %s đã đăng ký với nội dung khác; dùng revision mới' % key_text(key),
                      'Version %s is already registered with a different payload; use a new revision' % key_text(key))
        latest = reg['latest'].get(req['artifact_id'], 0)
        if req['revision'] != latest + 1:
            raise err('BAD_REVISION', 'Revision kế tiếp của %s là %d' % (req['artifact_id'], latest + 1), 'The next revision of %s is %d' % (req['artifact_id'], latest + 1))
        if req['predecessor'] is not None and req['predecessor'] != latest:
            raise err('BAD_PREDECESSOR', 'predecessor phải là revision %d của cùng ID' % latest, 'predecessor must be revision %d of the same ID' % latest)
        req['predecessor'] = latest or None
        owner = next((r for r in reg['artifacts'] if r['path'].casefold() == req['path'].casefold()), None)
        if owner is not None:
            raise err('PATH_REUSED', 'Path đã thuộc %s; mỗi phiên bản dùng file riêng' % key_text((owner['artifact_id'], owner['revision'])),
                      'The path already belongs to %s; each version uses its own file' % key_text((owner['artifact_id'], owner['revision'])))
        try:
            return _register_new(root, req, original, snapshot, wb, catalog, reg, opx, extra)
        except ContentError as e:
            if req['origin'] == 'P_CN1':
                raise unregistered(e)
            raise
    finally:
        wb.close()


def _register_new(root, req, original, snapshot, wb, catalog, reg, opx, extra):
    if req['origin'] == 'P_CN1':
        observed = pcn1_target(root, req)
    else:
        observed, why = observe_file(root, req['path'])
        if observed is None:
            raise file_error(req['path'], why)
        if req['sha256'] and req['sha256'] != observed['sha256']:
            raise err('FILE_DIGEST_MISMATCH', 'sha256 khai khác bytes file hiện tại', 'The declared sha256 differs from the current file bytes')
    keys = [(d['id'], d['revision']) for d in req['dependencies'] if d['kind'] == 'ARTIFACT']
    for d in req['dependencies']:
        if d['kind'] == 'ARTIFACT':
            row = reg['by_key'].get((d['id'], d['revision']))
            if row is None:
                raise err('UNKNOWN_ARTIFACT', 'Không có phiên bản %s@%d' % (d['id'], d['revision']), 'No such version %s@%d' % (d['id'], d['revision']))
            if row['sha256'] != d['sha256']:
                raise err('STALE_ARTIFACT_BASIS', 'Hash khai của %s@%d khác phiên bản đã đăng ký' % (d['id'], d['revision']), 'The declared hash of %s@%d differs from the registered version' % (d['id'], d['revision']))
    got = inspect_closure(root, reg, catalog, keys)
    inputs, sources, files, semantic = dict(got['inputs']), dict(got['sources']), dict(got['files']), {}
    for d in req['dependencies']:
        if d['kind'] == 'SOURCE':
            obs = observe_source(root, catalog, d['id'], d['revision'], d['sha256'])
            sources[obs['locator']] = obs['file_sha256']
            inputs[obs['path']] = source_input(obs)
            semantic[d['id']] = obs['semantic']
    if req['dependencies']:
        inputs[WORKBOOK_PATH] = {'path': WORKBOOK_PATH, 'digest': 'sha256:' + snapshot}
    if req['origin'] == 'P_CN1':
        check_prepared(sorted(inputs.values(), key=lambda i: i['path']), observed['inputs'])
    files[req['path']] = observed['sha256']
    stamp = now()
    arow = artifact_row(req, observed, stamp)
    rrows = []
    for d in req['dependencies']:
        rrows.append({'relation_key': relation_key('DERIVED_FROM', 'ARTIFACT', req['artifact_id'], req['revision'], d['kind'], d['id']),
                      'type': 'DERIVED_FROM', 'from_kind': 'ARTIFACT', 'from_id': req['artifact_id'], 'from_revision': req['revision'],
                      'to_kind': d['kind'], 'to_id': d['id'], 'to_revision': d['revision'], 'to_sha256': d['sha256'],
                      'recorded_by': req['recorded_by'], 'recorded_at': stamp})
        if d['kind'] == 'SOURCE':
            rrows[-1].update({snap: semantic[d['id']][k] for k, snap in library.SEMANTIC_FIELDS})

    def check(again):
        got_row = again['by_key'].get((req['artifact_id'], req['revision']))
        back = {(e['to_kind'], e['to_id'], e['to_revision'], e['to_sha256']) for e in again['deps'].get((req['artifact_id'], req['revision']), [])}
        if got_row is None or any(got_row[k] != (arow[k] if arow[k] is not None else got_row[k]) for k in ARTIFACT_HEADERS if k != 'predecessor') or \
                got_row['predecessor'] != arow['predecessor'] or back != {(d['kind'], d['id'], d['revision'], d['sha256']) for d in req['dependencies']} or \
                len(again['artifacts']) != len(reg['artifacts']) + 1 or len(again['edges']) != len(reg['edges']) + len(rrows):
            raise SyncError('Verification of the staged workbook failed: the version does not read back as written / '
                            'Kiểm bản sẽ ghi không đạt: phiên bản đọc lại không đúng như đã ghi')
    staged = stage(opx, original, wb, reg, [arow], rrows, check)
    observation = {'file': {'sha256': observed['sha256'], 'size': observed['size'], 'identity': observed['identity']},
                   'inputs': sorted(inputs.values(), key=lambda i: i['path']), 'dependencies': req['dependencies']}
    return dict(extra, action='record', snapshot=snapshot, bytes=staged, notices=library_sync.lossy_notices(original),
                observation=observation, sources=sources, files=files, current_basis_validated=True,
                publish_request=req if req['origin'] == 'P_CN1' else None,
                plan_token=library_sync.plan_token('artifact', req, snapshot, observation))


def relate_plan(root, request):
    """Xem trước thêm một quan hệ RELATED giữa hai nguồn theo SourceID ổn định. Không đọc bytes; không truyền stale.
    Preview adding one RELATED between two sources by stable SourceID. No bytes read; nothing propagates."""
    req = relate_request(request)
    original, snapshot, wb, catalog = open_writable(root)
    opx = library.need_openpyxl()
    try:
        history_ok(wb)
        reg = parse_registry(wb)
        for sid in (req['source_a'], req['source_b']):
            if len(catalog.get(sid) or []) != 1:
                raise err('UNKNOWN_SOURCE', 'SourceID không có hoặc trùng trong sổ: ' + sid, 'SourceID is missing or duplicated in the index: ' + sid)
        key = relation_key('RELATED', 'SOURCE', req['source_a'], None, 'SOURCE', req['source_b'])
        extra = {'relation': [req['source_a'], req['source_b']], 'sheets_created': not reg['present']}
        if any(e['relation_key'] == key for e in reg['related']):
            return dict(extra, action='already_recorded', snapshot=snapshot, bytes=None, notices=[], sources={}, files={}, observation=None,
                        plan_token=library_sync.plan_token('relate', req, snapshot, None), sheets_created=False)
        row = {'relation_key': key, 'type': 'RELATED', 'from_kind': 'SOURCE', 'from_id': req['source_a'], 'to_kind': 'SOURCE',
               'to_id': req['source_b'], 'recorded_by': req['recorded_by'], 'recorded_at': now()}

        def check(again):
            if key not in {e['relation_key'] for e in again['related']} or len(again['edges']) != len(reg['edges']) + 1 or \
                    len(again['artifacts']) != len(reg['artifacts']):
                raise SyncError('Verification of the staged workbook failed: the relation does not read back as written / '
                                'Kiểm bản sẽ ghi không đạt: quan hệ đọc lại không đúng như đã ghi')
        staged = stage(opx, original, wb, reg, [], [row], check)
        observation = {'sources': {sid: library.revision_of(catalog[sid][0].get('law_revision')) for sid in (req['source_a'], req['source_b'])}}
        return dict(extra, action='record', snapshot=snapshot, bytes=staged, notices=library_sync.lossy_notices(original),
                    observation=observation, sources={}, files={}, plan_token=library_sync.plan_token('relate', req, snapshot, observation))
    finally:
        wb.close()


def commit(root, make_plan, token, operation):
    """Ghi đúng bản đã xem qua library_sync.locked_commit: plan tính lại dưới khóa phải ra cùng token; recheck có kiểu cho nguồn và file kết quả.
    Write exactly what was previewed through library_sync.locked_commit, with a typed recheck of sources and result files."""
    token = library_sync.check_plan_token(token)

    def check(result):
        if result['plan_token'] != token:
            raise library_sync.PreviewMismatch('The workbook, the declared files/sources or the request changed after the preview; LAW did not change the index, '
                                               'and no new file transaction was created; preview again / '
                                               'Sổ, file/nguồn được khai hoặc yêu cầu đã đổi sau bản xem trước; LAW không sửa sổ, chưa tạo giao dịch file mới; '
                                               'xem trước lại.', 'plan')

    def recheck(result):
        for locator, sha in result['sources'].items():
            full, why = library.locate_source(root, locator)
            if full is None or library_sync.digest(full) != sha:
                raise SyncError('Source changed while preparing / Nguồn đổi khi chuẩn bị ghi; xem lại và chạy lại: ' + locator)
        for rel, sha in result['files'].items():
            obs, why = observe_file(root, rel)
            if obs is None or obs['sha256'] != sha:
                raise SyncError('A result file changed while preparing / File kết quả đổi khi chuẩn bị ghi; xem lại và chạy lại: ' + rel)
        if result.get('publish_request') is not None:
            try:
                observed = pcn1_target(root, result['publish_request'])
                if any(observed[k] != result['observation']['file'][k] for k in ('sha256', 'size', 'identity')):
                    raise err('TARGET_NOT_MATCHING', 'Bằng chứng hoặc identity file đã đổi trước ghi sổ',
                              'The file evidence or identity changed before the registry write')
            except ContentError as e:
                raise unregistered(e)

    return library_sync.locked_commit(root, make_plan, check=check, recheck=recheck, operation=operation)


def register_apply(root, request, expect_plan):
    return commit(root, lambda: register_plan(root, request), expect_plan, 'library-artifact')


def relate_apply(root, request, expect_plan):
    return commit(root, lambda: relate_plan(root, request), expect_plan, 'library-relation')


# ---------------------------------------------------------------- mở và sao chép / open and copy
def rebind(root, selection):
    """Gắn lựa chọn của client vào đúng phiên bản đã thấy: snapshot sổ + ID/revision + path + hash. (loaded, row, None) hoặc (None, None, lý do)."""
    if not isinstance(selection, dict) or set(selection) != {'snapshot', 'id', 'revision', 'path', 'sha256'}:
        return None, None, 'bad_selection'
    revision = selection['revision']
    if (not all(isinstance(selection[k], str) for k in ('snapshot', 'id', 'path', 'sha256')) or isinstance(revision, bool) or
            not isinstance(revision, int)):
        return None, None, 'bad_selection'
    loaded = load_book(root)
    if not loaded['ok']:
        return None, None, 'registry_malformed' if loaded['kind'] == 'registry_malformed' else 'index:' + loaded['kind']
    if loaded['snapshot'] != selection['snapshot']:
        return None, None, 'stale'
    row = loaded['registry']['by_key'].get((selection['id'], revision))
    if row is None:
        return None, None, 'stale'
    if row['path'] != selection['path'] or row['sha256'] != selection['sha256']:
        return None, None, 'changed'
    return loaded, row, None


def _launch(path):
    import current_draft
    if os.name != 'nt' and not os.environ.get(current_draft.OPEN_LOG_ENV):
        raise OSError('unsupported platform')
    current_draft._startfile(path)


def open_artifact(root, selection):
    """Mở đúng phiên bản đã đăng ký sau khi quan sát lại file đích ngay trước opener. OPENED nghĩa là đã yêu cầu hệ điều hành mở, không phải
    kiểm Office hay nghiệm thu. Bản có nguồn cũ vẫn mở nếu file còn đúng; file đổi/mất/liên kết/identity khác thì từ chối. Không ack."""
    loaded, row, why = rebind(root, selection)
    if why:
        return {'ok': False, 'reason': why}
    if not row['path'].lower().endswith(_open_ext()):
        return {'ok': False, 'reason': 'unsupported_type'}
    try:
        item = describe(loaded['registry'], loaded['catalog'], (row['artifact_id'], row['revision']))
    except ContentError as e:
        return {'ok': False, 'reason': e.code, 'message': str(e)}
    if row['origin'] == 'P_CN1':
        import current_draft as cd
        try:
            env = cd.Env(root)
            cand, _ = cd._candidate(env, row['publish_rid'])
            if cand is None:
                return {'ok': False, 'reason': 'evidence_unusable'}
            if (cand['path'], cand['digest'][len('sha256:'):], cand['size']) != (row['path'], row['sha256'], row['size']):
                return {'ok': False, 'reason': 'file_changed'}
            state, reason = cd._judge(cand, cd._facts(env, row['path']))
            if state != 'matched':
                return {'ok': False, 'reason': {'TARGET_MISSING': 'missing', 'TARGET_NOT_REGULAR': 'not_regular', 'PARENT_UNSAFE': 'link'}.get(reason, 'file_changed')}
            if '%s:%s' % (cand['identity']['volume_serial'], cand['identity']['file_index']) != row['publish_identity']:
                return {'ok': False, 'reason': 'identity_conflict'}
        except cd.Unavailable:
            return {'ok': False, 'reason': 'owner_unavailable'}
    else:
        obs, why = observe_file(root, row['path'])
        if obs is None:
            return {'ok': False, 'reason': why}
        if obs['sha256'] != row['sha256'] or obs['size'] != row['size']:
            return {'ok': False, 'reason': 'file_changed'}
    try:
        _launch(os.path.join(root, *row['path'].split('/')))
    except OSError:
        return {'ok': False, 'reason': 'launch_failed'}
    return {'ok': True, 'state': 'OPENED', 'path': row['path'], 'dependency_state': item['dependency_state'], 'bytes_observation': 'OBSERVED_BEFORE_OPEN'}


def copy_text(item, snapshot, lang='vi'):
    en = lang == 'en'
    lines = ['Result as input (from references/_INDEX.xlsx, index SHA-256 %s):' % snapshot if en
             else 'Kết quả dùng làm đầu vào (từ references/_INDEX.xlsx, SHA-256 sổ %s):' % snapshot,
             '- %s | %s | sha256 %s | %s | %s | %s' % (item['key'], item['path'], item['sha256'], item['kind'], item['origin'], item['title']),
             '%s %s%s' % ('Declared basis state:' if en else 'Trạng thái căn cứ đã khai:', item['dependency_state'],
                          (' (%s %s)' % ('superseded by revision' if en else 'đã có bản mới revision', item['superseded_by'])) if item['superseded_by'] else '')]
    for x in item['lineage']:
        lines.append('  %s %s | %s | %s %s' % ('•', x['key'], x['kind'], 'Limits:' if en else 'Giới hạn:', x['limits']))
    for d in item['dependencies']:
        lines.append('  DERIVED_FROM %s -> %s %s@%d | sha256 %s | %s (%s)' %
                     (d['from'], d['kind'], d['id'], d['revision'], d['sha256'], d['state'], d['reason']))
        if d['basis']:
            lines.append('    %s %s' % ('Recorded source basis:' if en else 'Căn cứ nguồn đã ghi:',
                                       library_sync.canon(d['basis'])))
    if item['unknown_basis']:
        lines.append(('Unknown basis (no declared source): ' if en else 'Chưa khai căn cứ: ') + ', '.join(item['unknown_basis']))
    if item['inference']:
        lines.append(('Carries INFERENCE labels: ' if en else 'Mang nhãn INFERENCE: ') + ', '.join(item['inference']))
    lines.append('Not verified: bytes_observation UNOBSERVED here. Run prepare right before use; this verifies nothing about the content.' if en
                 else 'Chưa kiểm: bytes_observation UNOBSERVED (chưa quan sát bytes ở đây). Chạy prepare ngay trước khi dùng; danh sách này không xác minh nội dung.')
    return '\n'.join(lines)


def copy_input(root, selection, lang='vi'):
    """Văn bản sao chép gắn đúng khóa/path/hash/snapshot và toàn bộ căn cứ, kind, giới hạn. Chỉ sao chép: không giao việc, không chọn Nháp, không ghi."""
    loaded, row, why = rebind(root, selection)
    if why:
        return {'ok': False, 'reason': why}
    try:
        item = describe(loaded['registry'], loaded['catalog'], (row['artifact_id'], row['revision']))
    except ContentError as e:
        return {'ok': False, 'reason': e.code, 'message': str(e)}
    return {'ok': True, 'text': copy_text(item, loaded['snapshot'], 'en' if lang == 'en' else 'vi')}


# ---------------------------------------------------------------- CLI
def summary(res):
    keep = ('action', 'artifact_key', 'relation', 'plan_token', 'snapshot', 'sheets_created', 'current_basis_validated', 'notices', 'protection')
    out = {k: res[k] for k in keep if k in res}
    out['index_changed'] = res.get('bytes') is not None
    return out


def print_items(res):
    print(L('%d/%d phiên bản kết quả khớp (chỉ đọc sổ; bytes: UNOBSERVED, chưa kiểm file).', '%d/%d artifact versions match (workbook only; bytes: UNOBSERVED, no file checked).') % (res['matched'], res['total']))
    if res['registry'] == 'absent':
        print(L('Sổ chưa có sheet "%s"/"%s": chưa đăng ký gì.', 'The workbook has no "%s"/"%s" sheets: nothing registered yet.') % (ARTIFACT_SHEET, RELATION_SHEET))
    for e in res['items']:
        print('- %s | %s | %s | %s' % (e['key'], e['kind'], e['origin'], e['title']))
        print('    %s | %s: %s | %s: %s' % (e['path'], L('Căn cứ', 'Basis'), e['dependency_state'], L('Bị thay bởi', 'Superseded by'), e['superseded_by'] or '-'))
        print('    %s: %s' % (L('Giới hạn', 'Limits'), e['limits']))
        if e['unknown_basis']:
            print('    %s: %s' % (L('Chưa khai căn cứ', 'Unknown basis'), ', '.join(e['unknown_basis'])))
        for d in e['diagnostics']:
            print('    ! %s %s %s: %s (%s)' % (d['from'], d['kind'], d['id'], d['state'], d['reason']))


def _option(rest, name, default=''):
    return library.option(rest, name, default)


def cli(argv):
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8', errors='replace')
    try:
        info = json.load(open(os.path.join(library.ROOT, '.agent', 'PACKAGE_INFO.json'), encoding='utf-8'))
        library.LANG = 'en' if info.get('language') == 'en' else 'vi'
    except (OSError, ValueError):
        pass
    if '--lang' in argv:
        i = argv.index('--lang')
        library.LANG = argv[i + 1] if i + 1 < len(argv) else library.LANG
        del argv[i:i + 2]
    root = os.path.abspath(_option(argv, '--project') or library.ROOT)
    if not argv or argv[0] in ('-h', '--help'):
        print(__doc__)
        return 0
    cmd, rest = argv[0], argv[1:]
    as_json = '--json' in rest
    rest = [a for a in rest if a != '--json']
    try:
        if cmd == 'find':
            artifact_id, source_id, revision = _option(rest, '--artifact-id'), _option(rest, '--source-id'), _option(rest, '--revision')
            res = find(root, ' '.join(rest), artifact_id, source_id, int(revision) if revision.isdigit() else None)
            if not res['ok']:
                if res['kind'] == 'openpyxl_missing':
                    library.need_openpyxl()
                print(res['message'], file=sys.stderr)
                return 1
            print(json.dumps(res, ensure_ascii=False) if as_json else '', end='\n' if as_json else '')
            if not as_json:
                print_items(res)
            return 0
        if cmd in ('register', 'relate'):
            sub = rest.pop(0) if rest else ''
            path, token = _option(rest, '--request'), _option(rest, '--expect-plan')
            if sub not in ('preview', 'apply') or not path or rest or (sub == 'apply' and not token):
                print(L('Cú pháp: %s preview|apply --request FILE [--expect-plan TOKEN] (apply cần --expect-plan)', 'Usage: %s preview|apply --request FILE [--expect-plan TOKEN] (apply needs --expect-plan)') % cmd, file=sys.stderr)
                return 2
            plan_fn, apply_fn = (register_plan, register_apply) if cmd == 'register' else (relate_plan, relate_apply)
            request = library_sync.load_request(path)
            res = apply_fn(root, request, token) if sub == 'apply' else plan_fn(root, request)
            out = dict(summary(res), applied=sub == 'apply' and res.get('bytes') is not None)
            if as_json:
                print(json.dumps(out, ensure_ascii=False))
            else:
                print({'record': L('Đã ghi.', 'Recorded.') if sub == 'apply' else L('Xem trước: sẽ ghi.', 'Preview: would be recorded.'),
                       'already_recorded': L('Đã ghi từ trước, không ghi lại (không xác nhận căn cứ hiện tại).', 'Already recorded; nothing written (this does not confirm the current basis).')}[res['action']])
                for k in ('artifact_key', 'plan_token'):
                    if k in out:
                        print('  %s: %s' % (k, out[k]))
                if sub == 'preview' and res.get('bytes') is not None:
                    print(L('Thêm apply --expect-plan TOKEN sau khi xem trước.', 'Add apply --expect-plan TOKEN after reviewing the preview.'))
            return 0
        if cmd == 'prepare':
            path = _option(rest, '--request')
            if not path or rest:
                print('Usage: prepare --request FILE [--json]', file=sys.stderr)
                return 2
            res = prepare(root, library_sync.load_request(path))
            print(json.dumps(res, ensure_ascii=False) if as_json else json.dumps(res, ensure_ascii=False, indent=2))
            return 0
        if cmd in ('open', 'copy'):
            artifact_id, revision, snap = _option(rest, '--artifact-id'), _option(rest, '--revision'), _option(rest, '--expect-snapshot')
            if not artifact_id or not revision.isdigit() or rest:
                print('Usage: %s --artifact-id ID --revision N [--expect-snapshot HEX]' % cmd, file=sys.stderr)
                return 2
            found = find(root, '', artifact_id, '', int(revision))
            if not found['ok'] or not found['items']:
                print(found.get('message') or L('Không có phiên bản này.', 'No such version.'), file=sys.stderr)
                return 1
            selection = dict(found['items'][0]['pick'], snapshot=snap or found['snapshot'])
            res = open_artifact(root, selection) if cmd == 'open' else copy_input(root, selection, library.LANG)
            if cmd == 'copy' and res['ok']:
                print(res['text'])
            else:
                print(json.dumps(res, ensure_ascii=False))
            return 0 if res['ok'] else 1
    except (ContentError, SyncError) as e:
        print(str(e), file=sys.stderr)
        return 1
    print(L('Lệnh không có: %s. Dùng: find | register | relate | prepare | open | copy', 'Unknown command: %s. Use: find | register | relate | prepare | open | copy') % cmd, file=sys.stderr)
    return 2


if __name__ == '__main__':
    sys.exit(cli(sys.argv[1:]))
