"""Real Excel (Office via COM) against the index writer. Dummy data in temp dirs. Not part of the unit-test suite."""
import hashlib
import os
import shutil
import subprocess
import sys
import tempfile
import time

LAW = r'D:\UNG DUNG AI\TOOL AI 2026\CVF-Workspace\local-agent-workspace'
SCRATCH = os.path.dirname(os.path.abspath(__file__))
PS1 = os.path.join(SCRATCH, 'excel_ops.ps1')
sys.path.insert(0, os.path.join(LAW, 'tools'))
import library_sync  # noqa: E402
import transaction  # noqa: E402
import openpyxl  # noqa: E402

TEMPLATE = os.path.join(LAW, 'skills/local/shared/source-library/assets/source-index-template.xlsx')


def sha(p):
    return hashlib.sha256(open(p, 'rb').read()).hexdigest()[:12]


def project():
    base = tempfile.mkdtemp(prefix='a06-excel-')
    os.makedirs(os.path.join(base, 'references', 'source-documents', 'team'))
    shutil.copyfile(TEMPLATE, os.path.join(base, 'references', '_INDEX.xlsx'))
    open(os.path.join(base, 'references', 'source-documents', 'team', 'one.txt'), 'w').write('one')
    return base


def sheet_index(index):
    wb = openpyxl.load_workbook(index)
    n = wb.sheetnames.index(library_sync.library.SHEET) + 1
    wb.close()
    return n


def ps(*args):
    return subprocess.Popen(['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', PS1] + list(args),
                            stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding='utf-8')


def excel_pids():
    out = subprocess.run(['tasklist', '/FI', 'IMAGENAME eq EXCEL.EXE', '/FO', 'CSV', '/NH'], capture_output=True, text=True).stdout
    return {line.split('","')[1] for line in out.splitlines() if 'EXCEL.EXE' in line}


before_pids = excel_pids()
try:
    # E1 (T14): a real Excel window has the workbook open while LAW writes
    base = project()
    index = os.path.join(base, 'references', '_INDEX.xlsx')
    original = sha(index)
    ready, release = os.path.join(base, 'ready.flag'), os.path.join(base, 'release.flag')
    holder = ps('-Mode', 'hold', '-Path', index, '-Sheet', str(sheet_index(index)), '-Ready', ready, '-Release', release)
    t0 = time.time()
    while not os.path.exists(ready) and time.time() - t0 < 60:
        time.sleep(0.3)
    print('E1 Excel has the workbook open:', os.path.exists(ready))
    try:
        library_sync.apply(base)
        print('E1 apply while Excel holds it: OK (written!)')
    except Exception as e:  # noqa: BLE001
        print('E1 apply while Excel holds it: %s -> %s' % (type(e).__name__, str(e).split(' / ')[0][:150]))
    print('E1 workbook unchanged:', sha(index) == original, '| open transactions:', transaction.incomplete(base))
    open(release, 'w').close()
    holder.communicate(timeout=60)
    time.sleep(1)
    library_sync.apply(base)
    print('E1 after Excel closed, same sync: OK, rows:', [c.value for c in openpyxl.load_workbook(index)[library_sync.library.SHEET]['A'] if c.value][-1:])

    # E3: the workbook LAW wrote opens in real Excel and reads back
    reader = ps('-Mode', 'read', '-Path', index, '-Sheet', str(sheet_index(index)))
    out, err = reader.communicate(timeout=90)
    print('E3 real Excel reads the LAW-written workbook:', out.strip() or err.strip()[:150])
    shutil.rmtree(base, ignore_errors=True)

    # E2 (T13): the user edits and saves in real Excel between LAW's read and LAW's write
    base = project()
    index = os.path.join(base, 'references', '_INDEX.xlsx')
    proposal = library_sync.plan(base)
    editor = ps('-Mode', 'edit', '-Path', index, '-Sheet', str(sheet_index(index)), '-Text', 'Sua trong Excel that')
    editor.communicate(timeout=90)
    saved = sha(index)
    print('E2 Excel saved the workbook after LAW read it:', saved != sha(TEMPLATE))
    tx = transaction.Transaction(base, 'library-sync')
    try:
        tx.add('references/_INDEX.xlsx', proposal['bytes'], expected_before=proposal['snapshot'])
        print('E2 LAW write: ACCEPTED (would overwrite the Excel save)')
    except transaction.TransactionError as e:
        print('E2 LAW write refused:', str(e).split(' / ')[0])
    print('E2 Excel-saved workbook untouched:', sha(index) == saved)
    library_sync.apply(base)  # a fresh sync on the Excel-saved file
    ws = openpyxl.load_workbook(index)[library_sync.library.SHEET]
    print('E2 fresh sync on the Excel-saved file keeps the edit and adds the source:',
          [(ws.cell(r, 1).value, ws.cell(r, 9).value) for r in range(5, 8) if ws.cell(r, 1).value or ws.cell(r, 9).value])
    shutil.rmtree(base, ignore_errors=True)
finally:
    leftovers = excel_pids() - before_pids
    for pid in leftovers:
        subprocess.run(['taskkill', '/F', '/PID', pid], capture_output=True)
    print('Excel processes started by this script and still running (killed):', sorted(leftovers))
