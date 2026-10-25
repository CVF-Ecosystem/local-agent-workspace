#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Hàm dùng chung cho python-docx: lề ô, nền ô, viền ô, đoạn văn, tiêu đề, bảng có hàng tiêu đề lặp.
Common python-docx helpers: cell margins/shading/borders, paragraphs, headings, tables with repeating header row.
Màu, phông, cỡ chữ là tham số (THEME), không gắn với tổ chức nào. Copy file vào thư mục làm việc hoặc import theo đường dẫn.
Colors/fonts/sizes are parameters (THEME); nothing organization-specific.

    import sys; sys.path.insert(0, 'skills/local/shared/office-documents/scripts')
    from docx_helpers import *
    doc = new_document(margins_cm=(2, 2, 3, 2))          # trên, dưới, trái, phải
    add_heading(doc, '1. Căn cứ', level=1)
    add_p(doc, 'Nội dung', align='justify', first_line_cm=1)
    t = add_table(doc, ['STT', 'Nội dung'], [['1', 'A'], ['2', 'B']], widths_cm=[1.5, 14])
    doc.save('out.docx')

Chỉ trong workspace LAW: văn bản hành chính theo Nghị định 30/2020/NĐ-CP: dùng new_nd30_document() thay cho new_document(); nó gọi cùng mã
với tools/vn_legal_docx.py (A4, lề, Times New Roman 13, cách đoạn 6 pt, số trang lề trên) nên không lệch thông số.
Inside LAW only, use new_nd30_document(): it calls the same code as tools/vn_legal_docx.py.
Sau khi dựng chạy: python tools/check_output.py out.docx
Kiểm: python docx_helpers.py --selftest   (tạo file thử, mở lại, đếm bảng/đoạn)
"""
import os
import sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

THEME = {
    'font': 'Times New Roman', 'size': 13, 'table_size': 11,
    'header_fill': 'D9D9D9', 'header_color': '000000', 'border_color': '000000', 'border_sz': 4,
    'heading_color': '000000',
}
ALIGN = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
         'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}


def _rgb(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_run_font(run, name=None, size=None, bold=None, italic=None, color=None):
    name = name or THEME['font']
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(a), name)                       # tiếng Việt/CJK cần eastAsia
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color: run.font.color.rgb = _rgb(color)


def new_document(margins_cm=(2, 2, 3, 2), font=None, size=None):
    """Tạo tài liệu A4 dọc, lề (trên, dưới, trái, phải), phông mặc định."""
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = (Cm(x) for x in margins_cm)
    st = d.styles['Normal']
    st.font.name = font or THEME['font']; st.font.size = Pt(size or THEME['size'])
    st.element.rPr.rFonts.set(qn('w:eastAsia'), font or THEME['font'])
    return d


def new_nd30_document(title=''):
    """Tài liệu theo thể thức Nghị định 30/2020/NĐ-CP, dùng chung mã với tools/vn_legal_docx.py (nguồn thông số duy nhất).
    Cần file này nằm trong workspace LAW (tìm tools/vn_legal_docx.py ở các thư mục cha); hàm này không có trong chế độ export độc lập. Khi chép riêng, dùng new_document() và native/mẫu đơn vị;
    new_document không tự cung cấp đủ cấu hình ND30. Lưu ý: số trang ở lề trên, trang đầu không hiện."""
    d = os.path.abspath(os.path.dirname(__file__))
    for _ in range(8):
        cand = os.path.join(d, 'tools', 'vn_legal_docx.py')
        if os.path.isfile(cand):
            if os.path.dirname(cand) not in sys.path:
                sys.path.insert(0, os.path.dirname(cand))
            import vn_legal_docx
            return vn_legal_docx.new_nd30_document(title)
        d = os.path.dirname(d)
    raise FileNotFoundError('new_nd30_document chỉ dùng trong workspace LAW; khi export dùng native/mẫu đơn vị. '
                            'hoặc chạy trong workspace LAW. / tools/vn_legal_docx.py not found above this file.')


def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    """Lề trong ô, đơn vị twip (dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    m = tcPr.find(qn('w:tcMar'))
    if m is None:
        m = OxmlElement('w:tcMar'); tcPr.append(m)
    for k, v in (('top', top), ('bottom', bottom), ('start', left), ('end', right)):
        e = m.find(qn('w:' + k))
        if e is None:
            e = OxmlElement('w:' + k); m.append(e)
        e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa')


def set_cell_background(cell, fill):
    """Nền ô (hex không #). Dùng shd val=clear + fill, không dùng solid (Word hiện đen khi thiếu color)."""
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:shd')):
        tcPr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_cell_borders(cell, sz=None, color=None, sides=('top', 'left', 'bottom', 'right'), val='single'):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders'); tcPr.append(b)
    for s in sides:
        e = b.find(qn('w:' + s))
        if e is None:
            e = OxmlElement('w:' + s); b.append(e)
        e.set(qn('w:val'), val); e.set(qn('w:sz'), str(sz or THEME['border_sz']))
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), color or THEME['border_color'])


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:tblHeader'); e.set(qn('w:val'), 'true'); trPr.append(e)


def keep_row_together(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:cantSplit'); e.set(qn('w:val'), 'true'); trPr.append(e)


def add_p(doc_or_cell, text='', align='left', bold=False, italic=False, size=None, color=None,
          first_line_cm=0, space_before=0, space_after=6, line=1.15, keep_next=False):
    p = doc_or_cell.add_paragraph()
    p.alignment = ALIGN[align]
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(space_before), Pt(space_after), line
    if first_line_cm: pf.first_line_indent = Cm(first_line_cm)
    pf.keep_with_next = keep_next
    if text:
        set_run_font(p.add_run(text), size=size or THEME['size'], bold=bold, italic=italic, color=color)
    return p


def add_heading(doc, text, level=1, size=None, color=None, align='left'):
    """Tiêu đề dùng định dạng trực tiếp (không phụ thuộc style Heading của mẫu) nhưng gắn outline level để có mục lục."""
    sizes = {1: THEME['size'] + 1, 2: THEME['size'], 3: THEME['size']}
    p = add_p(doc, text, align=align, bold=True, italic=(level == 3), size=size or sizes.get(level, THEME['size']),
              color=color or THEME['heading_color'], space_before=10 if level == 1 else 6, space_after=4, keep_next=True)
    ppr = p._p.get_or_add_pPr()
    ol = OxmlElement('w:outlineLvl'); ol.set(qn('w:val'), str(level - 1)); ppr.append(ol)
    return p


def style_table_header(table, fill=None, color=None, size=None):
    for c in table.rows[0].cells:
        set_cell_background(c, fill or THEME['header_fill'])
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                set_run_font(r, size=size or THEME['table_size'], bold=True, color=color or THEME['header_color'])
    set_repeat_header(table.rows[0])


def style_table_data_row(row, aligns=None, size=None):
    for i, c in enumerate(row.cells):
        for p in c.paragraphs:
            if aligns: p.alignment = ALIGN[aligns[i]]
            for r in p.runs:
                set_run_font(r, size=size or THEME['table_size'])
    keep_row_together(row)


def add_table(doc, headers, rows, widths_cm=None, aligns=None, header_fill=None):
    """Bảng có viền, hàng tiêu đề lặp lại ở trang sau, hàng không bị cắt đôi. rows: danh sách danh sách chuỗi."""
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = ''
        set_run_font(t.rows[0].cells[i].paragraphs[0].add_run(str(h)), size=THEME['table_size'], bold=True)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ''
            set_run_font(cells[i].paragraphs[0].add_run('' if v is None else str(v)), size=THEME['table_size'])
    for ri, row in enumerate(t.rows):
        for ci, c in enumerate(row.cells):
            set_cell_margins(c); set_cell_borders(c)
            if widths_cm: c.width = Cm(widths_cm[ci])
        if ri: style_table_data_row(row, aligns)
    style_table_header(t, fill=header_fill)
    return t


def _selftest():
    import tempfile, os
    d = new_document()
    add_heading(d, '1. Tiêu đề thử', 1)
    add_p(d, 'Đoạn văn thử có dấu tiếng Việt.', align='justify', first_line_cm=1)
    add_table(d, ['STT', 'Nội dung'], [['1', 'A'], ['2', 'B']], widths_cm=[1.5, 14], aligns=['center', 'left'])
    with tempfile.TemporaryDirectory() as td:
        f = os.path.join(td, 't.docx'); d.save(f)
        d2 = Document(f)
        assert len(d2.tables) == 1 and len(d2.tables[0].rows) == 3
        assert 'tblHeader' in d2.tables[0].rows[0]._tr.xml
    try:
        n = new_nd30_document('thử')
    except FileNotFoundError:
        print('OK: docx_helpers selftest (không thấy tools/vn_legal_docx.py nên bỏ qua new_nd30_document)')
        return
    sec = n.sections[0]
    assert abs(sec.page_width.cm - 21.0) < 0.05 and round(sec.left_margin.cm, 1) == 3.0 and 'PAGE' in sec.header._element.xml
    print('OK: docx_helpers selftest (kể cả new_nd30_document)')


if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    if '--selftest' in sys.argv:
        _selftest()
    else:
        print(__doc__)
