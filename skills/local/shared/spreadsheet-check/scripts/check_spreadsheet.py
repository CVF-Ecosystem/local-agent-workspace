#!/usr/bin/env python3
"""Kiểm tra nhanh file Excel/CSV trước khi lấy số liệu / Quick check of an Excel/CSV file before using its figures.
CHỈ ĐỌC, không sửa file gốc / READ-ONLY, never modifies the original file.

Dùng / Usage:
  python check_spreadsheet.py FILE [--lang vi|en] [--sheet NAME] [--khoa COL1,COL2] [--toi-da-vi-du N]
  (--key is accepted as an alias of --khoa; --max-examples as an alias of --toi-da-vi-du)

- CSV: Python 3.8+ (standard library) is enough. / CSV: chỉ cần thư viện chuẩn.
- XLSX/XLSM: needs openpyxl (pip install openpyxl). / cần openpyxl.
Output is grouped into three levels (impact, questions, notes). It reports findings only and never
concludes that data is right or wrong. / Chỉ nêu phát hiện, không kết luận đúng sai.
"""
import argparse, csv, io, re, statistics, sys, unicodedata
from collections import Counter, defaultdict

ERR = {'#REF!', '#DIV/0!', '#N/A', '#VALUE!', '#NAME?', '#NUM!', '#NULL!'}
TOTAL_WORDS = ('tổng', 'tong', 'total', 'cộng', 'cong', 'sum')

MSG = {
    'hidden': ('Sheet đang ẩn; xác nhận có dùng số liệu ở đây không.', 'Sheet is hidden; confirm whether the figures here should be used.'),
    'merged': ('Có %d vùng ô gộp; ô gộp dễ làm lệch cột khi đọc bằng máy.', '%d merged-cell ranges; merged cells easily misalign columns when read by machine.'),
    'nocache': ('%d/%d ô công thức chưa có giá trị đã tính (file chưa được mở và lưu trong Excel); số đọc ra có thể thiếu.', '%d/%d formula cells have no calculated value (the file was not opened and saved in Excel); figures read may be missing.'),
    'nodata': ('Gần như không có dữ liệu (dưới 2 dòng có nội dung).', 'Almost no data (fewer than 2 non-empty rows).'),
    'header': ('Dòng tiêu đề đoán là dòng %d của sheet; %d cột, %d dòng dữ liệu. Kiểm tra lại nếu đoán sai.', 'Header row guessed to be sheet row %d; %d columns, %d data rows. Check again if the guess is wrong.'),
    'duphdr': ('Tên cột trùng nhau: %s.', 'Duplicate column names: %s.'),
    'empty': ('%d/%d ô trống.', '%d/%d cells empty.'),
    'errs': ('%d ô lỗi công thức (%s).', '%d cells with formula errors (%s).'),
    'mixed': ('lẫn %d số thật và %d số lưu dạng chữ; phép cộng/lọc có thể bỏ sót.', 'mixes %d real numbers and %d numbers stored as text; sums/filters may miss some.'),
    'alltext': ('toàn bộ %d số đang lưu dạng chữ; cần đổi sang số trước khi tính.', 'all %d numbers are stored as text; convert them to numbers before calculating.'),
    'amb': ('%d giá trị dạng 1.234 mơ hồ giữa hai cách đọc: 1.234 là 1234 (dấu chấm ngăn cách nghìn) hoặc 1.234 là một phẩy hai trăm ba mươi bốn (dấu chấm thập phân); xác nhận định dạng số của file.', '%d values like 1.234 are ambiguous between two readings: 1.234 as 1234 (dot as thousands separator) or 1.234 as one point two three four (dot as decimal point); confirm the number format of the file.'),
    'totalalt': (' Tổng này tính giá trị mơ hồ theo cách 1234; nếu đọc theo cách còn lại (dấu chấm là dấu thập phân) thì tổng là %s (lệch %s).', ' This total reads ambiguous values as 1234; under the other reading (dot as decimal point) the total is %s (difference %s).'),
    'neg': ('%d giá trị âm (chỉ nêu, chưa kết luận sai).', '%d negative values (reported only, not concluded to be wrong).'),
    'outl': ('%d giá trị lệch xa so với phần còn lại (ví dụ %s); kiểm tra đơn vị/nhập liệu.', '%d values far from the rest (e.g. %s); check units/data entry.'),
    'textinnum': ('cột chủ yếu là số nhưng có %d ô chữ (%s); phép cộng/lọc sẽ bỏ qua các ô này.', 'column is mostly numbers but has %d text cell(s) (%s); sums/filters will skip them.'),
    'total': ('dòng tổng ghi %s nhưng tổng các dòng trên là %s (lệch %s).', 'total row says %s but the rows above add up to %s (difference %s).'),
    'datefmt': ('ngày viết nhiều định dạng (%s).', 'dates written in several formats (%s).'),
    'dateord': ('ngày dạng a/b/yyyy: xác nhận thứ tự ngày/tháng.', 'dates like a/b/yyyy: confirm day/month order.'),
    'variants': ('%d mục viết khác nhau nhưng có thể là một (hoa/thường, dấu, khoảng trắng): %s.', '%d entries are written differently but may be the same (case, accents, spaces): %s.'),
    'spaces': ('%d giá trị có khoảng trắng thừa đầu/cuối/giữa.', '%d values have extra leading/trailing/inner spaces.'),
    'nokey': ('Không thấy cột khóa đã chỉ định (%s); các cột có: %s.', 'Key column(s) not found (%s); columns present: %s.'),
    'dupkey': ('%d khóa trùng theo cột %s, ví dụ: %s.', '%d duplicate keys on column(s) %s, e.g.: %s.'),
    'duprows': ('%d dòng trùng hoàn toàn dòng khác (chưa chỉ định cột khóa; dùng --khoa để kiểm theo mã).', '%d rows are exact duplicates of other rows (no key column given; use --khoa to check by code).'),
    'col': ('Cột "%s": ', 'Column "%s": '),
    'title': ('Kiểm tra (chỉ đọc): %s | %d sheet: %s', 'Check (read-only): %s | %d sheet(s): %s'),
    'l1': ('ẢNH HƯỞNG ĐẾN SỐ LIỆU (xử lý hoặc xác nhận trước khi dùng)', 'AFFECTS THE FIGURES (fix or confirm before use)'),
    'l2': ('CẦN HỎI (định nghĩa/đơn vị chưa rõ)', 'TO ASK (definition/unit unclear)'),
    'l3': ('LƯU Ý (chưa ảnh hưởng kết quả)', 'NOTES (not affecting the result yet)'),
    'none': ('(không phát hiện trong phạm vi đã kiểm)', '(nothing found within the scope checked)'),
    'trace': ('Dấu vết: check_spreadsheet.py | file %s | sha256 %s | chạy %s', 'Trace: check_spreadsheet.py | file %s | sha256 %s | run %s'),
    'scope': ('Phạm vi: các kiểm tra trên là phát hiện máy móc, không chứng minh file đúng. Chưa kiểm: ý nghĩa nghiệp vụ, công thức phức tạp, liên kết giữa các sheet.', 'Scope: the checks above are mechanical findings and do not prove the file is correct. Not checked: business meaning, complex formulas, links between sheets.'),
    'ext': ('Chỉ hỗ trợ .csv .tsv .xlsx .xlsm (file .xls cũ: hãy lưu thành .xlsx rồi kiểm tra).', 'Only .csv .tsv .xlsx .xlsm are supported (for old .xls: save as .xlsx first).'),
    'noxl': ('Cần openpyxl để đọc Excel: pip install openpyxl (hoặc xuất sheet sang CSV rồi kiểm tra).', 'openpyxl is needed to read Excel: pip install openpyxl (or export the sheet to CSV and check that).'),
    'nosheet': ('Không thấy sheet "%s". Các sheet có: %s', 'Sheet "%s" not found. Sheets present: %s'),
    'colname': ('cột %d', 'column %d'),
    'dmy': ('d/m/yyyy hoặc m/d/yyyy', 'd/m/yyyy or m/d/yyyy'),
    'lv1': ('Ảnh hưởng', 'Impact'),
    'lv2': ('Cần hỏi', 'Ask'),
    'lv3': ('Lưu ý', 'Note'),
}
LANG = 0

def T(key):
    return MSG[key][LANG]

def fold(s):
    s = unicodedata.normalize('NFD', str(s)).replace('đ', 'd').replace('Đ', 'D')
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return re.sub(r'\s+', ' ', s).strip().lower()

def parse_number(s):
    """Trả (giá trị, ghi_chú). ghi_chú: 'ok', 'mơ hồ' (1.234 có thể là 1234 hoặc 1,234), hoặc None."""
    t = str(s).strip().replace(' ', ' ')
    t = re.sub(r'(?i)\s*(vnd|vnđ|đ|usd|\$|%)\s*$', '', t).strip()
    if not t or not re.fullmatch(r'[-+(]?\d[\d.,\s]*\)?', t):
        return None, None
    neg = t.startswith('-') or (t.startswith('(') and t.endswith(')'))
    t = t.strip('-+()').replace(' ', '')
    note = 'ok'
    if ',' in t and '.' in t:
        dec = ',' if t.rfind(',') > t.rfind('.') else '.'
        thou = '.' if dec == ',' else ','
        t = t.replace(thou, '').replace(dec, '.')
    elif ',' in t or '.' in t:
        sep = ',' if ',' in t else '.'
        parts = t.split(sep)
        if len(parts) > 2:
            t = ''.join(parts)
        elif len(parts[1]) == 3:
            note = 'mơ hồ'
            t = ''.join(parts)
        else:
            t = parts[0] + '.' + parts[1]
    try:
        v = float(t)
    except ValueError:
        return None, None
    return (-v if neg else v), note

def load(path, sheet):
    low = path.lower()
    if low.endswith(('.csv', '.tsv', '.txt')):
        raw = open(path, 'rb').read()
        for enc in ('utf-8-sig', 'cp1258', 'cp1252', 'latin-1'):
            try:
                text = raw.decode(enc); break
            except UnicodeDecodeError:
                continue
        dialect = csv.excel_tab if low.endswith('.tsv') else csv.Sniffer().sniff(text[:4096], delimiters=',;\t|') if text.strip() else csv.excel
        rows = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
        return [{'name': 'CSV', 'rows': rows, 'hidden': False, 'merged': 0, 'errors': [], 'formulas': 0, 'nocache': 0}]
    if not low.endswith(('.xlsx', '.xlsm')):
        sys.exit(T('ext'))
    try:
        import openpyxl
    except ImportError:
        sys.exit(T('noxl'))
    wb_v = openpyxl.load_workbook(path, data_only=True, read_only=False)
    wb_f = openpyxl.load_workbook(path, data_only=False, read_only=False)
    out = []
    for ws in wb_v.worksheets:
        if sheet and ws.title != sheet:
            continue
        wf = wb_f[ws.title]
        rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
        formulas = nocache = 0
        for rf, rv in zip(wf.iter_rows(), ws.iter_rows()):
            for cf, cv in zip(rf, rv):
                if isinstance(cf.value, str) and cf.value.startswith('='):
                    formulas += 1
                    if cv.value is None:
                        nocache += 1
        out.append({'name': ws.title, 'rows': rows, 'hidden': ws.sheet_state != 'visible',
                    'merged': len(ws.merged_cells.ranges), 'errors': [], 'formulas': formulas, 'nocache': nocache})
    if sheet and not out:
        sys.exit(T('nosheet') % (sheet, ', '.join(w.title for w in wb_v.worksheets)))
    return out

def fmt(x):
    s = '{:,.0f}'.format(x) if x == int(x) else '{:,.2f}'.format(x)
    if LANG == 0:  # tiếng Việt: dấu chấm ngăn cách nghìn, dấu phẩy thập phân
        s = s.replace(',', '\0').replace('.', ',').replace('\0', '.')
    return s

def is_blank(v):
    return v is None or (isinstance(v, str) and not v.strip())

def check_sheet(sh, keys, ex_n, R):
    name, rows = sh['name'], sh['rows']
    tag = '[%s] ' % name
    if sh['hidden']:
        R['Cần hỏi'].append(tag + T('hidden'))
    if sh['merged']:
        R['Lưu ý'].append(tag + T('merged') % sh['merged'])
    if sh['formulas'] and sh['nocache']:
        R['Ảnh hưởng'].append(tag + T('nocache') % (sh['nocache'], sh['formulas']))
    nums_row = [i + 1 for i, r in enumerate(rows) if any(not is_blank(c) for c in r)]
    rows = [r for r in rows if any(not is_blank(c) for c in r)]
    if len(rows) < 2:
        R['Cần hỏi'].append(tag + T('nodata'))
        return
    # dòng tiêu đề = dòng đầu có >=2 ô chữ không rỗng
    hdr_i = next((i for i, r in enumerate(rows[:15]) if sum(1 for c in r if isinstance(c, str) and c.strip()) >= 2), 0)
    header = [(str(c).strip() if not is_blank(c) else T('colname') % (j + 1)) for j, c in enumerate(rows[hdr_i])]
    data = rows[hdr_i + 1:]
    width = len(header)
    R['Lưu ý'].append(tag + T('header') % (nums_row[hdr_i], width, len(data)))
    dup_h = [h for h, n in Counter(header).items() if n > 1]
    if dup_h:
        R['Cần hỏi'].append(tag + T('duphdr') % ', '.join(dup_h))
    # tổng
    total_idx = None
    if data:
        first = next((str(c) for c in data[-1] if not is_blank(c)), '')
        if any(w in fold(first) for w in TOTAL_WORDS):
            total_idx = len(data) - 1
    body = data[:total_idx] if total_idx is not None else data
    for j, h in enumerate(header):
        col = [(r[j] if j < len(r) else None) for r in body]
        vals = [c for c in col if not is_blank(c)]
        empty = len(col) - len(vals)
        loc = tag + T('col') % h
        if col and empty:
            lv = 'Ảnh hưởng' if (h in keys or fold(h) in [fold(k) for k in keys]) else 'Lưu ý'
            if empty / len(col) > 0.3:
                lv = 'Cần hỏi' if lv == 'Lưu ý' else lv
            R[lv].append(loc + T('empty') % (empty, len(col)))
        errs = [c for c in vals if isinstance(c, str) and c.strip() in ERR]
        if errs:
            R['Ảnh hưởng'].append(loc + T('errs') % (len(errs), ', '.join(sorted(set(e.strip() for e in errs)))))
        nums, texts_num, amb, others, amb_vals = [], [], 0, [], []
        for c in vals:
            if isinstance(c, bool):
                others.append(c)
            elif isinstance(c, (int, float)):
                nums.append(float(c))
            elif isinstance(c, str):
                v, note = parse_number(c)
                if v is not None:
                    texts_num.append(v)
                    amb += (note == 'mơ hồ')
                    if note == 'mơ hồ': amb_vals.append(v)
                elif c.strip() not in ERR:
                    others.append(c)
            else:
                others.append(c)
        if texts_num and nums:
            R['Ảnh hưởng'].append(loc + T('mixed') % (len(nums), len(texts_num)))
        elif texts_num and not others and name != 'CSV':  # CSV: mọi ô đều là chữ, không có ý nghĩa cảnh báo
            R['Ảnh hưởng'].append(loc + T('alltext') % len(texts_num))
        if amb:
            R['Cần hỏi'].append(loc + T('amb') % amb)
        allnum = nums + texts_num
        numcol = bool(allnum) and (not others or (len(allnum) >= 2 and len(allnum) / len(vals) >= 0.6))
        txt = [str(c).strip() for c in others if isinstance(c, str)]
        if allnum and txt and numcol:
            R['Ảnh hưởng'].append(loc + T('textinnum') % (len(txt), ', '.join('"%s"' % x for x in sorted(set(txt))[:ex_n])))
        if numcol and len(allnum) >= 4 and (not others or len(txt) == len(others)):
            neg = sum(1 for x in allnum if x < 0)
            if neg:
                R['Lưu ý'].append(loc + T('neg') % neg)
            q = statistics.quantiles(allnum, n=4)
            iqr = q[2] - q[0]
            if iqr > 0:
                out = [x for x in allnum if x > q[2] + 3 * iqr or x < q[0] - 3 * iqr]
                if out:
                    R['Lưu ý'].append(loc + T('outl') % (len(out), ', '.join(fmt(x) for x in out[:ex_n])))
            if total_idx is not None and j < len(data[total_idx]):
                tv = data[total_idx][j]
                tn = float(tv) if isinstance(tv, (int, float)) and not isinstance(tv, bool) else parse_number(tv)[0] if isinstance(tv, str) else None
                if tn is not None and abs(tn - sum(allnum)) > max(0.5, abs(tn) * 1e-6):
                    msg = loc + T('total') % (fmt(tn), fmt(sum(allnum)), fmt(tn - sum(allnum)))
                    if amb_vals:
                        alt = sum(allnum) - sum(amb_vals) + sum(v / 1000.0 for v in amb_vals)
                        msg += T('totalalt') % (fmt(alt), fmt(tn - alt))
                    R['Ảnh hưởng'].append(msg)
        # ngày
        if others and not numcol and all(isinstance(c, str) for c in others):
            fm = Counter()
            for c in others:
                s = c.strip()
                if re.fullmatch(r'\d{1,2}/\d{1,2}/\d{4}', s): fm[T('dmy')] += 1
                elif re.fullmatch(r'\d{4}-\d{1,2}-\d{1,2}', s): fm['yyyy-mm-dd'] += 1
                elif re.fullmatch(r'\d{1,2}-\d{1,2}-\d{4}', s): fm['d-m-yyyy'] += 1
                elif re.fullmatch(r'\d{1,2}\.\d{1,2}\.\d{4}', s): fm['d.m.yyyy'] += 1
            if len(fm) > 1:
                R['Ảnh hưởng'].append(loc + T('datefmt') % ', '.join('%s: %d' % kv for kv in fm.items()))
            elif sum(fm.values()) == len(others) and fm and 'd/m/yyyy' in next(iter(fm)):
                R['Cần hỏi'].append(loc + T('dateord'))
        # ngày dạng datetime lẫn chữ
        # biến thể chính tả của cùng một mục (cột chữ, ít giá trị khác nhau)
        strs = [str(c) for c in others if isinstance(c, str)]
        if strs and not numcol and len(strs) == len(vals) and len(set(strs)) <= max(30, len(strs) // 3):
            groups = defaultdict(set)
            for s in set(strs):
                groups[fold(s)].add(s)
            var = [sorted(v) for v in groups.values() if len(v) > 1]
            if var:
                eg = '; '.join(' | '.join('"%s"' % x for x in v) for v in var[:ex_n])
                R['Ảnh hưởng'].append(loc + T('variants') % (len(var), eg))
            sp = [s for s in set(strs) if s != s.strip() or '  ' in s]
            if sp and not var:
                R['Lưu ý'].append(loc + T('spaces') % len(sp))
    # trùng
    kidx = [i for i, h in enumerate(header) if h in keys or fold(h) in [fold(k) for k in keys]]
    if keys and not kidx:
        R['Cần hỏi'].append(tag + T('nokey') % (', '.join(keys), ', '.join(header)))
    if kidx:
        c = Counter(tuple(fold(r[i]) if i < len(r) and not is_blank(r[i]) else '' for i in kidx) for r in body)
        d = [(k, n) for k, n in c.items() if n > 1 and any(k)]
        if d:
            eg = ', '.join('%s (x%d)' % ('/'.join(k), n) for k, n in d[:ex_n])
            R['Ảnh hưởng'].append(tag + T('dupkey') % (len(d), '+'.join(header[i] for i in kidx), eg))
    else:
        c = Counter(tuple('' if is_blank(x) else str(x).strip() for x in r) for r in body)
        d = sum(n - 1 for n in c.values() if n > 1)
        if d:
            R['Cần hỏi'].append(tag + T('duprows') % d)

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('file')
    ap.add_argument('--sheet')
    ap.add_argument('--lang', choices=['vi', 'en'], default='vi', help='ngôn ngữ kết quả / output language')
    ap.add_argument('--khoa', '--key', dest='khoa', default='', help='cột khóa/mã, cách nhau bằng dấu phẩy / key columns, comma-separated')
    ap.add_argument('--toi-da-vi-du', '--max-examples', dest='toi_da_vi_du', type=int, default=3)
    a = ap.parse_args()
    global LANG
    LANG = 0 if a.lang == 'vi' else 1
    keys = [k.strip() for k in a.khoa.split(',') if k.strip()]
    R = {'Ảnh hưởng': [], 'Cần hỏi': [], 'Lưu ý': []}
    sheets = load(a.file, a.sheet)
    for sh in sheets:
        check_sheet(sh, keys, a.toi_da_vi_du, R)
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    print(T('title') % (a.file, len(sheets), ', '.join(s['name'] for s in sheets)))
    labels = {'Ảnh hưởng': T('l1'), 'Cần hỏi': T('l2'), 'Lưu ý': T('l3')}
    for k in ('Ảnh hưởng', 'Cần hỏi', 'Lưu ý'):
        print('\n== %s ==' % labels[k])
        for line in R[k] or [T('none')]:
            print('- ' + line)
    print('\n' + T('scope'))
    import datetime, hashlib, os
    try:
        h = hashlib.sha256(open(a.file, 'rb').read()).hexdigest()[:12]
    except OSError:
        h = '?'
    print(T('trace') % (os.path.basename(a.file), h, datetime.datetime.now().strftime('%Y-%m-%d %H:%M')))

if __name__ == '__main__':
    main()
