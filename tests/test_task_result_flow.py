"""A-11 probe: task/result and the checkpoint before integration (T29, T30), on the existing owners, with fake data.

The checkpoint is NOT implemented here. The test extracts the python block marked `law-recipe: task-result-checkpoint` from roles/README.md, executes
it, and drives the real owners with it: the loader/binding of the v3 candidate checker, library.py/library_sync.py, publish_new.py (P-CN1) and
context_update.py (P-CTX1). If the guidance in the docs changes, this probe follows it; nothing in this file is product code.

Every write stays in a fresh temporary project. "Another writer" actions happen between the read and the write inside the test: this is a controlled
interleaving of a cooperative single-writer protocol, not a proof against every agent, the native Write tool or a kernel-level race. A matching
binding here means only "the declared digests match the current task"; it is not a PASS, a grant or user acceptance, and no live agent is involved.

B-03 (stop, timeout, late result, reopen, fallback, acceptance A/B) is probed in `Lifecycle` and `AcceptanceAB` below with the same recipe and the real owners
(publish_new status/observe/recover, context_update, library, current_draft, the v3 accept check). What is SIMULATED is named where it happens: a stopped or
timed-out host (a child process killed at a publish_new fault point, or just silence), quota/tool/provider availability and a Reviewer's absence (declared facts),
and Main's own judgement of the content (an explicit input). None of it proves a real host stopped, a real quota or a real fallback model.
"""
import contextlib
import copy
import hashlib
import importlib
import json
import os
import re
import shutil
import stat
import subprocess
import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'tools'))
sys.path.insert(0, str(ROOT / 'tests'))
import artifacts as art  # noqa: E402
import context_update as cu  # noqa: E402
import library_sync as ls  # noqa: E402
import publish_new as pn  # noqa: E402
import transaction as tx  # noqa: E402
import test_control_v3_workflow as wf  # noqa: E402  (only `wf.Loop`, the D4 builders, is used: it is not collected from this module)

try:
    import openpyxl
except ImportError:  # the source ledger needs it: skipped with a reason, never silently passed
    openpyxl = None

V3 = 'law.cvf.control/3'
EB = 'COOPERATIVE_WORKFLOW'
WORK = 'w-20261019-01'
MAIN = {'actor_id': 'main', 'role': 'main'}
WORKER = {'actor_id': 'worker-1', 'role': 'worker'}
SOURCE_REL = 'references/source-documents/team/quy-che-gia.txt'
TEMPLATE = ROOT / 'skills' / 'local' / 'shared' / 'source-library' / 'assets' / 'source-index-template.xlsx'
CAN_RUN = os.name == 'nt' and pn.volume_supported(tempfile.gettempdir())[0] and openpyxl is not None
STATE_FRAME = '\n'.join('## %s\n\n- (fixture)\n' % h for h in cu.STATE_HEADINGS)
MARKER = re.compile(r'```python\n# law-recipe: task-result-checkpoint\n(.*?)```', re.S)


def sha(data):
    return 'sha256:' + hashlib.sha256(data).hexdigest()


def att(locator='host:msg/0012'):
    return {'source_class': 'USER_CONFIRMATION', 'channel': 'USER_MESSAGE', 'evidence_locator': locator, 'recorded_by': dict(MAIN)}


def ref(kind, ref_id, locator, digit):
    return {'ref_id': ref_id, 'kind': kind, 'locator': locator, 'digest': 'sha256:' + digit * 64, 'attribution': att(locator)}


def recipe_source():
    text = (ROOT / 'roles' / 'README.md').read_text(encoding='utf-8')
    blocks = MARKER.findall(text)
    assert len(blocks) == 1, 'roles/README.md must hold exactly one law-recipe block'
    return blocks[0]


def recipe_from_docs():
    ns = {'__name__': 'law_recipe_from_docs'}
    exec(compile(recipe_source(), 'roles/README.md#law-recipe', 'exec'), ns)
    return types.SimpleNamespace(**ns)


SKIP_RUN = 'needs Windows NTFS local (P-CN1/P-CTX1 owners) and openpyxl: skipped, not passed'


@unittest.skipUnless(CAN_RUN, SKIP_RUN)
class FlowBase(unittest.TestCase):
    """Fixture and helpers shared by the A-11 probe (`Flow`) and the B-03 probe (`Lifecycle`); it holds no test of its own."""

    @classmethod
    def setUpClass(cls):
        cls.r = recipe_from_docs()
        cls.v3 = cls.r.load_v3(ROOT)

    def setUp(self):
        self.project = tempfile.mkdtemp(prefix='law-a11-')
        self.addCleanup(shutil.rmtree, self.project, True)
        os.makedirs(os.path.join(self.project, '.agent', 'runtime'))
        self.write('.agent/STATE.md', STATE_FRAME.encode('utf-8'))
        self.write('.agent/HANDOFF.md', b'# HANDOFF\n\nNone\n')
        os.makedirs(os.path.join(self.project, 'references', 'source-documents', 'team'))
        shutil.copyfile(str(TEMPLATE), self.abs('references/_INDEX.xlsx'))
        self.write(SOURCE_REL, b'quy che gia: muc gia v1\n')
        ls.apply(self.project)  # SRC-000001, revision 1: the ledger rows the checkpoint compares with
        self.ledger0 = self.read('references/_INDEX.xlsx')
        self.source0 = self.read(SOURCE_REL)

    # ------------------------------------------------------------ fixture helpers
    def abs(self, rel):
        return os.path.join(self.project, *rel.split('/'))

    def read(self, rel):
        with open(self.abs(rel), 'rb') as fh:
            return fh.read()

    def write(self, rel, data):
        path = self.abs(rel)
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, 'wb') as fh:
            fh.write(data)

    def tree(self):
        files, dirs = {}, []
        for d, subdirs, names in os.walk(self.project):
            dirs.extend(os.path.relpath(os.path.join(d, n), self.project) for n in subdirs)
            for n in names:
                p = os.path.join(d, n)
                with open(p, 'rb') as fh:
                    files[os.path.relpath(p, self.project)] = hashlib.sha256(fh.read()).hexdigest()
        return files, sorted(dirs)

    def slot(self, task_id, aid):
        return 'working/%s/%s/assignments/%s' % (WORK, task_id, aid)

    def task_rel(self, task_id):
        return 'working/%s/%s/TASK.json' % (WORK, task_id)

    def source_input(self, source_bytes=None):
        return {'path': SOURCE_REL, 'digest': sha(self.source0 if source_bytes is None else source_bytes), 'source_id': 'SRC-000001', 'revision': '1'}

    def task_record(self, task_id, rev, aid, statement='Tóm tắt mục giá', inputs=None, targets=None):
        slot = self.slot(task_id, aid)
        targets = targets if targets is not None else [slot + '/candidate/draft.txt']
        return {
            'schema_version': V3, 'record_type': 'task', 'evaluation_basis': EB, 'work_id': WORK, 'task_id': task_id, 'task_revision': rev,
            'proposed_by': dict(MAIN),
            'intent': {'refs': [ref('user_intent', 'intent-1', 'host:msg/0012', '5')], 'main_interpretation': 'Ban nhap ngan ve muc gia'},
            'authority': {'grant_refs': [ref('grant', 'grant-1', 'PROJECT.md#pham-vi', '1')], 'subject': dict(WORKER)},
            'inputs': [self.source_input()] if inputs is None else inputs,
            'effects': {'actions': [{'operation': 'create_new', 'tool': 'Write', 'targets': targets, 'write_zone': 'working'}],
                        'expected_outputs': list(targets), 'on_missing': 'RETURN_TO_MAIN'},
            'acceptance': {'requirements': [{'req_id': 'REQ-1', 'statement': statement, 'mandatory': True}]},
            'envelope': {'role': 'worker'},
        }

    def put_task(self, task):
        rel = self.task_rel(task['task_id'])
        self.write(rel, json.dumps(task, ensure_ascii=False, indent=2).encode('utf-8'))
        return rel

    def result_record(self, task, aid, body, candidate_rel=None):
        candidate_rel = candidate_rel or self.slot(task['task_id'], aid) + '/candidate/draft.txt'
        self.write(candidate_rel, body)
        outputs = [{'path': candidate_rel, 'digest': sha(body), 'role': 'ACTUAL_BYTES'}]
        binding = self.r.binding_stub(self.v3, task)  # Main computes it at dispatch; the role copies it and adds `outputs`
        binding['outputs'] = outputs
        return {'schema_version': V3, 'record_type': 'result', 'evaluation_basis': EB, 'result_id': 'r-%s-%s' % (task['task_id'], aid),
                'binding': binding, 'produced_by': dict(WORKER), 'artifacts': copy.deepcopy(outputs), 'evidence': [],
                'requirements': [{'req_id': 'REQ-1', 'mandatory': True, 'status': 'UNKNOWN'}],
                'gaps': ['No tool checked the semantic content of the draft.'], 'acceptance_states': ['CANDIDATE'],
                'state_set_by': dict(WORKER)}

    def assignment(self, task_id='t-a', rev=1, aid='a1', body=b'ban nhap muc gia, dua tren quy che v1\n', **task_over):
        """Main writes the current TASK, the role writes RESULT + candidate into the assignment slot. Returns (task_rel, result_rel, task)."""
        task = self.task_record(task_id, rev, aid, **task_over)
        task_rel = self.put_task(task)
        result = self.result_record(task, aid, body)
        result_rel = self.slot(task_id, aid) + '/RESULT.json'
        self.write(result_rel, json.dumps(result, ensure_ascii=False, indent=2).encode('utf-8'))
        return task_rel, result_rel, task

    def checkpoint(self, task_rel, result_rel, staged=None):
        return self.r.checkpoint(self.v3, self.project, task_rel, result_rel, staged=staged)

    def publish_request(self, cp, task_rel, result_rel, target='output/%s/bao-cao-tich-hop.txt' % WORK):
        task = cp['task']
        return {'work_id': task['work_id'], 'target': target, 'content_use': ['internal_draft'],
                'intent_ref': {k: v for k, v in task['intent']['refs'][0].items() if k != 'kind'},
                'main_interpretation': task['intent']['main_interpretation'],
                'requirements': [{'req_id': 'REQ-PUB-1', 'statement': 'Tích hợp đúng candidate đã checkpoint', 'mandatory': True}],
                'grant_refs': [{k: v for k, v in g.items() if k != 'kind'} for g in task['authority']['grant_refs']],
                'subject': dict(MAIN), 'inputs': self.r.publish_inputs(cp, task_rel, result_rel)}

    def new_rid(self):
        return os.urandom(16).hex()

    def edit_ledger(self, edit):
        path = self.abs('references/_INDEX.xlsx')
        wb = openpyxl.load_workbook(path)
        try:
            edit(wb['Danh mục'])
            wb.save(path)
        finally:
            wb.close()

    def assertRefusedBeforeMutation(self, res, before):
        self.assertEqual((res['outcome'], res['reason']), ('REFUSED', 'INPUT_CHANGED_SINCE_READ'), res)
        self.assertEqual(res['wrote'], [])
        self.assertEqual(self.tree(), before)  # nothing created, not even the runtime publish folder


@unittest.skipUnless(CAN_RUN, SKIP_RUN)
class Flow(FlowBase):
    # ------------------------------------------------------------ positive flow
    def test_two_assignment_slots_pass_the_checkpoint_and_publish_the_checked_bytes(self):
        a = self.assignment('t-a', 1, 'a1', body=b'ban nhap A\n')
        b = self.assignment('t-b', 1, 'a1', body=b'ban nhap B\n')
        shared = {'.agent/STATE.md', '.agent/HANDOFF.md', 'references/_INDEX.xlsx', SOURCE_REL}
        before = self.tree()
        cps = [self.checkpoint(t, r) for t, r, _ in (a, b)]
        self.assertEqual(self.tree(), before, 'a checkpoint reads only')
        for cp, (_, result_rel, task) in zip(cps, (a, b)):
            self.assertEqual(cp['problems'], [])
            slot = os.path.dirname(result_rel).replace('\\', '/')
            targets = [t for act in task['effects']['actions'] for t in act['targets']]
            self.assertTrue(targets and all(t.startswith(slot + '/') for t in targets))
            self.assertFalse(shared & set(targets))
            self.assertEqual(cp['result']['acceptance_states'], ['CANDIDATE'])
            self.assertEqual(sha(cp['candidate']), cp['result']['artifacts'][0]['digest'])
        slots = [os.path.dirname(a[1]), os.path.dirname(b[1])]
        self.assertNotEqual(slots[0], slots[1])
        self.assertFalse(any(x.startswith(y + '/') or y.startswith(x + '/') for x in slots for y in slots if x != y))
        # one writer: Main passes the very bytes the checkpoint compared
        cp, (task_rel, result_rel, task) = cps[0], a
        rid = self.new_rid()
        res = pn.publish(self.project, self.publish_request(cp, task_rel, result_rel), cp['candidate'], request_id=rid)
        self.assertEqual(res['outcome'], 'PUBLISHED', res)
        published = self.read('output/%s/bao-cao-tich-hop.txt' % WORK)
        self.assertEqual(published, cp['candidate'])
        self.assertEqual(sha(published), cp['result']['artifacts'][0]['digest'])
        prepared = pn.cj_loads(self.read('.agent/runtime/publish/%s.prepared.json' % rid))
        pub_task = prepared['task']
        self.assertEqual((pub_task['work_id'], pub_task['task_id'], pub_task['task_revision']), (task['work_id'], 'pcn1-' + rid, 1))
        bound = {i['path']: i['digest'] for i in pub_task['inputs']}
        self.assertEqual(bound[task_rel], sha(self.read(task_rel)))      # raw current TASK
        self.assertEqual(bound[result_rel], sha(self.read(result_rel)))  # raw RESULT
        self.assertEqual(bound[cp['result']['artifacts'][0]['path']], sha(published))
        self.assertEqual(bound['references/_INDEX.xlsx'], sha(self.read('references/_INDEX.xlsx')))  # ledger metadata was used
        self.assertEqual(json.loads(self.read(result_rel))['acceptance_states'], ['CANDIDATE'])  # publishing never turned it into an acceptance

    def test_small_job_needs_no_task_or_result(self):
        before = self.read('.agent/STATE.md')
        new = before.replace(b'(fixture)', b'cap nhat nho', 1)
        res = cu.update(self.project, 'state', new, expected_before=sha(before))
        self.assertEqual(res['result'], 'WRITTEN', res)
        self.assertFalse(os.path.exists(self.abs('working')), 'no TASK.json/RESULT.json is needed for a small job')

    # ------------------------------------------------------------ T30: requirement changed, source unchanged
    def test_task_revision_bump_voids_the_old_result_and_the_old_request(self):
        task_rel, result_rel, task1 = self.assignment('t-a', 1, 'a1')
        cp1 = self.checkpoint(task_rel, result_rel)
        self.assertEqual(cp1['problems'], [])
        request1, rid1, staged1 = self.publish_request(cp1, task_rel, result_rel), self.new_rid(), cp1['candidate']
        old_result = self.read(result_rel)
        # Main changes the requirement and re-dispatches into a new slot: TASK is rewritten with task_revision 2
        task2 = self.task_record('t-a', 2, 'a2', statement='Tóm tắt mục giá và nêu điểm chưa chắc')
        self.put_task(task2)
        stale = self.checkpoint(task_rel, result_rel)
        self.assertTrue(any(p.startswith('BINDING B-REF') for p in stale['problems']), stale['problems'])
        self.assertTrue(any(p.startswith('BINDING B-TASK') for p in stale['problems']), stale['problems'])
        self.assertEqual(self.read(result_rel), old_result, 'the old RESULT keeps the binding it was made with')
        self.assertEqual(json.loads(old_result)['binding']['task_revision'], 1)
        before = self.tree()
        self.assertRefusedBeforeMutation(pn.publish(self.project, request1, staged1, request_id=rid1), before)
        # the new assignment, bound to revision 2, passes and can be published
        result2 = self.result_record(task2, 'a2', b'ban nhap muc gia, co neu diem chua chac\n')
        result2_rel = self.slot('t-a', 'a2') + '/RESULT.json'
        self.write(result2_rel, json.dumps(result2, ensure_ascii=False, indent=2).encode('utf-8'))
        cp2 = self.checkpoint(task_rel, result2_rel)
        self.assertEqual(cp2['problems'], [])
        res = pn.publish(self.project, self.publish_request(cp2, task_rel, result2_rel), cp2['candidate'], request_id=self.new_rid())
        self.assertEqual(res['outcome'], 'PUBLISHED', res)

    # ------------------------------------------------------------ T30: source changed
    def test_source_changed_before_the_worker_returns_with_the_ledger_not_synced(self):
        task_rel, result_rel, _ = self.assignment('t-a', 1, 'a1')
        cp = self.checkpoint(task_rel, result_rel)
        self.assertEqual(cp['problems'], [])
        request, rid, staged = self.publish_request(cp, task_rel, result_rel), self.new_rid(), cp['candidate']
        self.write(SOURCE_REL, b'quy che gia: muc gia v2 (da sua)\n')  # the ledger is not synced: it still says revision 1
        late = self.checkpoint(task_rel, result_rel)
        self.assertIn('INPUT_CHANGED ' + SOURCE_REL, late['problems'])
        self.assertFalse([p for p in late['problems'] if p.startswith('LEDGER')], 'the ledger alone would not have noticed')
        before = self.tree()
        self.assertRefusedBeforeMutation(pn.publish(self.project, request, staged, request_id=rid), before)

    def test_missing_source_or_unreadable_ledger_is_not_treated_as_unchanged(self):
        task_rel, result_rel, _ = self.assignment('t-a', 1, 'a1')
        os.remove(self.abs(SOURCE_REL))
        self.assertIn('INPUT %s MISSING_OR_NOT_REGULAR' % SOURCE_REL, self.checkpoint(task_rel, result_rel)['problems'])
        self.write(SOURCE_REL, self.source0)
        self.write('references/_INDEX.xlsx', b'not a workbook')
        problems = self.checkpoint(task_rel, result_rel)['problems']
        self.assertIn('LEDGER_UNREADABLE unreadable', problems)

    def test_ledger_metadata_changes_while_the_source_bytes_stay_the_same(self):
        def bump_revision(ws):
            ws.cell(5, 22).value = 2

        def clear_id(ws):
            ws.cell(5, 1).value = 'SRC-000009'

        def duplicate_row(ws):
            for col in range(1, ws.max_column + 1):
                ws.cell(6, col).value = ws.cell(5, col).value
        cases = [('revision changed', bump_revision, 'LEDGER_METADATA_CHANGED SRC-000001'),
                 ('id no longer present', clear_id, 'LEDGER_ID_COUNT SRC-000001=0'),
                 ('id duplicated', duplicate_row, 'LEDGER_ID_COUNT SRC-000001=2')]
        for label, edit, expected in cases:
            with self.subTest(label):
                self.write('references/_INDEX.xlsx', self.ledger0)
                task_rel, result_rel, _ = self.assignment('t-a', 1, 'a1')
                cp = self.checkpoint(task_rel, result_rel)
                self.assertEqual(cp['problems'], [])
                request, rid, staged = self.publish_request(cp, task_rel, result_rel), self.new_rid(), cp['candidate']
                self.edit_ledger(edit)
                self.assertEqual(self.read(SOURCE_REL), self.source0)
                problems = self.checkpoint(task_rel, result_rel)['problems']
                self.assertIn(expected, problems)
                self.assertFalse([p for p in problems if p.startswith('INPUT')], 'the source bytes did not change')
                before = self.tree()  # the request that bound the old ledger snapshot cannot be reused
                self.assertRefusedBeforeMutation(pn.publish(self.project, request, staged, request_id=rid), before)

    def test_source_synced_to_a_new_revision_is_caught_by_bytes_and_by_the_ledger(self):
        task_rel, result_rel, _ = self.assignment('t-a', 1, 'a1')
        self.write(SOURCE_REL, b'quy che gia: muc gia v2\n')
        ls.apply(self.project)  # a legitimate sync: SRC-000001 is now revision 2
        problems = self.checkpoint(task_rel, result_rel)['problems']
        self.assertIn('INPUT_CHANGED ' + SOURCE_REL, problems)
        self.assertIn('LEDGER_METADATA_CHANGED SRC-000001', problems)

    # ------------------------------------------------------------ changed after the checkpoint
    def test_task_result_or_candidate_changed_after_the_checkpoint_voids_the_request(self):
        def edit_task(rel):
            self.write(rel, self.read(rel) + b'\n')  # even a byte of whitespace: the digest is of the raw bytes

        def edit_result(rel):
            self.write(rel, self.read(rel) + b'\n')

        def edit_candidate(rel):
            self.write(rel, self.read(rel) + b'sua sau checkpoint\n')
        for label, which, edit in (('TASK', 0, edit_task), ('RESULT', 1, edit_result), ('candidate', 2, edit_candidate)):
            with self.subTest(label):
                aid = 'a-' + label.lower()
                task_rel, result_rel, _ = self.assignment('t-' + label.lower(), 1, aid)
                cp = self.checkpoint(task_rel, result_rel)
                self.assertEqual(cp['problems'], [])
                request, rid, staged = self.publish_request(cp, task_rel, result_rel), self.new_rid(), cp['candidate']
                edit((task_rel, result_rel, cp['result']['artifacts'][0]['path'])[which])
                before = self.tree()
                self.assertRefusedBeforeMutation(pn.publish(self.project, request, staged, request_id=rid), before)

    def test_a_result_with_a_tampered_binding_or_a_claimed_acceptance_is_stopped(self):
        def drop_outputs(res):
            del res['binding']['outputs']

        def other_inputs(res):
            res['binding']['inputs'][0]['digest'] = 'sha256:' + 'e' * 64

        def other_task_digest(res):
            res['binding']['accepted_task_digest'] = 'sha256:' + 'd' * 64

        def claims_user_acceptance(res):
            res['acceptance_states'] = ['CANDIDATE', 'USER_ACCEPTED']
        cases = [('no outputs bound', drop_outputs, 'BINDING B-OUT-RESULT=NOT_BOUND'),
                 ('inputs differ from the task', other_inputs, 'BINDING B-INPUTS=MISMATCH'),
                 ('task digest differs', other_task_digest, 'BINDING B-TASK=MISMATCH'),
                 ('a worker claims acceptance', claims_user_acceptance, 'RECORD_REJECTED result SCHEMA_INVALID')]
        for label, tamper, expected in cases:
            with self.subTest(label):
                task_rel, result_rel, _ = self.assignment('t-' + str(abs(hash(label)) % 10 ** 6), 1, 'a1')
                result = json.loads(self.read(result_rel))
                tamper(result)
                self.write(result_rel, json.dumps(result, ensure_ascii=False, indent=2).encode('utf-8'))
                self.assertIn(expected, self.checkpoint(task_rel, result_rel)['problems'])

    # ------------------------------------------------------------ candidate -> staged
    def test_candidate_must_match_result_and_staged_must_be_the_checked_bytes(self):
        task_rel, result_rel, _ = self.assignment('t-a', 1, 'a1')
        good = self.checkpoint(task_rel, result_rel)
        self.assertEqual(self.checkpoint(task_rel, result_rel, staged=good['candidate'])['problems'], [])
        other = self.checkpoint(task_rel, result_rel, staged=good['candidate'] + b'x')
        self.assertEqual(other['problems'], ['STAGED_NOT_CANDIDATE'])
        candidate_rel = good['result']['artifacts'][0]['path']
        self.write(candidate_rel, good['candidate'] + b'doi sau khi Worker tra')
        changed = self.checkpoint(task_rel, result_rel, staged=good['candidate'])
        self.assertIn('CANDIDATE_CHANGED', changed['problems'])
        self.assertIn('STAGED_NOT_CANDIDATE', changed['problems'])
        self.assertIsNone(changed['candidate'])

    def test_the_writer_alone_does_not_relate_staged_bytes_to_the_candidate(self):
        # Pins a limitation: P-CN1 binds the candidate as an input but accepts any staged bytes; only the checkpoint relates them.
        task_rel, result_rel, _ = self.assignment('t-a', 1, 'a1')
        cp = self.checkpoint(task_rel, result_rel)
        wrong = cp['candidate'] + b'khong phai candidate\n'
        self.assertEqual(self.checkpoint(task_rel, result_rel, staged=wrong)['problems'], ['STAGED_NOT_CANDIDATE'])
        res = pn.publish(self.project, self.publish_request(cp, task_rel, result_rel), wrong, request_id=self.new_rid())
        self.assertEqual(res['outcome'], 'PUBLISHED', res)
        self.assertEqual(self.read('output/%s/bao-cao-tich-hop.txt' % WORK), wrong)

    def test_a_candidate_behind_a_junction_is_not_read(self):
        task_rel, result_rel, _ = self.assignment('t-a', 1, 'a1')
        candidate_dir = self.abs(self.slot('t-a', 'a1') + '/candidate')
        outside = self.abs('outside-the-slot')
        os.makedirs(outside)
        shutil.copy(os.path.join(candidate_dir, 'draft.txt'), outside)
        shutil.rmtree(candidate_dir)
        made = subprocess.run(['cmd', '/c', 'mklink', '/J', candidate_dir, outside], capture_output=True, text=True)
        if made.returncode != 0:
            self.skipTest('cannot create a junction here: %s' % made.stderr.strip())
        self.addCleanup(lambda: os.path.isdir(candidate_dir) and os.rmdir(candidate_dir))  # removes the junction, never its target
        problems = self.checkpoint(task_rel, result_rel)['problems']
        self.assertIn('CANDIDATE MISSING_OR_NOT_REGULAR', problems)

    def test_noncanonical_paths_are_rejected_before_reading_records_inputs_or_candidates(self):
        task_rel, result_rel, _ = self.assignment('t-path', 1, 'a1')
        for rel in ('../PROJECT.md', '/PROJECT.md', 'C:/PROJECT.md', task_rel.replace('/TASK.json', '/./TASK.json')):
            with self.subTest(raw_path=rel):
                self.assertEqual(self.r.read_file(self.project, rel), (None, 'PATH_UNSUPPORTED_FORM'))
        self.assertIn('RECORD_TASK PATH_UNSUPPORTED_FORM', self.checkpoint(task_rel.replace('/TASK.json', '/./TASK.json'), result_rel)['problems'])
        self.assertIn('RECORD_RESULT PATH_UNSUPPORTED_FORM', self.checkpoint(task_rel, result_rel.replace('/RESULT.json', '/./RESULT.json'))['problems'])
        bad_source = self.source_input()
        bad_source['path'] = SOURCE_REL.replace('/team/', '/team/../team/')
        task_rel, result_rel, _ = self.assignment('t-input', 1, 'a1', inputs=[bad_source])
        self.assertIn('INPUT %s PATH_UNSUPPORTED_FORM' % bad_source['path'], self.checkpoint(task_rel, result_rel)['problems'])
        # This prefix passes a naive slot.startswith test, but resolves outside that assignment.
        slot = self.slot('t-candidate', 'a1')
        escaped = slot + '/../../../outside-slot.txt'
        task = self.task_record('t-candidate', 1, 'a1', targets=[escaped])
        task_rel = self.put_task(task)
        result = self.result_record(task, 'a1', b'outside assignment bytes', candidate_rel=escaped)
        result_rel = slot + '/RESULT.json'
        self.write(result_rel, json.dumps(result).encode('utf-8'))
        cp = self.checkpoint(task_rel, result_rel)
        self.assertIn('TARGET_UNSUPPORTED_FORM', cp['problems'])
        self.assertIn('CANDIDATE PATH_UNSUPPORTED_FORM', cp['problems'])
        self.assertIsNone(cp['candidate'])

    def test_a_ledger_behind_a_junction_is_not_read_or_bound(self):
        task_rel, result_rel, _ = self.assignment('t-ledger-path', 1, 'a1')
        refs = self.abs('references')
        outside = self.abs('outside-references')
        os.rename(refs, outside)
        made = subprocess.run(['cmd', '/c', 'mklink', '/J', refs, outside], capture_output=True, text=True)
        if made.returncode != 0:
            self.skipTest('cannot create a junction here: %s' % made.stderr.strip())
        self.addCleanup(lambda: os.path.isdir(refs) and os.rmdir(refs))
        cp = self.checkpoint(task_rel, result_rel)
        self.assertIn('LEDGER_UNREADABLE unsafe_path', cp['problems'])
        self.assertIsNone(cp['ledger_snapshot'])

    def test_context_source_metadata_is_checked_and_the_ledger_snapshot_is_bound(self):
        task = self.task_record('t-context', 1, 'a1', inputs=[])
        task['context_refs'] = [self.source_input()]
        task_rel = self.put_task(task)
        result = self.result_record(task, 'a1', b'context source draft')
        result_rel = self.slot('t-context', 'a1') + '/RESULT.json'
        self.write(result_rel, json.dumps(result).encode('utf-8'))
        cp = self.checkpoint(task_rel, result_rel)
        self.assertEqual(cp['problems'], [])
        bound = {o['path']: o['digest'] for o in self.r.publish_inputs(cp, task_rel, result_rel)}
        self.assertEqual(bound[SOURCE_REL], sha(self.source0))
        self.assertEqual(bound.get('references/_INDEX.xlsx'), sha(self.ledger0))
        self.edit_ledger(lambda ws: setattr(ws.cell(5, 22), 'value', 2))
        self.assertIn('LEDGER_METADATA_CHANGED SRC-000001', self.checkpoint(task_rel, result_rel)['problems'])

    def test_generic_context_bytes_are_bound_into_the_publish_request(self):
        task = self.task_record('t-context-bytes', 1, 'a1', inputs=[])
        context_rel = 'working/context.txt'
        self.write(context_rel, b'context version one')
        task['context_refs'] = [{'path': context_rel, 'digest': sha(self.read(context_rel))}]
        task_rel = self.put_task(task)
        result = self.result_record(task, 'a1', b'context draft')
        result_rel = self.slot('t-context-bytes', 'a1') + '/RESULT.json'
        self.write(result_rel, json.dumps(result).encode('utf-8'))
        cp = self.checkpoint(task_rel, result_rel)
        self.assertEqual(cp['problems'], [])
        request = self.publish_request(cp, task_rel, result_rel)
        self.assertIn(task['context_refs'][0], request['inputs'])
        self.write(context_rel, b'context version two')
        self.assertIn('INPUT_CHANGED ' + context_rel, self.checkpoint(task_rel, result_rel)['problems'])
        self.assertRefusedBeforeMutation(pn.publish(self.project, request, cp['candidate'], request_id=self.new_rid()), self.tree())

    def test_a_result_that_targets_a_shared_file_is_stopped_at_the_checkpoint(self):
        for target in ('.agent/STATE.md', 'output/%s/bao-cao-tich-hop.txt' % WORK, 'references/_INDEX.xlsx'):
            with self.subTest(target):
                task_rel, result_rel, _ = self.assignment('t-a', 1, 'a1', targets=[target])
                self.assertIn('TARGET_OUTSIDE_SLOT', self.checkpoint(task_rel, result_rel)['problems'])

    # ------------------------------------------------------------ T29: owners refuse stale writes, current bytes kept
    def test_context_update_with_a_stale_expected_before_is_refused_and_keeps_the_other_writers_bytes(self):
        read_by_main = self.read('.agent/STATE.md')
        other_writer = read_by_main.replace(b'(fixture)', b'nguoi khac vua cap nhat', 1)
        self.write('.agent/STATE.md', other_writer)  # controlled interleaving between Main's read and Main's write
        mine = read_by_main.replace(b'(fixture)', b'ban cua Main', 1)
        for call in (cu.preview, cu.update):
            res = call(self.project, 'state', mine, expected_before=sha(read_by_main))
            self.assertEqual((res['result'], res['reason']), ('CONFLICT', 'EXPECTED_BEFORE_MISMATCH'), res)
            self.assertEqual(self.read('.agent/STATE.md'), other_writer)
        res = cu.update(self.project, 'state', mine, expected_before=sha(other_writer))
        self.assertEqual(res['result'], 'WRITTEN', res)

    def test_library_apply_with_a_stale_preview_is_refused_and_keeps_the_index(self):
        preview = ls.preview_binding(ls.plan(self.project))
        self.write('references/source-documents/team/phu-luc.txt', b'phu luc moi\n')  # another writer adds a source after the preview
        index_before = self.read('references/_INDEX.xlsx')
        with self.assertRaises(ls.PreviewMismatch) as caught:
            ls.apply(self.project, expected_preview=preview)
        self.assertEqual(caught.exception.reason, 'sources')
        self.assertEqual(self.read('references/_INDEX.xlsx'), index_before)
        fresh = ls.preview_binding(ls.plan(self.project))
        applied = ls.apply(self.project, expected_preview=fresh)
        self.assertEqual([c['action'] for c in applied['changes']], ['add'])

    def test_publish_to_an_existing_target_is_a_conflict_and_keeps_the_current_output(self):
        task_rel, result_rel, _ = self.assignment('t-a', 1, 'a1')
        cp = self.checkpoint(task_rel, result_rel)
        current = b'ban hien hanh do nguoi dung dat\n'
        self.write('output/%s/bao-cao-tich-hop.txt' % WORK, current)
        res = pn.publish(self.project, self.publish_request(cp, task_rel, result_rel), cp['candidate'], request_id=self.new_rid())
        self.assertEqual((res['outcome'], res['reason']), ('CONFLICT', 'TARGET_EXISTS'), res)
        self.assertEqual(self.read('output/%s/bao-cao-tich-hop.txt' % WORK), current)


# ---------------------------------------------------------------- B-03: stop, timeout, late result, reopen, fallback, content, acceptance A/B
CHILD_PUBLISH = '''import sys
tools, project, request_path, staged_path, rid = sys.argv[1:6]
sys.path.insert(0, tools)
import publish_new as pn
with open(request_path, 'rb') as fh:
    request = pn.cj_loads(fh.read())
with open(staged_path, 'rb') as fh:
    staged = fh.read()
print(pn.publish(project, request, staged, request_id=rid))
'''

# A fresh interpreter: it runs the recipe text taken from roles/README.md against the disk only. The regex reads the note the way Main would read its own
# STATE/HANDOFF; it is test-side, there is no parser or marker in the product.
CHILD_REOPEN = '''import json, os, re, sys
recipe, law, project, payload = sys.argv[1:5]
spec = json.loads(payload)
ns = {'__name__': 'law_recipe_reopen'}
with open(recipe, encoding='utf-8') as fh:
    exec(compile(fh.read(), 'recipe', 'exec'), ns)
v3 = ns['load_v3'](law)
out = {'pid': os.getpid(), 'results': {}}
task, _, why = ns['read_record'](v3, project, spec['task_rel'], 'task')
out['task_revision'] = task['task_revision'] if task else why
for name, rel in spec['results'].items():
    out['results'][name] = ns['checkpoint'](v3, project, spec['task_rel'], rel)['problems']
with open(os.path.join(project, '.agent', 'HANDOFF.md'), encoding='utf-8') as fh:
    note = fh.read()
out['note_says_writer_not_called'] = 'writer chưa được gọi' in note
found = re.search(r'request ID ([0-9a-f]{32}), đích (\\S+?);', note)
if found:
    out['reconcile'] = ns['reconcile_publish'](project, found.group(1), found.group(2))
print(json.dumps(out))
'''


@unittest.skipUnless(CAN_RUN, SKIP_RUN)
class Lifecycle(FlowBase):
    def setUp(self):
        super().setUp()
        self.scratch = tempfile.mkdtemp(prefix='law-b03-')  # request files and child scripts stay out of the project tree
        self.addCleanup(shutil.rmtree, self.scratch, True)

    # ------------------------------------------------------------ helpers
    def handoff_add(self, text):
        before = self.read('.agent/HANDOFF.md')
        return self.r.record_before_effect(self.project, 'handoff', before + text.encode('utf-8'), sha(before))

    def child(self, name, source, *args, fault=None):
        path = os.path.join(self.scratch, name)
        with open(path, 'w', encoding='utf-8') as fh:
            fh.write(source)
        env = dict(os.environ, PYTHONDONTWRITEBYTECODE='1', PYTHONUTF8='1')
        env.pop(pn.FAULT_ENV, None)
        if fault:
            env[pn.FAULT_ENV] = fault
        return subprocess.run([sys.executable, '-I', '-B', path] + list(args), capture_output=True, text=True, encoding='utf-8', env=env, timeout=120,
                              cwd=self.scratch)

    def no_writer(self):
        """Any call of a P-CN1 writer during the block fails the test: reconcile and the decision functions only read."""
        stack = contextlib.ExitStack()
        for name in ('publish', 'recover', 'discard_staged'):
            stack.enter_context(mock.patch.object(pn, name, side_effect=AssertionError('%s must not be called here' % name)))
        return stack

    def killed_publish(self, point):
        """The writer is really called in a child process that dies at a publish_new fault point: the caller never gets a response (SIMULATED timeout)."""
        task_rel, result_rel, _ = self.assignment('t-a', 1, 'a1')
        cp = self.checkpoint(task_rel, result_rel)
        self.assertEqual(cp['problems'], [])
        request, rid = self.publish_request(cp, task_rel, result_rel), self.new_rid()
        ok, res = self.handoff_add('Lần giao a1: t-a rev 1, slot a1; writer P-CN1, request ID %s, đích %s; writer đã được gọi, chưa nhận phản hồi.\n' % (rid, request['target']))
        self.assertTrue(ok, res)  # the mapping is on disk BEFORE the effect
        request_path, staged_path = os.path.join(self.scratch, 'request.json'), os.path.join(self.scratch, 'staged.bin')
        with open(request_path, 'wb') as fh:
            fh.write(pn.cj_dumps(request))
        with open(staged_path, 'wb') as fh:
            fh.write(cp['candidate'])
        proc = self.child('child_publish.py', CHILD_PUBLISH, str(ROOT / 'tools'), self.project, request_path, staged_path, rid, fault='pcn1-kill:' + point)
        self.assertEqual(proc.returncode, 70, proc.stdout + proc.stderr)
        return task_rel, result_rel, cp, request, rid

    def void(self, task_rel, result_rel, rev=2, aid='a1'):
        return self.r.void_assignment(self.v3, self.project, task_rel, self.task_record('t-a', rev, aid), result_rel)

    def facts(self, **over):
        base = {'voided': True, 'effect': 'NONE_ATTEMPTED', 'invocation': 'ENDED', 'grant_ok': True, 'capability': 'AVAILABLE', 'retry_left': 1, 'cost_ok': True, 'no_call_observed': True}  # SIMULATED live no-call observation; never inferred from an old note
        base.update(over)
        return base

    # ------------------------------------------------------------ stop before any effect
    def test_stop_before_any_effect_voids_the_result_and_the_prepared_request(self):
        task_rel, result_rel, _ = self.assignment('t-a', 1, 'a1')
        cp = self.checkpoint(task_rel, result_rel)
        self.assertEqual(cp['problems'], [], 'until the revision is raised the right RESULT still MATCHes: this is the gap the note must name')
        request, rid, staged = self.publish_request(cp, task_rel, result_rel), self.new_rid(), cp['candidate']
        ok, res = self.handoff_add('Lần giao a1: t-a rev 1, slot a1; writer P-CN1, request ID %s, đích %s; writer chưa được gọi. Đã yêu cầu dừng, host chưa xác nhận.\n'
                                   % (rid, request['target']))
        self.assertTrue(ok, res)
        candidate_rel = cp['result']['artifacts'][0]['path']
        old_result, old_candidate = self.read(result_rel), self.read(candidate_rel)
        voided = self.void(task_rel, result_rel)
        self.assertEqual(voided, {'voided': True, 'problems': [], 'host_stop': 'UNCONFIRMED'})
        problems = self.checkpoint(task_rel, result_rel)['problems']
        self.assertTrue(any(p.startswith('BINDING B-REF') for p in problems), problems)
        self.assertEqual((self.read(result_rel), self.read(candidate_rel)), (old_result, old_candidate), 'the old RESULT and candidate keep their bytes')
        before = self.tree()
        self.assertRefusedBeforeMutation(pn.publish(self.project, request, staged, request_id=rid), before)  # raw-TASK input guard of the request made before
        status = pn.status(self.project, rid)['items'][0]
        self.assertEqual((status['files'], status['terminal']), ([], None), 'no effect for that request ID')

    def test_failed_or_unconfirmed_revision_write_keeps_everything_on_hold(self):
        task_rel, result_rel, _ = self.assignment('t-a', 1, 'a1')
        raw = self.read(task_rel)
        os.chmod(self.abs(task_rel), stat.S_IREAD)  # the operating system refuses the replace (a real failure, not a flag set by the test)
        self.addCleanup(os.chmod, self.abs(task_rel), stat.S_IWRITE)
        res = self.void(task_rel, result_rel)
        self.assertEqual((res['voided'], res['problems']), (False, ['TASK_WRITE_FAILED']))
        self.assertEqual(self.read(task_rel), raw)
        self.assertFalse(os.path.exists(self.abs(task_rel) + '.new'), 'a failed write leaves no temporary file')
        self.assertEqual(self.checkpoint(task_rel, result_rel)['problems'], [], 'the old RESULT still MATCHes, so the recipe must say hold, not voided')
        self.assertEqual(self.r.next_step(self.facts(voided=res['voided'])), {'step': 'WAIT', 'reasons': ['ASSIGNMENT_NOT_VOIDED']})
        os.chmod(self.abs(task_rel), stat.S_IWRITE)
        res = self.void(task_rel, result_rel, rev=1)
        self.assertEqual((res['voided'], res['problems']), (False, ['REVISION_NOT_RAISED']))
        self.assertEqual(self.read(task_rel), raw)
        broken = self.task_record('t-a', 2, 'a1')
        del broken['intent']  # rejected before writing: a valid TASK is preserved
        res = self.r.void_assignment(self.v3, self.project, task_rel, broken, result_rel)
        self.assertFalse(res['voided'])
        self.assertEqual(res['problems'], ['NEW_TASK_REJECTED'])
        self.assertEqual(self.read(task_rel), raw)
        replace = os.replace

        def replace_then_corrupt(source, destination):
            replace(source, destination)
            self.write(task_rel, b'{}')  # controlled interleaving after write, before owner read-back

        with mock.patch.object(os, 'replace', side_effect=replace_then_corrupt):
            res = self.void(task_rel, result_rel)
        self.assertFalse(res['voided'])
        self.assertTrue(res['problems'][0].startswith('TASK_READBACK_MISMATCH'), res)
        self.assertTrue(self.checkpoint(task_rel, result_rel)['problems'], 'an unreadable TASK fails closed')

    def test_a_note_that_cannot_be_written_means_the_effect_has_not_started(self):
        before = self.read('.agent/HANDOFF.md')
        other = before + b'nguoi khac vua ghi\n'
        self.write('.agent/HANDOFF.md', other)  # controlled interleaving: another writer changed the note after Main read it
        ok, res = self.r.record_before_effect(self.project, 'handoff', before + b'ke hoach\n', sha(before))
        self.assertEqual((ok, res['result'], res['reason']), (False, 'CONFLICT', 'EXPECTED_BEFORE_MISMATCH'))
        self.assertEqual(self.read('.agent/HANDOFF.md'), other)
        with mock.patch.dict(os.environ, {tx.FAULT_ENV: 'commit:1'}):  # the replace lands, the transaction does not finish
            ok, res = self.r.record_before_effect(self.project, 'handoff', other + b'a\n', sha(other))
        self.assertEqual((ok, res['result'], res['reason']), (False, 'WRITTEN', 'COMPLETION_ERROR'), 'bytes written, yet not a clean note')
        current = self.read('.agent/HANDOFF.md')
        ok, res = self.r.record_before_effect(self.project, 'handoff', current + b'b\n', sha(current))
        self.assertEqual((ok, res['result'], res['reason']), (False, 'ERROR', 'TRANSACTION_UNFINISHED'))
        self.assertEqual(self.read('.agent/HANDOFF.md'), current)
        self.assertFalse(os.path.exists(self.abs('output')), 'the caller never reached the writer')

    # ------------------------------------------------------------ timeout, then a late RESULT
    def test_timeout_then_a_late_result_is_not_integrated_though_request_and_sources_did_not_change(self):
        task1 = self.task_record('t-a', 1, 'a1')
        task_rel = self.put_task(task1)  # dispatched: no RESULT yet
        ok, res = self.handoff_add('Lần giao a1: t-a rev 1, slot a1; invocation hết thời gian, chưa rõ đã kết thúc; chưa có RESULT.\n')
        self.assertTrue(ok, res)
        result_rel = self.slot('t-a', 'a1') + '/RESULT.json'
        voided = self.void(task_rel, result_rel)  # there is nothing to checkpoint yet: voided rests on the read-back and the changed binding
        self.assertEqual((voided['voided'], voided['problems']), (True, []))
        expected = json.loads(json.dumps(task1))
        expected['task_revision'] = 2
        self.assertEqual(json.loads(self.read(task_rel)), expected, 'only the revision moved: request and sources are unchanged')
        late = self.result_record(task1, 'a1', b'ban nhap den muon\n')  # carries the binding handed out at dispatch
        self.assertEqual(late['binding']['task_revision'], 1)
        self.write(result_rel, json.dumps(late, ensure_ascii=False, indent=2).encode('utf-8'))
        cp = self.checkpoint(task_rel, result_rel)
        self.assertTrue(any(p.startswith('BINDING B-REF') for p in cp['problems']), cp['problems'])
        self.assertFalse([p for p in cp['problems'] if p.startswith(('INPUT', 'LEDGER'))], 'the sources did not change')
        self.assertEqual(self.r.integration_gate(cp, {'REQ-1': 'MET'})['state'], 'HOLD')
        step = self.r.next_step(self.facts(invocation='UNKNOWN', retry_left=2))
        self.assertEqual(step, {'step': 'WAIT', 'reasons': ['INVOCATION_UNKNOWN']}, 'no retry while the invocation is UNKNOWN')

    # ------------------------------------------------------------ timeout after the effect
    def test_timeout_after_link_is_reconciled_by_request_id_and_is_never_published_twice(self):
        task_rel, result_rel, cp, request, rid = self.killed_publish('after_link')
        target = request['target']
        self.assertTrue(self.void(task_rel, result_rel)['voided'])
        out_dir = 'output/%s' % WORK
        listing, before = sorted(os.listdir(self.abs(out_dir))), self.tree()
        with self.no_writer():
            rec = self.r.reconcile_publish(self.project, rid, target)
            self.assertEqual((rec['decision'], rec['terminal']), ('NEEDS_RECOVER', None))
            self.assertIn('TARGET_EXISTS', rec['reasons'])
            self.assertEqual(self.r.next_step(self.facts(effect=rec['decision'])), {'step': 'WAIT', 'reasons': ['EFFECT_NEEDS_RECOVER']})
        self.assertEqual(self.tree(), before, 'reading the evidence changed nothing')
        self.assertEqual(self.read(target), cp['candidate'])
        recovered = pn.recover(self.project, rid)  # a WRITE (terminal evidence, staged copy removed): called on purpose, inside the simulated grant
        self.assertEqual((recovered['outcome'], recovered['recovered']), ('PUBLISHED', True))
        self.assertTrue(recovered['wrote'], 'recovery is declared as a write')
        with self.no_writer():
            rec = self.r.reconcile_publish(self.project, rid, target)
            self.assertEqual((rec['decision'], rec['terminal'], rec['reasons']), ('WRITTEN', 'published', []))
            self.assertEqual(self.r.next_step(self.facts(effect=rec['decision'])), {'step': 'NO_SECOND_WRITE', 'reasons': ['EFFECT_ALREADY_WRITTEN']})
        self.assertEqual(sorted(os.listdir(self.abs(out_dir))), listing, 'one output, no second file')
        self.assertEqual(self.read(target), cp['candidate'])
        self.assertTrue(self.checkpoint(task_rel, result_rel)['problems'], 'the written output does not make the old RESULT integrable under the new revision')
        retry = pn.publish(self.project, request, cp['candidate'], request_id=self.new_rid())  # a careless caller with a fresh ID: the old request is bound to the old TASK
        self.assertEqual((retry['outcome'], retry['reason']), ('REFUSED', 'INPUT_CHANGED_SINCE_READ'))
        self.assertEqual(sorted(os.listdir(self.abs(out_dir))), listing)

    def test_timeout_after_the_terminal_evidence_is_already_written_needs_no_recovery(self):
        task_rel, result_rel, cp, request, rid = self.killed_publish('after_published')
        target = request['target']
        with self.no_writer():
            rec = self.r.reconcile_publish(self.project, rid, target)
        self.assertEqual((rec['decision'], rec['terminal'], rec['reasons']), ('WRITTEN', 'published', ['CLEANUP_PENDING']))
        self.assertEqual(self.read(target), cp['candidate'])
        self.assertEqual(pn.recover(self.project, rid)['outcome'], 'NO_ACTION')  # only the leftover staged copy is cleaned: also a write
        with self.no_writer():
            self.assertEqual(self.r.reconcile_publish(self.project, rid, target)['reasons'], [])

    # ------------------------------------------------------------ evidence missing or contradictory
    def test_missing_or_contradictory_evidence_holds_and_writes_nothing(self):
        task_rel, result_rel, cp, request, rid = self.killed_publish('after_published')
        target, folder = request['target'], self.abs('.agent/runtime/publish')
        snapshot = {n: self.read('.agent/runtime/publish/' + n) for n in os.listdir(folder)}
        self.assertTrue(all(n.startswith(rid) for n in snapshot))

        def restore():
            for n in os.listdir(folder):
                os.remove(os.path.join(folder, n))
            for n, data in snapshot.items():
                self.write('.agent/runtime/publish/' + n, data)

        def drop_all():
            for n in snapshot:
                os.remove(os.path.join(folder, n))

        def drop_prepared():
            os.remove(os.path.join(folder, rid + '.prepared.json'))

        def second_terminal():
            self.write('.agent/runtime/publish/%s.conflict.json' % rid, snapshot[rid + '.published.json'])
        cases = [('all evidence removed while the output exists', drop_all, ['EVIDENCE_MISSING', 'TARGET_EXISTS_WITHOUT_EVIDENCE']),
                 ('prepared record missing', drop_prepared, ['PREPARED_MISSING']),
                 ('two terminal records', second_terminal, ['EVIDENCE_INCONSISTENT'])]
        for label, tamper, reasons in cases:
            with self.subTest(label):
                restore()
                tamper()
                before = self.tree()
                with self.no_writer():
                    rec = self.r.reconcile_publish(self.project, rid, target)
                    self.assertEqual(rec['decision'], 'HOLD')
                    self.assertLessEqual(set(reasons), set(rec['reasons']), rec)
                    self.assertEqual(self.r.next_step(self.facts(effect=rec['decision'])), {'step': 'WAIT', 'reasons': ['EFFECT_HOLD']})
                self.assertEqual(self.tree(), before, 'reading again lost nothing')
                self.assertEqual(self.read(target), cp['candidate'])
                self.assertEqual(self.read(cp['result']['artifacts'][0]['path']), cp['candidate'])
        with self.subTest('request ID never seen'):
            rec = self.r.reconcile_publish(self.project, self.new_rid(), 'output/%s/khac.txt' % WORK)
            self.assertEqual((rec['decision'], rec['reasons']), ('HOLD', ['EVIDENCE_MISSING']), 'no evidence is not "not written"')

    # ------------------------------------------------------------ reopen in a fresh process, reassign under the current revision
    def test_a_reopened_session_rereads_the_disk_and_only_the_new_slot_under_the_current_revision_passes(self):
        task_rel, result_rel, _ = self.assignment('t-a', 1, 'a1')
        cp = self.checkpoint(task_rel, result_rel)
        request, rid = self.publish_request(cp, task_rel, result_rel), self.new_rid()
        ok, res = self.handoff_add('Lần giao a1: t-a rev 1, slot a1; writer P-CN1, request ID %s, đích %s; writer có thể đã bắt đầu, chưa rõ kết quả. Đã yêu cầu dừng, host chưa xác nhận.\n'
                                   % (rid, request['target']))
        self.assertTrue(ok, res)
        self.write(request['target'], b'pre-existing output, not owned by this request')
        collision = pn.publish(self.project, request, cp['candidate'], request_id=rid)
        self.assertEqual(collision['outcome'], 'CONFLICT')  # real terminal proof before reopening, not an inference from an old note
        self.assertTrue(self.void(task_rel, result_rel)['voided'])
        ok, res = self.handoff_add('TASK t-a rev 2 đang giữ để reconcile, chưa dispatch; đối soát conflict evidence cho request ID %s.\n' % rid)
        self.assertTrue(ok, res)
        with open(os.path.join(self.scratch, 'recipe.py'), 'w', encoding='utf-8') as fh:
            fh.write(recipe_source())  # the text of the docs, not an object from this process

        def reopen(**results):
            payload = json.dumps({'task_rel': task_rel, 'results': results})
            proc = self.child('child_reopen.py', CHILD_REOPEN, os.path.join(self.scratch, 'recipe.py'), str(ROOT), self.project, payload)
            self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
            return json.loads(proc.stdout)
        old_bytes = self.read(result_rel)
        first = reopen(a1=result_rel)
        self.assertNotEqual(first['pid'], os.getpid())
        self.assertEqual(first['task_revision'], 2)
        self.assertTrue(any(p.startswith('BINDING B-REF') for p in first['results']['a1']), first['results'])
        self.assertEqual(first['reconcile']['decision'], 'NOT_PUBLISHED')
        self.assertEqual(first['reconcile']['reasons'], ['CONFLICT'])
        self.assertFalse(first['note_says_writer_not_called'])
        step = self.r.next_step(self.facts(effect=first['reconcile']['decision'], voided=True, no_call_observed=False))  # terminal proof from disk; host-ended remains SIMULATED

        self.assertEqual(step, {'step': 'REASSIGN', 'reasons': []})
        task3 = self.task_record('t-a', 3, 'a2')  # a new slot under a raised revision
        self.assertTrue(self.r.void_assignment(self.v3, self.project, task_rel, task3, result_rel)['voided'])
        result2 = self.result_record(task3, 'a2', b'ban nhap lan hai\n')
        result2_rel = self.slot('t-a', 'a2') + '/RESULT.json'
        self.write(result2_rel, json.dumps(result2, ensure_ascii=False, indent=2).encode('utf-8'))
        second = reopen(a1=result_rel, a2=result2_rel)
        self.assertEqual((second['task_revision'], second['results']['a2']), (3, []))
        self.assertTrue(second['results']['a1'], 'the old RESULT stays blocked')
        self.assertEqual(self.read(result_rel), old_bytes)
        self.write(SOURCE_REL, b'quy che gia: muc gia v3 (doi khi dong phien)\n')
        third = reopen(a2=result2_rel)
        self.assertIn('INPUT_CHANGED ' + SOURCE_REL, third['results']['a2'], 'a source changed while the session was closed is detected')

    # ------------------------------------------------------------ missing tool, quota or provider; one agent
    def test_missing_tool_quota_or_provider_waits_or_drafts_inside_the_grant_and_changes_nothing_else(self):
        task_rel, result_rel, _ = self.assignment('t-a', 1, 'a1')
        voided = self.void(task_rel, result_rel)
        self.assertTrue(voided['voided'])
        before = self.tree()
        reviewer_absent = 'MISSING'  # SIMULATED availability: nothing below asks a host, a tool, a quota, a provider or a reviewer
        cases = [({}, 'REASSIGN', []),
                 ({'capability': 'TOOL_MISSING'}, 'WAIT', ['CAPABILITY_TOOL_MISSING']),
                 ({'capability': 'QUOTA_EXHAUSTED', 'main_may_draft': True}, 'MAIN_DRAFT', ['CAPABILITY_QUOTA_EXHAUSTED']),
                 ({'capability': 'PROVIDER_UNAVAILABLE', 'main_may_draft': True}, 'MAIN_DRAFT', ['CAPABILITY_PROVIDER_UNAVAILABLE']),
                 ({'retry_left': None}, 'WAIT', ['NO_RETRY_OR_COST_BUDGET']),
                 ({'retry_left': 0, 'main_may_draft': True}, 'MAIN_DRAFT', ['NO_RETRY_OR_COST_BUDGET']),
                 ({'cost_ok': False}, 'WAIT', ['NO_RETRY_OR_COST_BUDGET']),
                 ({'invocation': 'RUNNING'}, 'WAIT', ['INVOCATION_RUNNING']),
                 ({'invocation': 'RUNNING', 'parallel_safe': True}, 'REASSIGN', []),
                 ({'grant_ok': False, 'main_may_draft': True}, 'WAIT', ['GRANT_NOT_COVERING'])]
        with mock.patch('socket.socket.connect', side_effect=AssertionError('no network in this probe')), self.no_writer():
            for over, step, reasons in cases:
                with self.subTest(over):
                    self.assertEqual(self.r.next_step(self.facts(**over)), {'step': step, 'reasons': reasons})
        self.assertEqual(self.tree(), before, 'no profile, native file, retry counter or other state was created')
        self.assertFalse(os.path.exists(self.abs('.agent/EXECUTION_PROFILE.json')) or os.path.exists(self.abs('.claude')))
        draft = self.assignment('t-d', 1, 'a1')  # a draft that passed the checkpoint, standing in for Main's own draft inside the grant
        cp = self.checkpoint(draft[0], draft[1])
        self.assertEqual(cp['problems'], [])
        held = self.r.integration_gate(cp, {'REQ-1': 'MET'}, review=reviewer_absent)
        self.assertEqual(held, {'state': 'HOLD', 'reasons': ['REVIEW_REQUIRED_BUT_MISSING']}, 'every requirement MET still waits for the missing Reviewer')
        self.assertEqual(self.r.integration_gate(cp, {'REQ-1': 'MET'}, review='PASSED')['state'], 'ACCEPTED_FOR_INTEGRATION')  # bound, veto-free review is a declared fact here

    # ------------------------------------------------------------ binding matches, content does not
    def test_a_matching_binding_does_not_mean_the_required_content_is_there(self):
        task = self.task_record('t-c', 1, 'a1')
        task['acceptance']['requirements'].append({'req_id': 'REQ-2', 'statement': 'Nêu điểm chưa chắc', 'mandatory': True})
        task_rel = self.put_task(task)
        result = self.result_record(task, 'a1', b'ban nhap muc gia, dua tren quy che v1\n')  # says nothing about what is uncertain
        result['requirements'].append({'req_id': 'REQ-2', 'mandatory': True, 'status': 'UNKNOWN'})
        result_rel = self.slot('t-c', 'a1') + '/RESULT.json'
        self.write(result_rel, json.dumps(result, ensure_ascii=False, indent=2).encode('utf-8'))
        raw = self.read(result_rel)
        cp = self.checkpoint(task_rel, result_rel)
        self.assertEqual(cp['problems'], [], 'bytes, binding and ledger all match')
        # Main's own reading of the content is an explicit INPUT here (SIMULATED judgement): no tool in this probe detects what is missing
        for review, reasons in (({'REQ-1': 'MET', 'REQ-2': 'UNKNOWN'}, ['REQUIREMENT REQ-2 UNKNOWN']), ({'REQ-1': 'MET'}, ['REQUIREMENT REQ-2 UNKNOWN']),
                                ({'REQ-1': 'MET', 'REQ-2': 'FAILED'}, ['REQUIREMENT REQ-2 FAILED']), ({'REQ-1': 'MET', 'REQ-2': 'BLOCKED'}, ['REQUIREMENT REQ-2 BLOCKED'])):
            with self.subTest(review):
                self.assertEqual(self.r.integration_gate(cp, review), {'state': 'HOLD', 'reasons': reasons})
        self.assertEqual(self.r.integration_gate(cp, {'REQ-1': 'MET', 'REQ-2': 'MET'}), {'state': 'ACCEPTED_FOR_INTEGRATION', 'reasons': []})
        self.assertEqual(self.read(result_rel), raw)
        self.assertEqual(json.loads(raw)['acceptance_states'], ['CANDIDATE'], 'the Worker RESULT stays a CANDIDATE whatever Main decides')
        self.assertEqual(self.r.integration_gate(self.checkpoint(task_rel, result_rel + '.missing'), {'REQ-1': 'MET', 'REQ-2': 'MET'})['state'], 'HOLD')


# A fresh interpreter that only reads the workbook through the B-02 owner; the opener is the owner's own TEST-ONLY log hook (nothing is launched).
CHILD_ARTIFACT_VIEW = '''import json, os, sys
tools, project, aid, rev = sys.argv[1], sys.argv[2], sys.argv[3], int(sys.argv[4])
sys.path.insert(0, tools)
sys.dont_write_bytecode = True
import artifacts
out = {'pid': os.getpid()}
found = artifacts.find(project, '', aid, revision=rev)
item = found['items'][0]
out.update(state=item['dependency_state'], limits=item['limits'], inference=item['inference'], unknown=item['unknown_basis'], superseded_by=item['superseded_by'],
           lineage=[[x['key'], x['limits']] for x in item['lineage']], bytes_observation=found['bytes_observation'], path=item['path'], sha256=item['sha256'])
try:
    prep = artifacts.prepare(project, {'inputs': [{'kind': 'ARTIFACT', 'id': aid, 'revision': rev}]})
    out['prepare'] = {'closure': prep['closure'], 'inputs': prep['inputs'], 'verifies_content': prep['verifies_content'], 'bytes_observation': prep['bytes_observation']}
except Exception as e:
    out['prepare_error'] = getattr(e, 'code', type(e).__name__)
out['open'] = artifacts.open_artifact(project, item['pick'])
print(json.dumps(out))
'''


@unittest.skipUnless(CAN_RUN, SKIP_RUN)
class IntegratedTrace(FlowBase):
    """M3/E04: ONE trace over owners that were accepted one by one: source (B-01 ledger + content) -> candidate with a declared basis (A-11 TASK/RESULT)
    -> checkpoint and a declared review -> P-CN1 publish -> B-02 registration -> a fresh interpreter reopens and reuses it -> the source changes
    (bytes, then a semantic event on the same bytes) -> the stale result is refused. Fake data, a temporary project, a .txt result.
    SIMULATED: the Reviewer verdict and Main's content judgement are declared inputs (no live agent). MOCK: the opener only logs the path through the owner's
    test hook; it proves nothing about an Office application. Nothing here is a PASS, a grant or user acceptance."""

    def setUp(self):
        super().setUp()
        self.scratch = tempfile.mkdtemp(prefix='law-e04-')
        self.addCleanup(shutil.rmtree, self.scratch, True)
        self.sid = 'SRC-000001'

    # ------------------------------------------------------------ helpers
    def src(self):
        return art.load_book(self.project)['catalog'][self.sid][0]

    def dep_source(self):
        rec = self.src()
        return {'kind': 'SOURCE', 'id': self.sid, 'revision': art.library.revision_of(rec['law_revision']), 'sha256': rec['law_sha']}

    def register(self, aid, rev, path, kind, deps, limits, origin='USER_FILE', rid=None, predecessor=None):
        req = {'artifact_id': aid, 'revision': rev, 'path': path, 'title': 'Kết quả ' + aid, 'kind': kind, 'limits': limits, 'origin': origin,
               'recorded_by': 'tester', 'dependencies': deps}
        if rid:
            req['rid'] = rid
        if predecessor:
            req['predecessor'] = predecessor
        res = art.register_apply(self.project, req, art.register_plan(self.project, req)['plan_token'])
        self.assertEqual(res['action'], 'record', res)
        return art.load_book(self.project)['registry']['by_key'][(aid, rev)]['sha256']

    def refused(self, code, fn, *args):
        with self.assertRaises(Exception) as caught:
            fn(*args)
        self.assertEqual(getattr(caught.exception, 'code', None), code, caught.exception)

    def prepare_inputs(self, *picks):
        prep = art.prepare(self.project, {'inputs': list(picks)})
        return prep, [i for i in prep['inputs'] if i['path'] != art.WORKBOOK_PATH]  # the ledger snapshot is bound by the checkpoint itself

    def publish(self, cp, task_rel, result_rel, target):
        request, rid = self.publish_request(cp, task_rel, result_rel, target=target), self.new_rid()
        res = pn.publish(self.project, request, cp['candidate'], request_id=rid)
        self.assertEqual(res['outcome'], 'PUBLISHED', res)
        return request, rid

    def view(self, aid, rev):
        """What a fresh interpreter sees through the B-02 owner; the project tree must not move."""
        log = os.path.join(self.scratch, 'open.log')
        path = os.path.join(self.scratch, 'view.py')
        with open(path, 'w', encoding='utf-8') as fh:
            fh.write(CHILD_ARTIFACT_VIEW)
        env = dict(os.environ, PYTHONDONTWRITEBYTECODE='1', PYTHONUTF8='1', LOCAL_AGENT_WORKSPACE_TEST_OPEN_LOG=log)
        before = self.tree()
        proc = subprocess.run([sys.executable, '-I', '-B', path, str(ROOT / 'tools'), self.project, aid, str(rev)], capture_output=True, text=True,
                              encoding='utf-8', env=env, timeout=120, cwd=self.scratch)
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        self.assertEqual(self.tree(), before, 'viewing, preparing and opening must not write, scan or regenerate anything')
        out = json.loads(proc.stdout)
        self.assertNotEqual(out['pid'], os.getpid())
        out['opened_paths'] = []
        if os.path.exists(log):
            with open(log, encoding='utf-8') as fh:
                out['opened_paths'] = fh.read().splitlines()
            os.remove(log)
        return out

    def state_with_decision(self, text):
        before = self.read('.agent/STATE.md')
        new = before.replace(b'## Decisions\n\n- (fixture)\n', ('## Decisions\n\n- %s\n' % text).encode('utf-8'))
        self.assertNotEqual(new, before)
        return self.r.record_before_effect(self.project, 'state', new, sha(before))

    # ------------------------------------------------------------ the trace
    def test_source_to_published_result_reopened_reused_and_refused_when_the_source_changes(self):
        # 1. B-01: the chosen source and what was read from it, with its limits
        rec = self.src()
        entry = {'entry_id': 'E-001', 'source_id': self.sid, 'source_revision': rec['law_revision'], 'source_sha256': rec['law_sha'], 'locator': 'dòng 1',
                 'read_scope': 'dòng 1 đã đọc', 'kind': 'EXTRACT', 'text': 'muc gia v1', 'limits': 'chỉ đọc dòng 1; dữ liệu giả', 'recorded_by': 'tester'}
        ls.content_apply(self.project, entry, ls.content_plan(self.project, entry)['plan_token'])
        # 2. an inference stays an INFERENCE with its own limits; a decision is written down BEFORE any effect (P-CTX1)
        self.write('working/e04/suy-luan.txt', '[INFERENCE] mức giá có thể tăng; suy luận của Main, chưa có nguồn xác nhận\n'.encode('utf-8'))
        inf_sha = self.register('ART-INF', 1, 'working/e04/suy-luan.txt', 'INFERENCE', [self.dep_source()], 'suy luận, không phải dữ kiện của nguồn')
        decision = 'QĐ-E04-1: dùng %s revision 1; suy luận chỉ mang nhãn INFERENCE, không thành nguồn.' % self.sid
        ok, res = self.state_with_decision(decision)
        self.assertTrue(ok, res)
        state_after_decision = self.read('.agent/STATE.md')
        # 3. A-11: a candidate whose TASK input set IS the B-02 closure (source + inference), then checkpoint and a declared review
        prep, inputs = self.prepare_inputs({'kind': 'ARTIFACT', 'id': 'ART-INF', 'revision': 1}, {'kind': 'SOURCE', 'id': self.sid})
        self.assertEqual(prep['inference'], ['ART-INF@1'])
        self.assertFalse(prep['verifies_content'])
        body_a = ('báo cáo mức giá; căn cứ %s revision 1 sha %s\n[INFERENCE] mức giá có thể tăng (ART-INF@1)\n' % (self.sid, rec['law_sha'][:12])).encode('utf-8')
        task_rel, result_rel, task_a = self.assignment('t-a', 1, 'a1', body=body_a, inputs=inputs)
        cp_a = self.checkpoint(task_rel, result_rel)
        self.assertEqual(cp_a['problems'], [])
        self.assertEqual(cp_a['ledger_snapshot'], prep['snapshot'])  # one workbook snapshot for both owners
        self.assertEqual(self.r.integration_gate(cp_a, {'REQ-1': 'MET'}, review='PASSED')['state'], 'ACCEPTED_FOR_INTEGRATION')  # SIMULATED verdict
        # 4. P-CN1 publishes the checked bytes; B-02 registers the PUBLISHED file with the same basis (check_prepared needs every closure input in the publish request)
        request_a, rid_a = self.publish(cp_a, task_rel, result_rel, 'output/%s/bao-cao-a.txt' % WORK)
        sha_a = self.register('ART-A', 1, request_a['target'], 'REPORT', [self.dep_source(), {'kind': 'ARTIFACT', 'id': 'ART-INF', 'revision': 1, 'sha256': inf_sha}],
                              'chỉ quy chế v1; suy luận chưa kiểm', origin='P_CN1', rid=rid_a)
        self.assertEqual(sha_a, hashlib.sha256(body_a).hexdigest())
        # 5. a fresh interpreter: same version, same basis, limits and inference kept, nothing scanned or rewritten, the opener only logged
        seen = self.view('ART-A', 1)
        self.assertEqual((seen['state'], seen['bytes_observation'], seen['superseded_by']), ('CURRENT_CATALOG_BASIS', 'UNOBSERVED', ''))
        self.assertEqual((seen['limits'], seen['inference'], seen['unknown']), ('chỉ quy chế v1; suy luận chưa kiểm', ['ART-INF@1'], []))
        self.assertEqual(sorted(seen['prepare']['closure']), ['ART-A@1', 'ART-INF@1'])
        self.assertEqual(seen['prepare']['inputs'], art.prepare(self.project, {'inputs': [{'kind': 'ARTIFACT', 'id': 'ART-A', 'revision': 1}]})['inputs'])
        self.assertEqual(seen['open']['state'], 'OPENED')
        self.assertEqual([os.path.normcase(p) for p in seen['opened_paths']], [os.path.normcase(self.abs(request_a['target']))])  # MOCK: logged, not launched
        # the result is a plain file readable without LAW, and so are the source and the basis it names
        with open(self.abs(request_a['target']), 'rb') as fh:
            outside = fh.read().decode('utf-8')
        with open(self.abs(SOURCE_REL), 'rb') as fh:
            self.assertIn('muc gia v1', fh.read().decode('utf-8'))
        self.assertIn('%s revision 1 sha %s' % (self.sid, rec['law_sha'][:12]), outside)
        self.assertEqual(self.read('.agent/STATE.md'), state_after_decision)  # the written decision survived the reopening
        # 6. reuse: ART-A is the INPUT of the next TASK; B-02 hands its whole closure to the checkpoint and to the publish request
        prep_b, inputs_b = self.prepare_inputs({'kind': 'ARTIFACT', 'id': 'ART-A', 'revision': 1})
        self.assertEqual(sorted(prep_b['closure']), ['ART-A@1', 'ART-INF@1'])
        self.assertEqual({i['path'] for i in inputs_b}, {request_a['target'], 'working/e04/suy-luan.txt', SOURCE_REL})
        body_b = b'resumen de ART-A rev 1 (dung lai, khong doc lai nguon)\n'
        task_b_rel, result_b_rel, _ = self.assignment('t-b', 1, 'b1', body=body_b, inputs=inputs_b)
        cp_b = self.checkpoint(task_b_rel, result_b_rel)
        self.assertEqual(cp_b['problems'], [])
        self.assertEqual(self.r.integration_gate(cp_b, {'REQ-1': 'MET'}, review='PASSED')['state'], 'ACCEPTED_FOR_INTEGRATION')  # its own gate: nothing is carried over from A
        request_b, rid_b = self.publish(cp_b, task_b_rel, result_b_rel, 'output/%s/tong-hop-b.txt' % WORK)
        self.register('ART-B', 1, request_b['target'], 'REPORT', [{'kind': 'ARTIFACT', 'id': 'ART-A', 'revision': 1, 'sha256': sha_a}], 'chỉ dùng ART-A rev 1', origin='P_CN1', rid=rid_b)
        reused = self.view('ART-B', 1)
        self.assertEqual((reused['state'], sorted(reused['inference'])), ('CURRENT_CATALOG_BASIS', ['ART-INF@1']))
        self.assertEqual({k: v for k, v in reused['lineage']}.get('ART-A@1'), 'chỉ quy chế v1; suy luận chưa kiểm')  # the limits of the ancestor travel with the result
        for rel in (result_rel, result_b_rel):
            self.assertEqual(json.loads(self.read(rel))['acceptance_states'], ['CANDIDATE'])  # neither the publish nor the open is user acceptance
        # 7. the source file changes, the ledger is not synced yet: the metadata view still says CURRENT (it never reads bytes), the bytes checks refuse
        self.write(SOURCE_REL, 'quy che gia: muc gia v2 (da doi)\n'.encode('utf-8'))
        self.assertEqual(self.view('ART-B', 1)['state'], 'CURRENT_CATALOG_BASIS')  # the limit of a metadata-only view, stated here on purpose
        self.refused('STALE_SOURCE_BASIS', art.prepare, self.project, {'inputs': [{'kind': 'ARTIFACT', 'id': 'ART-B', 'revision': 1}]})
        self.assertIn('INPUT_CHANGED ' + SOURCE_REL, self.checkpoint(task_b_rel, result_b_rel)['problems'])
        # 8. synced to revision 2: the ledger itself now marks every result built on revision 1; a request bound before the change is refused before any write
        ls.apply(self.project)
        rec2 = self.src()
        self.assertEqual(rec2['law_revision'], 2)
        for aid in ('ART-A', 'ART-B'):
            stale = self.view(aid, 1)
            self.assertEqual(stale['state'], 'NEEDS_REVIEW', aid)
            self.assertEqual(stale['prepare_error'], 'STALE_SOURCE_BASIS', aid)
            self.assertEqual(stale['open']['dependency_state'], 'NEEDS_REVIEW')  # the old file still opens, labelled; opening is not integrating
        before = self.tree()
        late = pn.publish(self.project, self.publish_request(cp_b, task_b_rel, result_b_rel, 'output/%s/muon.txt' % WORK), cp_b['candidate'], request_id=self.new_rid())
        self.assertEqual((late['outcome'], late['reason']), ('REFUSED', 'INPUT_CHANGED_SINCE_READ'))
        self.assertEqual(self.tree(), before)
        # 9. reconcile instead of retry: the TASK is voided with a raised revision, the inference is re-derived, a NEW version supersedes ART-A rev 1
        self.write('working/e04/suy-luan-2.txt', '[INFERENCE] mức giá v2; suy luận lại từ revision 2\n'.encode('utf-8'))
        inf2_sha = self.register('ART-INF', 2, 'working/e04/suy-luan-2.txt', 'INFERENCE', [self.dep_source()], 'suy luận lại, chưa kiểm', predecessor=1)
        prep2, inputs2 = self.prepare_inputs({'kind': 'ARTIFACT', 'id': 'ART-INF', 'revision': 2}, {'kind': 'SOURCE', 'id': self.sid})
        task2 = self.task_record('t-a', 2, 'a2', inputs=inputs2)
        voided = self.r.void_assignment(self.v3, self.project, task_rel, task2, result_rel)
        self.assertTrue(voided['voided'], voided)
        self.assertEqual(voided['host_stop'], 'UNCONFIRMED')
        body2 = ('báo cáo mức giá; căn cứ %s revision 2 sha %s\n[INFERENCE] mức giá v2 (ART-INF@2)\n' % (self.sid, rec2['law_sha'][:12])).encode('utf-8')
        result2 = self.result_record(task2, 'a2', body2)
        result2_rel = self.slot('t-a', 'a2') + '/RESULT.json'
        self.write(result2_rel, json.dumps(result2, ensure_ascii=False, indent=2).encode('utf-8'))
        self.assertTrue(self.checkpoint(task_rel, result_rel)['problems'], 'the old RESULT never matches the current TASK again')
        cp2 = self.checkpoint(task_rel, result2_rel)
        self.assertEqual(cp2['problems'], [])
        request2, rid2 = self.publish(cp2, task_rel, result2_rel, 'output/%s/bao-cao-a2.txt' % WORK)
        self.register('ART-A', 2, request2['target'], 'REPORT', [self.dep_source(), {'kind': 'ARTIFACT', 'id': 'ART-INF', 'revision': 2, 'sha256': inf2_sha}],
                      'chỉ quy chế v2', origin='P_CN1', rid=rid2, predecessor=1)
        self.assertEqual(self.view('ART-A', 2)['state'], 'CURRENT_CATALOG_BASIS')
        old = self.view('ART-A', 1)
        self.assertEqual(old['superseded_by'], 2)
        self.refused('UPSTREAM_SUPERSEDED', art.prepare, self.project, {'inputs': [{'kind': 'ARTIFACT', 'id': 'ART-B', 'revision': 1}]})  # B was built on A rev 1
        # 10. a semantic change with the SAME bytes (B-01 source-event) stales the new result as well
        ev = {'event_id': 'V-001', 'source_id': self.sid, 'expected_revision': 2, 'expected_sha256': rec2['law_sha'], 'mode': 'semantic_update',
              'changes': {'status': 'Hết hiệu lực'}, 'reason': 'thay thế bởi văn bản mới', 'evidence_ref': 'công văn giả 01', 'recorded_by': 'tester'}
        source_bytes = self.read(SOURCE_REL)
        done = ls.event_apply(self.project, ev, ls.event_plan(self.project, ev)['plan_token'])
        self.assertEqual((done['action'], done['revision_after']), ('event', 3))
        self.assertEqual(self.read(SOURCE_REL), source_bytes)  # the bytes did not move
        again = self.view('ART-A', 2)
        self.assertEqual((again['state'], again['prepare_error']), ('NEEDS_REVIEW', 'STALE_SOURCE_BASIS'))
        self.assertTrue(any(p.startswith('LEDGER_METADATA_CHANGED') for p in self.checkpoint(task_rel, result2_rel)['problems']))
        self.assertEqual(prep2['bytes_observation'], 'OBSERVED_AT_PREPARE')  # what prepare saw is a moment, not a standing fact


@unittest.skipUnless(CAN_RUN, SKIP_RUN)
class ReviewRegressions(FlowBase):
    """Reviewer: reject malformed replacements, review vetoes and mismatched publish evidence."""

    def test_invalid_new_task_is_refused_before_replacing_the_current_record(self):
        task_rel, result_rel, task = self.assignment()
        invalid = copy.deepcopy(task)
        invalid['task_revision'] += 1
        del invalid['acceptance']
        before = self.tree()
        out = self.r.void_assignment(self.v3, self.project, task_rel, invalid, result_rel)
        self.assertFalse(out['voided'])
        self.assertIn('NEW_TASK_REJECTED', out['problems'])
        self.assertEqual(self.tree(), before, 'schema validation must happen before replacement')

    def test_review_veto_or_unknown_never_allows_integration(self):
        task_rel, result_rel, _ = self.assignment()
        cp = self.checkpoint(task_rel, result_rel)
        self.assertEqual(cp['problems'], [])
        for verdict in ('DONE', 'FAILED', 'BLOCKED', 'UNKNOWN', 'ERROR', None):
            with self.subTest(verdict=verdict):
                self.assertEqual(self.r.integration_gate(cp, {'REQ-1': 'MET'}, review=verdict)['state'], 'HOLD')
        self.assertEqual(self.r.integration_gate(cp, {'REQ-1': 'MET'}, review='PASSED')['state'], 'ACCEPTED_FOR_INTEGRATION')

    def test_reconcile_requires_the_target_bound_to_the_request(self):
        task_rel, result_rel, _ = self.assignment()
        cp = self.checkpoint(task_rel, result_rel)
        request, rid = self.publish_request(cp, task_rel, result_rel), self.new_rid()
        self.assertEqual(pn.publish(self.project, request, cp['candidate'], request_id=rid)['outcome'], 'PUBLISHED')
        other = 'output/other.txt'
        self.write(other, cp['candidate'])
        before = self.tree()
        out = self.r.reconcile_publish(self.project, rid, other)
        self.assertEqual(out['decision'], 'HOLD')
        self.assertIn('REQUEST_TARGET_MISMATCH', out['reasons'])
        self.assertEqual(self.tree(), before)

    def test_reconcile_requires_published_file_identity_even_when_bytes_match(self):
        task_rel, result_rel, _ = self.assignment()
        cp = self.checkpoint(task_rel, result_rel)
        request, rid = self.publish_request(cp, task_rel, result_rel), self.new_rid()
        self.assertEqual(pn.publish(self.project, request, cp['candidate'], request_id=rid)['outcome'], 'PUBLISHED')
        path = request['target']
        before_identity = pn.identity_of(pn.file_info(self.abs(path)))
        self.write(path + '.replacement', cp['candidate'])
        os.replace(self.abs(path + '.replacement'), self.abs(path))
        self.assertNotEqual(pn.identity_of(pn.file_info(self.abs(path))), before_identity)
        out = self.r.reconcile_publish(self.project, rid, path)
        self.assertEqual(out['decision'], 'HOLD')
        self.assertIn('TARGET_IDENTITY_CHANGED', out['reasons'])

    def test_a_historical_no_call_note_does_not_authorize_retry(self):
        facts = {'voided': True, 'effect': 'NONE_ATTEMPTED', 'invocation': 'ENDED', 'grant_ok': True,
                 'capability': 'AVAILABLE', 'retry_left': 1, 'cost_ok': True}
        self.assertEqual(self.r.next_step(facts), {'step': 'WAIT', 'reasons': ['NO_CALL_NOT_OBSERVED']})
        facts['no_call_observed'] = True  # SIMULATED positive observation in the same live context, not a historical note
        self.assertEqual(self.r.next_step(facts), {'step': 'REASSIGN', 'reasons': []})

    def test_the_digest_observation_refuses_a_legacy_task_whose_layout_is_otherwise_valid(self):
        slot = self.slot('t-legacy', 'a1')
        cand, res = slot + '/candidate/draft.txt', slot + '/RESULT.json'
        task = self.task_record('t-legacy', 1, 'a1', targets=[cand, res])
        task['effects']['expected_outputs'] = [cand]
        task_rel = self.put_task(task)
        self.write(cand, b'ban nhap legacy')
        self.assertEqual(self.v3['office'].guard_task(task, lambda rel: self.r.read_file(self.project, rel)), [], 'the guard alone accepts a legacy task')
        before = self.tree()
        with mock.patch.object(art, 'observe_file') as spy:
            got = self.r.observe_candidate(self.v3, self.project, task_rel, cand, res, sha(self.read(task_rel)))
        spy.assert_not_called()
        self.assertEqual((got['observed'], got['observation'], got['problems']), (False, None, ['OFFICE_MARKER_AND_CONTEXT_REQUIRED']))
        self.assertEqual(self.tree(), before)


@unittest.skipUnless(wf.CAN_RUN, 'the P-CN1 owner supports Windows NTFS local only: skipped, not passed')
class AcceptanceAB(wf.Loop):
    """User acceptance of draft A is not acceptance of draft B. Real P-CN1 publishes, the existing v3 `accept` check, and current_draft for navigation."""

    def publish_draft(self, name, folder, staged):
        self.staged = staged
        rid = os.urandom(16).hex()
        return rid, self.publish(self.request(target='working/w-20261018-01/%s%s' % (folder, name)), rid)[1]

    def test_acceptance_of_draft_a_does_not_carry_to_draft_b_and_navigation_is_not_acceptance(self):
        try:
            import jsonschema  # noqa: F401  (the current_draft owner checks its pins with it)
        except ImportError:
            self.skipTest('jsonschema missing: the navigation part was not run')
        cd = importlib.import_module('current_draft')
        a = self.publish_draft('bao-cao-A.docx', '', b'PK ban A\n' * 30)
        b = self.publish_draft('bao-cao-B.docx', '', b'PK ban B\n' * 30)
        c = self.publish_draft('bao-cao-A.docx', 'b2/', b'PK ban A\n' * 30)  # C: the same name and the same bytes as A, in another folder
        self.assertEqual(len({a[1]['task_id'], b[1]['task_id'], c[1]['task_id']}), 3)
        self.assertEqual({a[1]['work_id'], b[1]['work_id'], c[1]['work_id']}, {'w-20261018-01'})
        bundle_a, _, _ = self.accept_bundle(*a)
        res_a = self.run_check(bundle_a, 'accept')
        self.assertEqual((res_a['disposition'], self.gate(res_a, 'REQUIREMENTS')['status']), ('PASS', 'USER_ACCEPTED_DECLARED'))
        accepted_a = copy.deepcopy(bundle_a['records'][3]['user_acceptance'])
        for label, draft in (('B', b), ('C: same name and bytes as A', c)):
            with self.subTest(label):
                unaccepted, _, _ = self.accept_bundle(*draft, acceptance=False)
                self.assert_not_pass(self.run_check(unaccepted, 'accept'))
                result = unaccepted['records'][3]
                result['acceptance_states'].append('USER_ACCEPTED')
                result['user_acceptance'] = copy.deepcopy(accepted_a)  # A's acceptance laid on this draft's own RESULT
                res = self.run_check(unaccepted, 'accept')
                self.assert_not_pass(res)
                self.assertEqual(self.gate(res, 'REQUIREMENTS')['status'], 'NOT_MET')
                self.assertIn('USER_ACCEPTANCE_STALE', res['reasons'])
        lines = ['# PROJECT STATE', '']
        for heading in cu.STATE_HEADINGS:
            lines += ['## ' + heading, '', '- Chưa chọn bản nào.' if heading == 'Current artifact' else 'Nội dung ' + heading, '']
        with open(os.path.join(self.project, '.agent', 'STATE.md'), 'wb') as fh:
            fh.write(('\n'.join(lines) + '\n').encode('utf-8'))
        opened = []
        with mock.patch.object(cd, '_startfile', lambda path: opened.append(path)):
            preview = cd.select_preview(self.project, b[0])
            selected = cd.select_apply(self.project, b[0], preview['state_sha256'], preview['selection_fingerprint'])
            shown = cd.open_draft(self.project, b[0])
        self.assertEqual((preview['result'], selected['result'], shown['result']), ('PREVIEW', 'WRITTEN', 'OPENED'))
        self.assertEqual(len(opened), 1)
        still, _, _ = self.accept_bundle(*b, acceptance=False)
        self.assert_not_pass(self.run_check(still, 'accept'))  # selected and opened, still not accepted
        self.assertNotIn('USER_ACCEPTED', still['records'][3]['acceptance_states'])
        again = self.run_check(self.accept_bundle(*a)[0], 'accept')
        self.assertEqual((again['disposition'], self.gate(again, 'REQUIREMENTS')['status']), ('PASS', 'USER_ACCEPTED_DECLARED'), 'A stays accepted, for A only')


class Docs(unittest.TestCase):
    def test_the_recipe_is_loaded_from_the_fixed_law_tree(self):
        text = (ROOT / 'roles' / 'README.md').read_text(encoding='utf-8')
        self.assertEqual(len(MARKER.findall(text)), 1)
        for needle in ('COOPERATIVE', 'Việc nhỏ không cần TASK/RESULT', 'task_revision', 'không phải PASS', 'Không ghi đè', 'B-03'):
            self.assertIn(needle, text)
        recipe = MARKER.findall(text)[0]
        self.assertFalse([q for q in ("'tests'", '"tests"', 'tests/') if q in recipe])  # nothing is imported from tests/ or a data root
        self.assertFalse([ln for ln in recipe.splitlines() if ln.lstrip().startswith(('import ', 'from ')) and 'tests' in ln])
        if openpyxl is None:
            self.skipTest('openpyxl missing: the package/library loading part was not run')
        r = recipe_from_docs()
        v3 = r.load_v3(ROOT)
        for name in ('pkg', 'loader', 'refs', 'binding', 'canon'):
            self.assertEqual(Path(v3[name].__file__).resolve().parent, (ROOT / '.cvf' / 'checks' / 'v3').resolve(), name)
        self.assertEqual(Path(v3['library'].__file__).resolve().parent, ROOT / 'tools')
        self.assertIsNotNone(importlib.import_module('law_cvf_check_v3'))

    def test_b03_section_and_functions_are_documented_and_the_recipe_controls_no_host(self):
        text = (ROOT / 'roles' / 'README.md').read_text(encoding='utf-8')
        recipe = MARKER.findall(text)[0]
        for name in ('record_before_effect', 'void_assignment', 'reconcile_publish', 'next_step', 'integration_gate'):
            self.assertRegex(recipe, r'(?m)^def %s\(' % name)
            self.assertIn('`%s`' % name, text.split('### Dừng, timeout, kết quả muộn và mở lại phiên (B-03)')[1].split('### Giới hạn cần nói rõ')[0])
        for needle in ('UNCONFIRMED', 'recover', 'Fallback một agent', 'user_acceptance', 'Stop, timeout, late results and reopening (B-03)'):
            self.assertIn(needle, text)
        # the recipe reads and writes project files through owners; it never stops, retires, exports or calls anything outside the project
        self.assertFalse([w for w in ('native_agents', 'subprocess', 'socket', 'urllib', 'os.kill', 'retire', 'taskkill') if w in recipe])
        self.assertNotIn('pn.publish(', recipe)
        self.assertNotRegex(recipe, r'publish_new\.(publish|recover|discard_staged)\(')


    def test_the_digest_observation_is_documented_and_the_helper_only_reads_through_the_owners(self):
        text = (ROOT / 'roles' / 'README.md').read_text(encoding='utf-8')
        recipe = MARKER.findall(text)[0]
        body = re.search(r'(?ms)^def observe_candidate\(.*?(?=^def )', recipe)
        self.assertIsNotNone(body, 'the helper lives in the recipe block, not in a second file')
        body = body.group(0)
        self.assertIn('artifacts.observe_file', body)
        self.assertIn('read_record(', body)
        self.assertIn('guard_task(', body)
        # reads only, through owners: no hashing of its own, no write, no ledger, no registry, no scan
        self.assertFalse([w for w in ('hashlib', 'sha(', "'wb'", "'xb'", '.write(', 'makedirs', 'os.replace', 'os.remove', 'rmtree', 'read_book', '_INDEX', 'listdir', 'walk(', 'glob',
                                       'register', 'prepare', 'relate') if w in body])
        for needle in ('Hai pha: Main quan sát digest, Worker ghi RESULT (V2-B)', 'Two phases for the digest (V2-B)', 'TOOL_OBSERVED', 'ACTUAL_BYTES', 'SendMessage', 'Agent.resume'):
            self.assertIn(needle, text)
        if openpyxl is None:
            self.skipTest('openpyxl missing: the executable part was not run')
        v3 = recipe_from_docs().load_v3(ROOT)
        self.assertEqual(Path(importlib.import_module('artifacts').__file__).resolve().parent, Path(v3['office'].__file__).resolve().parent)


class LargeInputCheckpoint(FlowBase):
    """Real owner regression: binary hashing is separate from bounded payload reads."""

    def large_assignment(self):
        path, context = 'references/large.bin', 'roles/large-context.bin'
        data = b'\x00\xff' * (1024 * 1024 + 123)
        self.write(path, data)
        self.write(context, data)
        t, r, task = self.assignment(inputs=[{'path': path, 'digest': sha(data)}])
        task['context_refs'] = [{'path': context, 'digest': sha(data)}]
        self.put_task(task)
        result = self.result_record(task, 'a1', b'candidate')
        self.write(r, json.dumps(result).encode())
        return t, r, path, context, data

    def test_large_binary_input_and_context_match_then_one_byte_change_is_detected(self):
        t, r, path, context, data = self.large_assignment()
        before = self.tree()
        self.assertEqual(self.checkpoint(t, r)['problems'], [])
        self.assertEqual(self.tree(), before)
        for changed in (path, context):
            with self.subTest(path=changed):
                self.write(changed, data[:-1] + b'x')
                self.assertEqual(self.checkpoint(t, r)['problems'], ['INPUT_CHANGED ' + changed])
                self.write(changed, data)

    def test_custom_byte_reader_is_not_bypassed_by_streaming(self):
        import task_checkpoint as tc
        t, r, path, _, _ = self.large_assignment()
        def refusing(project, rel):
            return (None, 'CALLER_REFUSED') if rel == path else tc.read_file(project, rel)
        with mock.patch.object(tc, 'file_digest', side_effect=AssertionError('must use caller')):
            cp = tc.checkpoint(self.v3, self.project, t, r, reader=refusing)
        self.assertIn('INPUT ' + path + ' CALLER_REFUSED', cp['problems'])

    def test_record_result_and_candidate_reads_remain_bounded(self):
        import task_checkpoint as tc
        t, r, _, _, _ = self.large_assignment()
        for rel, expected in ((t, 'RECORD_TASK TOO_LARGE'), (r, 'RECORD_RESULT TOO_LARGE'),
                              (self.slot('t-a', 'a1') + '/candidate/draft.txt', 'CANDIDATE TOO_LARGE')):
            with self.subTest(path=rel):
                old = self.read(rel)
                self.write(rel, b'x' * (2 * 1024 * 1024 + 1))
                self.assertIn(expected, tc.checkpoint(self.v3, self.project, t, r)['problems'])
                self.write(rel, old)

    def test_streaming_rejects_noncanonical_paths_and_junctions(self):
        import task_checkpoint as tc
        self.assertEqual(tc.file_digest(self.project, '../outside.bin'), (None, 0, 'PATH_UNSUPPORTED_FORM'))
        self.assertEqual(tc.file_digest(self.project, 'references/./outside.bin'), (None, 0, 'PATH_UNSUPPORTED_FORM'))
        outside = tempfile.mkdtemp(prefix='law-digest-outside-')
        self.addCleanup(shutil.rmtree, outside, True)
        Path(outside, 'secret.bin').write_bytes(b'secret')
        link = self.abs('references/linked')
        made = subprocess.run(['cmd', '/c', 'mklink', '/J', link, outside], capture_output=True, text=True)
        self.assertEqual(made.returncode, 0, made.stderr)
        self.addCleanup(lambda: os.path.isdir(link) and os.rmdir(link))
        self.assertEqual(tc.file_digest(self.project, 'references/linked/secret.bin'),
                         (None, 0, 'MISSING_OR_NOT_REGULAR'))


if __name__ == '__main__':
    unittest.main()
