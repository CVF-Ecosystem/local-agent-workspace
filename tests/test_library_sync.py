"""Focused source-index writer checks. All writes stay in temporary projects."""
import collections
import contextlib
import hashlib
import json
import os
from pathlib import Path
import shutil
import stat
import subprocess
import sys
import tempfile
import types
import unittest
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'tools'))
import library_sync
import result_schema
import transaction

try:
    import openpyxl
except ImportError:
    openpyxl = None


@unittest.skipUnless(openpyxl, 'openpyxl is needed for the source workbook')
class LibrarySync(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='law-source-sync-')
        self.addCleanup(self.tmp.cleanup)
        self.project = Path(self.tmp.name)
        self.ref = self.project / 'references'
        (self.ref / 'source-documents' / 'team').mkdir(parents=True)
        shutil.copyfile(ROOT / 'skills/local/shared/source-library/assets/source-index-template.xlsx',
                        self.ref / '_INDEX.xlsx')

    def source(self, name, content):
        path = self.ref / 'source-documents' / 'team' / name
        path.write_text(content, encoding='utf-8')
        return path

    def row(self):
        wb = openpyxl.load_workbook(self.ref / '_INDEX.xlsx', data_only=False)
        self.addCleanup(wb.close)
        ws = wb['Danh mục']
        return wb, [ws.cell(5, n).value for n in range(1, 26)]

    def test_preview_apply_idempotence_revision_and_history(self):
        path = self.source('report.txt', 'first')
        index = self.ref / '_INDEX.xlsx'
        before = hashlib.sha256(index.read_bytes()).hexdigest()
        proposal = library_sync.plan(str(self.project))
        self.assertEqual([x['action'] for x in proposal['changes']], ['add'])
        self.assertEqual(hashlib.sha256(index.read_bytes()).hexdigest(), before)
        library_sync.apply(str(self.project))
        self.assertIsNone(library_sync.apply(str(self.project))['bytes'])
        wb, values = self.row()
        self.assertEqual((values[0], values[21], values[24]), ('SRC-000001', 1, 'PRESENT'))
        wb.close()
        path.write_text('second', encoding='utf-8')
        changed = library_sync.apply(str(self.project))
        self.assertIn('changed', [x['action'] for x in changed['changes']])
        wb, values = self.row()
        self.assertEqual((values[0], values[13], values[21], values[24]),
                         ('SRC-000001', 'Cần đối chiếu', 2, 'NEEDS_REVIEW'))
        self.assertEqual(wb['LAW history'].cell(2, 2).value, 1)
        wb.close()
        self.assertIsNone(library_sync.apply(str(self.project))['bytes'])

    def test_missing_and_reappearing_file_keep_id(self):
        path = self.source('one.txt', 'same')
        library_sync.apply(str(self.project))
        path.unlink()
        self.assertEqual(library_sync.apply(str(self.project))['changes'][0]['action'], 'missing')
        _, values = self.row()
        self.assertEqual((values[0], values[24]), ('SRC-000001', 'MISSING'))
        path.write_text('same', encoding='utf-8')
        actions = library_sync.apply(str(self.project))['changes']
        self.assertIn('state', [x['action'] for x in actions])
        _, values = self.row()
        self.assertEqual((values[0], values[21], values[24]), ('SRC-000001', 1, 'PRESENT'))

    def test_same_bytes_at_new_path_do_not_prove_a_move(self):
        path = self.source('old.txt', 'same')
        library_sync.apply(str(self.project))
        path.unlink()
        library_sync.apply(str(self.project))
        self.source('new.txt', 'same')
        result = library_sync.apply(str(self.project))
        self.assertEqual([x['action'] for x in result['changes']], ['add'])
        wb, values = self.row()
        self.assertEqual((values[0], values[24]), ('SRC-000001', 'MISSING'))
        self.assertEqual(wb['Danh mục'].cell(6, 1).value, 'SRC-000002')
        wb.close()

    def test_user_notes_and_formula_survive(self):
        path = self.source('one.txt', 'same')
        wb = openpyxl.load_workbook(self.ref / '_INDEX.xlsx')
        ws = wb['Danh mục']
        ws.cell(5, 1, 'PRIVATE-42')
        ws.cell(5, 9, 'Tên do người dùng viết')
        ws.cell(5, 16, 'source-documents/team/one.txt')
        ws.cell(5, 21, '=1+2')
        wb.save(self.ref / '_INDEX.xlsx')
        wb.close()
        library_sync.apply(str(self.project))
        _, values = self.row()
        self.assertEqual((values[0], values[8], values[20]), ('PRIVATE-42', 'Tên do người dùng viết', '=1+2'))

    def test_snapshot_and_lock_block_conflicting_write(self):
        self.source('one.txt', 'same')
        proposal = library_sync.plan(str(self.project))
        index = self.ref / '_INDEX.xlsx'
        index.write_bytes(index.read_bytes() + b'changed-by-another-writer')
        tx = transaction.Transaction(str(self.project), 'source-sync-test')
        with self.assertRaises(transaction.TransactionError):
            tx.add('references/_INDEX.xlsx', proposal['bytes'], expected_before=proposal['snapshot'])
        with library_sync.writer_lock(str(self.project)):
            with self.assertRaises(library_sync.SyncError):
                library_sync.apply(str(self.project))

    def test_cli_preview_then_apply(self):
        self.source('one.txt', 'same')
        tools = self.project / 'tools'
        tools.mkdir()
        for name in ('library.py', 'library_sync.py', 'artifacts.py', 'transaction.py', 'result_schema.py'):
            shutil.copyfile(ROOT / 'tools' / name, tools / name)
        env = dict(os.environ, PYTHONUTF8='1', PYTHONDONTWRITEBYTECODE='1')
        for option, expected in (('', 'preview'), ('--apply', 'success')):
            args = [sys.executable, str(tools / 'library.py'), 'sync', '--json', '--result-json']
            if option:
                args.append(option)
            result = subprocess.run(args, cwd=self.project, env=env, capture_output=True,
                                    text=True, encoding='utf-8', timeout=30)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn('"status":"%s"' % expected, result.stdout)
            self.assertEqual(json.loads(result.stdout.splitlines()[0])['planned'], 1)


def make_dir_link(link, target):
    """A directory symlink, or a junction on Windows. False when this machine cannot create either."""
    try:
        os.symlink(str(target), str(link), target_is_directory=True)
        return True
    except (OSError, NotImplementedError):
        pass
    if os.name == 'nt':
        result = subprocess.run(['cmd', '/c', 'mklink', '/J', str(link), str(target)], capture_output=True)
        return result.returncode == 0 and os.path.isdir(str(link))
    return False


def drop_dir_link(link):
    """Remove a directory link without touching its target (rmdir for a junction, unlink for a symlink)."""
    try:
        os.rmdir(str(link))
    except OSError:
        os.unlink(str(link))


TEMPLATE_INDEX = ROOT / 'skills/local/shared/source-library/assets/source-index-template.xlsx'
SCOPE_LINE = '- Ngoại lệ thư mục nguồn cũ: %s\n'


@unittest.skipUnless(openpyxl, 'openpyxl is needed for the source workbook')
class LibraryScopeAndRename(unittest.TestCase):
    """A-05: old-source scope from PROJECT.md, README ownership, confirmed renames, links. Temporary projects only."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='law-source-scope-')
        self.addCleanup(self.tmp.cleanup)
        self.project = Path(self.tmp.name)
        self.ref = self.project / 'references'
        (self.ref / 'source-documents' / 'team').mkdir(parents=True)
        shutil.copyfile(ROOT / 'skills/local/shared/source-library/assets/source-index-template.xlsx',
                        self.ref / '_INDEX.xlsx')
        self.index = self.ref / '_INDEX.xlsx'

    def scope(self, value):
        (self.project / 'PROJECT.md').write_text('# PROJECT\n\n- Tên dự án: Thử\n' + SCOPE_LINE % value,
                                                  encoding='utf-8')

    def old(self, name='a.txt', content='old source', folder='Tai-lieu-cu'):
        path = self.project / folder / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding='utf-8')
        return path

    def team(self, name, content):
        path = self.ref / 'source-documents' / 'team' / name
        path.write_text(content, encoding='utf-8')
        return path

    def rows(self):
        wb = openpyxl.load_workbook(self.index, data_only=False)
        self.addCleanup(wb.close)
        ws = wb['Danh mục']
        out = [tuple(ws.cell(r, c).value for c in range(1, 26)) for r in range(5, ws.max_row + 1)
               if ws.cell(r, 1).value]
        return wb, out

    def digest(self):
        return hashlib.sha256(self.index.read_bytes()).hexdigest()

    def test_old_folder_is_scanned_only_with_the_scope_the_user_wrote(self):
        source = self.old()
        project_md = self.project / 'PROJECT.md'
        for value in ('Không có', 'None', '[Không có / đường dẫn và nhóm nguồn do người dùng chọn]'):
            self.scope(value)
            self.assertEqual(library_sync.plan(str(self.project))['changes'], [], value)
        self.assertFalse(project_md.read_text(encoding='utf-8').count('Tai-lieu-cu'))
        self.scope('Tai-lieu-cu')
        before = (hashlib.sha256(project_md.read_bytes()).hexdigest(), self.digest(),
                  hashlib.sha256(source.read_bytes()).hexdigest())
        preview = library_sync.plan(str(self.project))
        self.assertEqual([(c['action'], c['path']) for c in preview['changes']], [('add', '../Tai-lieu-cu/a.txt')])
        self.assertEqual(preview['scopes'], ['Tai-lieu-cu'])
        self.assertEqual(self.digest(), before[1], 'preview must not write the workbook')
        library_sync.apply(str(self.project))
        _, values = self.rows()
        self.assertEqual((values[0][0], values[0][15], values[0][24]), ('SRC-000001', '../Tai-lieu-cu/a.txt', 'PRESENT'))
        self.assertIn('Nguồn cũ giữ tại chỗ', values[0][20])
        self.assertEqual((hashlib.sha256(project_md.read_bytes()).hexdigest(), hashlib.sha256(source.read_bytes()).hexdigest()),
                         (before[0], before[2]), 'PROJECT.md and the old source are never written')
        self.assertIsNone(library_sync.apply(str(self.project))['bytes'], 'second run adds no rows')

    def test_old_source_rows_pass_the_existing_check_command(self):
        self.old()
        self.scope('Tai-lieu-cu')
        library_sync.apply(str(self.project))
        tools = self.project / 'tools'
        tools.mkdir()
        for name in ('library.py', 'library_sync.py', 'artifacts.py', 'transaction.py', 'result_schema.py'):
            shutil.copyfile(ROOT / 'tools' / name, tools / name)
        env = dict(os.environ, PYTHONUTF8='1', PYTHONDONTWRITEBYTECODE='1')
        result = subprocess.run([sys.executable, str(tools / 'library.py'), 'check', '--json'], cwd=self.project,
                                env=env, capture_output=True, text=True, encoding='utf-8', timeout=30)
        report = json.loads(result.stdout)
        self.assertEqual(report['errors'], [], result.stdout)
        self.assertNotIn('../Tai-lieu-cu/a.txt', ' '.join(report['unindexed']))

    def test_invalid_scope_blocks_the_sync_and_writes_nothing(self):
        self.old()
        before = self.digest()
        outside = Path(tempfile.gettempdir()).as_posix()
        for value in (outside, 'C:/Users', '../Tai-lieu-cu', 'references', 'tools', 'output', 'Khong-co-thu-muc',
                      'Tai-lieu-cu; [điền sau]', 'a//b', './'):
            with self.subTest(value=value):
                self.scope(value)
                with self.assertRaises(library_sync.SyncError):
                    library_sync.plan(str(self.project))
                with self.assertRaises(library_sync.SyncError):
                    library_sync.apply(str(self.project))
                self.assertEqual(self.digest(), before)

    def test_scope_removed_or_file_deleted_keeps_the_id_with_an_honest_state(self):
        source = self.old()
        self.scope('Tai-lieu-cu')
        library_sync.apply(str(self.project))
        self.scope('Không có')
        actions = [c['action'] for c in library_sync.apply(str(self.project))['changes']]
        self.assertEqual(actions, ['out_of_scope'])
        _, values = self.rows()
        self.assertEqual((values[0][0], values[0][24]), ('SRC-000001', 'OUT_OF_SCOPE'))
        self.scope('Tai-lieu-cu')
        library_sync.apply(str(self.project))
        _, values = self.rows()
        self.assertEqual((values[0][0], values[0][21], values[0][24]), ('SRC-000001', 1, 'PRESENT'))
        source.unlink()
        self.assertEqual([c['action'] for c in library_sync.apply(str(self.project))['changes']], ['missing'])
        _, values = self.rows()
        self.assertEqual(values[0][24], 'MISSING')

    def test_readme_ownership_not_basename(self):
        policies = self.ref / 'policies'
        policies.mkdir()
        shipped = policies / 'README.md'
        shipped.write_text('shipped guidance\n', encoding='utf-8')
        sha = hashlib.sha256(shipped.read_bytes()).hexdigest()
        (self.project / 'MANIFEST.sha256').write_text('%s  references/policies/README.md\n' % sha, encoding='utf-8')
        self.assertEqual(library_sync.plan(str(self.project))['changes'], [], 'package README is not a source')
        (self.project / 'MANIFEST.sha256').write_text('%s  en/references/policies/README.md\n' % sha, encoding='utf-8')
        self.assertEqual(library_sync.plan(str(self.project))['changes'], [], 'English package README is not a source')
        (self.project / 'MANIFEST.sha256').write_text('', encoding='utf-8')
        shipped.write_bytes(library_sync.library.group_readme('policies').encode('utf-8'))
        self.assertEqual(library_sync.plan(str(self.project))['changes'], [], 'README written by init is not a source')
        shipped.write_text('Quy chế do người dùng viết\n', encoding='utf-8')
        planned = library_sync.plan(str(self.project))['changes']
        self.assertEqual([(c['action'], c['path']) for c in planned], [('add', 'policies/README.md')])
        deeper = policies / 'hr'
        deeper.mkdir()
        (deeper / 'README.md').write_text('shipped guidance\n', encoding='utf-8')
        paths = [c['path'] for c in library_sync.plan(str(self.project))['changes']]
        self.assertIn('policies/hr/README.md', paths, 'a README below the group folder is a source')

    def test_confirmed_rename_keeps_the_id_and_user_data(self):
        old = self.team('old.txt', 'same bytes')
        library_sync.apply(str(self.project))
        wb = openpyxl.load_workbook(self.index)
        ws = wb['Danh mục']
        ws.cell(5, 9, 'Tên do người dùng viết')
        ws.cell(5, 21, '=1+2')
        wb.save(self.index)
        wb.close()
        old.unlink()
        self.team('moved.txt', 'same bytes')
        preview = library_sync.plan(str(self.project))
        self.assertEqual(sorted(c['action'] for c in preview['changes']), ['add', 'missing'], 'no automatic merge')
        self.assertEqual(len(preview['candidates']), 1)
        candidate = preview['candidates'][0]
        self.assertEqual((candidate['id'], candidate['from'], candidate['to'], candidate['ambiguous']),
                         ('SRC-000001', 'source-documents/team/old.txt', 'source-documents/team/moved.txt', False))
        before = self.digest()
        confirmed = library_sync.plan(str(self.project), [candidate['token']])
        self.assertEqual(self.digest(), before, 'preview does not write')
        self.assertEqual([c['action'] for c in confirmed['changes']], ['renamed'])
        self.assertTrue(confirmed['changes'][0]['same_content'])
        library_sync.apply(str(self.project), [candidate['token']])
        wb, values = self.rows()
        self.assertEqual(len(values), 1, 'no duplicate row')
        self.assertEqual((values[0][0], values[0][8], values[0][15], values[0][20], values[0][21], values[0][24]),
                         ('SRC-000001', 'Tên do người dùng viết', 'source-documents/team/moved.txt', '=1+2', 1, 'PRESENT'))
        self.assertNotIn('LAW history', wb.sheetnames, 'a pure move is not a content revision')
        self.assertIsNone(library_sync.apply(str(self.project))['bytes'])

    def test_hand_written_token_with_changed_content_is_accepted_but_never_offered_by_the_preview(self):
        """Unit check of the apply branch only. The preview offers tokens for equal hashes alone, so there is no
        preview-to-apply flow for a rename that also changes content; this token is hand-built, not preview evidence."""
        old = self.team('old.txt', 'first')
        library_sync.apply(str(self.project))
        old.unlink()
        new = self.team('new.txt', 'second version')
        self.assertEqual(library_sync.plan(str(self.project))['candidates'], [], 'the preview offers no such rename')
        token = 'SRC-000001=source-documents/team/new.txt#%s' % library_sync.digest(str(new))[:12]
        library_sync.apply(str(self.project), [token])
        wb, values = self.rows()
        self.assertEqual((values[0][0], values[0][15], values[0][21], values[0][24]),
                         ('SRC-000001', 'source-documents/team/new.txt', 2, 'NEEDS_REVIEW'))
        self.assertEqual(wb['LAW history'].cell(2, 2).value, 1)

    def test_copy_and_equal_hash_never_prove_a_move(self):
        self.team('old.txt', 'same bytes')
        library_sync.apply(str(self.project))
        copy = self.team('copy.txt', 'same bytes')
        preview = library_sync.plan(str(self.project))
        self.assertEqual(preview['candidates'], [], 'the old file still exists: this is a copy')
        token = 'SRC-000001=source-documents/team/copy.txt#%s' % library_sync.digest(str(copy))[:12]
        before = self.digest()
        with self.assertRaises(library_sync.SyncError):
            library_sync.apply(str(self.project), [token])
        self.assertEqual(self.digest(), before)
        library_sync.apply(str(self.project))
        _, values = self.rows()
        self.assertEqual([v[0] for v in values], ['SRC-000001', 'SRC-000002'], 'copy gets its own ID')

    def test_ambiguous_candidates_are_flagged_and_only_the_confirmed_one_moves(self):
        old = self.team('old.txt', 'same bytes')
        library_sync.apply(str(self.project))
        old.unlink()
        self.team('one.txt', 'same bytes')
        self.team('two.txt', 'same bytes')
        preview = library_sync.plan(str(self.project))
        self.assertEqual(len(preview['candidates']), 2)
        self.assertTrue(all(c['ambiguous'] for c in preview['candidates']))
        token = [c['token'] for c in preview['candidates'] if c['to'].endswith('one.txt')][0]
        library_sync.apply(str(self.project), [token])
        _, values = self.rows()
        self.assertEqual([(v[0], v[15]) for v in values],
                         [('SRC-000001', 'source-documents/team/one.txt'), ('SRC-000002', 'source-documents/team/two.txt')])

    def test_blocked_renames_write_nothing(self):
        old = self.team('old.txt', 'first')
        other = self.team('other.txt', 'other')
        library_sync.apply(str(self.project))
        old.unlink()
        new = self.team('new.txt', 'first')
        sha = library_sync.digest(str(new))[:12]
        good = 'SRC-000001=source-documents/team/new.txt#%s' % sha
        before = self.digest()
        cases = {
            'unknown id': 'SRC-000099=source-documents/team/new.txt#%s' % sha,
            'no confirmation code': 'SRC-000001=source-documents/team/new.txt',
            'short code': 'SRC-000001=source-documents/team/new.txt#abc',
            'wrong hash': 'SRC-000001=source-documents/team/new.txt#%s' % ('0' * 12),
            'path not scanned': 'SRC-000001=source-documents/team/nothing.txt#%s' % sha,
            'path owned by another id': 'SRC-000001=source-documents/team/other.txt#%s' % library_sync.digest(str(other))[:12],
            'old file still present': 'SRC-000002=source-documents/team/new.txt#%s' % sha,
        }
        for name, token in cases.items():
            with self.subTest(case=name):
                with self.assertRaises(library_sync.SyncError):
                    library_sync.apply(str(self.project), [token])
                self.assertEqual(self.digest(), before)
        with self.assertRaises(library_sync.SyncError):
            library_sync.apply(str(self.project), [good, good])
        new.write_text('changed after the preview', encoding='utf-8')
        with self.assertRaises(library_sync.SyncError):
            library_sync.apply(str(self.project), [good])
        self.assertEqual(self.digest(), before, 'a file changed after the preview blocks the confirmation')

    def test_links_do_not_extend_the_scope(self):
        outside = tempfile.TemporaryDirectory(prefix='law-outside-')
        self.addCleanup(outside.cleanup)
        secret = Path(outside.name) / 'secret.txt'
        secret.write_text('not a project source', encoding='utf-8')
        link = self.ref / 'source-documents' / 'team' / 'linked'
        if not make_dir_link(link, outside.name):
            self.skipTest('this machine cannot create a directory symlink or junction')
        self.team('inside.txt', 'inside')
        self.old()
        escape = self.project / 'Tai-lieu-cu' / 'escape'
        self.assertTrue(make_dir_link(escape, outside.name))
        self.scope('Tai-lieu-cu')
        planned = library_sync.plan(str(self.project))
        paths = [c['path'] for c in planned['changes']]
        self.assertEqual(sorted(paths), ['../Tai-lieu-cu/a.txt', 'source-documents/team/inside.txt'])
        self.assertTrue(any(reason == 'link' for _, reason in planned['skipped']))
        self.assertFalse(any('secret' in path for path in paths))
        top = self.project / 'Lien-ket'
        self.assertTrue(make_dir_link(top, outside.name))
        self.scope('Lien-ket')
        with self.assertRaises(library_sync.SyncError):
            library_sync.plan(str(self.project))

    def snapshot(self, folder):
        return {str(p.relative_to(folder)): hashlib.sha256(p.read_bytes()).hexdigest()
                for p in sorted(Path(folder).rglob('*')) if p.is_file()}

    def tmp_dir(self, prefix):
        folder = tempfile.TemporaryDirectory(prefix=prefix)
        self.addCleanup(folder.cleanup)
        return Path(folder.name)

    def test_references_junction_out_of_the_project_is_refused_before_any_write(self):
        target = self.tmp_dir('law-outside-refs-') / 'refs'
        (target / 'source-documents' / 'team').mkdir(parents=True)
        shutil.copyfile(TEMPLATE_INDEX, target / '_INDEX.xlsx')
        (target / 'source-documents' / 'team' / 'x.txt').write_text('outside source', encoding='utf-8')
        project = self.tmp_dir('law-linked-project-')
        link = project / 'references'
        if not make_dir_link(link, target):
            self.skipTest('this machine cannot create a directory symlink or junction')
        self.addCleanup(drop_dir_link, link)
        before = self.snapshot(target)
        calls = {'preview': lambda: library_sync.plan(str(project)),
                 'apply': lambda: library_sync.apply(str(project)),
                 'apply with a token': lambda: library_sync.apply(
                     str(project), ['SRC-000001=source-documents/team/x.txt#abcdef12'])}
        for name, call in calls.items():
            with self.subTest(call=name):
                with self.assertRaises(library_sync.SyncError):
                    call()
        self.assertEqual(self.snapshot(target), before, 'the workbook and source outside the project keep their bytes')
        self.assertFalse((project / '.agent').exists(), 'refused before the lock or any transaction was created')
        tools = project / 'tools'
        tools.mkdir()
        for name in ('library.py', 'library_sync.py', 'artifacts.py', 'transaction.py', 'result_schema.py'):
            shutil.copyfile(ROOT / 'tools' / name, tools / name)
        env = dict(os.environ, PYTHONUTF8='1', PYTHONDONTWRITEBYTECODE='1')
        for extra in ([], ['--apply']):
            result = subprocess.run([sys.executable, str(tools / 'library.py'), 'sync', '--json'] + extra, cwd=project,
                                    env=env, capture_output=True, text=True, encoding='utf-8', timeout=30)
            self.assertNotEqual(result.returncode, 0, result.stdout)
        self.assertEqual(self.snapshot(target), before, 'the CLI preview and apply write nothing outside either')
        self.assertFalse((project / '.agent').exists())

    def test_project_folder_reached_through_a_link_still_syncs(self):
        holder = self.tmp_dir('law-through-link-')
        project = holder / 'real' / 'proj'
        (project / 'references' / 'source-documents' / 'team').mkdir(parents=True)
        shutil.copyfile(TEMPLATE_INDEX, project / 'references' / '_INDEX.xlsx')
        (project / 'references' / 'source-documents' / 'team' / 'one.txt').write_text('one', encoding='utf-8')
        alias = holder / 'alias'
        if not make_dir_link(alias, holder / 'real'):
            self.skipTest('this machine cannot create a directory symlink or junction')
        self.addCleanup(drop_dir_link, alias)
        through = str(alias / 'proj')
        self.assertEqual([c['action'] for c in library_sync.plan(through)['changes']], ['add'])
        library_sync.apply(through)
        ws = openpyxl.load_workbook(project / 'references' / '_INDEX.xlsx')['Danh mục']
        self.assertEqual((ws.cell(5, 1).value, ws.cell(5, 16).value), ('SRC-000001', 'source-documents/team/one.txt'))

    def test_scope_removed_with_a_copy_never_offers_or_accepts_a_rename(self):
        original = self.old(content='old source')
        self.scope('Tai-lieu-cu')
        library_sync.apply(str(self.project))
        self.scope('Không có')
        copy = self.team('copy.txt', 'old source')
        preview = library_sync.plan(str(self.project))
        self.assertEqual(preview['candidates'], [], 'the original still exists and is only out of scope')
        self.assertEqual(sorted(c['action'] for c in preview['changes']), ['add', 'out_of_scope'])
        forged = 'SRC-000001=source-documents/team/copy.txt#%s' % library_sync.digest(str(copy))[:12]
        before = self.digest()
        for call in (library_sync.plan, library_sync.apply):
            with self.assertRaises(library_sync.SyncError):
                call(str(self.project), [forged])
            self.assertEqual(self.digest(), before, 'a hand-built token changes nothing')
        library_sync.apply(str(self.project))
        _, values = self.rows()
        self.assertEqual([(v[0], v[15], v[24]) for v in values],
                         [('SRC-000001', '../Tai-lieu-cu/a.txt', 'OUT_OF_SCOPE'),
                          ('SRC-000002', 'source-documents/team/copy.txt', 'PRESENT')])
        self.assertEqual(original.read_text(encoding='utf-8'), 'old source', 'the old source is never touched')

    def test_token_taken_before_the_scope_change_is_refused(self):
        original = self.old(content='old source')
        self.scope('Tai-lieu-cu')
        library_sync.apply(str(self.project))
        original.unlink()
        self.team('moved.txt', 'old source')
        offered = library_sync.plan(str(self.project))['candidates']
        self.assertEqual([(c['id'], c['from']) for c in offered], [('SRC-000001', '../Tai-lieu-cu/a.txt')])
        self.scope('Không có')
        before = self.digest()
        self.assertEqual(library_sync.plan(str(self.project))['candidates'], [])
        with self.assertRaises(library_sync.SyncError):
            library_sync.apply(str(self.project), [offered[0]['token']])
        self.assertEqual(self.digest(), before, 'the token from the earlier scope no longer applies')

    def test_move_out_of_an_old_source_scope_keeps_the_id_in_one_plan(self):
        original = self.old(content='old source')
        self.scope('Tai-lieu-cu')
        library_sync.apply(str(self.project))
        original.unlink()
        self.team('moved.txt', 'old source')
        _, values = self.rows()
        self.assertEqual(values[0][24], 'PRESENT', 'no earlier sync has written MISSING')
        preview = library_sync.plan(str(self.project))
        self.assertEqual(sorted(c['action'] for c in preview['changes']), ['add', 'missing'], 'no automatic merge')
        self.assertEqual(len(preview['candidates']), 1)
        token = preview['candidates'][0]['token']
        self.assertEqual([c['action'] for c in library_sync.plan(str(self.project), [token])['changes']], ['renamed'])
        library_sync.apply(str(self.project), [token])
        _, values = self.rows()
        self.assertEqual([(v[0], v[15], v[24]) for v in values],
                         [('SRC-000001', 'source-documents/team/moved.txt', 'PRESENT')])

    def test_present_package_file_is_not_a_move(self):
        guide = self.team('guide.txt', 'shipped text')
        library_sync.apply(str(self.project))
        manifest = '%s  references/source-documents/team/guide.txt\n' % library_sync.digest(str(guide))
        (self.project / 'MANIFEST.sha256').write_text(manifest, encoding='utf-8')
        copy = self.team('copy.txt', 'shipped text')
        self.assertEqual(library_sync.plan(str(self.project))['candidates'], [],
                         'the old path still holds a file; it is only skipped as a package file')
        token = 'SRC-000001=source-documents/team/copy.txt#%s' % library_sync.digest(str(copy))[:12]
        before = self.digest()
        with self.assertRaises(library_sync.SyncError):
            library_sync.apply(str(self.project), [token])
        self.assertEqual(self.digest(), before)

    def test_old_path_skipped_as_a_link_is_not_a_move(self):
        folder = self.ref / 'source-documents' / 'team' / 'sub'
        folder.mkdir()
        (folder / 'a.txt').write_text('same bytes', encoding='utf-8')
        library_sync.apply(str(self.project))
        elsewhere = self.tmp_dir('law-link-target-')
        (elsewhere / 'a.txt').write_text('same bytes', encoding='utf-8')
        shutil.rmtree(folder)
        if not make_dir_link(folder, elsewhere):
            self.skipTest('this machine cannot create a directory symlink or junction')
        self.addCleanup(drop_dir_link, folder)
        copy = self.team('copy.txt', 'same bytes')
        preview = library_sync.plan(str(self.project))
        self.assertTrue(any(reason == 'link' for _, reason in preview['skipped']))
        self.assertEqual(preview['candidates'], [], 'the old path still resolves to a file behind a link')
        token = 'SRC-000001=source-documents/team/copy.txt#%s' % library_sync.digest(str(copy))[:12]
        before = self.digest()
        with self.assertRaises(library_sync.SyncError):
            library_sync.apply(str(self.project), [token])
        self.assertEqual(self.digest(), before)

    def test_cli_preview_prints_the_token_and_apply_accepts_it(self):
        old = self.team('old.txt', 'same bytes')
        library_sync.apply(str(self.project))
        old.unlink()
        self.team('moved.txt', 'same bytes')
        tools = self.project / 'tools'
        tools.mkdir()
        for name in ('library.py', 'library_sync.py', 'artifacts.py', 'transaction.py', 'result_schema.py'):
            shutil.copyfile(ROOT / 'tools' / name, tools / name)
        env = dict(os.environ, PYTHONUTF8='1', PYTHONDONTWRITEBYTECODE='1')

        def run(*extra):
            return subprocess.run([sys.executable, str(tools / 'library.py'), 'sync', '--json'] + list(extra),
                                  cwd=self.project, env=env, capture_output=True, text=True, encoding='utf-8', timeout=30)

        preview = run()
        self.assertEqual(preview.returncode, 0, preview.stderr)
        token = json.loads(preview.stdout.splitlines()[0])['candidates'][0]['token']
        before = self.digest()
        bad = run('--apply', '--rename', 'SRC-000001=source-documents/team/moved.txt')
        self.assertNotEqual(bad.returncode, 0)
        self.assertEqual(self.digest(), before)
        good = run('--apply', '--rename', token)
        self.assertEqual(good.returncode, 0, good.stderr + good.stdout)
        _, values = self.rows()
        self.assertEqual([(v[0], v[15]) for v in values], [('SRC-000001', 'source-documents/team/moved.txt')])



def tree(folder):
    """Relative path -> sha256 for every file under folder (empty dict when the folder is absent)."""
    out = {}
    for path in sorted(Path(folder).rglob('*')):
        if path.is_file():
            out[str(path.relative_to(folder))] = hashlib.sha256(path.read_bytes()).hexdigest()
    return out


def only_open(items):
    assert len(items) == 1, items
    return items[0]


def hold_like_excel(path):
    """Open a file the way Excel holds an open workbook: others may read it, nobody may write or replace it.

    Windows only. This is an operating-system sharing lock, not Excel itself; see the HANDOFF for the real-Excel run.
    """
    import ctypes
    from ctypes import wintypes
    k32 = ctypes.WinDLL('kernel32', use_last_error=True)
    k32.CreateFileW.restype = wintypes.HANDLE
    k32.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
                                wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
    k32.CloseHandle.argtypes = [wintypes.HANDLE]
    handle = k32.CreateFileW(str(path), 0x80000000, 0x1, None, 3, 0x80, None)  # GENERIC_READ, share READ only, OPEN_EXISTING
    if handle in (None, ctypes.c_void_p(-1).value):
        raise OSError(ctypes.get_last_error(), 'could not open the workbook with a read-only share mode')
    return lambda: k32.CloseHandle(handle)


@unittest.skipUnless(openpyxl, 'openpyxl is needed for the source workbook')
class LibraryRuntimeAndBoundary(unittest.TestCase):
    """A-06: runtime links, the check-to-write boundary, and T13-T16 conflict/recovery for the index writer.

    Temporary projects only. Junctions are real; a real symlink cannot be created on every machine (Windows error 1314)
    and is reported as skipped, never as passed. Locks and disk errors here are simulations, not Excel.
    """

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='law-a06-')
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        self.project = self.base / 'proj'
        self.ref = self.project / 'references'
        (self.ref / 'source-documents' / 'team').mkdir(parents=True)
        shutil.copyfile(TEMPLATE_INDEX, self.ref / '_INDEX.xlsx')
        self.index = self.ref / '_INDEX.xlsx'
        (self.ref / 'source-documents' / 'team' / 'one.txt').write_text('one', encoding='utf-8')
        self.original = self.index.read_bytes()

    def outside(self, name='outside'):
        folder = self.base / name
        folder.mkdir()
        (folder / 'sentinel.txt').write_text('keep', encoding='utf-8')
        return folder

    def link(self, where, target):
        if not make_dir_link(where, target):
            self.skipTest('this machine cannot create a directory symlink or junction')
        self.addCleanup(drop_dir_link, where)

    def rows(self):
        wb = openpyxl.load_workbook(self.index)
        self.addCleanup(wb.close)
        ws = wb['Danh mục']
        return [tuple(ws.cell(r, c).value for c in (1, 9, 16, 25)) for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value]

    def open_transactions(self):
        return transaction.incomplete(str(self.project))

    # ---- P1: runtime components that are links ----
    def make_project(self, name):
        project = self.base / name
        (project / 'references' / 'source-documents' / 'team').mkdir(parents=True)
        shutil.copyfile(TEMPLATE_INDEX, project / 'references' / '_INDEX.xlsx')
        (project / 'references' / 'source-documents' / 'team' / 'one.txt').write_text('one', encoding='utf-8')
        return project

    def test_runtime_links_are_refused_before_any_lock_or_journal(self):
        cases = ('.agent', '.agent/runtime', '.agent/runtime/library', '.agent/runtime/transactions')
        for number, name in enumerate(cases):
            with self.subTest(link=name):
                project = self.make_project('runtime-%d' % number)
                where = project.joinpath(*name.split('/'))
                where.parent.mkdir(parents=True, exist_ok=True)
                target = self.outside('outside-%d' % number)
                self.link(where, target)
                before = tree(target)
                index_before = (project / 'references' / '_INDEX.xlsx').read_bytes()
                with self.assertRaises(library_sync.SyncError):
                    library_sync.apply(str(project))
                self.assertEqual(tree(target), before, 'no lock, journal or staging was created outside the project')
                self.assertEqual((project / 'references' / '_INDEX.xlsx').read_bytes(), index_before)

    def test_runtime_link_cli_refuses_and_status_reports_it(self):
        target = self.outside()
        self.link(self.project / '.agent', target)
        tools = self.project / 'tools'
        tools.mkdir()
        for name in ('library.py', 'library_sync.py', 'artifacts.py', 'transaction.py', 'result_schema.py'):
            shutil.copyfile(ROOT / 'tools' / name, tools / name)
        env = dict(os.environ, PYTHONUTF8='1', PYTHONDONTWRITEBYTECODE='1')
        before = tree(target)

        def run(*args):
            return subprocess.run([sys.executable] + list(args), cwd=self.project, env=env, capture_output=True,
                                  text=True, encoding='utf-8', timeout=60)
        self.assertNotEqual(run('tools/library.py', 'sync', '--apply').returncode, 0)
        status = run('tools/transaction.py', 'status')
        self.assertEqual(status.returncode, 2, status.stdout + status.stderr)
        self.assertEqual(tree(target), before)
        self.assertEqual(self.index.read_bytes(), self.original)

    # ---- check-to-write boundary ----
    def swap_references_hook(self, outside):
        """Test hook: at the moment the folder is held and checked, try to turn references into a link to outside."""
        state = {'done': False, 'blocked': None, 'swapped': False}

        def hook(point, path):
            if point != 'commit' or os.path.basename(path) != '_INDEX.xlsx' or state['done']:
                return
            state['done'] = True
            try:
                os.rename(str(self.ref), str(self.ref) + '.real')
            except OSError as e:
                state['blocked'] = e
                return
            if not make_dir_link(self.ref, outside):
                os.rename(str(self.ref) + '.real', str(self.ref))
                state['blocked'] = OSError('could not create the link')
                return
            state['swapped'] = True
        return hook, state

    def restore_references(self):
        if os.path.isdir(str(self.ref) + '.real'):
            drop_dir_link(self.ref)

    def test_swapping_references_between_check_and_write_is_blocked_while_the_folder_is_held(self):
        if os.name != 'nt':
            self.skipTest('the held-folder protection is implemented for Windows only; see test_swap_is_refused_when_no_folder_can_be_held')
        outside = self.outside()
        shutil.copyfile(self.index, outside / '_INDEX.xlsx')  # same bytes, same hash: expected_before alone would accept it
        before = tree(outside)
        hook, state = self.swap_references_hook(outside)
        with mock.patch.object(transaction, '_HOOK', hook):
            library_sync.apply(str(self.project))
        self.assertTrue(state['done'])
        self.assertIsInstance(state['blocked'], PermissionError, 'the swap itself was refused by the operating system')
        self.assertFalse(state['swapped'])
        self.assertEqual(tree(outside), before, 'the identical workbook outside the project was not written')
        self.assertEqual([r[0] for r in self.rows()], ['SRC-000001'], 'the write landed in the real references')

    def test_swap_is_refused_when_no_folder_can_be_held(self):
        outside = self.outside()
        shutil.copyfile(self.index, outside / '_INDEX.xlsx')
        before = tree(outside)
        hook, state = self.swap_references_hook(outside)
        self.addCleanup(self.restore_references)
        with mock.patch.object(transaction, 'PIN_ENABLED', False), mock.patch.object(transaction, '_HOOK', hook):
            with self.assertRaises(library_sync.SyncError):
                library_sync.apply(str(self.project))
        self.assertTrue(state['swapped'], 'without a held folder the swap does happen; the re-check must catch it')
        self.assertEqual(tree(outside), before, 'the identical workbook outside the project was not written')
        self.assertEqual((Path(str(self.ref) + '.real') / '_INDEX.xlsx').read_bytes(), self.original,
                         'the real workbook was not half written either')

    # ---- T13: edited outside the LAW lock ----
    def test_workbook_saved_by_the_user_after_planning_is_never_overwritten(self):
        def user_saves(point, path):
            if point != 'commit':
                return
            wb = openpyxl.load_workbook(self.index)
            ws = wb['Danh mục']
            ws.cell(5, 1, 'PRIVATE-42')
            ws.cell(5, 9, 'Dòng do người dùng thêm')
            wb.save(self.index)
            wb.close()
        with mock.patch.object(transaction, '_HOOK', user_saves):
            with self.assertRaises(library_sync.SyncError):
                library_sync.apply(str(self.project))
        saved = self.index.read_bytes()
        self.assertEqual(self.rows()[0][:2], ('PRIVATE-42', 'Dòng do người dùng thêm'))
        (item,) = self.open_transactions()
        result = transaction.rollback(str(self.project), item['id'])
        self.assertEqual(result['code'], 1, 'a user edit made after the plan is kept, not restored over')
        self.assertEqual(self.index.read_bytes(), saved)
        self.assertEqual(self.open_transactions(), [], 'rollback_partial does not block new work')
        library_sync.apply(str(self.project))
        rows = self.rows()
        self.assertEqual(rows[0][:2], ('PRIVATE-42', 'Dòng do người dùng thêm'), 'the user row survives the next sync')
        self.assertIn('SRC-000001', [r[0] for r in rows])

    # ---- T14: file open or locked ----
    def test_index_held_open_like_excel_leaves_it_unchanged_and_leaves_no_open_transaction(self):
        if os.name != 'nt':
            self.skipTest('the sharing-lock simulation uses the Windows file-sharing API')
        release = hold_like_excel(self.index)
        try:
            with self.assertRaises(library_sync.SyncError) as ctx:
                library_sync.apply(str(self.project))
            self.assertIn('Chưa ghi gì', str(ctx.exception))
        finally:
            release()
        self.assertEqual(self.index.read_bytes(), self.original)
        self.assertEqual(self.open_transactions(), [], 'nothing was written, so nothing is left to resume')
        library_sync.apply(str(self.project))
        self.assertEqual([r[0] for r in self.rows()], ['SRC-000001'], 'the same sync works once the file is released')

    def test_read_only_index_is_refused_and_unchanged(self):
        if os.name != 'nt':
            self.skipTest('replacing over a read-only file is refused by Windows; on POSIX the folder permissions decide')
        os.chmod(self.index, 0o444)
        self.addCleanup(os.chmod, self.index, 0o666)
        with self.assertRaises(library_sync.SyncError):
            library_sync.apply(str(self.project))
        self.assertEqual(self.index.read_bytes(), self.original)
        self.assertEqual(self.open_transactions(), [])

    # ---- T15: invalid workbook, disk and permission errors ----
    def test_invalid_workbooks_are_refused_without_writing_or_leaving_a_journal(self):
        def not_a_workbook():
            self.index.write_bytes(b'this is not an xlsx file')

        def missing_sheet():
            wb = openpyxl.load_workbook(TEMPLATE_INDEX)
            wb['Danh mục'].title = 'Sheet khác'
            wb.save(self.index)
            wb.close()

        def missing_columns():
            wb = openpyxl.load_workbook(TEMPLATE_INDEX)
            wb['Danh mục'].cell(library_sync.library.HEADER_ROW, 1, 'Cột khác')
            wb.save(self.index)
            wb.close()
        for name, make in (('not a workbook', not_a_workbook), ('missing sheet', missing_sheet),
                           ('missing columns', missing_columns)):
            with self.subTest(case=name):
                make()
                before = self.index.read_bytes()
                with self.assertRaises(library_sync.SyncError):
                    library_sync.apply(str(self.project))
                self.assertEqual(self.index.read_bytes(), before)
                self.assertEqual(self.open_transactions(), [])
                self.assertFalse(list((self.project / '.agent' / 'runtime').glob('transactions/*')),
                                 'no journal for a write that never started')

    def test_simulated_disk_and_permission_errors_write_nothing(self):
        """Simulations (patched disk_usage and write probe), not a full disk or a locked-down account."""
        usage = collections.namedtuple('usage', 'total used free')
        with mock.patch.object(transaction.shutil, 'disk_usage', lambda path: usage(100, 100, 10)):
            with self.assertRaises(library_sync.SyncError) as ctx:
                library_sync.apply(str(self.project))
        self.assertIn('Thiếu dung lượng', str(ctx.exception))
        def refuse(project):
            raise OSError(13, 'Permission denied')
        with mock.patch.object(transaction, 'probe_writable', refuse):
            with self.assertRaises(library_sync.SyncError):
                library_sync.apply(str(self.project))
        self.assertEqual(self.index.read_bytes(), self.original)
        self.assertEqual(self.open_transactions(), [])
        library_sync.apply(str(self.project))
        self.assertEqual([r[0] for r in self.rows()], ['SRC-000001'], 'the same sync works once the fault is gone')

    def x14_workbook(self):
        """The template with an x14 data-validation block in every sheet, the shape real Excel writes when it re-saves the file."""
        import io as _io
        import zipfile as _zip
        ext = (b'<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" '
               b'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" '
               b'xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main"><x14:dataValidations count="1">'
               b'<x14:dataValidation type="list"><x14:formula1><xm:f>Sheet2!$A$2:$A$5</xm:f></x14:formula1>'
               b'<xm:sqref>B5:B204</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>')
        wb = openpyxl.load_workbook(self.index)  # real Excel leaves no classic validation behind, only the x14 block
        for ws in wb.worksheets:
            ws.data_validations.dataValidation = []
        wb.save(self.index)
        wb.close()
        out = _io.BytesIO()
        with _zip.ZipFile(self.index) as zin, _zip.ZipFile(out, 'w', _zip.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.startswith('xl/worksheets/') and item.filename.endswith('.xml'):
                    data = data.replace(b'</worksheet>', ext + b'</worksheet>')
                zout.writestr(item, data)
        self.index.write_bytes(out.getvalue())

    def x14_count(self):
        import zipfile as _zip
        with _zip.ZipFile(self.index) as z:
            return sum(z.read(n).count(b'x14:dataValidation ') for n in z.namelist() if n.startswith('xl/worksheets/'))

    def test_apply_refuses_before_any_transaction_when_it_would_drop_x14_drop_downs(self):
        """A-06 x14 known loss: preview stays readable; apply stops before Transaction, staging and backup, workbook bytes unchanged."""
        import warnings
        self.x14_workbook()
        before = self.index.read_bytes()
        self.assertGreater(self.x14_count(), 0)
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            preview = library_sync.plan(str(self.project))
        self.assertEqual(preview['notices'], ['x14_data_validation'])
        self.assertIsNotNone(preview['bytes'], 'the preview is still produced')
        self.assertEqual(self.index.read_bytes(), before, 'preview does not touch the file')
        created = []
        real = transaction.Transaction
        def spy(*args, **kw):
            created.append(args)
            return real(*args, **kw)
        with mock.patch.object(transaction, 'Transaction', spy):
            with self.assertRaises(library_sync.SyncError) as ctx:
                library_sync.apply(str(self.project))
        self.assertIn('x14', str(ctx.exception))
        self.assertEqual(created, [], 'no Transaction was even created')
        self.assertEqual(self.index.read_bytes(), before, 'workbook bytes and hash are unchanged')
        self.assertEqual(self.open_transactions(), [])
        runtime = self.project / '.agent' / 'runtime'
        self.assertEqual(sorted(os.listdir(str(runtime))) if runtime.exists() else [], ['library'] if runtime.exists() else [],
                         'only the lock folder: no journal, staging or backup')
        self.assertGreater(self.x14_count(), 0, 'the x14 validation is still in the file')

    def test_direct_cli_apply_with_x14_exits_nonzero_and_reports_nothing_written(self):
        import json as _json
        self.x14_workbook()
        before = self.index.read_bytes()
        preview = self.cli('tools/library.py', 'sync')
        self.assertEqual(preview.returncode, 0, preview.stdout + preview.stderr)
        self.assertIn('x14', preview.stdout)
        self.assertEqual(self.index.read_bytes(), before)
        result = self.cli('tools/library.py', 'sync', '--apply', '--result-json')
        self.assertNotEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertEqual(self.index.read_bytes(), before)
        self.assertNotIn('Đã cập nhật sổ nguồn', result.stdout)
        self.assertNotIn('Source index updated', result.stdout)
        self.assertIn('x14', result.stderr)
        _, envelope = result_schema.split_output(result.stdout)
        self.assertEqual(envelope['status'], 'blocked')
        self.assertFalse(envelope['changes'])
        self.assertEqual(self.open_transactions(), [])

    def test_a_scan_that_cannot_finish_refuses_apply_instead_of_assuming_no_x14(self):
        import warnings
        before = self.index.read_bytes()
        import types
        import zipfile as _zip
        def cannot_open(*args, **kw):
            raise OSError('boom')
        broken = types.SimpleNamespace(ZipFile=cannot_open, BadZipFile=_zip.BadZipFile)  # only the x14 scan sees this stand-in
        with mock.patch.object(library_sync, 'zipfile', broken):
            with warnings.catch_warnings():
                warnings.simplefilter('ignore')
                preview = library_sync.plan(str(self.project))
            self.assertEqual(preview['notices'], ['x14_scan_incomplete'])
            created = []
            real = transaction.Transaction
            with mock.patch.object(transaction, 'Transaction', lambda *a, **k: created.append(a) or real(*a, **k)):
                with self.assertRaises(library_sync.SyncError):
                    library_sync.apply(str(self.project))
        self.assertEqual(created, [])
        self.assertEqual(self.index.read_bytes(), before)
        self.assertEqual(self.open_transactions(), [])

    def test_plain_workbook_still_applies_and_a_repeat_is_a_no_op(self):
        self.assertEqual(library_sync.plan(str(self.project))['notices'], [], 'a plain workbook raises no notice')
        changed = library_sync.apply(str(self.project))
        self.assertIsNotNone(changed['bytes'])
        self.assertEqual([r[0] for r in self.rows()], ['SRC-000001'])
        after = self.index.read_bytes()
        again = library_sync.apply(str(self.project))
        self.assertIsNone(again['bytes'])
        self.assertEqual(self.index.read_bytes(), after, 'the repeat writes nothing')

    def test_current_x14_workbook_is_rechecked_after_preview_but_a_no_op_still_writes_nothing(self):
        library_sync.apply(str(self.project))
        self.assertIsNone(library_sync.plan(str(self.project))['bytes'], 'the earlier clean preview has no changes')
        self.x14_workbook()  # Excel saves an extension after that preview.
        before = self.index.read_bytes()
        with mock.patch.object(transaction, 'Transaction', side_effect=AssertionError('must not create a transaction')):
            no_op = library_sync.apply(str(self.project))
            self.assertIsNone(no_op['bytes'], 'an x14 workbook needs no rewrite when the catalog is unchanged')
            self.assertEqual(no_op['protection'], 'not_needed')
            (self.ref / 'source-documents' / 'team' / 'two.txt').write_text('two', encoding='utf-8')
            with mock.patch.object(library_sync, 'lossy_notices', wraps=library_sync.lossy_notices) as scan:
                with self.assertRaises(library_sync.SyncError):
                    library_sync.apply(str(self.project))
            scan.assert_called_once_with(before)
        self.assertEqual(self.index.read_bytes(), before)
        self.assertEqual(self.open_transactions(), [])

    def test_direct_cli_scan_failure_is_blocked_before_transaction(self):
        before = self.index.read_bytes()
        # Inject only the scanner's ZipFile failure in a fresh CLI process; openpyxl still reads the real workbook.
        code = '\n'.join((
            'import runpy, sys, types',
            'sys.path.insert(0, "tools")',
            'import library_sync, transaction',
            'def cannot_open(*args, **kw): raise RuntimeError("scanner failed")',
            'def no_transaction(*args, **kw): raise AssertionError("transaction created")',
            'library_sync.zipfile = types.SimpleNamespace(ZipFile=cannot_open, BadZipFile=library_sync.zipfile.BadZipFile)',
            'transaction.Transaction = no_transaction',
            'sys.argv = ["tools/library.py", "sync", "--apply", "--result-json"]',
            'runpy.run_path("tools/library.py", run_name="__main__")',
        ))
        result = self.cli('-B', '-W', 'error::ResourceWarning', '-c', code)
        self.assertNotEqual(result.returncode, 0, result.stdout + result.stderr)
        _, envelope = result_schema.split_output(result.stdout)
        self.assertEqual(envelope['status'], 'blocked', result.stdout + result.stderr)
        self.assertFalse(envelope['changes'])
        self.assertIn('x14', result.stderr)
        self.assertNotIn('Đã cập nhật sổ nguồn', result.stdout)
        self.assertNotIn('Source index updated', result.stdout)
        self.assertEqual(self.index.read_bytes(), before)
        self.assertEqual(self.open_transactions(), [])
        self.assertEqual(sorted(os.listdir(str(self.project / '.agent' / 'runtime'))), ['library'])

    # ---- A-06 corrective: a journal state that cannot be determined must stop the write ----
    def journal_dir(self, project, name, text=None, journal=None):
        """Plant .agent/runtime/transactions/<name>/journal.json with raw text or a JSON value (fake data)."""
        folder = project / '.agent' / 'runtime' / 'transactions' / name
        folder.mkdir(parents=True)
        if text is None and journal is not None:
            text = json.dumps(journal)
        if text is not None:
            (folder / 'journal.json').write_text(text, encoding='utf-8')
        return folder

    def valid_journal(self, status, states=('done',)):
        return {'schema': transaction.SCHEMA, 'id': 'fake', 'operation': 'fixture', 'status': status,
                'ops': [{'path': 'references/x.txt', 'state': s} for s in states]}

    def unknown_journals(self):
        good = self.valid_journal('committed')
        return {
            'invalid_json': dict(text='{not json'),
            'empty_file': dict(text=''),
            'json_list_not_object': dict(text='[]'),
            'wrong_schema': dict(journal=dict(good, schema='someone-else/v9')),
            'unknown_status': dict(journal=dict(good, status='half-done')),
            'status_missing': dict(journal={k: v for k, v in good.items() if k != 'status'}),
            'ops_not_a_list': dict(journal=dict(good, ops='nope')),
            'op_not_an_object': dict(journal=dict(good, ops=['done'])),
            'op_state_unknown': dict(journal=dict(good, ops=[{'path': 'a', 'state': 'maybe'}])),
            'folder_without_journal': dict(),
        }

    def guard_snapshot(self, project):
        return (tree(project / '.agent' / 'runtime' / 'transactions'), (project / 'references' / '_INDEX.xlsx').read_bytes(),
                tree(project / 'references' / 'source-documents'))

    def test_a_journal_that_cannot_be_interpreted_blocks_apply_before_plan_or_transaction(self):
        for number, (name, spec) in enumerate(self.unknown_journals().items()):
            with self.subTest(journal=name):
                project = self.make_project('unknown-%d' % number)
                self.journal_dir(project, 'tx-bad', **spec)
                before = self.guard_snapshot(project)
                self.assertTrue(library_sync.plan(str(project))['changes'], 'the fixture has a real change to write')
                with mock.patch.object(library_sync, 'plan', side_effect=AssertionError('planned past the guard')), \
                        mock.patch.object(transaction, 'Transaction', side_effect=AssertionError('transaction created')):
                    with self.assertRaises(library_sync.SyncError) as caught:
                        library_sync.apply(str(project))
                self.assertIn('Cannot tell whether', str(caught.exception), 'the guard itself must be the cause, not another blocker')
                self.assertEqual(self.guard_snapshot(project), before, 'workbook, sources and old journal keep their bytes')

    def test_simulated_read_and_listing_errors_block_apply_too(self):
        # Mocks: a real unreadable journal or folder listing is hard to make portably, so the errors are injected.
        real_listdir, real_lstat = os.listdir, os.lstat
        def listing_fails(path='.'):
            if str(path).replace('\\', '/').endswith('runtime/transactions'):
                raise PermissionError('simulated listing failure')
            return real_listdir(path)
        def inspection_fails(path, *a, **k):
            if str(path).replace('\\', '/').endswith('runtime/transactions'):
                raise PermissionError('simulated inspection failure')
            return real_lstat(path, *a, **k)
        cases = {
            'journal_read_error': (transaction, '_load', mock.Mock(side_effect=OSError('simulated read failure'))),
            'folder_listing_error': (transaction.os, 'listdir', listing_fails),
            'folder_inspection_error': (transaction.os, 'lstat', inspection_fails),
        }
        for number, (name, (owner, attr, replacement)) in enumerate(cases.items()):
            with self.subTest(simulated=name):
                project = self.make_project('simulated-%d' % number)
                self.journal_dir(project, 'tx-old', journal=self.valid_journal('committed'))
                before = self.guard_snapshot(project)
                with mock.patch.object(library_sync, 'plan', side_effect=AssertionError('planned past the guard')), \
                        mock.patch.object(transaction, 'Transaction', side_effect=AssertionError('transaction created')), \
                        mock.patch.object(owner, attr, replacement):
                    with self.assertRaises(library_sync.SyncError) as caught:
                        library_sync.apply(str(project))
                self.assertIn('Cannot tell whether', str(caught.exception))
                self.assertEqual(self.guard_snapshot(project), before)

    def test_valid_unfinished_journals_still_block_with_their_own_message(self):
        for status in ('prepared', 'committing'):
            with self.subTest(status=status):
                project = self.make_project('open-' + status)
                self.journal_dir(project, 'tx-open', journal=self.valid_journal(status, ('pending', 'done')))
                before = self.guard_snapshot(project)
                with self.assertRaises(library_sync.SyncError) as caught:
                    library_sync.apply(str(project))
                self.assertIn('Unfinished file transaction', str(caught.exception))
                self.assertNotIn('Cannot tell whether', str(caught.exception))
                self.assertEqual(self.guard_snapshot(project), before)

    def test_absent_or_terminal_journals_do_not_block_a_normal_apply(self):
        cases = {'absent': None, 'committed': 'committed', 'rolled_back': 'rolled_back', 'rollback_partial': 'rollback_partial'}
        for name, status in cases.items():
            with self.subTest(journals=name):
                project = self.make_project('terminal-' + name)
                old = None
                if status:
                    old = self.journal_dir(project, 'tx-old', journal=self.valid_journal(status, ('done', 'skipped_conflict')))
                    old_bytes = (old / 'journal.json').read_bytes()
                done = library_sync.apply(str(project))
                self.assertIsNotNone(done['bytes'], 'the change was written')
                self.assertIsNone(library_sync.apply(str(project))['bytes'], 'the repeat is a no-op')
                if old:
                    self.assertEqual((old / 'journal.json').read_bytes(), old_bytes, 'a terminal journal is left alone')
                self.assertEqual(transaction.incomplete(str(project)), [])

    def test_preview_stays_read_only_and_works_with_an_unreadable_journal(self):
        project = self.make_project('preview-bad')
        self.journal_dir(project, 'tx-bad', text='{not json')
        before = self.guard_snapshot(project)
        result = library_sync.plan(str(project))
        self.assertTrue(result['changes'])
        self.assertEqual(self.guard_snapshot(project), before)

    def test_direct_cli_apply_with_an_unreadable_journal_is_blocked_and_changes_nothing(self):
        for name, spec in {'invalid_json': dict(text='{not json'),
                           'unknown_status': dict(journal=dict(self.valid_journal('committed'), status='half-done'))}.items():
            with self.subTest(journal=name):
                project = self.make_project('cli-' + name)
                self.journal_dir(project, 'tx-bad', **spec)
                before = self.guard_snapshot(project)
                self.project = project  # cli() runs inside self.project
                result = self.cli('-B', '-W', 'error::ResourceWarning', 'tools/library.py', 'sync', '--apply', '--result-json')
                self.assertNotEqual(result.returncode, 0, result.stdout + result.stderr)
                _, envelope = result_schema.split_output(result.stdout)
                self.assertEqual(envelope['status'], 'blocked', result.stdout + result.stderr)
                self.assertFalse(envelope['changes'])
                self.assertIn('Cannot tell whether', result.stderr + result.stdout)
                self.assertIn('no new file transaction was created', result.stderr + result.stdout)
                self.assertNotIn('so nothing was written', result.stderr + result.stdout)
                self.assertNotIn('Đã cập nhật sổ nguồn', result.stdout)
                self.assertNotIn('Source index updated', result.stdout)
                self.assertEqual(self.guard_snapshot(project), before)

    def test_apply_names_the_level_of_protection_it_actually_had(self):
        expected = 'held' if os.name == 'nt' else 'recheck_only'
        self.assertEqual(library_sync.apply(str(self.project))['protection'], expected)
        self.assertEqual(library_sync.apply(str(self.project))['protection'], 'not_needed', 'nothing to write, nothing claimed')
        (self.ref / 'source-documents' / 'team' / 'two.txt').write_text('two', encoding='utf-8')
        with mock.patch.object(transaction, 'PIN_ENABLED', False):
            self.assertEqual(library_sync.apply(str(self.project))['protection'], 'recheck_only',
                             'without folder handles the Windows level of protection is never claimed')

    # ---- A-06 corrective: the lock leaf itself (index.lock) must be one plain single-link file before it is byte-locked ----
    def leaf_path(self, project):
        return Path(project) / '.agent' / 'runtime' / 'library' / 'index.lock'

    def lock_folder(self, project):
        folder = self.leaf_path(project).parent
        folder.mkdir(parents=True, exist_ok=True)
        return folder

    def lock_state(self, project):
        """What a refusal must leave alone: workbook, sources, journals, and the lock leaf's own bytes."""
        lock = self.leaf_path(project)
        return (self.guard_snapshot(project), lock.read_bytes() if lock.is_file() else None)

    @contextlib.contextmanager
    def spy_lock_file(self, project, on_open=None):
        """Record os.open flags for index.lock, the descriptors opened and those closed. Delegates every call to the real os."""
        lock = os.path.normcase(os.path.abspath(str(self.leaf_path(project))))
        spy = types.SimpleNamespace(flags=[], opened=[], closed=[])
        real_open, real_close = os.open, os.close

        def watched_open(path, flags, *args, **kw):
            mine = os.path.normcase(os.path.abspath(str(path))) == lock
            if mine:
                spy.flags.append(flags)
                if on_open:
                    on_open(len(spy.flags), flags)
            fd = real_open(path, flags, *args, **kw)
            if mine:
                spy.opened.append(fd)
            return fd

        def watched_close(fd):
            spy.closed.append(fd)
            return real_close(fd)

        with mock.patch.object(os, 'open', watched_open), mock.patch.object(os, 'close', watched_close):
            yield spy

    def refuses_before_the_byte_lock(self, project, expect):
        """apply() must raise SyncError naming expect, without byte-locking, planning or creating a transaction."""
        with mock.patch.object(library_sync, 'take_os_lock', side_effect=AssertionError('byte-locked a refused leaf')), \
                mock.patch.object(library_sync, 'plan', side_effect=AssertionError('planned past the lock guard')), \
                mock.patch.object(transaction, 'Transaction', side_effect=AssertionError('transaction created')):
            with self.assertRaises(library_sync.SyncError) as caught:
                library_sync.apply(str(project))
        self.assertIn('index.lock', str(caught.exception))
        self.assertIn(expect, str(caught.exception))

    def plant_hardlinked_lock(self, project, name):
        outside = self.outside(name)
        sentinel = outside / 'sentinel.txt'
        folder = self.lock_folder(project)
        try:
            os.link(str(sentinel), str(folder / 'index.lock'))
        except (OSError, NotImplementedError, AttributeError) as e:
            self.skipTest('this machine cannot create a hard link: %s' % e)
        return sentinel

    def test_a_hard_link_lock_leaf_to_a_file_outside_is_refused_before_the_byte_lock(self):
        # Real NTFS hard link: the outside file is the same file as index.lock (link count 2). Before this fix it was byte-locked.
        project = self.make_project('hardlink')
        sentinel = self.plant_hardlinked_lock(project, 'outside-hl')
        outside_before = sentinel.read_bytes()
        self.assertEqual(os.lstat(str(self.leaf_path(project))).st_nlink, 2, 'fixture really is a two-name file')
        before = self.lock_state(project)
        seen = []
        real_refuse = library_sync.refuse_lock_leaf
        read_alias = 'import sys; print(open(sys.argv[1], "rb").read(1).hex())'

        def refusal_hook(fd, en, vi):
            # At the moment of refusal another process can still read byte 0 of the outside file through the alias:
            # nothing byte-locked it. (A held byte lock makes that read fail with PermissionError.)
            done = subprocess.run([sys.executable, '-c', read_alias, str(sentinel)], capture_output=True, text=True, timeout=30)
            seen.append((done.returncode, done.stdout.strip(), done.stderr.strip()))
            return real_refuse(fd, en, vi)

        with mock.patch.object(library_sync, 'refuse_lock_leaf', refusal_hook), self.spy_lock_file(project) as spy:
            self.refuses_before_the_byte_lock(project, '2 names')
        self.assertEqual(seen, [(0, outside_before[:1].hex(), '')], 'the alias was readable, so no byte lock was held at refusal')
        self.assertEqual(sentinel.read_bytes(), outside_before, 'the outside file keeps its bytes')
        self.assertEqual(os.lstat(str(self.leaf_path(project))).st_nlink, 2, 'the link is not removed, replaced or repaired')
        self.assertEqual(self.lock_state(project), before)
        self.assertEqual(spy.flags, [], 'a leaf that is already a two-name file is refused on its metadata, before any open')

    def test_direct_cli_apply_with_a_hard_link_lock_leaf_is_blocked_and_changes_nothing(self):
        project = self.make_project('hardlink-cli')
        sentinel = self.plant_hardlinked_lock(project, 'outside-hl-cli')
        outside_before = sentinel.read_bytes()
        before = self.lock_state(project)
        self.project = project  # cli() runs inside self.project
        result = self.cli('-B', '-W', 'error::ResourceWarning', 'tools/library.py', 'sync', '--apply', '--result-json')
        self.assertNotEqual(result.returncode, 0, result.stdout + result.stderr)
        _, envelope = result_schema.split_output(result.stdout)
        self.assertEqual(envelope['status'], 'blocked', result.stdout + result.stderr)
        self.assertFalse(envelope['changes'])
        self.assertIn('index.lock', result.stderr + result.stdout)
        self.assertNotIn('Đã cập nhật sổ nguồn', result.stdout)
        self.assertNotIn('Source index updated', result.stdout)
        self.assertEqual(sentinel.read_bytes(), outside_before)
        self.assertEqual(self.lock_state(project), before)

    def make_file_symlink(self, link, target):
        try:
            os.symlink(str(target), str(link))
        except (OSError, NotImplementedError) as e:
            code = getattr(e, 'winerror', None) or getattr(e, 'errno', None)
            self.skipTest('UNAVAILABLE here, not tested (file symlink cannot be created, error %s): a real file-symlink proof needs '
                          'Developer Mode or the privilege; a hard link is not a substitute' % code)
        self.addCleanup(lambda: os.path.lexists(str(link)) and os.unlink(str(link)))

    def test_a_real_file_symlink_lock_leaf_to_an_existing_outside_file_is_refused(self):
        project = self.make_project('symlink-existing')
        outside = self.outside('outside-sl-existing') / 'sentinel.txt'
        self.make_file_symlink(self.lock_folder(project) / 'index.lock', outside)
        outside_before = outside.read_bytes()
        self.refuses_before_the_byte_lock(project, 'symlink')
        self.assertEqual(outside.read_bytes(), outside_before)

    def test_a_real_file_symlink_lock_leaf_to_an_absent_target_creates_nothing(self):
        project = self.make_project('symlink-absent')
        target = self.outside('outside-sl-absent') / 'not-there.txt'
        self.make_file_symlink(self.lock_folder(project) / 'index.lock', target)
        self.refuses_before_the_byte_lock(project, 'index.lock')
        self.assertFalse(target.exists(), 'O_CREAT must not follow the link and create the outside file')

    def fake_stat(self, **override):
        fields = dict(st_mode=stat.S_IFREG | 0o600, st_nlink=1, st_dev=7, st_ino=11, st_file_attributes=0x20)
        fields.update(override)
        return types.SimpleNamespace(**fields)

    def test_simulated_unusable_leaf_metadata_is_refused_before_the_byte_lock(self):
        # Simulation: os.lstat is replaced for index.lock only. These are not real symlinks, reparse points or devices.
        cases = {
            'symlink_mode': (dict(st_mode=stat.S_IFLNK | 0o777), 'symlink'),
            'reparse_attribute': (dict(st_file_attributes=0x400 | 0x20), 'reparse'),
            'directory': (dict(st_mode=stat.S_IFDIR | 0o700), 'regular'),
            'not_regular': (dict(st_mode=stat.S_IFCHR | 0o600), 'regular'),
            'two_links': (dict(st_nlink=2), '2 names'),
            'zero_links': (dict(st_nlink=0), 'link count'),
            'links_unknown': (dict(st_nlink=None), 'link count'),
            'links_boolean': (dict(st_nlink=True), 'link count'),
            'file_id_zero': (dict(st_ino=0), 'identity'),
            'file_id_missing': (dict(st_ino=None), 'identity'),
            'device_zero': (dict(st_dev=0), 'identity'),
            'device_boolean': (dict(st_dev=True), 'identity'),
        }
        for number, (name, (override, expect)) in enumerate(cases.items()):
            with self.subTest(metadata=name):
                project = self.make_project('meta-%d' % number)
                lock = os.path.normcase(os.path.abspath(str(self.leaf_path(project))))
                real_lstat = os.lstat
                fake = self.fake_stat(**override)

                def lstat(path, *a, **k):
                    if os.path.normcase(os.path.abspath(str(path))) == lock:
                        return fake
                    return real_lstat(path, *a, **k)

                before = self.lock_state(project)
                with mock.patch.object(os, 'lstat', lstat), self.spy_lock_file(project) as spy:
                    self.refuses_before_the_byte_lock(project, expect)
                self.assertEqual(spy.flags, [], 'nothing was opened for a leaf whose metadata was refused')
                self.assertEqual(self.lock_state(project), before)

    def test_simulated_metadata_error_on_the_leaf_is_refused(self):
        project = self.make_project('meta-error')
        lock = os.path.normcase(os.path.abspath(str(self.leaf_path(project))))
        real_lstat = os.lstat

        def lstat(path, *a, **k):
            if os.path.normcase(os.path.abspath(str(path))) == lock:
                raise PermissionError('simulated metadata failure')
            return real_lstat(path, *a, **k)

        with mock.patch.object(os, 'lstat', lstat), self.spy_lock_file(project) as spy:
            self.refuses_before_the_byte_lock(project, 'metadata')
        self.assertEqual(spy.flags, [])

    def test_a_new_lock_is_created_exclusively_and_an_old_one_is_opened_without_create_or_truncate(self):
        project = self.make_project('flags')
        binary = getattr(os, 'O_BINARY', 0)
        with self.spy_lock_file(project) as spy:
            with library_sync.writer_lock(str(project)):
                pass
        self.assertEqual(spy.flags, [os.O_RDWR | os.O_CREAT | os.O_EXCL | binary], 'absent leaf: exclusive create')
        self.assertTrue(set(spy.opened) <= set(spy.closed))
        lock = self.leaf_path(project)
        lock.write_bytes(b'424242 old process id\n')  # a lock file from the previous version: bytes kept, never trusted
        with self.spy_lock_file(project) as spy:
            with library_sync.writer_lock(str(project)):
                pass  # the byte lock covers the first byte, so the bytes are read back after release
        self.assertEqual(spy.flags, [os.O_RDWR | binary], 'existing leaf: no O_CREAT, no O_TRUNC')
        self.assertEqual(lock.read_bytes(), b'424242 old process id\n', 'old bytes survive')
        self.assertEqual(os.lstat(str(lock)).st_nlink, 1)

    def test_controlled_race_a_leaf_that_appears_before_create_is_not_followed_by_a_create(self):
        # Hook, not a kernel race: the leaf is planted just before the exclusive create, which must then fail and re-check.
        project = self.make_project('race-appears')
        lock = self.leaf_path(project)
        binary = getattr(os, 'O_BINARY', 0)
        self.lock_folder(project)

        def plant(number, flags):
            if number == 1:
                lock.write_bytes(b'planted')

        with self.spy_lock_file(project, plant) as spy:
            with library_sync.writer_lock(str(project)):
                pass
        self.assertEqual(spy.flags, [os.O_RDWR | os.O_CREAT | os.O_EXCL | binary, os.O_RDWR | binary],
                         'the failed exclusive create is followed by a checked open without O_CREAT')
        self.assertEqual(lock.read_bytes(), b'planted')

    def test_controlled_race_a_hard_link_that_appears_before_create_is_refused_not_locked(self):
        project = self.make_project('race-appears-link')
        sentinel = self.outside('outside-race') / 'sentinel.txt'
        before_outside = sentinel.read_bytes()
        lock = self.leaf_path(project)
        self.lock_folder(project)

        def plant(number, flags):
            if number == 1:
                try:
                    os.link(str(sentinel), str(lock))
                except (OSError, NotImplementedError) as e:
                    self.skipTest('this machine cannot create a hard link: %s' % e)

        with self.spy_lock_file(project, plant) as spy:
            self.refuses_before_the_byte_lock(project, '2 names')
        self.assertEqual(len(spy.flags), 1, 'only the exclusive create was tried; the re-check refused on metadata, with no second open')
        self.assertTrue(spy.flags[0] & os.O_EXCL)
        self.assertEqual(spy.opened, [], 'the failed create opened nothing')
        self.assertEqual(sentinel.read_bytes(), before_outside)

    def test_controlled_race_a_leaf_swapped_before_the_existing_open_is_refused(self):
        # Hook, not a kernel race: a different file takes the leaf's place between the metadata check and the open.
        project = self.make_project('race-swap')
        lock = self.leaf_path(project)
        with library_sync.writer_lock(str(project)):
            pass
        spare = lock.with_name('spare.tmp')
        spare.write_bytes(b'other file')

        def swap(number, flags):
            if number == 1:
                os.replace(str(spare), str(lock))

        with self.spy_lock_file(project, swap) as spy:
            self.refuses_before_the_byte_lock(project, 'not the one that was opened')
        self.assertTrue(set(spy.opened) <= set(spy.closed))
        self.assertEqual(lock.read_bytes(), b'other file', 'the swapped-in file is left as found')

    def test_controlled_race_a_real_second_name_before_existing_open_is_refused_and_closed(self):
        # Controlled hook with a real hard link: lstat saw one name, but the subsequently opened descriptor has two.
        project = self.make_project('race-existing-alias')
        lock = self.leaf_path(project)
        with library_sync.writer_lock(str(project)):
            pass
        lock.write_bytes(b'old process id\n')
        alias = self.outside('outside-existing-alias') / 'lock-alias.txt'
        before = self.lock_state(project)
        seen = []
        real_refuse = library_sync.refuse_lock_leaf

        def plant(number, flags):
            if number == 1:
                try:
                    os.link(str(lock), str(alias))
                except (OSError, NotImplementedError) as e:
                    self.skipTest('this machine cannot create a hard link: %s' % e)

        def refusal_hook(fd, en, vi):
            self.assertIsNotNone(fd, 'the descriptor, not the earlier one-name metadata, caused this refusal')
            self.assertEqual(os.fstat(fd).st_nlink, 2, 'real descriptor metadata sees the new alias')
            read_alias = 'import sys\nwith open(sys.argv[1], "rb") as f: print(f.read(1).hex())'
            done = subprocess.run([sys.executable, '-B', '-c', read_alias, str(alias)],
                                  capture_output=True, text=True, timeout=30)
            seen.append((done.returncode, done.stdout.strip(), done.stderr.strip()))
            return real_refuse(fd, en, vi)

        with self.spy_lock_file(project, plant) as spy, mock.patch.object(library_sync, 'refuse_lock_leaf', refusal_hook):
            self.refuses_before_the_byte_lock(project, '2 names')
        self.assertEqual(seen, [(0, b'o'.hex(), '')], 'the alias remains readable at refusal before any byte lock')
        self.assertEqual(spy.flags, [os.O_RDWR | getattr(os, 'O_BINARY', 0)])
        self.assertEqual(len(spy.opened), 1)
        self.assertEqual(spy.closed.count(spy.opened[0]), 1, 'the refused descriptor is closed exactly once')
        self.assertEqual(self.lock_state(project), before)
        self.assertEqual(alias.read_bytes(), b'old process id\n')

    def test_simulated_descriptor_checks_are_refused_before_and_after_the_byte_lock(self):
        # Simulation: os.fstat is replaced for the lock descriptor only (a real swap of an open file fails on Windows).
        cases = {
            'descriptor_has_two_links': (dict(st_nlink=2), 'before', '2 names'),
            'descriptor_file_id_zero': (dict(st_ino=0), 'before', 'identity'),
            'descriptor_is_another_file': (dict(st_ino=999999), 'before', 'not the one that was opened'),
            'second_name_appears_after_the_lock': (dict(st_nlink=2), 'after', '2 names'),
            'identity_changes_after_the_lock': (dict(st_ino=999999), 'after', 'not the one that was opened'),
        }
        for number, (name, (override, when, expect)) in enumerate(cases.items()):
            with self.subTest(check=name):
                project = self.make_project('fd-%d' % number)
                with library_sync.writer_lock(str(project)):
                    pass
                state = types.SimpleNamespace(locked=False)
                real_fstat, real_take = os.fstat, library_sync.take_os_lock
                with self.spy_lock_file(project) as spy:
                    def take(fd):
                        real_take(fd)
                        state.locked = True

                    def fstat(fd):
                        st = real_fstat(fd)
                        if fd in spy.opened and (state.locked == (when == 'after')):
                            fields = {k: getattr(st, k) for k in ('st_mode', 'st_nlink', 'st_dev', 'st_ino')}
                            fields['st_file_attributes'] = getattr(st, 'st_file_attributes', 0)
                            fields.update(override)
                            return types.SimpleNamespace(**fields)
                        return st

                    with mock.patch.object(library_sync, 'take_os_lock', take), mock.patch.object(os, 'fstat', fstat), \
                            mock.patch.object(library_sync, 'plan', side_effect=AssertionError('planned past the lock guard')):
                        with self.assertRaises(library_sync.SyncError) as caught:
                            library_sync.apply(str(project))
                self.assertIn(expect, str(caught.exception))
                self.assertEqual(state.locked, when == 'after', 'refused before the byte lock, or after it with the lock released')
                self.assertTrue(set(spy.opened) <= set(spy.closed))
                self.assertEqual(spy.closed.count(spy.opened[0]), 1, 'the descriptor is closed exactly once')
                with library_sync.writer_lock(str(project)):
                    pass  # not left locked

    # ---- T16: interruption, resume, rollback ----
    def cli(self, *args, fault=None):
        tools = self.project / 'tools'
        if not tools.exists():
            tools.mkdir()
            for name in ('library.py', 'library_sync.py', 'artifacts.py', 'transaction.py', 'result_schema.py'):
                shutil.copyfile(ROOT / 'tools' / name, tools / name)
        env = dict(os.environ, PYTHONUTF8='1', PYTHONDONTWRITEBYTECODE='1')
        if fault:
            env[transaction.FAULT_ENV] = fault
        return subprocess.run([sys.executable] + list(args), cwd=self.project, env=env, capture_output=True,
                              text=True, encoding='utf-8', timeout=60)

    def kill_during_write(self):
        result = self.cli('tools/library.py', 'sync', '--apply', fault='kill:1')
        self.assertEqual(result.returncode, 70, result.stdout + result.stderr)
        return only_open(self.open_transactions())['id']

    def test_killed_write_can_be_resumed_and_the_next_sync_is_not_blocked_by_the_dead_writers_lock(self):
        txid = self.kill_during_write()
        self.assertTrue((self.project / '.agent' / 'runtime' / 'library' / 'index.lock').exists(), 'the lock file is never deleted')
        self.assertEqual([r[0] for r in self.rows()], ['SRC-000001'], 'the workbook was replaced before the kill and reads back')
        result = transaction.resume(str(self.project), txid)
        self.assertEqual(result['code'], 0)
        self.assertEqual(self.open_transactions(), [])
        again = self.cli('tools/library.py', 'sync', '--apply')
        self.assertEqual(again.returncode, 0, again.stdout + again.stderr)
        self.assertEqual([r[0] for r in self.rows()], ['SRC-000001'], 'no duplicate row')

    def test_killed_write_can_be_rolled_back_to_the_exact_original(self):
        txid = self.kill_during_write()
        result = transaction.rollback(str(self.project), txid)
        self.assertEqual(result['code'], 0)
        self.assertEqual(self.index.read_bytes(), self.original, 'byte-identical original restored')
        self.assertEqual(self.rows(), [])
        again = self.cli('tools/library.py', 'sync', '--apply')
        self.assertEqual(again.returncode, 0, again.stdout + again.stderr)
        self.assertEqual([r[0] for r in self.rows()], ['SRC-000001'])

    def test_rollback_after_a_kill_keeps_what_the_user_saved_in_between(self):
        txid = self.kill_during_write()
        wb = openpyxl.load_workbook(self.index)
        wb['Danh mục'].cell(6, 1, 'PRIVATE-7')
        wb['Danh mục'].cell(6, 9, 'Sửa sau sự cố')
        wb.save(self.index)
        wb.close()
        saved = self.index.read_bytes()
        result = transaction.rollback(str(self.project), txid)
        self.assertEqual(result['code'], 1)
        self.assertEqual(self.index.read_bytes(), saved, 'the user edit is not overwritten by the restore')
        self.assertIn(('PRIVATE-7', 'Sửa sau sự cố'), [r[:2] for r in self.rows()])
        self.assertTrue(any((self.project / '.agent' / 'runtime' / 'transactions' / txid / 'backup').rglob('*.xlsx')),
                        'the original stays in backup for manual recovery')

    def test_faults_before_the_write_leave_nothing_open_and_do_not_block(self):
        for fault in ('stage:1', 'backup:1'):
            with self.subTest(fault=fault):
                with mock.patch.dict(os.environ, {transaction.FAULT_ENV: fault}):
                    with self.assertRaises(library_sync.SyncError):
                        library_sync.apply(str(self.project))
                self.assertEqual(self.index.read_bytes(), self.original)
                self.assertEqual(self.open_transactions(), [])
        library_sync.apply(str(self.project))
        self.assertEqual([r[0] for r in self.rows()], ['SRC-000001'])

    def lock_path(self):
        return self.project / '.agent' / 'runtime' / 'library' / 'index.lock'

    def test_a_held_lock_is_never_taken_removed_or_rewritten_by_another_writer(self):
        with library_sync.writer_lock(str(self.project)):
            for attempt in range(3):
                with self.assertRaises(library_sync.SyncError):
                    with library_sync.writer_lock(str(self.project)):
                        self.fail('a second writer got in')
            with self.assertRaises(library_sync.SyncError):
                library_sync.apply(str(self.project))
            self.assertTrue(self.lock_path().exists(), 'the failed attempts did not remove the lock')
        with library_sync.writer_lock(str(self.project)):
            pass  # released cleanly, no error from either side

    def test_two_writers_that_both_find_an_old_lock_file_let_exactly_one_in(self):
        """The reviewer's race (both see a dead owner, one takes over, the other deletes the new lock) has no step left to race on."""
        import threading
        import time
        self.lock_path().parent.mkdir(parents=True)
        self.lock_path().write_text('999999', encoding='ascii')  # what an older version left behind: a dead process id
        for round_number in range(15):
            barrier = threading.Barrier(2)
            outcome = []

            def writer(name):
                barrier.wait()
                try:
                    with library_sync.writer_lock(str(self.project)):
                        outcome.append((name, 'in'))
                        time.sleep(0.3)
                except library_sync.SyncError:
                    outcome.append((name, 'busy'))
                except Exception as e:  # noqa: BLE001 - a cleanup error would show up here
                    outcome.append((name, repr(e)))
            threads = [threading.Thread(target=writer, args=(n,)) for n in ('A', 'B')]
            [x.start() for x in threads]
            [x.join() for x in threads]
            self.assertEqual(sorted(o[1] for o in outcome), ['busy', 'in'], (round_number, outcome))
        self.assertTrue(self.lock_path().exists())

    def test_lock_of_a_live_process_is_respected_and_a_killed_one_is_released(self):
        import time
        code = ('import sys, time\nsys.path.insert(0, %r)\nimport library_sync\n'
                'with library_sync.writer_lock(sys.argv[1]):\n    print("held", flush=True)\n    time.sleep(60)\n') % str(ROOT / 'tools')
        holder = subprocess.Popen([sys.executable, '-c', code, str(self.project)], stdout=subprocess.PIPE, text=True)
        self.addCleanup(holder.stdout.close)
        self.addCleanup(lambda: holder.poll() is None and holder.kill())
        self.assertEqual(holder.stdout.readline().strip(), 'held')
        with self.assertRaises(library_sync.SyncError):
            library_sync.apply(str(self.project))
        self.assertEqual(self.index.read_bytes(), self.original, 'a live writer keeps the index to itself')
        holder.kill()
        holder.wait()
        started = time.time()
        while True:  # the operating system releases the lock when the process is gone; allow it a moment
            try:
                library_sync.apply(str(self.project))
                break
            except library_sync.SyncError:
                self.assertLess(time.time() - started, 10, 'the lock of a killed process was never released')
                time.sleep(0.1)
        self.assertEqual([r[0] for r in self.rows()], ['SRC-000001'])

    def test_the_process_id_in_an_old_lock_file_is_never_trusted(self):
        """A live-looking id in the file means nothing: only the lock the operating system holds counts (a reused id cannot block or free it)."""
        self.lock_path().parent.mkdir(parents=True)
        self.lock_path().write_text(str(os.getpid()), encoding='ascii')
        library_sync.apply(str(self.project))
        self.assertEqual([r[0] for r in self.rows()], ['SRC-000001'])


@unittest.skipUnless(openpyxl, 'openpyxl is needed for the source workbook')
class LibraryPreviewGuard(unittest.TestCase):
    """A-07: the optional expected_preview of apply(). None keeps the CLI behaviour; anything else must match the plan recomputed
    under the lock, before any Transaction. It is a cooperative precondition, not a grant or a receipt. Temporary projects only."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='law-preview-guard-')
        self.addCleanup(self.tmp.cleanup)
        self.project = Path(self.tmp.name)
        self.ref = self.project / 'references'
        (self.ref / 'source-documents' / 'team').mkdir(parents=True)
        shutil.copyfile(TEMPLATE_INDEX, self.ref / '_INDEX.xlsx')
        self.index = self.ref / '_INDEX.xlsx'
        self.source = self.ref / 'source-documents' / 'team' / 'a.txt'
        self.source.write_text('first', encoding='utf-8')

    def digest(self):
        return hashlib.sha256(self.index.read_bytes()).hexdigest()

    def binding(self):
        return library_sync.preview_binding(library_sync.plan(str(self.project)))

    def refuse_before_transaction(self, expected, reason):
        before = self.digest()
        with mock.patch.object(transaction, 'Transaction', side_effect=AssertionError('stopped before any Transaction')):
            with self.assertRaises(library_sync.PreviewMismatch) as caught:
                library_sync.apply(str(self.project), (), expected)
        self.assertEqual(caught.exception.reason, reason)
        self.assertIsInstance(caught.exception, library_sync.SyncError)
        if reason != 'invalid':
            self.assertIn('LAW did not change the index or source files', str(caught.exception))
            self.assertIn('no new file transaction was created', str(caught.exception))
            self.assertNotIn('nothing was written', str(caught.exception))
        self.assertEqual(self.digest(), before)

    def test_binding_is_json_stable_and_carries_no_bytes(self):
        binding = self.binding()
        self.assertEqual(set(binding), {'snapshot', 'sources', 'scopes'})
        self.assertEqual(json.loads(json.dumps(binding)), binding)
        self.assertEqual(binding['snapshot'], self.digest())
        self.assertNotIn('bytes', json.dumps(binding))

    def test_a_matching_preview_applies_through_the_same_writer_and_none_keeps_the_cli_behaviour(self):
        result = library_sync.apply(str(self.project), (), self.binding())
        self.assertIsNotNone(result['bytes'])
        self.assertEqual(result['protection'] in ('held', 'recheck_only'), True)
        self.source.write_text('second', encoding='utf-8')
        again = library_sync.apply(str(self.project))  # không truyền gì: như CLI hiện hành / nothing passed: as the CLI does today
        self.assertIn('changed', [c['action'] for c in again['changes']])

    def test_malformed_previews_are_refused_before_the_lock_or_any_read(self):
        binding = self.binding()
        bad = [{}, [], 'x', 0, False, dict(binding, extra=1), {'snapshot': binding['snapshot']}, dict(binding, snapshot='short'),
               dict(binding, snapshot=None), dict(binding, sources=None), dict(binding, sources={'a.txt': 'zz'}), dict(binding, sources={1: '0' * 64}),
               dict(binding, scopes=('Cu',)), dict(binding, scopes=[None]), dict(binding, snapshot=binding['snapshot'].upper())]
        for value in bad:
            with self.subTest(value=str(value)[:60]):
                self.refuse_before_transaction(value, 'invalid')
        self.assertFalse((self.project / '.agent').exists())  # chưa có khóa hay nhật ký / no lock, no journal

    def test_each_kind_of_drift_is_named_and_refused_even_when_the_new_plan_is_a_no_op(self):
        binding = self.binding()
        self.source.write_text('changed after the preview', encoding='utf-8')
        self.refuse_before_transaction(binding, 'sources')
        self.source.write_text('first', encoding='utf-8')
        self.index.write_bytes(self.index.read_bytes() + b'edited by someone')
        self.refuse_before_transaction(binding, 'snapshot')
        shutil.copyfile(TEMPLATE_INDEX, self.index)
        (self.project / 'Tai-lieu-cu').mkdir()
        (self.project / 'PROJECT.md').write_text('# PROJECT\n\n' + SCOPE_LINE % 'Tai-lieu-cu', encoding='utf-8')
        self.refuse_before_transaction(binding, 'scopes')
        (self.project / 'PROJECT.md').unlink()
        library_sync.apply(str(self.project))  # một người ghi khác áp dụng đúng thay đổi đó: plan mới là no-op / another writer applied it: the new plan is a no-op
        self.assertIsNone(library_sync.plan(str(self.project))['bytes'])
        self.refuse_before_transaction(binding, 'snapshot')

    def test_the_guard_runs_before_the_x14_and_journal_refusals_and_those_still_hold(self):
        binding = self.binding()
        before = self.digest()
        with mock.patch.object(library_sync, 'lossy_notices', return_value=['x14_data_validation']), \
                mock.patch.object(transaction, 'Transaction', side_effect=AssertionError('no Transaction')):  # mô phỏng / simulated
            with self.assertRaisesRegex(library_sync.SyncError, 'x14'):
                library_sync.apply(str(self.project), (), self.binding())
        with mock.patch.object(transaction, 'incomplete', side_effect=ValueError('bad journal')), \
                mock.patch.object(transaction, 'Transaction', side_effect=AssertionError('no Transaction')):  # mô phỏng / simulated
            with self.assertRaisesRegex(library_sync.SyncError, 'journal'):
                library_sync.apply(str(self.project), (), binding)
        self.assertEqual(self.digest(), before)

    def test_cli_sync_never_sends_an_expected_preview(self):
        seen = {}
        real = library_sync.apply

        def spy(root, renames=(), expected_preview=None):
            seen['expected'] = expected_preview
            return real(root, renames, expected_preview)
        import library
        with mock.patch.object(library_sync, 'apply', spy), mock.patch.object(library, 'ROOT', str(self.project)), \
                contextlib.redirect_stdout(types.SimpleNamespace(write=lambda s: len(s), flush=lambda: None)):
            library.cmd_sync({'apply': True, 'json': True, 'rename': []})
        self.assertIsNone(seen['expected'])


if __name__ == '__main__':
    unittest.main()
