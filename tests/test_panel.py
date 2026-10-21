"""Panel checks: operations history, status data, transaction recovery, diagnostics bundle, merge comparison.

Every test works in temporary folders. The panel's ROOT is patched to a temporary project and its TOOLS to a
temporary copy of the package, so the real repository (and its .agent/runtime) is never written.
"""
import contextlib
import getpass
import hashlib
import io
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from unittest.mock import patch
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'tools'))
import app
import library
import library_sync
import merge_center
import ops_log
import result_schema
import transaction as tx

try:
    import openpyxl
except ImportError:
    openpyxl = None

FAULT = tx.FAULT_ENV
RUNTIME = '.agent/runtime/'
PKG_IGNORE = shutil.ignore_patterns('.git', '__pycache__', 'dist', 'roadmap', 'tests', 'feedback')
ALLOWED_RECORD_FIELDS = {'time', 'action', 'status', 'code', 'files', 'transaction', 'duration_ms'}


def write(root, rel, text):
    p = Path(root) / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    if isinstance(text, bytes):
        p.write_bytes(text)
    else:
        p.write_text(text, encoding='utf-8')
    return p


def snapshot(root):
    """Mọi file trừ dữ liệu giao dịch / every file except transaction runtime data."""
    out = {}
    root = Path(root)
    for p in root.rglob('*'):
        rel = p.relative_to(root).as_posix()
        if p.is_file() and '__pycache__' not in p.parts and not rel.startswith(RUNTIME):
            out[rel] = p.read_bytes()
    return out


def manifest(root):
    lines = []
    for p in sorted(Path(root).rglob('*')):
        if p.is_file() and p.name != 'MANIFEST.sha256':
            lines.append(hashlib.sha256(p.read_bytes().replace(b'\r\n', b'\n')).hexdigest() + '  ' + p.relative_to(root).as_posix())
    write(root, 'MANIFEST.sha256', '\n'.join(lines) + '\n')


PKG_TMP = None
PKG = None  # temporary copy of the package; the tools run from here, so the real repository is not touched


def setUpModule():
    global PKG_TMP, PKG
    PKG_TMP = tempfile.TemporaryDirectory(prefix='panel-pkg-')
    PKG = Path(PKG_TMP.name) / 'package'
    shutil.copytree(ROOT, PKG, ignore=PKG_IGNORE)
    shutil.rmtree(str(PKG / '.agent' / 'runtime'), ignore_errors=True)  # dữ liệu vận hành của máy phát triển không vào bản sao / dev-machine runtime data never enters the copy
    [f.unlink() for f in (PKG / 'assets').rglob('*') if f.is_file() and f.name != 'README.md']  # logo, font của máy phát triển / dev-machine branding and fonts


def tearDownModule():
    PKG_TMP.cleanup()


def run_script(script, *args, env=None):
    e = dict(os.environ, PYTHONUTF8='1', PYTHONDONTWRITEBYTECODE='1')
    e.pop(FAULT, None)
    e.update(env or {})
    return subprocess.run([sys.executable, str(PKG / 'tools' / script), *map(str, args)],
                          cwd=str(PKG), env=e, capture_output=True, encoding='utf-8', timeout=300)


class PanelBase(unittest.TestCase):
    """Temp project + patched ROOT and TOOLS (the temp package copy)."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='panel-check-')
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.project = self.root / 'project'
        self.project.mkdir()
        self.patch(app, 'ROOT', str(self.project))
        self.patch(app, 'TOOLS', str(PKG / 'tools'))

    def patch(self, obj, name, value):
        p = patch.object(obj, name, value)
        p.start()
        self.addCleanup(p.stop)

    def upgrade_pair(self, name):
        """Gói cũ và mới; năm thao tác khi nâng cấp / old and new packages; five ops on upgrade."""
        old, new = self.root / (name + '-old'), self.root / (name + '-new')
        for r in (old, new):
            write(r, 'VERSION', '1.6.0\n')
            write(r, 'untouched.txt', 'old package')
            write(r, 'conflict.txt', 'old package')
            write(r, 'kept.txt', 'old package')
            write(r, 'PROJECT.md', 'private context')
            manifest(r)
        write(old, '.agent/PACKAGE_INFO.json', json.dumps({'package': 'local-agent-workspace', 'version': '1.6.0', 'language': 'vi'}))
        write(old, 'conflict.txt', 'user change')
        write(old, 'kept.txt', 'user change')
        write(new, 'untouched.txt', 'new package')
        write(new, 'conflict.txt', 'new package')
        write(new, 'added.txt', 'new file')
        manifest(new)
        return old, new

    def interrupt_upgrade(self, old, new, fault='commit:1'):
        """Nâng cấp bị dừng giữa chừng bằng biến kiểm thử / an upgrade interrupted by the test-only fault variable."""
        r = run_script('upgrade_package.py', '--project', old, '--from', new, '--apply', env={FAULT: fault})
        self.assertEqual(r.returncode, 1, r.stdout + r.stderr)
        items = tx.incomplete(str(old))
        self.assertEqual(len(items), 1, items)
        return items[0]['id']


class OperationsHistory(PanelBase):
    def test_records_hold_only_allowed_fields_and_no_secrets(self):
        user = getpass.getuser()
        secret_project = 'Dự án BÍ MẬT 4471'
        secret_file = 'SECRET-PERSON Nguyen Van An.docx.new'
        secret_query = 'NGUYEN-VAN-AN-QUERY-7733'
        app.run_api({'action': 'merge_state', 'params': {'path': secret_file, 'state': 'ready'}})
        app.run_api({'action': 'library_find', 'params': {'query': secret_query}})  # A-07: reading the library is not logged
        app.run_api({'action': 'init', 'params': {'name': secret_project, 'lang': 'vi'}})
        app.run_api({'action': 'check_output', 'params': {'files': [str(self.project / secret_file)]}})
        app.run_api({'action': 'tx_resume', 'params': {'txid': 'bad id!'}})
        text = (self.project / '.agent' / 'runtime' / 'operations.jsonl').read_text(encoding='utf-8')
        self.assertEqual(len(text.splitlines()), 4)
        self.assertNotIn('library_find', text)
        for line in text.splitlines():
            rec = json.loads(line)
            self.assertEqual(set(rec), ALLOWED_RECORD_FIELDS)
        for forbidden in (secret_project, secret_file, 'SECRET-PERSON', secret_query, 'Nguyen Van An',
                          str(self.project), str(self.project).replace('\\', '/'), os.path.expanduser('~'),
                          app.TOKEN, user):
            if len(forbidden) >= 2:
                self.assertNotIn(forbidden.lower(), text.lower())
        self.assertIn('"status": "blocked"', text)

    def test_unknown_action_name_is_recorded_as_unknown(self):
        app.run_api({'action': 'my_private_action_name', 'params': {}})
        rec = ops_log.read(str(self.project))[-1]
        self.assertEqual(rec['action'], 'unknown')
        self.assertEqual(rec['status'], 'blocked')
        self.assertNotIn('my_private', (self.project / '.agent' / 'runtime' / 'operations.jsonl').read_text(encoding='utf-8'))

    def test_cap_keeps_only_newest_records(self):
        extra = 25
        for i in range(ops_log.MAX_RECORDS + extra):
            ops_log.record(str(self.project), 'check_output', 'success', 0, i)
        records = ops_log.read(str(self.project))
        self.assertEqual(len(records), ops_log.MAX_RECORDS)
        self.assertEqual(records[0]['files'], extra)
        self.assertEqual(records[-1]['files'], ops_log.MAX_RECORDS + extra - 1)

    def test_clear_removes_only_the_history_file(self):
        ops_log.record(str(self.project), 'check_output', 'success', 0)
        keep = {
            '.agent/runtime/transactions/t1/journal.json': '{"schema": "x"}',
            '.agent/STATE.md': 'state text',
            'output/report.txt': 'document text',
        }
        for rel, text in keep.items():
            write(self.project, rel, text)
        before = {rel: (self.project / rel).read_bytes() for rel in keep}
        res = app.action('history_clear', {})
        self.assertEqual(res['result']['status'], 'success')
        self.assertFalse((self.project / '.agent' / 'runtime' / 'operations.jsonl').exists())
        for rel, data in before.items():
            self.assertEqual((self.project / rel).read_bytes(), data)
        self.assertEqual(ops_log.read(str(self.project)), [])

    def test_status_reports_history_and_latest_outcomes(self):
        ops_log.record(str(self.project), 'check_output', 'success', 0)
        ops_log.record(str(self.project), 'upgrade', 'partial', 1, 2, 'tid-1')
        ops_log.record(str(self.project), 'check_package', 'error', -1)
        with patch.object(app, 'pack_info', return_value={'items': []}):
            status = app.status()
        self.assertEqual([h['action'] for h in status['history']], ['check_package', 'upgrade', 'check_output'])
        summary = status['history_summary']
        self.assertEqual(summary['last_success']['action'], 'check_output')
        self.assertEqual(summary['last_failure']['action'], 'check_package')
        self.assertEqual(summary['last_unfinished']['transaction'], 'tid-1')


class StatusData(PanelBase):
    def test_status_keeps_old_fields_and_adds_new_ones(self):
        with patch.object(app, 'pack_info', return_value={'items': []}):
            status = app.status()
        for key in ('version', 'initialized', 'workspace_kind', 'info', 'project_name', 'folder', 'root', 'lang', 'state',
                    'merges', 'skills', 'packs', 'counts', 'is_source', 'release_source', 'template_folder', 'dests',
                    'picker', 'python', 'platform', 'transactions', 'components', 'history', 'history_summary'):
            self.assertIn(key, status)
        self.assertEqual(status['transactions'], [])
        self.assertEqual(status['history'], [])
        self.assertEqual({c['feature'] for c in status['components']}, {'python', 'library', 'xlsx', 'picker', 'controlled'})
        self.assertTrue(all(isinstance(c['usable'], bool) for c in status['components']))

    def test_status_lists_sources_results_and_purpose_newest_first(self):
        write(self.project, 'references/source-documents/a.txt', 'a')
        write(self.project, 'references/README.md', 'scaffold')
        write(self.project, 'output/old.md', 'old')
        write(self.project, 'output/new.md', 'new')
        os.utime(str(self.project / 'output' / 'old.md'), (1700000000, 1700000000))
        write(self.project, 'PROJECT.md', '- Tên dự án: X\n- Mục đích và người đọc: Biên bản cho tổ\n- Đầu ra mong đợi: [Tài liệu]\n')
        with patch.object(app, 'pack_info', return_value={'items': []}):
            status = app.status()
        self.assertEqual([f['path'] for f in status['sources']['items']], ['references/source-documents/a.txt'])
        self.assertEqual([f['name'] for f in status['outputs']['items']], ['new.md', 'old.md'])
        self.assertEqual((status['purpose'], status['expected']), ('Biên bản cho tổ', None))
        self.assertEqual(status['parent'], os.path.dirname(app.ROOT))

    def test_project_open_accepts_only_another_initialised_project(self):
        for bad in ('', str(self.project / 'nope'), str(app.ROOT)):
            with self.assertRaises(app.BadInput):
                app.open_project(bad, 'folder')

    def test_missing_optional_components_are_reported_not_usable(self):
        with patch.object(app, '_has_module', return_value=False), patch.object(app, 'tk_available', return_value=False):
            by_feature = {c['feature']: c for c in app.components()}
        self.assertFalse(by_feature['library']['usable'])
        self.assertFalse(by_feature['xlsx']['usable'])
        self.assertFalse(by_feature['picker']['usable'])
        self.assertEqual(by_feature['library']['needs'], 'openpyxl')

    def test_the_controlled_flow_readiness_is_its_own_component_and_never_hides_the_others(self):
        unavailable = {'usable': False, 'reason': 'DEPENDENCY_MISSING', 'python': sys.executable, 'version': '3.11.9', 'jsonschema': None, 'fs': None}
        import current_draft
        with patch.object(current_draft, 'readiness', return_value=unavailable):
            by_feature = {c['feature']: c for c in app.components()}
        self.assertEqual((by_feature['controlled']['usable'], by_feature['controlled']['reason']), (False, 'DEPENDENCY_MISSING'))
        self.assertEqual(by_feature['controlled']['python_path'], sys.executable)  # the exact interpreter the panel runs on
        self.assertTrue(by_feature['python']['usable'])
        with patch.object(current_draft, 'readiness', side_effect=RuntimeError('boom')):
            broken = {c['feature']: c for c in app.components()}['controlled']
        self.assertEqual((broken['usable'], broken['reason']), (False, 'READINESS_UNKNOWN'))  # unmeasured is unusable, never guessed

    def test_status_lists_an_interrupted_transaction(self):
        old, new = self.upgrade_pair('status')
        txid = self.interrupt_upgrade(old, new)
        with patch.object(app, 'ROOT', str(old)), patch.object(app, 'pack_info', return_value={'items': []}):
            status = app.status()
        self.assertEqual(len(status['transactions']), 1)
        item = status['transactions'][0]
        self.assertEqual((item['id'], item['operation'], item['status'], item['done']), (txid, 'upgrade', 'committing', 1))

    def test_status_branding_from_assets(self):
        work = Path(tempfile.mkdtemp(prefix='test-brand-'))
        try:
            write(work, 'VERSION', '1.8.0\n')
            write(work, 'PROJECT.md', '- Tên dự án: Demo\n')
            write(work, 'assets/branding.json', json.dumps({'title': 'Đơn vị Mẫu', 'subtitle': 'Agent Workspace', 'logo': 'logo.png'}))
            write(work, 'assets/logo.png', b'\x89PNG\r\n\x1a\nfakeimage')
            with patch.object(app, 'ROOT', str(work)), patch.object(app, 'pack_info', return_value={'items': []}):
                st = app.status()
            self.assertIsNotNone(st['branding'])
            self.assertEqual(st['branding']['title'], 'Đơn vị Mẫu')
            self.assertEqual(st['branding']['subtitle'], 'Agent Workspace')
            self.assertTrue(st['branding']['logo'].startswith('data:image/png;base64,'))
        finally:
            shutil.rmtree(work, ignore_errors=True)


class TransactionRecovery(PanelBase):
    def test_bad_transaction_ids_are_rejected_before_any_tool_runs(self):
        bad = ['../x', 'a/b', 'a.b', '', '   ', 'a b', 'x;rm', 'abc\n', 'a' * 81, '..\\x']
        for txid in bad:
            for act in ('tx_resume', 'tx_rollback'):
                with self.subTest(txid=txid, action=act):
                    with self.assertRaises(app.BadInput):
                        app.action(act, {'txid': txid})
                    res = app.run_api({'action': act, 'params': {'txid': txid}})
                    self.assertEqual(res[0], 200)
                    self.assertEqual(res[1]['result']['status'], 'blocked')
                    self.assertTrue(res[1]['input_error'])

    def test_resume_completes_interrupted_upgrade_like_a_clean_apply(self):
        clean_old, clean_new = self.upgrade_pair('clean')
        r = run_script('upgrade_package.py', '--project', clean_old, '--from', clean_new, '--apply')
        self.assertEqual(r.returncode, 0, r.stderr)
        clean = snapshot(clean_old)
        old, new = self.upgrade_pair('broken')
        txid = self.interrupt_upgrade(old, new)
        with patch.object(app, 'ROOT', str(old)):
            res = app.action('tx_resume', {'txid': txid, 'lang': 'en'})
        self.assertEqual(res['code'], 0, res['out'])
        self.assertEqual(res['result']['code'], res['code'])
        self.assertEqual(res['result']['status'], 'success')
        self.assertEqual(res['result']['transaction'], txid)
        self.assertEqual(snapshot(old), clean)
        self.assertEqual(tx.incomplete(str(old)), [])

    def test_rollback_restores_the_original_files(self):
        old, new = self.upgrade_pair('rb')
        before = snapshot(old)
        txid = self.interrupt_upgrade(old, new, 'commit:3')
        with patch.object(app, 'ROOT', str(old)):
            res = app.action('tx_rollback', {'txid': txid})
        self.assertEqual((res['code'], res['result']['status'], res['result']['code']), (0, 'success', 0))
        self.assertEqual(snapshot(old), before)
        self.assertEqual(tx.incomplete(str(old)), [])

    def test_rollback_keeps_a_user_edit_and_reports_next_steps(self):
        old, new = self.upgrade_pair('edit')
        txid = self.interrupt_upgrade(old, new, 'commit:3')
        write(old, 'untouched.txt', 'USER EDIT')
        with patch.object(app, 'ROOT', str(old)):
            res = app.run_api({'action': 'tx_rollback', 'params': {'txid': txid}})
            payload = res[1]
        self.assertEqual(payload['result']['status'], 'partial')
        self.assertEqual(payload['code'], 1)
        self.assertEqual(payload['result']['code'], payload['code'])
        self.assertEqual(payload['result']['transaction'], txid)
        self.assertTrue(payload['result']['next_actions'])
        self.assertEqual((old / 'untouched.txt').read_text(encoding='utf-8'), 'USER EDIT')
        history = ops_log.read(str(old))
        self.assertEqual((history[-1]['status'], history[-1]['transaction']), ('partial', txid))


class Diagnostics(PanelBase):
    def test_bundle_contains_no_paths_user_name_or_document_text(self):
        old, new = self.upgrade_pair('diag')
        write(old, 'output/secret-report.txt', 'CONFIDENTIAL BODY TEXT 8842 Nguyen Van An')
        write(old, 'references/legal/law.txt', 'LEGAL BODY TEXT 5519')
        write(old, 'skills/local/packs/demo/SKILL.md.new', 'pack body')
        write(old, 'SECRET-PERSON-NAME.new', 'named file')
        txid = self.interrupt_upgrade(old, new)
        with patch.object(app, 'ROOT', str(old)), patch.object(app, 'pack_info', return_value={'items': []}):
            res = app.action('diagnostics', {})
            self.assertEqual(res['result']['status'], 'success')
            rel = res['result']['files'][0]['path']
            self.assertTrue(rel.startswith('.agent/runtime/diagnostics/') and rel.endswith('.zip'), rel)
            archive = old / rel
            self.assertTrue(archive.is_file())
            with zipfile.ZipFile(str(archive)) as z:
                self.assertEqual(z.namelist(), ['diagnostics.json'])
                text = z.read('diagnostics.json').decode('utf-8')
        data = json.loads(text)
        self.assertEqual(data['transactions'][0]['id'], txid)
        pendings = [item['pending'] for item in data['merges']['items']]
        self.assertIn('skills/local/packs/demo/SKILL.md.new', pendings)
        self.assertNotIn('SECRET-PERSON-NAME.new', pendings)  # not package-managed: not listed at all
        self.assertEqual(data['merges']['count'], len(merge_center.scan(str(old))))
        for forbidden in ('CONFIDENTIAL BODY TEXT', 'LEGAL BODY TEXT', 'secret-report', 'Nguyen Van An', 'SECRET-PERSON-NAME',
                          'pack body', 'named file'):
            self.assertNotIn(forbidden, text)
        home = os.path.expanduser('~')
        user = getpass.getuser()
        for forbidden in (str(old), str(old).replace('\\', '/'), home, home.replace('\\', '/'), user):
            if len(forbidden) >= 2:
                self.assertNotIn(forbidden.lower(), text.lower())
        self.assertIsNone(re.search(r'[A-Za-z]:[\\/]', text), 'absolute Windows path found')
        self.assertNotIn('/Users/', text)
        self.assertNotIn('/home/', text)

    def test_bundle_exports_only_safe_component_metadata_and_keeps_local_paths(self):
        import current_draft
        paths = (r'C:\HostedToolCache\windows\Python\3.12.10\x64\python.exe',
                 'C:/SharedTools/Python/python.exe',
                 r'C:\\HostedToolCache\\windows\\Python\\python.exe',
                 '/opt/hostedtoolcache/Python/3.12.10/x64/bin/python',
                 str(self.project / 'python.exe'), os.path.join(os.path.expanduser('~'), 'python.exe'))
        readiness = {'usable': True, 'reason': None, 'jsonschema': '4.23.0', 'fs': 'NTFS'}
        with patch.object(current_draft, 'readiness', return_value=readiness), \
                patch.object(app, 'run_tool', return_value={'result': {}}):
            for path in paths:
                with self.subTest(path=path), patch.object(app.sys, 'executable', path):
                    local = app.components()
                    controlled = next(row for row in local if row['feature'] == 'controlled')
                    self.assertEqual(controlled['python_path'], path)
                    # Future metadata must not add paths, nested payloads, or mistyped values to exports.
                    controlled.update(executable=path, install_root=path, details={'path': path})
                    rows = local + [{'feature': path, 'needs': [path], 'version': {'path': path},
                                     'reason': path, 'jsonschema': path, 'fs': path, 'usable': path}]
                    with patch.object(app, 'components', return_value=rows):
                        rel = app.diagnostics_bundle()
                    with zipfile.ZipFile(str(self.project / rel)) as archive:
                        self.assertEqual(archive.namelist(), ['diagnostics.json'])
                        text = archive.read('diagnostics.json').decode('utf-8')
                    exported = json.loads(text)['components']
                    safe = next(row for row in exported if row.get('feature') == 'controlled')
                    self.assertEqual(safe, {'feature': 'controlled', 'needs': 'jsonschema>=4.18',
                                           'usable': True, 'reason': None, 'version': app.sys.version.split()[0],
                                           'jsonschema': '4.23.0', 'fs': 'NTFS'})
                    self.assertEqual(exported[-1], {})
                    self.assertNotIn(path, text)
                    self.assertNotIn(json.dumps(path)[1:-1], text)
                    for key in ('python_path', 'executable', 'install_root', 'details'):
                        self.assertNotIn(key, text)
                    self.assertEqual(controlled['python_path'], path)

    def test_bundle_is_written_only_under_runtime_diagnostics(self):
        before = snapshot(self.project)
        with patch.object(app, 'pack_info', return_value={'items': []}):
            res = app.action('diagnostics', {})
        self.assertTrue(res['result']['files'][0]['path'].startswith('.agent/runtime/diagnostics/'))
        after = snapshot(self.project)
        self.assertEqual(before, after)


class MergeComparison(PanelBase):
    def setUp(self):
        super().setUp()
        root = self.project
        write(root, 'managed.txt', 'line one\nline two\nline three\nline four\nline five\n')
        write(root, 'managed.txt.new', 'line one\nline 2 changed\nline three\nline four\nline five\nline six\n')
        write(root, 'same.txt', 'identical text\n')
        write(root, 'same.txt.new', 'identical text\n')
        write(root, 'binary.bin', b'\x00\x01\x02')
        write(root, 'binary.bin.new', b'\x00\x05\x06')
        write(root, 'notutf.txt', 'plain text\n')
        write(root, 'notutf.txt.new', b'\xff\xfe\xfd\x00')
        write(root, 'big.txt', 'x' * 500 + '\n')
        write(root, 'big.txt.new', 'y' * 500 + '\n')
        write(root, 'missing.txt.new', 'only new\nsecond line\n')
        write(root, 'user-note.txt.new', 'not managed')
        manifest_lines = ['0' * 64 + '  ' + name for name in ('managed.txt', 'same.txt', 'binary.bin', 'notutf.txt',
                                                              'big.txt', 'missing.txt')]
        write(root, 'MANIFEST.sha256', '\n'.join(manifest_lines) + '\n')

    def test_text_diff_has_hunks_with_line_numbers(self):
        d = merge_center.diff(str(self.project), 'managed.txt.new')
        self.assertFalse(d['identical'] or d['binary'] or d['too_large'] or d['current_missing'])
        self.assertEqual((d['added'], d['removed']), (2, 1))
        self.assertEqual(len(d['hunks']), 1)
        hunk = d['hunks'][0]
        removed = [x for x in hunk['lines'] if x['op'] == 'removed']
        added = [x for x in hunk['lines'] if x['op'] == 'added']
        self.assertEqual([(x['old'], x['text']) for x in removed], [(2, 'line two')])
        self.assertEqual([(x['new'], x['text']) for x in added], [(2, 'line 2 changed'), (6, 'line six')])
        self.assertTrue(any(x['op'] == 'context' and x['text'] == 'line one' for x in hunk['lines']))

    def test_flags_identical_binary_undecodable_too_large_and_missing(self):
        root = str(self.project)
        self.assertTrue(merge_center.diff(root, 'same.txt.new')['identical'])
        self.assertEqual(merge_center.diff(root, 'same.txt.new')['hunks'], [])
        self.assertTrue(merge_center.diff(root, 'binary.bin.new')['binary'])
        self.assertTrue(merge_center.diff(root, 'notutf.txt.new')['binary'])
        self.assertTrue(merge_center.diff(root, 'big.txt.new', max_bytes=100)['too_large'])
        self.assertEqual(merge_center.diff(root, 'big.txt.new', max_bytes=100)['hunks'], [])
        missing = merge_center.diff(root, 'missing.txt.new')
        self.assertTrue(missing['current_missing'])
        self.assertEqual(missing['added'], 2)
        self.assertEqual(missing['removed'], 0)

    def test_unmanaged_paths_and_traversal_are_refused(self):
        root = str(self.project)
        for rel in ('user-note.txt.new', '../outside.txt.new', 'a/../../outside.txt.new', 'C:/managed.txt.new',
                    'managed.txt', '', 'missing-entirely.txt.new'):
            with self.subTest(rel=rel):
                with self.assertRaises(ValueError):
                    merge_center.diff(root, rel)
                with self.assertRaises(app.BadInput):
                    app.action('merge_diff', {'path': rel})

    def test_diff_is_read_only_and_stores_nothing(self):
        before = snapshot(self.project)
        merge_center.diff(str(self.project), 'managed.txt.new')
        res = app.action('merge_diff', {'path': 'managed.txt.new'})
        self.assertEqual(res['code'], 0)
        self.assertEqual(res['result']['diff']['path'], 'managed.txt')
        self.assertEqual(snapshot(self.project), before)
        self.assertFalse((self.project / '.agent' / 'MERGE_STATUS.json').exists())

    def test_set_state_still_refuses_unmanaged_files(self):
        with self.assertRaises(ValueError):
            merge_center.set_state(str(self.project), 'user-note.txt.new', 'ready')
        self.assertEqual(merge_center.set_state(str(self.project), 'managed.txt.new', 'reviewing')['state'], 'reviewing')


class EnvelopeCodes(PanelBase):
    """Each structured tool, called through app.action, returns an envelope whose code is the exit code.
    Runs against a temporary copy of the package, so the tools write only inside it."""

    def test_structured_tools_return_matching_envelopes(self):
        work = self.root / 'work'
        work.mkdir()
        sample = write(work, 'sample.txt', 'A complete text with enough content for the check.')
        write(work, 'adopt-target/AGENTS.md', 'own instructions')
        old, new = self.upgrade_pair('env')
        self.patch(app, 'ROOT', str(PKG))
        calls = [
            ('library_init', {}),
            ('library_stats', {}),
            ('library_check', {}),
            ('library_find', {'query': 'luat'}),
            ('pack_list', {}),
            ('check_package', {}),
            ('check_package', {'verify': True}),
            ('check_output', {'files': [str(sample)]}),
            ('install_skills', {'target': 'custom', 'dest': str(work / 'skills-out')}),
            ('init', {'name': 'Demo', 'lang': 'en', 'apply': False}),
            ('adopt', {'project': str(work / 'adopt-target'), 'name': 'Demo'}),
            ('upgrade_other', {'project': str(old)}),
            ('tx_resume', {'txid': '20000101-000000-deadbeef'}),
        ]
        for action_name, params in calls:
            with self.subTest(action=action_name, params=params):
                res = app.action(action_name, params)
                env = res['result']
                self.assertEqual(env['schema'], result_schema.SCHEMA)
                self.assertEqual(env['code'], res['code'])
                self.assertFalse(any(u.get('reason') == 'invalid_result_schema' for u in env['unchecked']))
        stats = app.action('library_stats', {})
        self.assertEqual(stats['result']['status'], 'success', stats['out'])
        self.assertEqual(json.loads(stats['out'])['index'], True)
        found = app.action('library_find', {'query': 'luat'})
        self.assertEqual(json.loads(found['out'])['query'], 'luat')
        self.assertEqual(app.action('check_output', {'files': [str(sample)]})['result']['status'], 'success')


# ------------------------------------------------------------------ A-07: thư viện dùng chung / the shared library
TEMPLATE_INDEX = ROOT / 'skills/local/shared/source-library/assets/source-index-template.xlsx'
LAW_COLUMNS = ('LAW revision', 'LAW SHA-256', 'LAW observed', 'LAW file state')


def tree_state(root):
    """Mọi file kể cả runtime (băm), để chứng minh một thao tác đọc không đổi gì / every file incl. runtime, hashed."""
    out = {}
    root = Path(root)
    for p in sorted(root.rglob('*')):
        rel = p.relative_to(root).as_posix()
        out[rel] = hashlib.sha256(p.read_bytes()).hexdigest() if p.is_file() else 'dir'
    return out


def make_index(project, rows, law=True):
    """Sổ từ mẫu với các dòng giả; law=False là sổ cũ không có cột LAW / the template with synthetic rows; law=False is an older index."""
    ref = Path(project) / 'references'
    ref.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(str(TEMPLATE_INDEX), str(ref / '_INDEX.xlsx'))
    wb = openpyxl.load_workbook(str(ref / '_INDEX.xlsx'))
    ws = wb['Danh mục']
    cols = {}
    for key, prefix in library.KEYS.items():
        for i in range(1, ws.max_column + 1):
            if str(ws.cell(4, i).value or '').startswith(prefix) and i not in cols.values():
                cols[key] = i
                break
    if law:
        for label in LAW_COLUMNS:
            ws.cell(4, ws.max_column + 1, label)
        cols.update({'law_revision': 22, 'law_sha': 23, 'law_observed': 24, 'law_state': 25})
    for n, rec in enumerate(rows, 5):
        for key, value in rec.items():
            ws.cell(n, cols[key], value)
    wb.save(str(ref / '_INDEX.xlsx'))
    wb.close()
    return ref / '_INDEX.xlsx'


def row(ident, rev, path, **kw):
    rec = {'id': ident, 'group': 'Tài liệu nguồn', 'number': kw.pop('number', ''), 'title': kw.pop('title', 'Tài liệu ' + (ident or 'không mã')),
           'status': kw.pop('status', 'Còn hiệu lực'), 'fulltext': kw.pop('fulltext', 'Có'), 'verify': kw.pop('verify', 'Cần đối chiếu'),
           'file': path, 'law_revision': rev, 'law_state': kw.pop('state', 'PRESENT')}
    rec.update(kw)
    return {k: v for k, v in rec.items() if v is not None}


def dir_link(link, target):
    """Liên kết thư mục (symlink hoặc junction Windows); False nếu máy này không tạo được / a directory link; False when not creatable."""
    try:
        os.symlink(str(target), str(link), target_is_directory=True)
        return True
    except (OSError, NotImplementedError):
        pass
    if os.name == 'nt':
        done = subprocess.run(['cmd', '/c', 'mklink', '/J', str(link), str(target)], capture_output=True)
        return done.returncode == 0 and os.path.isdir(str(link))
    return False


def drop_link(link):
    try:
        os.rmdir(str(link))
    except OSError:
        os.unlink(str(link))


@unittest.skipUnless(openpyxl, 'openpyxl is needed for the source workbook')
class LibraryReadApi(PanelBase):
    """A-07: browse/search/profile/select/open read the saved workbook only. Fake data in a temporary project."""

    def setUp(self):
        super().setUp()
        self.ref = self.project / 'references'
        write(self.project, 'references/legal/dat-dai/luat.pdf', 'synthetic law')
        write(self.project, 'references/policies/quy-che/quy-che.docx', 'synthetic policy')
        make_index(self.project, [
            row('SRC-000001', 1, 'legal/dat-dai/luat.pdf', group='Pháp luật', number='31/2024/QH15', title='Luật Đất đai'),
            row('SRC-000002', 3, 'policies/quy-che/quy-che.docx', group='Nội bộ', number='QC-01', title='Quy chế nội bộ',
                verify='Đã đối chiếu toàn văn'),
            row('SRC-000003', 1, 'legal/khong-co/mat.pdf', group='Pháp luật', title='Văn bản mất file', state='MISSING'),
        ])
        self.opened = []
        self.patch(app, 'launch', lambda path: self.opened.append(path))

    def browse(self, query='', group=''):
        return app.library_api('/api/library/browse', {'query': query, 'group': group})

    def pick(self, rec, snapshot):
        return {'snapshot': snapshot, 'row': rec['row'], 'id': rec['id'], 'revision': rec['revision'], 'locator': rec['locator']}

    def rewrite(self, mutate):
        wb = openpyxl.load_workbook(str(self.ref / '_INDEX.xlsx'))
        mutate(wb['Danh mục'])
        wb.save(str(self.ref / '_INDEX.xlsx'))
        wb.close()

    # -- (1) cùng sổ, cùng ID/revision/trạng thái với CLI / the same workbook, same identity as the CLI
    def test_blank_keyword_browses_everything_with_identity_status_and_locator(self):
        res = self.browse()
        self.assertTrue(res['ok'])
        self.assertEqual((res['total'], res['matched'], res['truncated']), (3, 3, False))
        first = res['rows'][0]
        self.assertEqual((first['id'], first['revision'], first['identity'], first['file_state'], first['status'], first['verify']),
                         ('SRC-000001', 1, 'ok', 'PRESENT', 'Còn hiệu lực', 'Cần đối chiếu'))
        self.assertEqual((first['locator'], first['file'], first['fulltext_declared']),
                         ('legal/dat-dai/luat.pdf', 'references/legal/dat-dai/luat.pdf', True))
        self.assertEqual(res['groups'], ['Pháp luật', 'Nội bộ'])  # nhóm lấy từ dữ liệu / groups come from the data
        self.assertEqual(len(res['snapshot']), 64)
        self.assertEqual(res['summary'], {'rows': 3, 'declared_fulltext': 3, 'not_declared': 0, 'needs_check': 2, 'identified': 3})

    def test_accent_insensitive_search_and_group_filter_run_on_the_saved_workbook(self):
        self.assertEqual([r['id'] for r in self.browse('dat dai')['rows']], ['SRC-000001'])
        self.assertEqual([r['id'] for r in self.browse('', 'Nội bộ')['rows']], ['SRC-000002'])
        self.assertEqual([r['id'] for r in self.browse('quy che', 'noi bo')['rows']], ['SRC-000002'])
        none = self.browse('khong-co-tu-nay-dau')
        self.assertTrue(none['ok'])
        self.assertEqual((none['matched'], none['rows'], none['total']), (0, [], 3))  # lọc ra 0 hit, sổ vẫn có dòng / filter-empty, not index-empty

    def test_groups_default_to_the_four_current_groups_when_the_data_names_none(self):
        make_index(self.project, [])
        res = self.browse()
        self.assertTrue(res['ok'])
        self.assertEqual(res['groups'], ['Pháp luật', 'Nội bộ', 'Mẫu biểu', 'Tài liệu nguồn'])
        self.assertEqual((res['total'], res['rows']), (0, []))  # sổ hợp lệ 0 dòng / a valid index with no rows

    def test_cli_find_and_panel_browse_agree_on_id_revision_status_and_file_state(self):
        book = library.read_book(str(self.project))  # cùng bộ phân tích: CLI find và browse đọc qua read_book / one reader behind both
        cli = {r['_row']: library.show(r) for r in book['data']}
        for rec in self.browse()['rows']:
            shown = cli[rec['row']]
            self.assertEqual((rec['id'], rec['revision'], rec['status'], rec['verify'], rec['file_state'], rec['title']),
                             (shown['id'], shown['revision'], shown['status'], shown['verify'], shown['file_state'], shown['title']))

    def test_legacy_workbook_without_id_and_law_columns_is_viewable_and_invents_nothing(self):
        make_index(self.project, [row(None, None, 'legal/dat-dai/luat.pdf', group='Pháp luật', title='Luật cũ', state=None)], law=False)
        res = self.browse()
        self.assertTrue(res['ok'], res)
        rec = res['rows'][0]
        self.assertEqual((rec['id'], rec['revision'], rec['identity'], rec['file_state'], rec['sha12']), ('', None, 'no_id', '', ''))
        self.assertEqual(res['summary']['identified'], 0)
        wb = openpyxl.load_workbook(str(self.ref / '_INDEX.xlsx'))
        self.assertIsNone(wb['Danh mục']['V4'].value)  # không ép cột LAW vào sổ cũ để xem / no LAW column is forced into an older index
        wb.close()

    # -- (2) đọc không phát sinh tác động / reads cause no effect
    def test_browse_search_profile_select_leave_the_tree_untouched_and_call_no_scan_or_writer(self):
        before = tree_state(self.project)
        boom = AssertionError('a read path touched a scan, a writer or a provider')
        res = self.browse()
        recs = res['rows']
        with patch.object(library_sync, 'plan', side_effect=boom), patch.object(library_sync, 'apply', side_effect=boom), \
                patch.object(library_sync, 'sources', side_effect=boom), patch.object(library_sync, 'digest', side_effect=boom), \
                patch.object(library, 'library_files', side_effect=boom), patch.object(app, 'run_tool', side_effect=boom), \
                patch.object(tx, 'Transaction', side_effect=boom), patch('os.walk', side_effect=boom), \
                patch.object(ops_log, 'record', side_effect=boom) as logger:
            self.browse()
            self.browse('luat', 'Pháp luật')
            self.browse('khong-co')
            out = app.library_api('/api/library/select', {'selections': [self.pick(r, res['snapshot']) for r in recs[:2]]})
            self.assertTrue(out['ok'], out)
            logger.assert_not_called()
        self.assertEqual(tree_state(self.project), before)  # nguồn, sổ, STATE, runtime, nhật ký: không đổi / sources, index, STATE, runtime, log: unchanged
        self.assertFalse((self.project / '.agent').exists())

    def test_library_find_and_stats_through_the_panel_api_write_no_log_line(self):
        for name, params in (('library_find', {'query': 'luat'}), ('library_stats', {})):
            app.run_api({'action': name, 'params': params})
        self.assertFalse((self.project / '.agent' / 'runtime' / 'operations.jsonl').exists())
        app.run_api({'action': 'library_check', 'params': {}})  # kiểm tra vẫn là một thao tác có ghi / check stays a logged action
        self.assertEqual([r['action'] for r in ops_log.read(str(self.project))], ['library_check'])

    # -- (3) ma trận lỗi / the error matrix
    def test_error_matrix_tells_apart_missing_unreadable_structure_and_valid_empty(self):
        index = self.ref / '_INDEX.xlsx'
        index.unlink()
        self.assertEqual(self.browse()['kind'], 'no_index')
        index.write_bytes(b'this is not a zip workbook')
        self.assertEqual(self.browse()['kind'], 'unreadable')

        wb = openpyxl.Workbook()
        wb.active.title = 'Sheet khác'
        wb.save(str(index))
        self.assertEqual(self.browse()['kind'], 'missing_sheet')

        wb = openpyxl.Workbook()
        wb.active.title = 'Danh mục'
        wb.active['A1'] = 'chỉ có dòng 1'
        wb.save(str(index))
        self.assertEqual(self.browse()['kind'], 'no_header')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Danh mục'
        for i, name in enumerate(('Mã / ID', 'Trích yếu / Title'), 1):
            ws.cell(4, i, name)
        wb.save(str(index))
        res = self.browse()
        self.assertEqual((res['ok'], res['kind']), (False, 'missing_columns'))
        self.assertIn('Nhóm', res['message'])

        make_index(self.project, [])
        valid_empty = self.browse()
        self.assertTrue(valid_empty['ok'])
        self.assertEqual(valid_empty['total'], 0)

    def test_simulated_permission_error_is_denied_not_no_index(self):
        real_open = open

        def denied(path, *a, **k):
            if str(path).endswith('_INDEX.xlsx'):
                raise PermissionError(13, 'simulated: access denied')  # mô phỏng, không phải quyền thật / a simulation, not a real ACL
            return real_open(path, *a, **k)
        with patch.object(library, 'open', denied, create=True):
            res = self.browse()
        self.assertEqual((res['ok'], res['kind']), (False, 'denied'))

    def test_missing_openpyxl_is_not_checked_not_empty(self):
        with patch.dict(sys.modules, {'openpyxl': None}):
            res = self.browse()
        self.assertEqual((res['ok'], res['kind']), (False, 'openpyxl_missing'))
        self.assertIn('openpyxl', ' '.join(res['next_actions']))

    def test_cli_load_stops_on_the_structure_the_panel_calls_a_failure(self):
        wb = openpyxl.Workbook()
        wb.active.title = 'Danh mục'
        wb.active['A1'] = 'x'
        wb.save(str(self.ref / '_INDEX.xlsx'))
        with patch.object(library, 'ROOT', str(self.project)):
            with self.assertRaises(SystemExit) as caught:
                library.load()
        self.assertNotIn(caught.exception.code, (0, None))  # không bao giờ thành công-rỗng / never success-empty

    # -- (4) mở / chọn: bản ghi cũ, trùng, thiếu danh tính, chèn locator / stale, duplicate, legacy, locator injection
    def test_select_copies_identity_and_reminds_the_agent_to_check(self):
        res = self.browse()
        out = app.library_api('/api/library/select', {'selections': [self.pick(r, res['snapshot']) for r in res['rows'][:2]], 'lang': 'vi'})
        self.assertTrue(out['ok'])
        self.assertEqual(out['count'], 2)
        for needle in ('SRC-000001', 'revision 1', 'references/legal/dat-dai/luat.pdf', 'Còn hiệu lực', 'Cần đối chiếu',
                       'SRC-000002', 'revision 3', 'Đã đối chiếu toàn văn', 'PRESENT', 'đối chiếu ID, revision và đường dẫn', 'không xác minh nội dung'):
            self.assertIn(needle, out['text'])
        self.assertIn(res['snapshot'][:12], out['text'])
        en = app.library_api('/api/library/select', {'selections': [self.pick(res['rows'][0], res['snapshot'])], 'lang': 'en'})
        self.assertIn('verifies nothing about the content', en['text'])

    def test_select_refuses_missing_duplicate_and_unrecorded_identity_and_copies_nothing(self):
        make_index(self.project, [
            row('SRC-000001', 1, 'legal/dat-dai/luat.pdf', group='Pháp luật'),
            row('DUP-1', 1, 'legal/a.pdf', group='Pháp luật'), row('DUP-1', 2, 'legal/b.pdf', group='Pháp luật'),
            row(None, None, 'legal/c.pdf', group='Pháp luật', title='Dòng cũ không ID', state=None),
            row('SRC-000009', None, 'legal/d.pdf', group='Pháp luật', title='Chưa có revision'),
        ])
        res = self.browse()
        by_title = {r['title']: r for r in res['rows']}
        self.assertEqual({r['identity'] for r in res['rows']}, {'ok', 'dup_id', 'no_id', 'no_revision'})
        out = app.library_api('/api/library/select', {'selections': [self.pick(r, res['snapshot']) for r in res['rows']]})
        self.assertFalse(out['ok'])
        self.assertNotIn('text', out)
        self.assertEqual(sorted(i['reason'] for i in out['items']), ['dup_id', 'dup_id', 'no_id', 'no_revision'])
        ok_only = app.library_api('/api/library/select', {'selections': [self.pick(by_title['Tài liệu SRC-000001'], res['snapshot'])]})
        self.assertTrue(ok_only['ok'])
        self.assertEqual(self.browse('Dòng cũ')['rows'][0]['row'], by_title['Dòng cũ không ID']['row'])  # xem vẫn được / still viewable

    def test_stale_changed_and_forged_selections_are_refused_before_any_open(self):
        res = self.browse()
        good = self.pick(res['rows'][0], res['snapshot'])
        cases = [(dict(good, snapshot='0' * 64), 'stale'), (dict(good, id='SRC-999999'), 'changed'), (dict(good, revision=9), 'changed'),
                 (dict(good, locator='legal/khac.pdf'), 'changed'), (dict(good, row=99), 'stale'), (dict(good, row='5'), 'bad_selection'),
                 (dict(good, row=True), 'bad_selection'), ({}, 'bad_selection'), (None, 'bad_selection'), ('x', 'bad_selection'),
                 (dict(good, revision=True), 'bad_selection')]
        for sel, why in cases:
            with self.subTest(selection=str(sel)[:60]):
                out = app.library_api('/api/library/open', {'selection': sel})
                self.assertEqual((out['ok'], out['reason']), (False, why))
        self.assertEqual(self.opened, [])
        self.rewrite(lambda ws: ws.cell(5, 9, 'Trích yếu đã đổi'))  # sổ đổi sau khi tải bảng / the workbook changes after the table loaded
        self.assertEqual(app.library_api('/api/library/open', {'selection': good})['reason'], 'stale')
        self.assertEqual(app.library_api('/api/library/select', {'selections': [good]})['items'][0]['reason'], 'stale')
        self.assertEqual(self.opened, [])

    def test_open_hands_exactly_the_rebound_source_to_the_opener_and_nothing_else(self):
        res = self.browse()
        out = app.library_api('/api/library/open', {'selection': self.pick(res['rows'][0], res['snapshot']), 'path': '../../secret.txt'})
        self.assertEqual(out, {'ok': True})
        self.assertEqual([os.path.normcase(p) for p in self.opened],
                         [os.path.normcase(str(self.ref / 'legal' / 'dat-dai' / 'luat.pdf'))])
        self.assertFalse((self.project / '.agent' / 'runtime' / 'operations.jsonl').exists())  # mở không ghi lịch sử / no history line

    def test_open_launch_failure_is_reported_not_hidden(self):
        res = self.browse()
        with patch.object(app, 'launch', side_effect=OSError('no handler')):
            out = app.library_api('/api/library/open', {'selection': self.pick(res['rows'][0], res['snapshot'])})
        self.assertEqual((out['ok'], out['reason']), (False, 'launch_failed'))

    def test_open_does_not_strip_a_second_references_prefix_from_the_bound_locator(self):
        nested = self.ref / 'references' / 'legal' / 'a.txt'
        write(self.project, 'references/references/legal/a.txt', 'the file shown by the row')
        write(self.project, 'references/legal/a.txt', 'a different file')
        make_index(self.project, [row('SRC-000001', 1, 'references/references/legal/a.txt')])
        res = self.browse()
        rec = res['rows'][0]
        self.assertEqual((rec['locator'], rec['file']), ('references/legal/a.txt', 'references/references/legal/a.txt'))
        out = app.library_api('/api/library/open', {'selection': self.pick(rec, res['snapshot'])})
        self.assertEqual(out, {'ok': True})
        self.assertEqual([os.path.normcase(p) for p in self.opened], [os.path.normcase(str(nested))])

    def test_open_refuses_formats_outside_the_managed_source_extensions_before_launch(self):
        # Inert dummy bytes only: these paths are never passed to an OS opener.
        paths = ['legal/inert' + suffix for suffix in ('.bat', '.cmd', '.exe', '.lnk', '.url')]
        for path in paths:
            write(self.project, 'references/' + path, 'inert review fixture')
        make_index(self.project, [row('SRC-%06d' % n, 1, path) for n, path in enumerate(paths, 1)])
        res = self.browse()
        for rec in res['rows']:
            with self.subTest(locator=rec['locator']):
                out = app.library_api('/api/library/open', {'selection': self.pick(rec, res['snapshot'])})
                self.assertEqual(out, {'ok': False, 'reason': 'unsupported_type'})
        self.assertEqual(self.opened, [])

    def test_locator_injection_in_the_workbook_is_refused_by_the_path_rules(self):
        evil = ['C:/Windows/win.ini', '/etc/passwd', '../../outside.txt', 'legal/../../x.pdf', '../Cu/../../x', 'legal//x.pdf',
                'legal/[tên].pdf', '../only', 'legal/./x.pdf']
        make_index(self.project, [row('SRC-%06d' % n, 1, p, group='Pháp luật') for n, p in enumerate(evil, 1)])
        res = self.browse()
        self.assertEqual(len(res['rows']), len(evil))
        for rec in res['rows']:
            with self.subTest(locator=rec['locator']):
                out = app.library_api('/api/library/open', {'selection': self.pick(rec, res['snapshot'])})
                self.assertEqual((out['ok'], out['reason']), (False, 'bad_locator'))
        self.assertEqual(self.opened, [])

    def test_the_index_itself_and_hidden_or_temp_files_are_not_openable_sources(self):
        write(self.project, 'references/legal/.an/x.pdf', 'hidden')
        write(self.project, 'references/legal/~$tam.docx', 'office lock file')
        make_index(self.project, [row('SRC-000001', 1, '_INDEX.xlsx', group='Pháp luật'), row('SRC-000002', 1, 'legal/.an/x.pdf', group='Pháp luật'),
                                  row('SRC-000003', 1, 'legal/~$tam.docx', group='Pháp luật')])
        res = self.browse()
        for rec in res['rows']:
            with self.subTest(locator=rec['locator']):
                out = app.library_api('/api/library/open', {'selection': self.pick(rec, res['snapshot'])})
                self.assertEqual((out['ok'], out['reason']), (False, 'out_of_scope'))
        self.assertEqual(self.opened, [])

    # -- nguồn mất / quyền / ngoài phạm vi / liên kết tách khỏi "không khai toàn văn" / source states are not "no full text"
    def test_source_states_are_told_apart_at_open_time(self):
        write(self.project, 'references/legal/dat-dai/thu-muc.pdf/inside.txt', 'a directory named like a file')
        write(self.project, 'Tai-lieu-cu/ho-so.pdf', 'old source')
        write(self.project, 'Khac/ngoai.pdf', 'outside the declared scope')
        write(self.project, 'PROJECT.md', '- Ngoại lệ thư mục nguồn cũ: Tai-lieu-cu\n')
        make_index(self.project, [
            row('SRC-000001', 1, 'legal/khong-co/mat.pdf', group='Pháp luật', title='mất', state='MISSING'),
            row('SRC-000002', 1, 'legal/dat-dai/thu-muc.pdf', group='Pháp luật', title='thư mục'),
            row('SRC-000003', 1, '../Tai-lieu-cu/ho-so.pdf', group='Tài liệu nguồn', title='nguồn cũ trong phạm vi'),
            row('SRC-000004', 1, '../Khac/ngoai.pdf', group='Tài liệu nguồn', title='ngoài phạm vi', state='OUT_OF_SCOPE'),
            row('SRC-000005', 1, 'legal/dat-dai/luat.pdf', group='Pháp luật', title='không khai toàn văn', fulltext='Không'),
            row('SRC-000006', 1, None, group='Pháp luật', title='chưa ghi đường dẫn', fulltext='Không'),
        ])
        res = self.browse()
        got = {}
        for rec in res['rows']:
            self.opened.clear()
            out = app.library_api('/api/library/open', {'selection': self.pick(rec, res['snapshot'])})
            got[rec['title']] = (out.get('reason'), bool(self.opened))
        self.assertEqual(got, {'mất': ('missing', False), 'thư mục': ('not_regular', False), 'nguồn cũ trong phạm vi': (None, True),
                               'ngoài phạm vi': ('out_of_scope', False), 'không khai toàn văn': (None, True),
                               'chưa ghi đường dẫn': ('no_locator', False)})
        self.assertFalse(next(r for r in res['rows'] if r['title'] == 'không khai toàn văn')['fulltext_declared'])
        self.assertEqual(next(r for r in res['rows'] if r['title'] == 'mất')['file_state'], 'MISSING')  # đã ghi, không phải kiểm lúc này / recorded, not checked now

    def test_old_source_is_refused_when_the_project_scope_does_not_name_it(self):
        write(self.project, 'Tai-lieu-cu/ho-so.pdf', 'old source')
        make_index(self.project, [row('SRC-000001', 1, '../Tai-lieu-cu/ho-so.pdf', group='Tài liệu nguồn')])
        res = self.browse()
        sel = self.pick(res['rows'][0], res['snapshot'])
        self.assertEqual(app.library_api('/api/library/open', {'selection': sel})['reason'], 'out_of_scope')  # PROJECT.md không có ngoại lệ
        write(self.project, 'PROJECT.md', '- Ngoại lệ thư mục nguồn cũ: [Không có / thư mục]\n')
        self.assertEqual(app.library_api('/api/library/open', {'selection': sel})['reason'], 'out_of_scope')  # chỗ trống: không rõ thì không mở
        self.assertEqual(self.opened, [])

    def test_simulated_permission_error_at_open_is_denied_not_missing(self):
        res = self.browse()
        real = os.lstat
        target = os.path.normcase(str(self.ref / 'legal' / 'dat-dai' / 'luat.pdf'))

        def lstat(path, *a, **k):
            if os.path.normcase(str(path)) == target:
                raise PermissionError(13, 'simulated: access denied')  # mô phỏng / a simulation
            return real(path, *a, **k)
        with patch.object(os, 'lstat', lstat):
            out = app.library_api('/api/library/open', {'selection': self.pick(res['rows'][0], res['snapshot'])})
        self.assertEqual((out['ok'], out['reason']), (False, 'denied'))
        self.assertEqual(self.opened, [])

    def test_a_link_in_the_path_is_refused_when_this_machine_can_make_one(self):
        outside = self.root / 'outside'
        write(outside, 'luat.pdf', 'outside the project')
        link = self.ref / 'legal' / 'lien-ket'
        if not dir_link(link, outside):
            self.skipTest('this machine cannot create a symlink or junction (WinError 1314 or similar)')
        self.addCleanup(drop_link, link)
        make_index(self.project, [row('SRC-000001', 1, 'legal/lien-ket/luat.pdf', group='Pháp luật')])
        res = self.browse()
        out = app.library_api('/api/library/open', {'selection': self.pick(res['rows'][0], res['snapshot'])})
        self.assertEqual((out['ok'], out['reason']), (False, 'link'))
        self.assertEqual(self.opened, [])


@unittest.skipUnless(openpyxl, 'openpyxl is needed for the source workbook')
class ArtifactLinksApi(PanelBase):
    """B-02: the dedicated read routes for related sources, results, exact-version open and copy. Fake data in a temporary project."""

    def setUp(self):
        super().setUp()
        import artifacts
        self.artifacts = artifacts
        self.root_s = str(self.project)
        write(self.project, 'references/source-documents/a.txt', 'alpha')
        write(self.project, 'working/r1.txt', 'report one')
        write(self.project, 'references/_INDEX.xlsx', (ROOT / 'skills/local/shared/source-library/assets/source-index-template.xlsx').read_bytes())
        library_sync.apply(self.root_s)
        rec = artifacts.load_book(self.root_s)['catalog']['SRC-000001'][0]
        req = {'artifact_id': 'R1', 'revision': 1, 'path': 'working/r1.txt', 'title': 'Báo cáo giả', 'kind': 'INFERENCE', 'limits': 'giả định X',
               'origin': 'USER_FILE', 'recorded_by': 'tester',
               'dependencies': [{'kind': 'SOURCE', 'id': 'SRC-000001', 'revision': 1, 'sha256': rec['law_sha']}]}
        artifacts.register_apply(self.root_s, req, artifacts.register_plan(self.root_s, req)['plan_token'])
        self.opened = []
        self.patch(artifacts, '_launch', lambda path: self.opened.append(path))
        book = library.read_book(self.root_s)
        self.source_pick = {'snapshot': book['snapshot'], 'row': 5, 'id': 'SRC-000001', 'revision': 1, 'locator': 'source-documents/a.txt'}

    def links(self, sel=None):
        return app.library_api('/api/library/links', {'selection': sel if sel is not None else self.source_pick})

    def test_links_return_the_results_with_exact_picks_and_the_page_can_open_and_copy_them(self):
        res = self.links()
        self.assertTrue(res['ok'], res)
        item = res['results'][0]
        self.assertEqual((item['artifact_id'], item['revision'], item['kind'], item['dependency_state'], item['inference'], item['via_direct']),
                         ('R1', 1, 'INFERENCE', 'CURRENT_CATALOG_BASIS', ['R1@1'], True))
        out = app.library_api('/api/library/artifact-open', {'selection': item['pick']})
        self.assertEqual((out['ok'], out['state']), (True, 'OPENED'))
        self.assertEqual(self.opened, [str(self.project / 'working' / 'r1.txt')])
        copied = app.library_api('/api/library/artifact-copy', {'selection': item['pick'], 'lang': 'en'})
        self.assertTrue(copied['ok'])
        self.assertIn('giả định X', copied['text'])
        self.assertIn('INFERENCE', copied['text'])

    def test_reads_leave_the_tree_untouched_and_call_no_scan_writer_log_or_provider(self):
        item = self.links()['results'][0]
        before = tree_state(self.project)
        boom = AssertionError('a read path touched a scan, a writer, the log or a provider')
        with patch.object(library_sync, 'plan', side_effect=boom), patch.object(library_sync, 'apply', side_effect=boom), \
                patch.object(library_sync, 'sources', side_effect=boom), patch.object(library_sync, 'locked_commit', side_effect=boom), \
                patch.object(library, 'library_files', side_effect=boom), patch.object(app, 'run_tool', side_effect=boom), \
                patch.object(library_sync, 'digest', side_effect=boom), patch.object(library, 'locate_source', side_effect=boom), \
                patch.object(tx, 'Transaction', side_effect=boom), patch('os.walk', side_effect=boom), \
                patch.object(ops_log, 'record', side_effect=boom) as logger:
            self.assertTrue(self.links()['ok'])
            self.assertTrue(app.library_api('/api/library/artifact-copy', {'selection': item['pick']})['ok'])
            logger.assert_not_called()
        self.assertEqual(tree_state(self.project), before)

    def test_open_may_hash_the_chosen_file_but_nothing_else_and_writes_nothing(self):
        item = self.links()['results'][0]
        before = tree_state(self.project)
        with patch.object(library_sync, 'plan', side_effect=AssertionError('scan')), patch.object(library_sync, 'sources', side_effect=AssertionError('scan')), \
                patch.object(ops_log, 'record', side_effect=AssertionError('log')):
            self.assertTrue(app.library_api('/api/library/artifact-open', {'selection': item['pick']})['ok'])
        self.assertEqual(tree_state(self.project), before)

    def test_stale_changed_and_malformed_selections_are_refused_with_a_reason(self):
        item = self.links()['results'][0]
        for action, body, reason in (
                ('links', {'selection': dict(self.source_pick, snapshot='9' * 64)}, 'stale'),
                ('links', {'selection': dict(self.source_pick, id='SRC-000009')}, 'changed'),
                ('links', {'selection': 'nope'}, 'bad_selection'),
                ('artifact-open', {'selection': dict(item['pick'], snapshot='9' * 64)}, 'stale'),
                ('artifact-open', {'selection': dict(item['pick'], sha256='8' * 64)}, 'changed'),
                ('artifact-open', {'selection': None}, 'bad_selection'),
                ('artifact-copy', {'selection': dict(item['pick'], revision=7)}, 'stale')):
            out = app.library_api('/api/library/' + action, body)
            self.assertEqual((out['ok'], out['reason']), (False, reason), (action, body))
        self.assertEqual(self.opened, [])
        write(self.project, 'working/r1.txt', 'edited by hand')
        self.assertEqual(app.library_api('/api/library/artifact-open', {'selection': item['pick']}), {'ok': False, 'reason': 'file_changed'})
        self.assertEqual(self.opened, [])

    def test_a_malformed_registry_is_reported_and_the_unknown_action_is_not_served(self):
        wb = openpyxl.load_workbook(str(self.project / 'references' / '_INDEX.xlsx'))
        wb['LAW artifacts']['A1'] = 'broken'
        wb.save(str(self.project / 'references' / '_INDEX.xlsx'))
        wb.close()
        out = self.links()
        self.assertEqual((out['ok'], out['reason']), (False, 'index:registry_malformed'))
        self.assertIsNone(app.library_api('/api/library/artifact-nothing', {}))


class LibrarySyncApi(PanelBase):
    """A-07: preview/apply of the source index through the one library_sync writer, bound to the preview that was shown."""

    def setUp(self):
        super().setUp()
        self.ref = self.project / 'references'
        (self.ref / 'source-documents' / 'team').mkdir(parents=True)
        shutil.copyfile(str(TEMPLATE_INDEX), str(self.ref / '_INDEX.xlsx'))
        self.index = self.ref / '_INDEX.xlsx'
        self.source('a.txt', 'first')

    def source(self, name, content):
        return write(self.ref, 'source-documents/team/' + name, content)

    def index_hash(self):
        return hashlib.sha256(self.index.read_bytes()).hexdigest()

    def preview(self):
        res = app.action('library_sync', {})
        self.assertEqual(res['code'], 0, res)
        return res['result']

    def apply(self, expected):
        return app.action('library_sync', {'apply': True, 'expected_preview': expected})

    def no_transaction(self):
        return patch.object(tx, 'Transaction', side_effect=AssertionError('a refused apply must stop before any Transaction'))

    def test_preview_writes_nothing_and_binds_snapshot_sources_and_scopes_without_bytes(self):
        before = tree_state(self.project)
        env = self.preview()
        self.assertEqual((env['status'], env['changes']), ('preview', None))
        self.assertEqual(tree_state(self.project), before)
        exp = env['expected_preview']
        self.assertEqual(set(exp), {'snapshot', 'sources', 'scopes'})
        self.assertEqual(exp['snapshot'], self.index_hash())
        self.assertEqual(list(exp['sources']), ['source-documents/team/a.txt'])
        self.assertEqual(exp['scopes'], [])
        self.assertEqual(json.loads(json.dumps(exp)), exp)  # ổn định qua JSON / stable through JSON
        self.assertEqual([r['action'] for r in env['rows']], ['add'])

    def test_apply_goes_through_the_common_writer_and_the_next_preview_is_a_no_op(self):
        exp = self.preview()['expected_preview']
        with patch.object(library_sync, 'apply', wraps=library_sync.apply) as spy:
            res = self.apply(exp)
        spy.assert_called_once_with(str(self.project), (), exp)
        self.assertEqual((res['code'], res['result']['status'], res['result']['changes']), (0, 'success', True))
        self.assertEqual(res['result']['files'], [{'path': 'references/_INDEX.xlsx', 'action': 'update'}])
        wb = openpyxl.load_workbook(str(self.index))
        self.assertEqual(wb['Danh mục'].cell(5, 1).value, 'SRC-000001')
        wb.close()
        again = self.preview()
        self.assertEqual((again['status'], again['changes']), ('success', False))
        self.assertNotIn('expected_preview', again)

    def test_apply_without_a_valid_preview_is_refused_before_lock_and_transaction(self):
        exp = self.preview()['expected_preview']
        bad = [None, [], {}, 'x', 7, dict(exp, extra=1), {k: v for k, v in exp.items() if k != 'scopes'}, dict(exp, snapshot='abc'),
               dict(exp, snapshot=7), dict(exp, sources=[]), dict(exp, sources={'a': 'not-a-hash'}), dict(exp, sources={3: 'a' * 64}),
               dict(exp, scopes='Cu'), dict(exp, scopes=[1]), dict(exp, snapshot=exp['snapshot'].upper())]
        before = self.index_hash()
        for value in bad:
            with self.subTest(expected=str(value)[:50]), self.no_transaction():
                res = self.apply(value)
                self.assertEqual((res['code'], res['result']['status'], res['result']['reason']), (1, 'blocked', 'PREVIEW_INVALID'))
        with self.no_transaction():
            res = app.action('library_sync', {'apply': True})  # thiếu hẳn: không phải hành vi CLI / missing entirely is not the CLI behaviour
            self.assertEqual(res['result']['reason'], 'PREVIEW_INVALID')
        self.assertEqual(self.index_hash(), before)
        self.assertFalse((self.project / '.agent').exists())  # chưa có khóa hay nhật ký / no lock, no journal

    def reset(self):
        shutil.copyfile(str(TEMPLATE_INDEX), str(self.index))
        for extra in ('a.txt', 'b.txt'):
            (self.ref / 'source-documents' / 'team' / extra).unlink(missing_ok=True)
        self.source('a.txt', 'first')

    def test_a_changed_index_source_or_scope_makes_the_preview_stale_before_any_transaction(self):
        edits = {
            'index': lambda: self.index.write_bytes(self.index.read_bytes() + b'x'),
            'source changed': lambda: self.source('a.txt', 'second'),
            'source added': lambda: self.source('b.txt', 'new'),
            'source removed': lambda: (self.ref / 'source-documents' / 'team' / 'a.txt').unlink(),
            'scope only': lambda: ((self.project / 'Tai-lieu-cu').mkdir(exist_ok=True),
                                   write(self.project, 'PROJECT.md', '- Ngoại lệ thư mục nguồn cũ: Tai-lieu-cu\n')),  # thư mục rỗng: chỉ phạm vi đổi
        }
        for name, edit in edits.items():
            with self.subTest(change=name):
                self.reset()
                (self.project / 'PROJECT.md').unlink(missing_ok=True)
                exp = self.preview()['expected_preview']
                edit()
                before = self.index_hash()
                with self.no_transaction():
                    res = self.apply(exp)
                self.assertEqual((res['code'], res['result']['status'], res['result']['reason']), (1, 'blocked', 'PREVIEW_STALE'), res['out'])
                self.assertEqual(self.index_hash(), before)
                if name == 'scope only':
                    self.assertIn('PROJECT.md', res['out'])

    def test_a_plan_that_became_a_no_op_is_still_stale(self):
        exp = self.preview()['expected_preview']
        library_sync.apply(str(self.project))  # một người ghi khác đã áp dụng đúng thay đổi đó / another writer applied the same change
        self.assertIsNone(library_sync.plan(str(self.project))['bytes'])
        before = self.index_hash()
        with self.no_transaction():
            res = self.apply(exp)
        self.assertEqual((res['code'], res['result']['status'], res['result']['reason']), (1, 'blocked', 'PREVIEW_STALE'))
        self.assertEqual(self.index_hash(), before)

    def test_x14_and_unreadable_journal_refusals_still_stop_the_panel_apply(self):
        exp = self.preview()['expected_preview']
        before = self.index_hash()
        with patch.object(library_sync, 'lossy_notices', return_value=['x14_data_validation']), self.no_transaction():  # mô phỏng / simulated
            res = self.apply(self.preview()['expected_preview'])
        self.assertEqual((res['code'], res['result']['status']), (1, 'blocked'))
        self.assertIn('x14', res['out'])
        with patch.object(tx, 'incomplete', side_effect=ValueError('bad journal')), self.no_transaction():  # mô phỏng / simulated
            res = self.apply(exp)
        self.assertEqual((res['code'], res['result']['status']), (1, 'blocked'))
        self.assertIn('journal', res['out'])
        self.assertEqual(self.index_hash(), before)

    def test_missing_openpyxl_gives_not_checked_for_preview(self):
        with patch.dict(sys.modules, {'openpyxl': None}), contextlib.redirect_stderr(io.StringIO()):
            res = app.action('library_sync', {})
        self.assertEqual((res['code'], res['result']['status']), (2, 'not_checked'))

    def test_sync_is_a_logged_action_with_no_names_or_paths_in_the_record(self):
        env = app.run_api({'action': 'library_sync', 'params': {}})[1]['result']
        app.run_api({'action': 'library_sync', 'params': {'apply': True, 'expected_preview': env['expected_preview']}})
        records = ops_log.read(str(self.project))
        self.assertEqual([(r['action'], r['status']) for r in records], [('library_sync', 'preview'), ('library_sync', 'success')])
        text = (self.project / '.agent' / 'runtime' / 'operations.jsonl').read_text(encoding='utf-8')
        self.assertNotIn('a.txt', text)
        for rec in records:
            self.assertEqual(set(rec), ALLOWED_RECORD_FIELDS)


# ------------------------------------------------------------------ D-01b: xem trước và thay một tài liệu / preview and replace one guide
GUIDES = merge_center.REPLACE_ALLOWLIST
OTHER_MANAGED = ('docs/other-notes.md', 'tools/helper.py', 'AGENTS.md', 'docs/HUONG_DAN_SKILLS_VI.html', '.agent/STATE.md')
HEX_A, HEX_B = 'a' * 64, 'b' * 64


def sha(data):
    return hashlib.sha256(data).hexdigest()


def guide_current(rel):
    return ('# Hướng dẫn %s\r\ndòng A cũ\r\ndòng B\r\n' % rel).encode('utf-8')  # CRLF: EOL của người dùng / the user's EOL


def guide_new(rel):
    return ('# Hướng dẫn %s\ndòng A mới\ndòng B\ndòng C thêm\n' % rel).encode('utf-8')  # LF


def build_project(root):
    """Dự án giả: tám tài liệu có .new khác bytes/EOL, file quản lý khác, dấu mốc của người dùng, MANIFEST giả.
    Fake project: eight guides with a different .new (bytes and EOL), other managed files, user markers, a fake MANIFEST."""
    root = Path(root)
    for rel in GUIDES:
        write(root, rel, guide_current(rel))
        write(root, rel + '.new', guide_new(rel))
    for rel in OTHER_MANAGED:
        write(root, rel, 'current ' + rel + '\n')
        write(root, rel + '.new', 'new ' + rel + '\n')
    write(root, 'skills/local/packs/demo/SKILL.md', 'pack current\n')
    write(root, 'skills/local/packs/demo/SKILL.md.new', 'pack new\n')
    write(root, 'output/report.txt', 'USER OUTPUT MARKER\n')
    write(root, 'references/source.txt', 'SOURCE MARKER\n')
    write(root, 'working/notes.txt', 'WORKING MARKER\n')
    write(root, 'unmanaged.md', 'not in the manifest\n')
    write(root, 'unmanaged.md.new', 'not in the manifest\n')
    lines = ['0' * 64 + '  ' + name for name in GUIDES + OTHER_MANAGED]
    write(root, 'MANIFEST.sha256', '\n'.join(lines) + '\n')
    return root


def user_markers(root):
    return {rel: (Path(root) / rel).read_bytes() for rel in ('output/report.txt', 'references/source.txt', 'working/notes.txt',
                                                              'MANIFEST.sha256')}


class MergeReplaceBase(PanelBase):
    def setUp(self):
        super().setUp()
        build_project(self.project)
        self.opened = []
        self.patch(app, 'launch', lambda path: self.opened.append(path))  # chép đường dẫn không được mở gì / copying a path must open nothing

    guide = GUIDES[2]  # docs/HUONG_DAN_SKILLS_VI.md

    def pv(self, rel=None, root=None):
        return merge_center.preview_replace(str(root or self.project), (rel or self.guide) + '.new')

    def go(self, rel=None, pv=None, root=None):
        rel = rel or self.guide
        pv = pv or self.pv(rel, root)
        return merge_center.apply_replace(str(root or self.project), rel + '.new', pv['current_sha256'], pv['candidate_sha256'])

    def raw(self, rel, root=None):
        return (Path(root or self.project) / rel).read_bytes()

    def runtime_exists(self):
        return (self.project / '.agent' / 'runtime').exists()

    def archive_ids(self, root=None):
        base = Path(root or self.project) / '.agent' / 'merge-backups'
        return sorted(p.name for p in base.iterdir()) if base.is_dir() else []

    def assert_nothing_written(self, before):
        self.assertEqual(tree_state(self.project), before)
        self.assertFalse(self.runtime_exists())
        self.assertEqual(self.archive_ids(), [])


class MergeReplacePreview(MergeReplaceBase):
    def test_preview_describes_the_snapshot_and_writes_nothing(self):
        before = tree_state(self.project)
        pv = self.pv()
        self.assertEqual(pv['schema'], merge_center.PREVIEW_SCHEMA)
        self.assertEqual((pv['path'], pv['pending'], pv['eligible'], pv['reason']), (self.guide, self.guide + '.new', True, None))
        self.assertEqual(pv['current_sha256'], sha(self.raw(self.guide)))  # hash thô, gồm cả CRLF / raw hash, CRLF included
        self.assertEqual(pv['candidate_sha256'], sha(self.raw(self.guide + '.new')))
        self.assertTrue(pv['pending_stays'] and not pv['html_rebuilt'])
        self.assertEqual((pv['diff']['added'], pv['diff']['removed']), (2, 1))
        self.assertEqual(tree_state(self.project), before)
        self.assertFalse(self.runtime_exists())
        self.assertFalse((self.project / '.agent' / 'MERGE_STATUS.json').exists())
        self.assertEqual(self.archive_ids(), [])

    def test_hash_and_diff_come_from_the_same_buffers(self):
        current, candidate = self.raw(self.guide), self.raw(self.guide + '.new')
        reads = []
        real = merge_center._capture

        def counting(root, rel):
            reads.append(rel)
            return real(root, rel)
        with patch.object(merge_center, '_capture', counting):
            pv = self.pv()
        self.assertEqual(reads, [self.guide, self.guide + '.new'])  # mỗi file đọc đúng một lần / each file read once
        self.assertEqual((pv['current_sha256'], pv['candidate_sha256']), (sha(current), sha(candidate)))

    def test_unsupported_and_invalid_paths_are_refused_before_any_read(self):
        before = tree_state(self.project)
        cases = {
            'unsupported': ['docs/other-notes.md.new', 'tools/helper.py.new', 'AGENTS.md.new', 'docs/HUONG_DAN_SKILLS_VI.html.new',
                            '.agent/STATE.md.new', 'skills/local/packs/demo/SKILL.md.new', 'unmanaged.md.new',
                            ' docs/HUONG_DAN_SKILLS_VI.md.new', 'docs/HUONG_DAN_SKILLS_VI.md.new.new'],
            'invalid_path': ['', '/docs/HUONG_DAN_SKILLS_VI.md.new', 'C:/docs/HUONG_DAN_SKILLS_VI.md.new', '//server/share/x.md.new',
                             '../docs/HUONG_DAN_SKILLS_VI.md.new', 'docs/../docs/HUONG_DAN_SKILLS_VI.md.new', './docs/HUONG_DAN_SKILLS_VI.md.new',
                             'docs\\HUONG_DAN_SKILLS_VI.md.new', 'docs//HUONG_DAN_SKILLS_VI.md.new', 'docs/HUONG_DAN_SKILLS_VI.md.new ',
                             'docs/HUONG_DAN_SKILLS_VI.md.new.', 'docs /HUONG_DAN_SKILLS_VI.md.new',
                             'docs/HUONG_DAN_SKILLS_VI.md:stream.new', 'docs/HUONG_DAN_SKILLS_VI.md.new:$DATA',
                             'docs/huong_dan_skills_vi.md.new', 'DOCS/HUONG_DAN_SKILLS_VI.md.new', 'docs/HUONG_DAN_SKILLS_VI.md',
                             'docs/HUONG_DAN_SKILLS_VI.md.NEW', 'docs/HUONG_DAN_SKILLS_VI.md\n.new', None, 5, ['x'], {'a': 1}],
        }
        for reason, paths in cases.items():
            for path in paths:
                with self.subTest(path=path):
                    pv = merge_center.preview_replace(str(self.project), path)
                    self.assertFalse(pv['eligible'])
                    self.assertEqual(pv['reason'], reason)
                    self.assertIsNone(pv['current_sha256'])
                    res = merge_center.apply_replace(str(self.project), path, HEX_A, HEX_B)
                    self.assertEqual((res['outcome'], res['reason']), ('blocked', reason))
        self.assertEqual(tree_state(self.project), before)
        self.assertFalse(self.runtime_exists())

    def test_allowlist_pair_must_still_be_managed(self):
        write(self.project, 'MANIFEST.sha256', '0' * 64 + '  other.txt\n')
        pv = self.pv()
        self.assertEqual((pv['eligible'], pv['reason']), (False, 'not_managed'))

    def test_content_limits_make_a_pair_ineligible(self):
        write(self.project, self.guide, b'\x00\x01\x02')
        self.assertEqual(self.pv()['reason'], 'binary')
        write(self.project, self.guide, 'ok\n')
        write(self.project, self.guide + '.new', b'\xff\xfe\xfd')
        self.assertEqual(self.pv()['reason'], 'binary')  # không phải UTF-8 / not UTF-8
        write(self.project, self.guide + '.new', b'text\n\x00 hidden after the first block\n')
        self.assertEqual(self.pv()['reason'], 'binary')  # NUL ở bất kỳ vị trí nào / NUL anywhere
        write(self.project, self.guide + '.new', 'line\n' * 50)
        with patch.object(merge_center, 'DIFF_MAX_BYTES', 100):
            self.assertEqual(self.pv()['reason'], 'too_large')
        write(self.project, self.guide, 'a\n' * 40)
        write(self.project, self.guide + '.new', 'b\n' * 40)
        with patch.object(merge_center, 'DIFF_MAX_ROWS', 5):
            pv = self.pv()
        self.assertEqual((pv['eligible'], pv['reason'], pv['diff']['truncated']), (False, 'truncated', True))

    def test_missing_files_and_read_errors_are_told_apart(self):
        (self.project / self.guide).unlink()
        self.assertEqual(self.pv()['reason'], 'current_missing')
        write(self.project, self.guide, guide_current(self.guide))
        (self.project / (self.guide + '.new')).unlink()
        self.assertEqual(self.pv()['reason'], 'candidate_missing')
        write(self.project, self.guide + '.new', guide_new(self.guide))
        real = os.lstat

        def deny(path, *a, **k):
            if str(path).replace('\\', '/').endswith(self.guide):
                raise PermissionError(13, 'denied')
            return real(path, *a, **k)
        with patch.object(merge_center.os, 'lstat', deny):
            pv = self.pv()
        self.assertEqual((pv['eligible'], pv['reason']), (False, 'read_error'))  # không coi là thiếu file / not read as a missing file
        (self.project / self.guide).unlink()
        (self.project / self.guide).mkdir()
        self.assertEqual(self.pv()['reason'], 'not_regular')

    def test_identical_pair_is_ineligible_but_reported(self):
        write(self.project, self.guide + '.new', self.raw(self.guide))
        pv = self.pv()
        self.assertEqual((pv['eligible'], pv['identical'], pv['reason']), (False, True, 'identical'))

    def test_unfinished_or_unreadable_journal_makes_the_pair_ineligible(self):
        folder = self.project / '.agent' / 'runtime' / 'transactions' / '20260101-000000-aaaaaaaa'
        journal = {'schema': tx.SCHEMA, 'id': folder.name, 'operation': 'upgrade', 'status': 'committing', 'ops': []}
        write(folder, 'journal.json', json.dumps(journal))
        pv = self.pv()
        self.assertEqual((pv['eligible'], pv['reason']), (False, 'transaction_pending'))
        self.assertEqual(pv['unfinished'][0]['id'], folder.name)
        write(folder, 'journal.json', '{not json')
        self.assertEqual(self.pv()['reason'], 'journal_unreadable')  # không bỏ qua bằng non-strict / not skipped like non-strict

    def test_links_at_the_pair_archive_and_runtime_are_refused(self):
        outside = self.root / 'outside'
        (outside / 'en').mkdir(parents=True)
        for rel in ('docs/en/NEW_PROJECT_GUIDE.md', 'docs/en/NEW_PROJECT_GUIDE.md.new'):
            write(outside / 'en', rel.rsplit('/', 1)[1], b'outside bytes\n')
        shutil.rmtree(str(self.project / 'docs' / 'en'))
        link = self.project / 'docs' / 'en'
        if not dir_link(link, outside / 'en'):
            self.skipTest('this machine cannot create a directory link (symlink or junction)')
        try:
            outside_before = tree_state(outside)
            guide = 'docs/en/NEW_PROJECT_GUIDE.md'
            pv = self.pv(guide)
            self.assertEqual((pv['eligible'], pv['reason']), (False, 'unsafe_link'))
            res = merge_center.apply_replace(str(self.project), guide + '.new', HEX_A, HEX_B)
            self.assertEqual((res['outcome'], res['reason']), ('blocked', 'unsafe_link'))
            self.assertEqual(tree_state(outside), outside_before)
        finally:
            drop_link(link)

    def test_archive_and_runtime_links_are_refused(self):
        for name, reason in (('merge-backups', 'archive_unsafe'), ('runtime', 'runtime_unsafe')):
            with self.subTest(link=name):
                outside = self.root / ('out-' + name)
                outside.mkdir()
                (self.project / '.agent').mkdir(exist_ok=True)
                link = self.project / '.agent' / name
                if not dir_link(link, outside):
                    self.skipTest('this machine cannot create a directory link (symlink or junction)')
                try:
                    pv = self.pv()
                    self.assertEqual((pv['eligible'], pv['reason']), (False, reason))
                    res = self.go(pv=pv)
                    self.assertEqual((res['outcome'], res['reason']), ('blocked', reason))
                    self.assertEqual(list(outside.iterdir()), [])
                    self.assertEqual(self.raw(self.guide), guide_current(self.guide))
                finally:
                    drop_link(link)


class MergeReplaceApply(MergeReplaceBase):
    def test_every_allowlisted_guide_is_replaced_with_both_versions_kept(self):
        for n, rel in enumerate(GUIDES):
            with self.subTest(guide=rel):
                root = build_project(self.root / ('p%d' % n))
                old, new = guide_current(rel), guide_new(rel)
                self.assertNotEqual(old.count(b'\r\n'), new.count(b'\r\n'))
                pv = self.pv(rel, root)
                res = self.go(rel, pv, root)
                self.assertEqual((res['outcome'], res['reason']), ('replaced', None), res)
                self.assertEqual(self.raw(rel, root), new)  # bytes và EOL của .new / bytes and EOL of the .new
                self.assertEqual(self.raw(rel + '.new', root), new)
                arc = res['archive']
                self.assertEqual(self.raw(arc['current'], root), old)
                self.assertEqual(self.raw(arc['incoming'], root), new)

    def test_success_keeps_new_manifest_status_markers_and_html(self):
        merge_center.set_state(str(self.project), self.guide + '.new', 'reviewing')
        html = write(self.project, 'docs/HUONG_DAN_SKILLS_VI.html', '<html>built earlier</html>').read_bytes()
        status = self.raw('.agent/MERGE_STATUS.json')
        markers = user_markers(self.project)
        others = {rel: self.raw(rel) for rel in OTHER_MANAGED + tuple(o + '.new' for o in OTHER_MANAGED)}
        old, new = self.raw(self.guide), self.raw(self.guide + '.new')
        pending_before = merge_center.scan(str(self.project))
        res = self.go()
        self.assertEqual(res['outcome'], 'replaced')
        self.assertEqual(self.raw(self.guide), new)
        self.assertEqual(self.raw(self.guide + '.new'), new)
        self.assertEqual(self.raw('.agent/MERGE_STATUS.json'), status)
        self.assertEqual(user_markers(self.project), markers)
        self.assertEqual(self.raw('docs/HUONG_DAN_SKILLS_VI.html'), html)
        self.assertEqual({rel: self.raw(rel) for rel in others}, others)
        pending_after = merge_center.scan(str(self.project))
        self.assertEqual(pending_after, pending_before)  # vẫn pending, trạng thái giữ nguyên / still pending, state kept
        self.assertTrue(res['pending_stays'] and not res['html_rebuilt'])
        arc = res['archive']
        self.assertEqual(arc['dir'], '.agent/merge-backups/' + arc['id'])
        self.assertEqual(sorted(p.name for p in (self.project / arc['dir']).iterdir()), ['current.txt', 'incoming.txt', 'record.json'])
        self.assertEqual((self.raw(arc['current']), self.raw(arc['incoming'])), (old, new))
        record = json.loads(self.raw(arc['record']).decode('utf-8'))
        self.assertEqual(record, {'schema': merge_center.RECORD_SCHEMA, 'record_version': 1, 'path': self.guide,
                                  'current_sha256': sha(old), 'incoming_sha256': sha(new)})
        self.assertEqual((arc['current_sha256'], arc['incoming_sha256']), (sha(old), sha(new)))
        text = self.raw(arc['record']).decode('utf-8')
        self.assertNotIn(str(self.project), text)
        self.assertNotIn('Hướng dẫn', text)  # không có nội dung / no content
        self.assertEqual(tx.incomplete(str(self.project)), [])

    def test_the_four_operations_share_one_transaction_and_replace_comes_last(self):
        res = self.go()
        journal = json.loads(self.raw('.agent/runtime/transactions/%s/journal.json' % res['txid']).decode('utf-8'))
        self.assertEqual(journal['status'], 'committed')
        ops = journal['ops']
        arc = res['archive']
        self.assertEqual([o['path'] for o in ops], [arc['current'], arc['incoming'], arc['record'], self.guide])
        self.assertEqual([o['action'] for o in ops], ['create_exclusive'] * 3 + ['write'])
        self.assertEqual([o['before'] is None for o in ops], [True, True, True, False])
        self.assertEqual(ops[3]['before'], sha(guide_current(self.guide)))

    def test_old_version_can_be_recovered_by_hand_into_a_new_file(self):
        old = self.raw(self.guide)
        res = self.go()
        recovered = self.project / 'working' / 'HUONG_DAN_SKILLS_VI.recovered.md'
        shutil.copyfile(str(self.project / res['archive']['current']), str(recovered))
        self.assertEqual(recovered.read_bytes(), old)  # đúng từng byte gồm CRLF / byte for byte, CRLF included
        self.assertNotEqual(self.raw(self.guide), old)

    def test_identical_pair_is_a_noop_without_runtime_or_archive(self):
        write(self.project, self.guide + '.new', self.raw(self.guide))
        before = tree_state(self.project)
        pv = self.pv()
        res = self.go(pv=dict(pv, current_sha256=sha(self.raw(self.guide)), candidate_sha256=sha(self.raw(self.guide))))
        self.assertEqual((res['outcome'], res['reason']), ('noop', 'identical'))
        self.assertEqual(tree_state(self.project), before)
        self.assertFalse(self.runtime_exists())
        self.assertEqual([m['pending'] for m in merge_center.scan(str(self.project)) if m['path'] == self.guide], [self.guide + '.new'])

    def test_identical_pair_still_refuses_an_unreadable_journal(self):
        write(self.project, self.guide + '.new', self.raw(self.guide))
        pv = self.pv()
        write(self.project, '.agent/runtime/transactions/broken/journal.json', '{')
        before = tree_state(self.project)
        result = self.go(pv=pv)
        self.assertEqual((result['outcome'], result['reason']), ('blocked', 'journal_unreadable'))
        self.assertEqual(tree_state(self.project), before)

    def test_forged_or_malformed_hashes_never_pass_the_server_checks(self):
        before = tree_state(self.project)
        pv = self.pv()
        for cur, cand in ((HEX_A, HEX_B), (pv['current_sha256'], HEX_B), (HEX_A, pv['candidate_sha256']),
                          (pv['current_sha256'].upper(), pv['candidate_sha256']), (None, None), ('', ''), (5, [1])):
            with self.subTest(hashes=str(cur)[:6]):
                res = merge_center.apply_replace(str(self.project), self.guide + '.new', cur, cand)
                self.assertIn(res['outcome'], ('stale', 'blocked'))
        self.assertEqual(tree_state(self.project), before)

    def test_ineligible_pairs_are_not_applied_even_with_their_real_hashes(self):
        before = tree_state(self.project)
        for rel in OTHER_MANAGED + ('unmanaged.md', 'skills/local/packs/demo/SKILL.md'):
            with self.subTest(path=rel):
                res = merge_center.apply_replace(str(self.project), rel + '.new', sha(self.raw(rel)), sha(self.raw(rel + '.new')))
                self.assertEqual((res['outcome'], res['reason']), ('blocked', 'unsupported'))
        self.assertEqual(tree_state(self.project), before)

    def test_current_or_candidate_changed_after_preview_is_stale_and_writes_nothing(self):
        for which in ('current', 'candidate'):
            with self.subTest(changed=which):
                root = build_project(self.root / ('stale-' + which))
                pv = self.pv(root=root)
                rel = self.guide + ('' if which == 'current' else '.new')
                edit = b'USER EDIT AFTER PREVIEW\n'
                write(root, rel, edit)
                before = tree_state(root)
                res = self.go(pv=pv, root=root)
                self.assertEqual((res['outcome'], res['reason']), ('stale', 'stale_preview'))
                self.assertEqual(tree_state(root), before)
                self.assertEqual(self.raw(rel, root), edit)
                self.assertFalse((root / '.agent' / 'runtime').exists())
                self.assertEqual(self.archive_ids(root), [])

    def test_archive_collision_blocks_without_overwrite_or_another_id(self):
        write(self.project, '.agent/merge-backups/fixed/current.txt', 'EXISTING ARCHIVE MARKER')
        calls = []

        def fixed():
            calls.append(1)
            return 'fixed'
        before = tree_state(self.project)
        with patch.object(merge_center, '_new_id', fixed):
            res = self.go()
        self.assertEqual((res['outcome'], res['reason']), ('blocked', 'archive_collision'))
        self.assertEqual(len(calls), 1)  # không vòng đổi id / no retry loop
        self.assertEqual(tree_state(self.project), before)
        self.assertFalse(self.runtime_exists())

    def test_unfinished_transaction_blocks_apply(self):
        pv = self.pv()
        folder = self.project / '.agent' / 'runtime' / 'transactions' / '20260101-000000-bbbbbbbb'
        write(folder, 'journal.json', json.dumps({'schema': tx.SCHEMA, 'id': folder.name, 'operation': 'upgrade',
                                                   'status': 'prepared', 'ops': []}))
        before = tree_state(self.project)
        res = self.go(pv=pv)
        self.assertEqual((res['outcome'], res['reason'], res['unfinished']), ('blocked', 'transaction_pending', [folder.name]))
        self.assertEqual(tree_state(self.project), before)

    def test_candidate_changed_before_the_run_is_refused(self):
        pv = self.pv()

        def hook(point):
            if point == 'captured':
                write(self.project, self.guide + '.new', b'CHANGED RIGHT AFTER CAPTURE\n')
        with patch.object(merge_center, '_HOOK', hook):
            res = self.go(pv=pv)
        self.assertEqual((res['outcome'], res['reason']), ('stale', 'candidate_changed'))
        self.assertEqual(self.raw(self.guide), guide_current(self.guide))
        self.assertEqual(self.raw(self.guide + '.new'), b'CHANGED RIGHT AFTER CAPTURE\n')
        self.assertFalse(self.runtime_exists())
        self.assertEqual(self.archive_ids(), [])

    def test_candidate_changed_late_keeps_the_live_file_and_only_the_snapshot_is_used(self):
        later = b'LATER VERSION OF THE NEW COPY\n'
        snapshot_bytes = self.raw(self.guide + '.new')
        for point, patcher in (('before_run', merge_center), ('stage', tx)):
            with self.subTest(point=point):
                root = build_project(self.root / ('late-' + point))
                pv = self.pv(root=root)
                done = []

                def change(*args):
                    if not done and args[0] == point:
                        done.append(1)
                        write(root, self.guide + '.new', later)
                with patch.object(patcher, '_HOOK', change):
                    res = self.go(pv=pv, root=root)
                self.assertEqual(done, [1])
                self.assertEqual(res['outcome'], 'replaced', res)
                self.assertEqual(self.raw(self.guide, root), snapshot_bytes)  # đúng snapshot đã xem / exactly the previewed snapshot
                self.assertEqual(self.raw(res['archive']['incoming'], root), snapshot_bytes)  # không pha hai lần đọc / no mix of two reads
                self.assertEqual(self.raw(self.guide + '.new', root), later)  # file live mới được giữ / the live file is kept
                self.assertTrue(res['candidate_live_changed'])

    def test_nothing_is_written_when_the_failure_comes_before_the_first_file(self):
        for fault in ('stage:2', 'backup:1'):
            with self.subTest(fault=fault):
                root = build_project(self.root / ('early-' + fault.replace(':', '-')))
                with patch.dict(os.environ, {FAULT: fault}):
                    res = self.go(root=root)
                self.assertEqual((res['outcome'], res['reason']), ('error', 'write_failed'), res)
                self.assertEqual(self.raw(self.guide, root), guide_current(self.guide))
                self.assertEqual(self.archive_ids(root), [])
                self.assertEqual(tx.incomplete(str(root)), [])  # không chặn việc sau / does not block later work

    def test_fault_after_the_first_archive_file_is_reported_as_unfinished_and_resume_completes(self):
        old, new = self.raw(self.guide), self.raw(self.guide + '.new')
        with patch.dict(os.environ, {FAULT: 'commit:1'}):
            res = self.go()
        self.assertEqual((res['outcome'], res['reason']), ('partial', 'transaction_unfinished'))
        self.assertEqual((res['done'], res['total'], res['current_state']), (1, 4, 'unchanged'))
        self.assertEqual(res['archive']['saved'], {'current': res['archive']['current']})
        self.assertEqual(self.raw(res['archive']['current']), old)
        self.assertEqual(self.raw(self.guide), old)
        self.assertEqual([i['id'] for i in tx.incomplete(str(self.project))], [res['txid']])
        out = tx.resume(str(self.project), res['txid'])
        self.assertEqual(out['status'], 'success', out)
        self.assertEqual(self.raw(self.guide), new)
        self.assertEqual((self.raw(res['archive']['current']), self.raw(res['archive']['incoming'])), (old, new))
        self.assertEqual(json.loads(self.raw(res['archive']['record']).decode('utf-8'))['incoming_sha256'], sha(new))
        self.assertEqual(self.raw(self.guide + '.new'), new)
        self.assertEqual(tx.incomplete(str(self.project)), [])

    def test_fault_before_replace_keeps_both_versions_and_rollback_leaves_current_intact(self):
        old, new = self.raw(self.guide), self.raw(self.guide + '.new')
        with patch.dict(os.environ, {FAULT: 'commit:3'}):
            res = self.go()
        self.assertEqual((res['outcome'], res['reason'], res['current_state']), ('partial', 'transaction_unfinished', 'unchanged'))
        arc = res['archive']
        self.assertEqual(sorted(arc['saved']), ['current', 'incoming', 'record'])
        self.assertEqual((self.raw(arc['current']), self.raw(arc['incoming'])), (old, new))
        self.assertEqual(self.raw(self.guide), old)
        out = tx.rollback(str(self.project), res['txid'])
        self.assertEqual(out['status'], 'success', out)
        self.assertEqual(self.raw(self.guide), old)
        self.assertEqual(self.raw(self.guide + '.new'), new)
        self.assertEqual(tx.incomplete(str(self.project)), [])

    def test_user_edit_during_the_replace_is_kept_by_rollback(self):
        edit = 'USER EDIT DURING THE REPLACE\n'.encode('utf-8')

        def hook(point, path):
            if point == 'commit' and path.replace('\\', '/').endswith('/' + self.guide):
                write(self.project, self.guide, edit)
        with patch.object(tx, '_HOOK', hook):
            res = self.go()
        self.assertEqual((res['outcome'], res['reason'], res['current_state']), ('partial', 'transaction_unfinished', 'unknown'))
        self.assertEqual(self.raw(self.guide), edit)
        self.assertEqual(sorted(res['archive']['saved']), ['current', 'incoming', 'record'])
        out = tx.rollback(str(self.project), res['txid'])
        self.assertEqual(out['status'], 'partial', out)
        self.assertEqual(self.raw(self.guide), edit)  # rollback không đè bản người dùng sửa / rollback does not overwrite the user's edit
        self.assertEqual(self.raw(self.guide + '.new'), guide_new(self.guide))


class MergeReplaceHttp(MergeReplaceBase):
    def serve(self):
        srv = app.Server(('127.0.0.1', 0), app.Handler)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        self.addCleanup(lambda: (srv.shutdown(), srv.server_close()))
        self.patch(app, 'TOKEN', 'test-token')
        return 'http://127.0.0.1:%d' % srv.server_address[1]

    def post(self, base, path, body, token='test-token'):
        headers = {'Content-Type': 'application/json'}
        if token:
            headers['X-Panel-Token'] = token
        # token sai: gửi thân rỗng, vì bảng điều khiển từ chối trước khi đọc thân (thân chưa đọc làm Windows ngắt kết nối)
        # bad token: send an empty body, because the panel refuses before reading the body (unread bytes make Windows reset the connection)
        data = json.dumps(body).encode('utf-8') if token == 'test-token' else b''
        req = urllib.request.Request(base + path, data=data, method='POST', headers=headers)
        try:
            with urllib.request.urlopen(req, timeout=30) as r:
                return r.status, json.loads(r.read().decode('utf-8'))
        except urllib.error.HTTPError as e:
            return e.code, json.loads(e.read().decode('utf-8') or '{}')

    def test_preview_route_is_read_only_checks_the_token_and_logs_nothing(self):
        base = self.serve()
        before = tree_state(self.project)
        for token in (None, 'wrong'):
            self.assertEqual(self.post(base, '/api/merge/preview', {'path': self.guide + '.new'}, token)[0], 403)
        code, body = self.post(base, '/api/merge/preview', {'path': self.guide + '.new'})
        self.assertEqual(code, 200)
        pv = body['replace_preview']
        self.assertEqual((body['status'], body['changes'], pv['eligible'], pv['path']), ('success', False, True, self.guide))
        self.assertEqual(self.post(base, '/api/merge/preview', {'path': '../x.new'})[1]['status'], 'blocked')
        self.assertEqual(tree_state(self.project), before)
        self.assertFalse(self.runtime_exists())  # không runtime, không operations.jsonl / no runtime, no operations.jsonl
        self.assertEqual(ops_log.read(str(self.project)), [])

    def test_replace_route_needs_the_preview_hashes_and_is_not_in_the_operations_log(self):
        base = self.serve()
        code, body = self.post(base, '/api/merge/preview', {'path': self.guide + '.new'})
        pv = body['replace_preview']
        self.assertEqual(self.post(base, '/api/merge/replace', {'path': self.guide + '.new'}, None)[0], 403)
        code, blocked = self.post(base, '/api/merge/replace', {'path': self.guide + '.new'})
        self.assertEqual((blocked['status'], blocked['replace']['reason']), ('blocked', 'bad_request'))
        self.assertEqual(self.raw(self.guide), guide_current(self.guide))
        code, ok = self.post(base, '/api/merge/replace', {'path': self.guide + '.new', 'current_sha256': pv['current_sha256'],
                                                          'candidate_sha256': pv['candidate_sha256']})
        self.assertEqual((code, ok['status'], ok['code'], ok['changes']), (200, 'success', 0, True))
        self.assertEqual(ok['replace']['outcome'], 'replaced')
        self.assertEqual([f['action'] for f in ok['files']], ['replaced', 'created', 'created', 'created'])
        self.assertEqual(self.raw(self.guide), guide_new(self.guide))
        self.assertFalse((self.project / '.agent' / 'runtime' / 'operations.jsonl').exists())
        self.assertEqual(self.opened, [])

    def test_envelope_statuses_follow_the_outcome(self):
        cases = {'noop': ('success', 0, False), 'stale': ('blocked', 1, False), 'error': ('error', 1, False),
                 'partial': ('partial', 1, None)}
        for outcome, (status_value, code, changes) in cases.items():
            with self.subTest(outcome=outcome):
                fake = {'outcome': outcome, 'reason': 'x', 'path': self.guide, 'pending': self.guide + '.new', 'txid': 'T1'}
                with patch.object(merge_center, 'apply_replace', return_value=fake):
                    env = app.merge_api('/api/merge/replace', {})
                self.assertEqual((env['status'], env['code'], env['changes']), (status_value, code, changes))
        self.assertIsNone(app.merge_api('/api/merge/other', {}))

    def test_keep_and_copy_never_reach_the_state_writer_or_the_opener(self):
        with patch.object(merge_center, 'set_state', side_effect=AssertionError('state write')):
            app.merge_api('/api/merge/preview', {'path': self.guide + '.new'})
            app.merge_api('/api/merge/replace', {'path': self.guide + '.new', 'current_sha256': HEX_A, 'candidate_sha256': HEX_B})
        self.assertEqual(self.opened, [])
        self.assertFalse((self.project / '.agent' / 'MERGE_STATUS.json').exists())


UI_FILE = ROOT / 'tools' / 'app_ui.html'
HARNESS = r"""
'use strict';
const calls = [], toasts = [], copied = []; let closed = 0, refreshed = 0;
const box = {innerHTML: ''}, els = {}, RESP = {};
function $(id) {
  if (id === 'sh-diff') return {querySelector: function () { return box; }, addEventListener: function () {}, insertAdjacentHTML: function () {}};
  return els[id] || (els[id] = {innerHTML: '', textContent: '', addEventListener: function () {}});
}
function t(k) { return k === 'mp_partial' ? 'mp_partial:%s' : k; }
function esc(s) { return String(s == null ? '' : s).replace(/[&<>"']/g, function (c) { return {'&': '&amp;', '<': '&lt;', '>': '&gt;', '"': '&quot;', "'": '&#39;'}[c]; }); }
function ico() { return ''; }
const drawn = [];
function drawDiff(d) { drawn.push(d); }
function copyText(x) { copied.push(x); }
function toast(m) { toasts.push(m); }
function closeSheet() { closed++; mpReset(); }
function refresh() { refreshed++; return Promise.resolve(); }
function api(p, b) { calls.push([p, JSON.parse(JSON.stringify(b))]); const r = RESP[p]; return Promise.resolve(typeof r === 'function' ? r(b) : r); }
%BLOCK%
const PV = {schema: 'local-agent-workspace/merge-replace-preview-v1', path: 'docs/a.md', pending: 'docs/a.md.new', eligible: true,
            reason: null, current_sha256: 'a'.repeat(64), candidate_sha256: 'b'.repeat(64)};
const applies = () => calls.filter(c => c[0] === '/api/merge/replace');
(async function () {
  const out = {};
  await mpAct('use'); out.no_preview_calls = calls.length;
  RESP['/api/merge/preview'] = {replace_preview: Object.assign({}, PV, {eligible: false, reason: 'binary'})};
  await mpLoad('docs/a.md.new', 'docs/a.md'); out.inelig_html = box.innerHTML; out.inelig_mp = MP;
  await mpAct('use'); out.inelig_applies = applies().length;
  RESP['/api/merge/preview'] = {replace_preview: Object.assign({}, PV, {schema: 'forged'})};
  await mpLoad('docs/a.md.new', 'docs/a.md'); out.forged_mp = MP;
  RESP['/api/merge/preview'] = {replace_preview: Object.assign({}, PV, {pending: 'docs/other.md.new'})};
  await mpLoad('docs/a.md.new', 'docs/a.md'); out.other_pair_mp = MP;
  RESP['/api/merge/preview'] = {replace_preview: Object.assign({}, PV, {reason: 'unsupported', eligible: false})};
  await mpLoad('docs/a.md.new', 'docs/a.md'); out.unsupported_html = box.innerHTML;
  out.drawn_before = drawn.length;
  RESP['/api/merge/preview'] = {replace_preview: Object.assign({}, PV, {diff: {added: 2, removed: 1, hunks: []}})};
  await mpLoad('docs/a.md.new', 'docs/a.md'); out.drawn_after = drawn.slice();
  RESP['/api/merge/preview'] = {replace_preview: PV};
  calls.length = 0;
  await mpLoad('docs/a.md.new', 'docs/a.md');
  out.eligible_html = box.innerHTML; out.after_load = calls.map(c => c[0]);
  RESP['/api/merge/replace'] = {replace: {outcome: 'replaced', archive: {current: '.agent/merge-backups/x/current.txt', incoming: '.agent/merge-backups/x/incoming.txt'}, candidate_live_changed: true}};
  await mpAct('use');
  out.after_use = calls.map(c => c); out.done_html = box.innerHTML; out.done_refresh = refreshed;
  await mpAct('use'); out.after_second_use = applies().length;
  // keep: no server call, no state write
  calls.length = 0; closed = 0;
  await mpLoad('docs/a.md.new', 'docs/a.md');
  await mpAct('keep'); out.keep_calls = calls.map(c => c[0]); out.keep_closed = closed; out.keep_toast = toasts.slice(); out.keep_mp = MP;
  await mpAct('use'); out.keep_then_use = applies().length;
  // manual: shows the paths, copies on request, never calls the server
  calls.length = 0;
  await mpLoad('docs/a.md.new', 'docs/a.md'); calls.length = 0;
  await mpAct('manual'); out.manual_html = box.innerHTML; out.manual_calls = calls.length; out.manual_mp = MP;
  await mpAct('copy', 'docs/a.md'); out.copied = copied.slice();
  // a slow older preview must not replace a newer one
  let slow;
  RESP['/api/merge/preview'] = () => new Promise(r => { slow = r; });
  const first = mpLoad('docs/a.md.new', 'docs/a.md');
  RESP['/api/merge/preview'] = {replace_preview: Object.assign({}, PV, {eligible: false, reason: 'identical'})};
  await mpLoad('docs/a.md.new', 'docs/a.md');
  slow({replace_preview: PV}); await first; out.race_mp = MP;
  // every outcome has a message and a stable status; refresh only when something may have been written
  const outcomes = {noop: {outcome: 'noop'}, stale: {outcome: 'stale'}, collision: {outcome: 'blocked', reason: 'archive_collision'},
                    pending: {outcome: 'blocked', reason: 'transaction_pending'}, partial: {outcome: 'partial', reason: 'transaction_unfinished', txid: 'TX-1'},
                    unproven: {outcome: 'partial', reason: 'verify_mismatch', txid: 'TX-2'}, error: {outcome: 'error'}, net: null};
  out.outcomes = {};
  for (const k of Object.keys(outcomes)) {
    RESP['/api/merge/preview'] = {replace_preview: PV};
    RESP['/api/merge/replace'] = outcomes[k] === null ? () => Promise.reject(new Error('offline')) : {replace: outcomes[k]};
    await mpLoad('docs/a.md.new', 'docs/a.md'); const r0 = refreshed;
    await mpAct('use'); out.outcomes[k] = {html: box.innerHTML, refreshed: refreshed - r0};
  }
  console.log(JSON.stringify(out));
})();
"""


def ui_block():
    text = UI_FILE.read_text(encoding='utf-8')
    return text[text.index('// D01B-BEGIN'):text.index('// D01B-END')]


def ui_dictionary():
    text = UI_FILE.read_text(encoding='utf-8')
    block = text[text.index('var D={'):text.index('\n};', text.index('var D={'))]
    return {m.group(1): (m.group(2), m.group(3)) for m in re.finditer(
        r"([A-Za-z_][A-Za-z0-9_]*)\s*:\s*\[\s*'((?:[^'\\]|\\.)*)'\s*,\s*'((?:[^'\\]|\\.)*)'\s*\]", block)}


@unittest.skipUnless(shutil.which('node'), 'node is needed to run the panel script outside a browser')
class MergeReplaceUiScript(unittest.TestCase):
    """Chạy đúng khối JS của bảng merge trong node với api giả: chứng minh luồng, không chứng minh giao diện trực quan.
    Runs the real merge-panel script block in node with a fake api: it proves the flow, not the visual look in a browser."""

    @classmethod
    def setUpClass(cls):
        with tempfile.TemporaryDirectory(prefix='merge-ui-') as tmp:
            script = Path(tmp) / 'harness.js'
            script.write_text(HARNESS.replace('%BLOCK%', ui_block()), encoding='utf-8')
            done = subprocess.run(['node', str(script)], capture_output=True, encoding='utf-8', timeout=60)
        assert done.returncode == 0, done.stdout + done.stderr
        cls.out = json.loads(done.stdout.strip().splitlines()[-1])

    def test_apply_click_does_nothing_without_an_eligible_preview(self):
        o = self.out
        self.assertEqual(o['no_preview_calls'], 0)
        self.assertIsNone(o['inelig_mp'])
        self.assertNotIn('data-mp="use"', o['inelig_html'])
        self.assertEqual(o['inelig_applies'], 0)
        self.assertIsNone(o['forged_mp'])  # schema lạ không thành preview / a foreign schema is not a preview
        self.assertIsNone(o['other_pair_mp'])  # preview của cặp khác không dùng được / a preview of another pair is not usable
        self.assertNotIn('data-mp="use"', o['unsupported_html'])
        self.assertIn('mp_unsupported', o['unsupported_html'])

    def test_loading_a_preview_does_not_apply_and_the_click_sends_the_previewed_hashes_once(self):
        o = self.out
        self.assertEqual((o['drawn_before'], o['drawn_after']), (0, [{'added': 2, 'removed': 1, 'hunks': []}]))  # diff hiển thị lấy từ chính preview / the shown diff is the preview's own
        self.assertEqual(o['after_load'], ['/api/merge/preview'])  # mở diff/preview không tự áp / opening never applies
        for key in ('mp_will_cur', 'mp_will_save', 'mp_will_pend', 'mp_will_html'):
            self.assertIn(key, o['eligible_html'])
        self.assertIn('data-mp="use"', o['eligible_html'])
        self.assertEqual(o['after_use'][1], ['/api/merge/replace', {'path': 'docs/a.md.new', 'current_sha256': 'a' * 64, 'candidate_sha256': 'b' * 64}])
        self.assertEqual(o['after_second_use'], 1)  # preview đã dùng, không gửi lần hai / the preview is spent
        self.assertIn('mp_done', o['done_html'])
        self.assertIn('mp_live_changed', o['done_html'])
        self.assertIn('.agent/merge-backups/x/current.txt', o['done_html'])
        self.assertEqual(o['done_refresh'], 1)

    def test_keep_and_manual_never_call_the_server_or_a_state_writer(self):
        o = self.out
        self.assertEqual(o['keep_calls'], ['/api/merge/preview'])  # chỉ lần xem trước, không merge_state/replace / only the preview
        self.assertEqual((o['keep_closed'], o['keep_toast'], o['keep_mp']), (1, ['mp_kept'], None))
        self.assertEqual(o['keep_then_use'], 0)  # sau khi giữ, bấm dùng không gửi gì / after keeping, a use click sends nothing
        self.assertEqual(o['manual_calls'], 0)
        self.assertIsNone(o['manual_mp'])
        self.assertIn('docs/a.md', o['manual_html'])
        self.assertIn('docs/a.md.new', o['manual_html'])
        self.assertEqual(o['copied'], ['docs/a.md'])

    def test_an_older_slow_preview_cannot_replace_a_newer_one(self):
        self.assertIsNone(self.out['race_mp'])

    def test_late_legacy_comparison_cannot_rebind_another_file_or_reopen_a_closed_sheet(self):
        text = UI_FILE.read_text(encoding='utf-8')
        legacy = text[text.index('function openDiff('):text.index('function drawDiff(')]
        prefix = HARNESS.split('(async function () {')[0].replace('%BLOCK%', ui_block())
        script = prefix + legacy + r"""
function openSheet(title) { $('sh-t').textContent = title; }
const oldDollar = $;
$ = function(id) { const x = oldDollar(id); x.setAttribute = function() {}; return x; };
(async function() {
  let resolveOld;
  RESP['/api/run'] = () => new Promise(r => { resolveOld = r; });
  const old = openDiff('docs/a.md.new', 'docs/a.md');
  RESP['/api/run'] = {result: {diff: {path: 'docs/b.md'}}};
  RESP['/api/merge/preview'] = {replace_preview: Object.assign({}, PV, {path: 'docs/b.md', pending: 'docs/b.md.new'})};
  await openDiff('docs/b.md.new', 'docs/b.md');
  resolveOld({result: {diff: {path: 'docs/a.md'}}}); await old;
  const selected = MP.pending, title = $('sh-t').textContent;
  if (selected !== 'docs/b.md.new' || !title.endsWith('docs/b.md') || drawn.some(d => d.path === 'docs/a.md')) throw Error('stale comparison rebound selection');
  RESP['/api/run'] = () => new Promise(r => { resolveOld = r; });
  const closing = openDiff('docs/a.md.new', 'docs/a.md'); closeSheet();
  const count = calls.length;
  resolveOld({result: {diff: {path: 'docs/a.md'}}}); await closing;
  if (MP !== null || calls.length !== count) throw Error('closed comparison loaded a preview');
  console.log(JSON.stringify({ok: true}));
})();
"""
        done = subprocess.run(['node', '-e', script], capture_output=True, encoding='utf-8', timeout=15)
        self.assertEqual(done.returncode, 0, done.stdout + done.stderr)
        self.assertTrue(json.loads(done.stdout)['ok'])

    def test_every_outcome_has_its_own_message_and_only_possible_writes_refresh(self):
        res = self.out['outcomes']
        expect = {'noop': ('mp_noop', 0), 'stale': ('mp_stale', 0), 'collision': ('mp_collision', 0), 'pending': ('mp_r_pending', 0),
                  'partial': ('mp_partial:TX-1', 1), 'unproven': ('mp_unproven', 1), 'error': ('mp_err', 1), 'net': ('mp_unknown', 1)}
        for key, (needle, refreshed) in expect.items():
            with self.subTest(outcome=key):
                self.assertIn(needle, res[key]['html'])
                self.assertEqual(res[key]['refreshed'], refreshed)
                self.assertNotIn('mp_done', res[key]['html'])  # không báo thay xong khi chưa thay / never says replaced when it was not

    def test_script_text_has_no_launcher_no_state_writer_and_one_apply_path(self):
        block = ui_block()
        for forbidden in ('/api/open', 'openPath', 'merge_state', 'launch', 'setTimeout', 'startfile'):
            self.assertNotIn(forbidden, block)
        self.assertEqual(sorted(set(re.findall(r"api\('(/api/[a-z/]+)'", block))), ['/api/merge/preview', '/api/merge/replace'])
        self.assertEqual(len(re.findall(r'mpApply\(', block)), 2)  # định nghĩa + mpAct('use') / the definition and the one dispatch
        text = UI_FILE.read_text(encoding='utf-8')
        table = text[text.index('function drawMerges'):text.index('// merge comparison')]
        self.assertNotIn('data-open', table)  # bảng merge không còn mở mặc định / no default opener in the merge table
        self.assertIn('data-cp', table)


class MergeReplaceUiTexts(unittest.TestCase):
    FORBIDDEN = ('Reviewed', 'reviewed', 'Resolved', 'resolved', 'đã giải quyết', 'Đã giải quyết', 'Đã hợp nhất', 'Merged', 'atomic', 'nguyên tử', 'Undo', 'Hoàn tác:')

    def keys(self):
        block = ui_block()
        return sorted(set(re.findall(r"'(mp_[a-z_]+)'", block)) | set(re.findall(r'\b(mp_[a-z_]+):\[', UI_FILE.read_text(encoding='utf-8'))) |
                      {'mg_cp_current', 'mg_cp_new'})

    def test_every_key_has_both_languages_and_no_false_claims(self):
        d = ui_dictionary()
        dynamic = {'mp_r_'}  # khóa ghép động được kiểm bên dưới / keys built at run time are checked below
        for key in self.keys():
            if key in dynamic:
                continue
            with self.subTest(key=key):
                self.assertIn(key, d)
                vi, en = d[key]
                self.assertTrue(vi.strip() and en.strip())
                for bad in self.FORBIDDEN:
                    for text in (vi, en):
                        self.assertIsNone(re.search(r'(?<![A-Za-z])' + re.escape(bad) + r'(?![A-Za-z])', text), (key, bad))

    def test_reason_and_failure_messages_say_what_happened_and_the_next_step(self):
        d = ui_dictionary()
        for key in ('mp_r_current_missing', 'mp_r_candidate_missing', 'mp_r_binary', 'mp_r_too_large', 'mp_r_truncated', 'mp_r_unsafe', 'mp_r_other'):
            self.assertIn('Không có gì bị thay đổi', d[key][0], key)
            self.assertIn('Nothing was changed', d[key][1], key)
        for key in ('mp_pv_fail', 'mp_r_pending', 'mp_stale', 'mp_collision', 'mp_blocked', 'mp_partial', 'mp_err', 'mp_unknown'):
            self.assertIn('Bước tiếp', d[key][0], key)
            self.assertIn('Next step', d[key][1], key)
        for key in ('mp_stale', 'mp_collision', 'mp_blocked'):
            self.assertIn('chưa ghi gì', d[key][0].lower())
            self.assertIn('nothing was written', d[key][1].lower())
        self.assertIn('%s', d['mp_partial'][0])
        self.assertIn('%s', d['mp_partial'][1])

    def test_the_decision_names_the_four_consequences_in_both_languages(self):
        d = ui_dictionary()
        for key, vi_part, en_part in (('mp_will_cur', 'bị thay', 'replaced'), ('mp_will_save', 'cả hai bản được lưu', 'both versions are stored'),
                                      ('mp_will_pend', 'danh sách chờ', 'pending list'), ('mp_will_html', 'HTML', 'HTML')):
            self.assertIn(vi_part, d[key][0])
            self.assertIn(en_part, d[key][1])
        self.assertIn('.new', d['mp_use'][0] + d['mp_use'][1])
        self.assertIn('không tự hợp nhất', d['mg_safe'][0])
        self.assertIn('never merges automatically', d['mg_safe'][1])


# ---------------------------------------------------------------- D-01c: lưu .new khỏi danh sách chờ / archive the .new, keep the current file
windows_only = unittest.skipUnless(os.name == 'nt', 'the no-replace rename is only measured on Windows')
FINISH_ID_RE = r'finish-[0-9a-f]{32}'


def raise_error(message):
    raise RuntimeError(message)


class MergeFinishBase(MergeReplaceBase):
    def fpv(self, rel=None, root=None):
        return merge_center.preview_finish(str(root or self.project), (rel or self.guide) + '.new')

    def fin(self, rel=None, pv=None, root=None):
        rel = rel or self.guide
        pv = pv or self.fpv(rel, root)
        return merge_center.apply_finish(str(root or self.project), rel + '.new', pv['current_sha256'], pv['candidate_sha256'], pv['finish_id'])

    def obs(self, pv, rel=None, root=None, **over):
        args = dict(current_sha256=pv['current_sha256'], candidate_sha256=pv['candidate_sha256'], finish_id=pv['finish_id'])
        args.update(over)
        return merge_center.observe_finish(str(root or self.project), (rel or self.guide) + '.new', args['current_sha256'],
                                           args['candidate_sha256'], args['finish_id'])

    def arc(self, pv, name):
        return '.agent/merge-backups/%s/%s' % (pv['finish_id'], name)

    def listed(self, rel=None, root=None):
        return [m['pending'] for m in merge_center.scan(str(root or self.project)) if m['path'] == (rel or self.guide)]

    def assert_obs(self, pv, state, reason=None, **over):
        before = tree_state(self.project)
        res = self.obs(pv, **over)
        self.assertEqual(tree_state(self.project), before)  # quan sát không đổi snapshot toàn cây / observing leaves the whole tree as it was
        self.assertEqual((res['state'], res['reason']), (state, reason), res)
        return res

    def hook(self, point, action):
        done = []

        def hook(p):
            if p == point and not done:
                done.append(1)
                action()
        return patch.object(merge_center, '_HOOK', hook)


class MergeFinishPreview(MergeFinishBase):
    @windows_only
    def test_preview_describes_the_pair_and_writes_nothing(self):
        before = tree_state(self.project)
        pv = self.fpv()
        self.assertEqual(pv['schema'], merge_center.FINISH_PREVIEW_SCHEMA)
        self.assertEqual((pv['path'], pv['pending'], pv['eligible'], pv['reason'], pv['identical']), (self.guide, self.guide + '.new', True, None, False))
        self.assertRegex(pv['finish_id'], FINISH_ID_RE)
        self.assertEqual((pv['current_sha256'], pv['candidate_sha256']), (sha(self.raw(self.guide)), sha(self.raw(self.guide + '.new'))))  # hash thô, gồm CRLF
        self.assertEqual((pv['diff']['added'], pv['diff']['removed']), (2, 1))
        base = '.agent/merge-backups/' + pv['finish_id']
        self.assertEqual(pv['archive'], {'id': pv['finish_id'], 'dir': base, 'current': base + '/current.txt', 'incoming': base + '/incoming.txt',
                                         'record': base + '/record.json', 'retired': base + '/pending-retired.txt'})
        self.assertTrue(pv['current_stays'] and pv['pending_moves'] and not pv['manifest_updated'] and not pv['html_rebuilt'])
        self.assertEqual(tree_state(self.project), before)
        self.assertFalse(self.runtime_exists())
        self.assertEqual(self.archive_ids(), [])
        self.assertNotEqual(self.fpv()['finish_id'], pv['finish_id'])  # mỗi lần xem trước một ID mới / a new id per preview

    def test_hash_and_diff_come_from_the_same_buffers(self):
        reads = []
        real = merge_center._capture

        def counting(root, rel):
            reads.append(rel)
            return real(root, rel)
        with patch.object(merge_center, '_capture', counting):
            pv = self.fpv()
        self.assertEqual(reads, [self.guide, self.guide + '.new'])
        self.assertEqual((pv['current_sha256'], pv['candidate_sha256']), (sha(self.raw(self.guide)), sha(self.raw(self.guide + '.new'))))

    @windows_only
    def test_identical_pair_can_finish_while_replace_keeps_its_noop(self):
        write(self.project, self.guide + '.new', self.raw(self.guide))
        pv = self.fpv()
        self.assertEqual((pv['eligible'], pv['reason'], pv['identical'], pv['diff']), (True, None, True, None))
        self.assertRegex(pv['finish_id'], FINISH_ID_RE)
        replace = self.pv()
        self.assertEqual((replace['eligible'], replace['reason']), (False, 'identical'))

    def test_unsupported_and_invalid_paths_are_refused_without_an_id(self):
        before = tree_state(self.project)
        for path, reason in (('docs/other-notes.md.new', 'unsupported'), ('tools/helper.py.new', 'unsupported'), ('roles/scout/agent.md.new', 'unsupported'),
                             ('skills/local/packs/demo/SKILL.md.new', 'unsupported'), ('../docs/HUONG_DAN_SKILLS_VI.md.new', 'invalid_path'),
                             ('docs/HUONG_DAN_SKILLS_VI.md:stream.new', 'invalid_path'), ('docs/huong_dan_skills_vi.md.new', 'invalid_path'),
                             ('C:/docs/HUONG_DAN_SKILLS_VI.md.new', 'invalid_path'), ('//server/share/x.md.new', 'invalid_path'),
                             ('docs\\HUONG_DAN_SKILLS_VI.md.new', 'invalid_path'), (None, 'invalid_path'), (5, 'invalid_path')):
            with self.subTest(path=path):
                pv = merge_center.preview_finish(str(self.project), path)
                self.assertEqual((pv['eligible'], pv['reason'], pv['finish_id']), (False, reason, None))
        self.assertEqual(tree_state(self.project), before)

    def test_content_and_environment_problems_block_the_preview(self):
        def blocked(reason):
            pv = self.fpv()
            self.assertEqual((pv['eligible'], pv['reason'], pv['finish_id'], pv['archive']), (False, reason, None, None))
        (self.project / (self.guide + '.new')).write_bytes(b'bin\x00ary')
        blocked('binary')
        write(self.project, self.guide + '.new', b'x' * (merge_center.DIFF_MAX_BYTES + 1))
        blocked('too_large')
        write(self.project, self.guide + '.new', ''.join('b%d\n' % i for i in range(5000)))
        write(self.project, self.guide, ''.join('a%d\n' % i for i in range(5000)))
        blocked('truncated')  # không hiện đủ diff thì không cho xác nhận / no full diff, no confirmation
        (self.project / (self.guide + '.new')).unlink()
        blocked('candidate_missing')
        write(self.project, self.guide + '.new', 'n\n')
        (self.project / self.guide).unlink()
        blocked('current_missing')
        write(self.project, self.guide, 'c\n')
        write(self.project, '.agent/runtime/transactions/broken/journal.json', '{')
        blocked('journal_unreadable')
        shutil.rmtree(str(self.project / '.agent' / 'runtime'))
        folder = self.project / '.agent' / 'runtime' / 'transactions' / '20260101-000000-bbbbbbbb'
        write(folder, 'journal.json', json.dumps({'schema': tx.SCHEMA, 'id': folder.name, 'operation': 'upgrade', 'status': 'committing', 'ops': []}))
        blocked('transaction_pending')
        shutil.rmtree(str(self.project / '.agent' / 'runtime'))
        self.assertEqual(self.archive_ids(), [])

    def test_platform_and_volume_limits_are_reported_not_worked_around(self):
        with patch.object(merge_center, '_windows_supported', lambda: False):
            pv = self.fpv()
        self.assertEqual((pv['eligible'], pv['reason'], pv['finish_id']), (False, 'os_unsupported', None))
        self.assertIsNotNone(pv['diff'])  # vẫn xem và chép đường dẫn được / the diff and the paths stay usable
        with patch.object(merge_center, '_windows_supported', lambda: True), patch.object(merge_center, '_same_volume', lambda *a: False):
            pv = self.fpv()
        self.assertEqual((pv['eligible'], pv['reason'], pv['finish_id']), (False, 'filesystem_unsupported', None))
        good = self.fpv() if os.name == 'nt' else dict(self.pv(), finish_id=None)
        with patch.object(merge_center, '_windows_supported', lambda: False), patch.object(merge_center.os, 'rename', wraps=os.rename) as spy:
            res = merge_center.apply_finish(str(self.project), self.guide + '.new', good['current_sha256'], good['candidate_sha256'], 'finish-' + 'a' * 32)
        spy.assert_not_called()
        self.assertEqual((res['outcome'], res['reason']), ('blocked', 'os_unsupported'))
        self.assertFalse(self.runtime_exists())
        self.assertEqual(self.archive_ids(), [])

    def test_links_in_the_pair_or_the_archive_folder_are_refused(self):
        outside = self.root / 'outside'
        write(outside, 'x.md', 'outside\n')
        names = sorted(p.name for p in outside.iterdir())
        cases = []
        shutil.rmtree(str(self.project / 'docs' / 'en'))
        if dir_link(self.project / 'docs' / 'en', outside):
            write(outside, 'SKILLS_GUIDE.md', 'o\n')
            write(outside, 'SKILLS_GUIDE.md.new', 'n\n')
            names = sorted(p.name for p in outside.iterdir())
            cases.append(('docs/en/SKILLS_GUIDE.md', 'unsafe_link', self.project / 'docs' / 'en'))
        link = self.project / '.agent' / 'merge-backups'
        link.parent.mkdir(parents=True, exist_ok=True)
        if dir_link(link, outside):
            cases.append((self.guide, 'archive_unsafe', link))
        if not cases:
            self.skipTest('this machine cannot create directory links')
        for rel, reason, path in cases:
            with self.subTest(reason=reason):
                pv = self.fpv(rel)
                self.assertEqual((pv['eligible'], pv['reason'], pv['finish_id']), (False, reason, None))
                drop_link(path)
        self.assertEqual(sorted(p.name for p in outside.iterdir()), names)


@windows_only
class MergeFinishApply(MergeFinishBase):
    def test_every_guide_finishes_with_current_kept_and_pending_archived_for_both_pair_kinds(self):
        for n, rel in enumerate(GUIDES):
            for identical in (False, True):
                with self.subTest(guide=rel, identical=identical):
                    root = build_project(self.root / ('g%d%d' % (n, identical)))
                    if identical:
                        write(root, rel + '.new', self.raw(rel, root))
                    cur, cand = self.raw(rel, root), self.raw(rel + '.new', root)
                    others = {o: self.raw(o + '.new', root) for o in OTHER_MANAGED}
                    markers = user_markers(root)
                    pv = self.fpv(rel, root)
                    self.assertTrue(pv['eligible'])
                    res = self.fin(rel, pv, root)
                    self.assertEqual((res['outcome'], res['reason'], res['phase'], res['snapshot_saved'], res['pending_moved']),
                                     ('finished', None, 'committed', True, True), res)
                    self.assertEqual(self.raw(rel, root), cur)  # current nguyên raw bytes/EOL / current untouched
                    self.assertFalse((root / (rel + '.new')).exists())
                    self.assertEqual(self.raw(self.arc(pv, 'pending-retired.txt'), root), cand)  # nguyên bytes / bytes intact
                    self.assertEqual((self.raw(self.arc(pv, 'current.txt'), root), self.raw(self.arc(pv, 'incoming.txt'), root)), (cur, cand))
                    self.assertEqual(sorted(p.name for p in (root / '.agent' / 'merge-backups' / pv['finish_id']).iterdir()),
                                     ['current.txt', 'incoming.txt', 'pending-retired.txt', 'record.json'])
                    self.assertNotIn(rel + '.new', [m['pending'] for m in merge_center.scan(str(root))])
                    self.assertEqual({o: self.raw(o + '.new', root) for o in OTHER_MANAGED}, others)
                    self.assertEqual(user_markers(root), markers)
                    self.assertEqual(tx.incomplete(str(root)), [])

    def test_manifest_status_html_and_other_pending_rows_are_kept(self):
        merge_center.set_state(str(self.project), self.guide + '.new', 'reviewing')
        html = write(self.project, 'docs/HUONG_DAN_SKILLS_VI.html', '<html>built earlier</html>').read_bytes()
        status, markers = self.raw('.agent/MERGE_STATUS.json'), user_markers(self.project)
        others = [m['pending'] for m in merge_center.scan(str(self.project)) if m['path'] != self.guide]
        res = self.fin()
        self.assertEqual(res['outcome'], 'finished')
        self.assertEqual(self.raw('.agent/MERGE_STATUS.json'), status)  # không cần sửa MERGE_STATUS / MERGE_STATUS is not edited
        self.assertEqual(self.raw('docs/HUONG_DAN_SKILLS_VI.html'), html)
        self.assertEqual(user_markers(self.project), markers)
        self.assertEqual([m['pending'] for m in merge_center.scan(str(self.project))], others)
        self.assertTrue(res['current_stays'] and not res['html_rebuilt'] and not res['manifest_updated'])

    def test_record_and_journal_hold_only_the_three_snapshot_files(self):
        old, new = self.raw(self.guide), self.raw(self.guide + '.new')
        pv = self.fpv()
        res = self.fin(pv=pv)
        record_bytes = self.raw(self.arc(pv, 'record.json'))
        self.assertEqual(record_bytes, merge_center._record_bytes(pv['finish_id'], self.guide, sha(old), sha(new), res['txid']))
        record = json.loads(record_bytes.decode('utf-8'))
        self.assertEqual(sorted(record), ['archive_txid', 'current_sha256', 'finish_id', 'incoming_sha256', 'path', 'record_version', 'schema'])
        self.assertEqual(record['schema'], merge_center.FINISH_RECORD_SCHEMA)
        text = record_bytes.decode('utf-8')
        self.assertNotIn(str(self.project), text)
        self.assertNotIn('Hướng dẫn', text)
        journal = json.loads(self.raw('.agent/runtime/transactions/%s/journal.json' % res['txid']).decode('utf-8'))
        self.assertEqual((journal['status'], journal['operation']), ('committed', merge_center.FINISH_OPERATION))
        self.assertEqual([(o['path'], o['action'], o['before'] is None) for o in journal['ops']],
                         [(self.arc(pv, n), 'create_exclusive', True) for n in ('current.txt', 'incoming.txt', 'record.json')])
        self.assertNotIn(self.guide, [o['path'] for o in journal['ops']])  # current và .new ngoài ops / neither is an op
        self.assertNotIn(self.guide + '.new', [o['path'] for o in journal['ops']])

    def test_replace_then_finish_needs_a_new_preview_of_the_new_pair(self):
        old_preview = self.fpv()
        rep = self.go()
        self.assertEqual(rep['outcome'], 'replaced')
        self.assertEqual(self.listed(), [self.guide + '.new'])  # replace vẫn để .new như cũ / replace still leaves the .new
        stale = self.fin(pv=old_preview)
        self.assertEqual((stale['outcome'], stale['reason'], stale['stale']), ('stale', 'stale_preview', {'current': True, 'candidate': False}))
        self.assertEqual(self.listed(), [self.guide + '.new'])
        fresh = self.fpv()
        self.assertTrue(fresh['identical'] and fresh['eligible'])
        self.assertNotEqual(fresh['finish_id'], old_preview['finish_id'])
        res = self.fin(pv=fresh)
        self.assertEqual(res['outcome'], 'finished')
        self.assertEqual(self.raw(self.guide), guide_new(self.guide))
        self.assertEqual(self.listed(), [])

    def test_manual_edit_or_a_new_candidate_after_replace_shows_the_new_pair(self):
        self.go()
        write(self.project, self.guide, b'MY OWN EDIT AFTER REPLACE\r\n')
        write(self.project, self.guide + '.new', b'A NEWER CANDIDATE\n')
        pv = self.fpv()
        self.assertEqual((pv['current_sha256'], pv['candidate_sha256']), (sha(b'MY OWN EDIT AFTER REPLACE\r\n'), sha(b'A NEWER CANDIDATE\n')))
        self.assertFalse(pv['identical'])
        self.assertIn('MY OWN EDIT AFTER REPLACE', json.dumps(pv['diff'], ensure_ascii=False))
        res = self.fin(pv=pv)
        self.assertEqual(res['outcome'], 'finished')
        self.assertEqual(self.raw(self.guide), b'MY OWN EDIT AFTER REPLACE\r\n')
        self.assertEqual(self.raw(self.arc(pv, 'pending-retired.txt')), b'A NEWER CANDIDATE\n')

    def test_each_stale_hash_is_reported_on_its_own_and_nothing_is_written(self):
        for which in ('current', 'candidate'):
            with self.subTest(changed=which):
                root = build_project(self.root / ('stale-' + which))
                pv = self.fpv(root=root)
                rel = self.guide + ('' if which == 'current' else '.new')
                write(root, rel, b'EDIT AFTER PREVIEW\n')
                before = tree_state(root)
                res = self.fin(pv=pv, root=root)
                self.assertEqual((res['outcome'], res['reason'], res['stale']),
                                 ('stale', 'stale_preview', {'current': which == 'current', 'candidate': which == 'candidate'}))
                self.assertEqual(tree_state(root), before)
                self.assertFalse((root / '.agent' / 'runtime').exists())

    def test_malformed_ids_and_hashes_never_pass_and_write_nothing(self):
        pv = self.fpv()
        before = tree_state(self.project)
        good = pv['finish_id']
        for fid in (None, '', 5, ['x'], good.upper(), good[:-1], good + '0', good.replace('finish-', 'finis-'), 'finish-' + 'g' * 32, good + '/../x',
                    '../' + good, '.agent/merge-backups/' + good, good + '\n', ' ' + good):
            with self.subTest(finish_id=fid):
                res = merge_center.apply_finish(str(self.project), self.guide + '.new', pv['current_sha256'], pv['candidate_sha256'], fid)
                self.assertEqual((res['outcome'], res['reason']), ('blocked', 'bad_request'))
                seen = merge_center.observe_finish(str(self.project), self.guide + '.new', pv['current_sha256'], pv['candidate_sha256'], fid)
                self.assertEqual(seen['state'], 'refused')
        for cur, cand in ((HEX_A, HEX_B), (None, None), ('', ''), (5, [1]), (pv['current_sha256'].upper(), pv['candidate_sha256'])):
            res = merge_center.apply_finish(str(self.project), self.guide + '.new', cur, cand, good)
            self.assertIn(res['outcome'], ('stale', 'blocked'))
        self.assertEqual(tree_state(self.project), before)

    def test_ineligible_pairs_are_not_finished_even_with_their_real_hashes(self):
        before = tree_state(self.project)
        # D-01e: AGENTS.md có finish riêng nên không còn là ca ngoài scope của finish; replace vẫn từ chối nó (MergeReplaceApply) / AGENTS.md has its own finish scope now; replace still refuses it
        for rel in tuple(x for x in OTHER_MANAGED if x != 'AGENTS.md') + ('unmanaged.md', 'skills/local/packs/demo/SKILL.md'):
            with self.subTest(path=rel):
                res = merge_center.apply_finish(str(self.project), rel + '.new', sha(self.raw(rel)), sha(self.raw(rel + '.new')), 'finish-' + 'c' * 32)
                self.assertEqual((res['outcome'], res['reason']), ('blocked', 'unsupported'))
        self.assertEqual(tree_state(self.project), before)

    def test_archive_collision_and_unreadable_journal_block_before_any_write(self):
        pv = self.fpv()
        write(self.project, '.agent/merge-backups/%s/other.txt' % pv['finish_id'], 'SQUATTER')
        before = tree_state(self.project)
        res = self.fin(pv=pv)
        self.assertEqual((res['outcome'], res['reason']), ('blocked', 'archive_collision'))
        self.assertEqual(tree_state(self.project), before)  # không ghi đè, không đổi ID, không dùng lại thư mục / no overwrite, no new id, no reuse
        self.assertEqual(self.listed(), [self.guide + '.new'])
        pv2 = self.fpv()
        write(self.project, '.agent/runtime/transactions/broken/journal.json', '{')
        before = tree_state(self.project)
        res = self.fin(pv=pv2)
        self.assertEqual((res['outcome'], res['reason']), ('blocked', 'journal_unreadable'))
        self.assertEqual(tree_state(self.project), before)

    def test_a_locked_candidate_keeps_the_pending_item_and_the_committed_archive(self):
        handles = []
        pv = self.fpv()
        with self.hook('finish_before_rename', lambda: handles.append(open(str(self.project / (self.guide + '.new')), 'rb'))):
            res = self.fin(pv=pv)
        [h.close() for h in handles]
        self.assertEqual((res['outcome'], res['reason'], res['snapshot_saved'], res['pending_moved']), ('still_pending', 'rename_failed', True, False), res)
        self.assertEqual(self.raw(self.guide + '.new'), guide_new(self.guide))
        self.assertEqual(self.listed(), [self.guide + '.new'])
        self.assertFalse((self.project / self.arc(pv, 'pending-retired.txt')).exists())
        self.assertEqual(self.raw(self.arc(pv, 'incoming.txt')), guide_new(self.guide))
        self.assertEqual(self.raw(self.guide), guide_current(self.guide))
        self.assert_obs(pv, 'still_pending')

    def test_an_occupied_slot_found_by_the_precheck_is_kept_and_no_rename_is_attempted(self):
        pv = self.fpv()
        slot = self.project / self.arc(pv, 'pending-retired.txt')
        with patch.object(merge_center.os, 'rename', wraps=os.rename) as spy, self.hook('finish_after_run', lambda: slot.write_bytes(b'SQUATTER')):
            res = self.fin(pv=pv)
        spy.assert_not_called()
        self.assertEqual((res['outcome'], res['reason'], res['detail']), ('conflict', 'both_present', None))
        self.assertEqual((slot.read_bytes(), self.raw(self.guide + '.new')), (b'SQUATTER', guide_new(self.guide)))

    def test_a_slot_taken_at_the_rename_is_refused_by_the_os_and_kept(self):
        pv = self.fpv()
        slot = self.project / self.arc(pv, 'pending-retired.txt')
        with self.hook('finish_before_rename', lambda: slot.write_bytes(b'SQUATTER LATE')):
            res = self.fin(pv=pv)
        self.assertEqual((res['outcome'], res['reason'], res['detail']), ('conflict', 'both_present', 'FileExistsError'), res)
        self.assertEqual(slot.read_bytes(), b'SQUATTER LATE')  # không đè, không xóa / not overwritten, not deleted
        self.assertEqual(self.raw(self.guide + '.new'), guide_new(self.guide))

    def test_candidate_changed_just_before_the_rename_keeps_the_moved_bytes(self):
        later = b'CANDIDATE CHANGED RIGHT BEFORE THE RENAME\n'
        pv = self.fpv()
        with self.hook('finish_before_rename', lambda: write(self.project, self.guide + '.new', later)):
            res = self.fin(pv=pv)
        self.assertEqual((res['outcome'], res['reason'], res['pending_moved']), ('conflict', 'moved_bytes_differ', True), res)
        self.assertEqual(self.raw(self.arc(pv, 'pending-retired.txt')), later)  # giữ đúng bytes đã chuyển, không move-back / kept, not moved back
        self.assertFalse((self.project / (self.guide + '.new')).exists())
        self.assertEqual(self.raw(self.arc(pv, 'incoming.txt')), guide_new(self.guide))  # snapshot là bản đã xem / the snapshot is what was previewed
        self.assertNotEqual(self.obs(pv)['state'], 'finished')

    def test_candidate_or_current_changed_before_the_precheck_stops_the_rename(self):
        for target, reason in (('.new', 'candidate_changed'), ('', 'current_changed')):
            with self.subTest(target=target):
                root = build_project(self.root / ('pre-' + reason))
                pv = self.fpv(root=root)
                with patch.object(merge_center.os, 'rename', wraps=os.rename) as spy, \
                        self.hook('finish_after_run', lambda: write(root, self.guide + target, b'CHANGED AFTER THE ARCHIVE COMMIT\n')):
                    res = self.fin(pv=pv, root=root)
                spy.assert_not_called()
                self.assertEqual((res['outcome'], res['reason'], res['snapshot_saved'], res['pending_moved']), ('conflict', reason, True, False), res)
                self.assertEqual(self.raw(self.guide + target, root), b'CHANGED AFTER THE ARCHIVE COMMIT\n')
                self.assertEqual(self.listed(root=root), [self.guide + '.new'])

    def test_candidate_changed_after_the_recheck_but_before_the_run_is_caught_by_the_precheck(self):
        later = b'CANDIDATE CHANGED BEFORE THE ARCHIVE RUN\n'
        pv = self.fpv()
        with patch.object(merge_center.os, 'rename', wraps=os.rename) as spy, \
                self.hook('finish_before_run', lambda: write(self.project, self.guide + '.new', later)):
            res = self.fin(pv=pv)
        spy.assert_not_called()
        self.assertEqual((res['outcome'], res['reason'], res['snapshot_saved'], res['pending_moved']), ('conflict', 'candidate_changed', True, False), res)
        self.assertEqual(self.raw(self.arc(pv, 'incoming.txt')), guide_new(self.guide))  # snapshot là bản đã xem, không phải bản đổi muộn / the snapshot is the previewed copy
        self.assertEqual(self.raw(self.guide + '.new'), later)

    def test_current_changed_late_after_the_move_is_not_reported_finished(self):
        pv = self.fpv()
        edit = b'CURRENT EDITED RIGHT BEFORE THE RENAME\n'
        with self.hook('finish_before_rename', lambda: write(self.project, self.guide, edit)):
            res = self.fin(pv=pv)
        self.assertEqual((res['outcome'], res['reason']), ('conflict', 'current_changed'), res)
        self.assertEqual(self.raw(self.guide), edit)  # current không bị thao tác này ghi / current is never written by this
        self.assertEqual(self.raw(self.arc(pv, 'pending-retired.txt')), guide_new(self.guide))

    def test_error_after_the_move_is_still_finished_because_the_files_decide(self):
        pv = self.fpv()
        with self.hook('finish_after_rename', lambda: raise_error('after move')):
            res = self.fin(pv=pv)
        self.assertEqual((res['outcome'], res['reason'], res['detail']), ('finished', 'observed_after_error', 'RuntimeError'))
        root = build_project(self.root / 'os-error')
        pv2 = self.fpv(root=root)
        real = os.rename

        def move_then_fail(src, dst):
            real(src, dst)
            raise OSError('lost the response after the move')
        with patch.object(merge_center.os, 'rename', move_then_fail):
            res = self.fin(pv=pv2, root=root)
        self.assertEqual((res['outcome'], res['reason']), ('finished', 'observed_after_error'))
        self.assertFalse((root / (self.guide + '.new')).exists())

    def test_hook_error_before_the_rename_leaves_the_item_pending(self):
        pv = self.fpv()
        with self.hook('finish_before_rename', lambda: raise_error('before move')):
            res = self.fin(pv=pv)
        self.assertEqual((res['outcome'], res['reason'], res['snapshot_saved']), ('still_pending', 'rename_failed', True))
        self.assertEqual(self.listed(), [self.guide + '.new'])

    def test_platform_or_volume_lost_after_the_archive_commit_does_not_rename(self):
        for name in ('os', 'volume'):
            with self.subTest(lost=name):
                root = build_project(self.root / ('lost-' + name))
                pv = self.fpv(root=root)
                started = []

                def lose():
                    started.append(patch.object(merge_center, '_windows_supported', lambda: False) if name == 'os'
                                   else patch.object(merge_center, '_same_volume', lambda *a: False))
                    started[-1].start()
                try:
                    with patch.object(merge_center.os, 'rename', wraps=os.rename) as spy, self.hook('finish_after_run', lose):
                        res = self.fin(pv=pv, root=root)
                finally:
                    [p.stop() for p in started]
                spy.assert_not_called()
                self.assertEqual((res['outcome'], res['reason'], res['snapshot_saved']),
                                 ('still_pending', 'os_unsupported' if name == 'os' else 'filesystem_unsupported', True))
                self.assertTrue((root / (self.guide + '.new')).exists())

    def test_fault_before_the_first_file_writes_nothing(self):
        for fault in ('stage:1', 'stage:3'):  # backup:N không chạy ở đây vì ba thao tác đều là file mới / backup:N never fires here, all three ops create new files
            with self.subTest(fault=fault):
                root = build_project(self.root / ('early-' + fault.replace(':', '-')))
                with patch.dict(os.environ, {FAULT: fault}):
                    res = self.fin(root=root)
                self.assertEqual((res['outcome'], res['reason']), ('error', 'write_failed'), res)
                self.assertEqual(self.archive_ids(root), [])
                self.assertTrue((root / (self.guide + '.new')).exists())
                self.assertEqual(tx.incomplete(str(root)), [])

    def test_cleanup_failure_after_committed_archive_reports_saved_and_never_renames(self):
        pv = self.fpv()
        with patch.object(tx.Transaction, '_cleanup', side_effect=OSError('cleanup after commit')):
            with patch.object(merge_center.os, 'rename', wraps=os.rename) as rename:
                result = self.fin(pv=pv)
        self.assertEqual((result['outcome'], result['phase'], result['reason']),
                         ('still_pending', 'committed', 'archive_committed_after_error'))
        self.assertTrue(result['snapshot_saved'])
        self.assertFalse(result['pending_moved'])
        rename.assert_not_called()
        self.assertEqual(self.raw(self.guide), guide_current(self.guide))
        self.assertEqual(self.raw(self.guide + '.new'), guide_new(self.guide))
        self.assertEqual(self.raw(self.arc(pv, 'current.txt')), guide_current(self.guide))
        self.assertEqual(self.raw(self.arc(pv, 'incoming.txt')), guide_new(self.guide))
        self.assertEqual(self.obs(pv)['state'], 'still_pending')
        self.assertEqual(tx.incomplete(str(self.project)), [])
        self.assertEqual(tx.rollback(str(self.project), result['txid'])['status'], 'blocked')

    def test_fault_after_the_first_snapshot_keeps_the_pending_item_and_resume_does_not_rename(self):
        pv = self.fpv()
        with patch.dict(os.environ, {FAULT: 'commit:1'}), patch.object(merge_center.os, 'rename', wraps=os.rename) as spy:
            res = self.fin(pv=pv)
        spy.assert_not_called()
        self.assertEqual((res['outcome'], res['reason'], res['phase'], res['done'], res['total']), ('partial', 'transaction_unfinished', 'unfinished', 1, 3), res)
        self.assertEqual((res['snapshot_saved'], res['pending_moved']), (False, False))
        self.assertEqual(self.raw(self.guide + '.new'), guide_new(self.guide))
        self.assertEqual(self.raw(self.guide), guide_current(self.guide))
        seen = self.assert_obs(pv, 'partial', 'archive_incomplete')  # record chưa ghi nên chưa biết txid, nhưng thấy giao dịch dở / no record yet, the open transaction is still listed
        self.assertEqual(seen['unfinished'], [res['txid']])
        out = tx.resume(str(self.project), res['txid'])
        self.assertEqual(out['status'], 'success', out)
        self.assertTrue((self.project / (self.guide + '.new')).exists())  # resume chỉ hoàn tất archive / resume only completes the archive
        self.assertFalse((self.project / self.arc(pv, 'pending-retired.txt')).exists())
        self.assert_obs(pv, 'still_pending')

    def test_fault_after_all_three_snapshots_but_before_committed_never_renames_and_rollback_keeps_pending(self):
        pv = self.fpv()
        with patch.dict(os.environ, {FAULT: 'commit:3'}), patch.object(merge_center.os, 'rename', wraps=os.rename) as spy:
            res = self.fin(pv=pv)
        spy.assert_not_called()
        self.assertEqual((res['outcome'], res['reason'], res['phase']), ('partial', 'transaction_unfinished', 'unfinished'), res)
        for name in ('current.txt', 'incoming.txt', 'record.json'):
            self.assertTrue((self.project / self.arc(pv, name)).exists())
        self.assertEqual(json.loads(self.raw('.agent/runtime/transactions/%s/journal.json' % res['txid']).decode('utf-8'))['status'], 'committing')
        self.assertEqual(self.raw(self.guide + '.new'), guide_new(self.guide))
        self.assert_obs(pv, 'partial', 'transaction_unfinished')
        out = tx.rollback(str(self.project), res['txid'])
        self.assertEqual(out['status'], 'success', out)
        self.assertEqual(self.raw(self.guide + '.new'), guide_new(self.guide))  # rollback gỡ snapshot, pending vẫn ở chỗ cũ / pending stays live
        self.assertEqual(self.raw(self.guide), guide_current(self.guide))
        self.assertNotEqual(self.obs(pv)['state'], 'finished')

    def test_a_commit_that_cannot_be_proven_does_not_rename(self):
        pv = self.fpv()
        with patch.object(merge_center, '_journal_info', lambda *a: ('unreadable', False)), \
                patch.object(merge_center.os, 'rename', wraps=os.rename) as spy:
            res = self.fin(pv=pv)
        spy.assert_not_called()
        self.assertEqual((res['outcome'], res['reason'], res['snapshot_saved'], res['pending_moved']), ('partial', 'verify_mismatch', False, False))
        self.assertTrue((self.project / (self.guide + '.new')).exists())


@windows_only
class MergeFinishObserve(MergeFinishBase):
    def test_observation_table_over_one_finished_item(self):
        pv = self.fpv()
        self.assertEqual(self.assert_obs(pv, 'not_started')['phase'], 'none')  # chưa làm gì / nothing done yet
        self.fin(pv=pv)
        seen = self.assert_obs(pv, 'finished')
        self.assertEqual((seen['phase'], seen['observed']['source']['state'], seen['observed']['slot']['sha256']), ('committed', 'absent', sha(guide_new(self.guide))))
        self.assertNotIn('Hướng dẫn', json.dumps(seen, ensure_ascii=False))  # trạng thái, hash và đường dẫn tương đối, không nội dung / status, hashes, relative paths only
        self.assertNotIn(str(self.project), json.dumps(seen))
        slot, pending = self.project / self.arc(pv, 'pending-retired.txt'), self.project / (self.guide + '.new')
        current, journal = self.project / self.guide, self.project / '.agent' / 'runtime' / 'transactions' / seen['txid'] / 'journal.json'
        keep = {p: p.read_bytes() for p in (slot, current, journal, self.project / self.arc(pv, 'incoming.txt'), self.project / self.arc(pv, 'record.json'))}

        def restore():
            for path, data in keep.items():
                path.write_bytes(data)
            if pending.exists():
                pending.unlink()
            self.assert_obs(pv, 'finished')
        write(self.project, self.guide + '.new', b'A NEW COPY APPEARED LATER\n')
        self.assert_obs(pv, 'conflict', 'both_present')  # .new mới hiện lại, không bị record cũ che / a new .new is not hidden by the old record
        self.assertEqual(self.listed(), [self.guide + '.new'])
        restore()
        slot.write_bytes(b'OTHER BYTES')
        self.assert_obs(pv, 'conflict', 'moved_bytes_differ')
        restore()
        slot.rename(slot.with_name('moved-away.bak'))
        self.assert_obs(pv, 'unknown', 'both_absent')
        slot.with_name('moved-away.bak').rename(slot)
        restore()
        current.write_bytes(b'CURRENT CHANGED LATER\n')
        self.assert_obs(pv, 'conflict', 'current_changed')
        current.unlink()
        self.assert_obs(pv, 'conflict', 'current_missing')
        restore()
        for name in ('incoming.txt', 'record.json'):
            (self.project / self.arc(pv, name)).write_bytes(b'TAMPERED')
            self.assert_obs(pv, 'conflict', 'archive_unverified')
            restore()
        self.assert_obs(pv, 'conflict', 'archive_unverified', candidate_sha256=HEX_B)  # binding khác không mượn record / another binding cannot borrow the record
        journal.write_bytes(b'{')
        self.assert_obs(pv, 'unknown', 'unreadable_or_unsafe')
        restore()
        flipped = json.loads(keep[journal].decode('utf-8'))
        flipped['status'] = 'committing'
        journal.write_bytes(json.dumps(flipped).encode('utf-8'))
        self.assertNotEqual(self.obs(pv)['state'], 'finished')
        restore()

    def test_a_link_in_place_of_the_archive_folder_is_unknown_never_finished(self):
        pv = self.fpv()
        self.fin(pv=pv)
        folder = self.project / '.agent' / 'merge-backups' / pv['finish_id']
        away = folder.with_name('real-folder')
        folder.rename(away)
        if not dir_link(folder, away):
            away.rename(folder)
            self.skipTest('this machine cannot create directory links')
        try:
            self.assert_obs(pv, 'unknown', 'unreadable_or_unsafe')
        finally:
            drop_link(folder)
            away.rename(folder)
        self.assert_obs(pv, 'finished')

    def test_an_error_after_the_move_is_confirmed_by_observation_without_a_receipt(self):
        pv = self.fpv()
        real = os.rename

        def move_then_fail(src, dst):
            real(src, dst)
            raise OSError('response lost')
        with patch.object(merge_center.os, 'rename', move_then_fail):
            self.fin(pv=pv)
        self.assert_obs(pv, 'finished')
        self.assertEqual(self.raw('.agent/runtime/transactions/%s/journal.json' % self.obs(pv)['txid']).count(b'"committed"'), 1)

    def test_forged_inputs_are_refused_and_other_files_are_never_located(self):
        pv = self.fpv()
        before = tree_state(self.project)
        for path in ('docs/other-notes.md.new', 'roles/scout/agent.md.new', '../x.new', 'docs/huong_dan_skills_vi.md.new', None):
            seen = merge_center.observe_finish(str(self.project), path, pv['current_sha256'], pv['candidate_sha256'], pv['finish_id'])
            self.assertEqual(seen['state'], 'refused')
        for fid in ('finish-' + 'e' * 32, ):  # ID đúng định dạng nhưng chưa từng được phát: không phải lỗi, chỉ là chưa có gì / well-formed but never issued: nothing there
            self.assertEqual(merge_center.observe_finish(str(self.project), self.guide + '.new', pv['current_sha256'], pv['candidate_sha256'], fid)['state'], 'not_started')
        self.assertEqual(tree_state(self.project), before)


class MergeFinishHttp(MergeFinishBase):
    serve = MergeReplaceHttp.serve
    post = MergeReplaceHttp.post

    def routes(self):
        return ('/api/merge/finish-preview', '/api/merge/finish', '/api/merge/finish-observe')

    def test_every_finish_route_checks_the_token_and_the_host(self):
        import http.client
        base = self.serve()
        port = int(base.rsplit(':', 1)[1])
        before = tree_state(self.project)
        body = {'path': self.guide + '.new', 'current_sha256': HEX_A, 'candidate_sha256': HEX_B, 'finish_id': 'finish-' + 'a' * 32}
        for route in self.routes():
            for token in (None, 'wrong'):
                self.assertEqual(self.post(base, route, body, token)[0], 403, route)
            conn = http.client.HTTPConnection('127.0.0.1', port, timeout=30)
            try:
                # Host, like the token, is rejected before the body is read. Unread request bytes can reset the socket on Windows.
                conn.request('POST', route, b'', {'Host': 'evil.example', 'X-Panel-Token': 'test-token', 'Content-Type': 'application/json'})
                response = conn.getresponse()
                self.assertEqual(response.status, 403, route)
                self.assertEqual(json.loads(response.read().decode('utf-8')), {'error': 'forbidden'})
            finally:
                conn.close()
        self.assertEqual(tree_state(self.project), before)

    @windows_only
    def test_preview_and_observe_are_read_only_and_not_logged(self):
        base = self.serve()
        before = tree_state(self.project)
        code, body = self.post(base, '/api/merge/finish-preview', {'path': self.guide + '.new'})
        pv = body['finish_preview']
        self.assertEqual((code, body['status'], body['changes'], pv['eligible'], pv['path']), (200, 'success', False, True, self.guide))
        self.assertEqual(self.post(base, '/api/merge/finish-preview', {'path': '../x.new'})[1]['status'], 'blocked')
        code, seen = self.post(base, '/api/merge/finish-observe', {'path': self.guide + '.new', 'current_sha256': pv['current_sha256'],
                                                                   'candidate_sha256': pv['candidate_sha256'], 'finish_id': pv['finish_id']})
        self.assertEqual((seen['status'], seen['changes'], seen['finish_observation']['state']), ('success', False, 'not_started'))
        bad = self.post(base, '/api/merge/finish-observe', {'path': self.guide + '.new', 'finish_id': 'nope'})[1]
        self.assertEqual((bad['status'], bad['code'], bad['finish_observation']['state']), ('blocked', 2, 'refused'))
        self.assertEqual(tree_state(self.project), before)
        self.assertFalse(self.runtime_exists())  # không runtime, không operations.jsonl / no runtime, no operations.jsonl
        self.assertEqual(ops_log.read(str(self.project)), [])

    @windows_only
    def test_finish_route_runs_the_flow_and_is_not_in_the_operations_log(self):
        base = self.serve()
        pv = self.post(base, '/api/merge/finish-preview', {'path': self.guide + '.new'})[1]['finish_preview']
        pair = {'path': self.guide + '.new', 'current_sha256': pv['current_sha256'], 'candidate_sha256': pv['candidate_sha256']}
        self.assertEqual(self.post(base, '/api/merge/finish', dict(pair, finish_id=pv['finish_id']), None)[0], 403)
        code, blocked = self.post(base, '/api/merge/finish', pair)
        self.assertEqual((blocked['status'], blocked['finish']['reason']), ('blocked', 'bad_request'))
        self.assertTrue((self.project / (self.guide + '.new')).exists())
        code, ok = self.post(base, '/api/merge/finish', dict(pair, finish_id=pv['finish_id']))
        self.assertEqual((code, ok['status'], ok['code'], ok['changes'], ok['finish']['outcome']), (200, 'success', 0, True, 'finished'))
        self.assertEqual([f['action'] for f in ok['files']], ['created', 'created', 'created', 'moved_from_' + self.guide + '.new'])
        self.assertFalse((self.project / (self.guide + '.new')).exists())
        self.assertEqual(self.raw(self.guide), guide_current(self.guide))
        seen = self.post(base, '/api/merge/finish-observe', dict(pair, finish_id=pv['finish_id']))[1]
        self.assertEqual(seen['finish_observation']['state'], 'finished')
        self.assertFalse((self.project / '.agent' / 'runtime' / 'operations.jsonl').exists())
        self.assertEqual(self.opened, [])

    def test_envelope_statuses_and_changes_follow_what_really_happened(self):
        arc = merge_center.finish_locators('finish-' + 'f' * 32)
        cases = {'finished': ({'snapshot_saved': True, 'archive': arc}, ('success', 0, True)),
                 'still_pending': ({'snapshot_saved': True, 'archive': arc}, ('partial', 1, True)),
                 'conflict': ({'snapshot_saved': True, 'archive': arc}, ('partial', 1, True)),
                 'partial': ({'snapshot_saved': False}, ('partial', 1, None)),
                 'unknown': ({'snapshot_saved': False}, ('partial', 1, None)),
                 'stale': ({}, ('blocked', 1, False)), 'blocked': ({}, ('blocked', 1, False)), 'error': ({}, ('error', 1, False))}
        for outcome, (extra, (status_value, code, changes)) in cases.items():
            with self.subTest(outcome=outcome):
                fake = dict({'outcome': outcome, 'reason': 'x', 'path': self.guide, 'pending': self.guide + '.new', 'txid': 'T1'}, **extra)
                with patch.object(merge_center, 'apply_finish', return_value=fake):
                    env = app.merge_api('/api/merge/finish', {})
                self.assertEqual((env['status'], env['code'], env['changes']), (status_value, code, changes))
                self.assertEqual(len(env['files']), 4 if outcome == 'finished' else (3 if extra.get('snapshot_saved') else 0))
        self.assertIsNone(app.merge_api('/api/merge/finish-other', {}))

    def test_finish_never_reaches_the_state_writer_or_the_opener(self):
        with patch.object(merge_center, 'set_state', side_effect=AssertionError('state write')):
            for name in ('finish-preview', 'finish', 'finish-observe'):
                app.merge_api('/api/merge/' + name, {'path': self.guide + '.new', 'current_sha256': HEX_A, 'candidate_sha256': HEX_B,
                                                     'finish_id': 'finish-' + 'a' * 32})
        self.assertEqual(self.opened, [])
        self.assertFalse((self.project / '.agent' / 'MERGE_STATUS.json').exists())


HARNESS_C = r"""
'use strict';
const calls = [], toasts = [], copied = [], diffs = [], boxes = {}, RESP = {}, els = {}; let closed = 0, refreshed = 0;
function adopt(html) { const m = /id="([^"]+)"/.exec(html); if (m) boxes['#' + m[1]] = {innerHTML: '', insertAdjacentHTML: function (pos, h) { adopt(h); }}; }
const sheet = {listeners: [], querySelector: function (sel) { return boxes[sel] || null; }, insertAdjacentHTML: function (pos, h) { adopt(h); },
               addEventListener: function (type, fn) { sheet.listeners.push(fn); }};
function $(id) { return id === 'sh-diff' ? sheet : (els[id] || (els[id] = {innerHTML: '', textContent: '', addEventListener: function () {}})); }
function t(k) { return (k === 'mf_partial' || k === 'mf_unproven') ? k + ':%s' : k; }
function esc(s) { return String(s == null ? '' : s).replace(/[&<>"']/g, function (c) { return {'&': '&amp;', '<': '&lt;', '>': '&gt;', '"': '&quot;', "'": '&#39;'}[c]; }); }
function ico() { return ''; }
function drawDiff(d) { diffs.push(d); Object.keys(boxes).forEach(function (k) { delete boxes[k]; }); }  // innerHTML = ... wipes the boxes
function copyText(x) { copied.push(x); }
function toast(m) { toasts.push(m); }
function closeSheet() { closed++; mpReset(); }
function refresh() { refreshed++; return Promise.resolve(); }
function api(p, b) { calls.push([p, JSON.parse(JSON.stringify(b))]); const r = RESP[p]; return Promise.resolve(typeof r === 'function' ? r(b) : r); }
%BLOCKB%
%BLOCKC%
const PVS = 'local-agent-workspace/merge-replace-preview-v1', FPS = 'local-agent-workspace/merge-finish-preview-v1', FOS = 'local-agent-workspace/merge-finish-observe-v1';
const A = 'a'.repeat(64), B = 'b'.repeat(64), C = 'c'.repeat(64), D = 'd'.repeat(64), FID = 'finish-' + '1'.repeat(32), FID2 = 'finish-' + '2'.repeat(32);
const PV = {schema: PVS, path: 'docs/a.md', pending: 'docs/a.md.new', eligible: true, reason: null, current_sha256: A, candidate_sha256: B};
function archive(id) { const base = '.agent/merge-backups/' + id; return {id: id, dir: base, current: base + '/current.txt', incoming: base + '/incoming.txt', record: base + '/record.json', retired: base + '/pending-retired.txt'}; }
function FPV(id, cur, cand, extra) { return Object.assign({schema: FPS, path: 'docs/a.md', pending: 'docs/a.md.new', eligible: true, reason: null, identical: false,
  current_sha256: cur, candidate_sha256: cand, diff: {added: 2, removed: 1, hunks: []}, finish_id: id, archive: archive(id)}, extra || {}); }
const fcalls = () => calls.filter(c => c[0] === '/api/merge/finish');
const html = k => boxes['#' + k] ? boxes['#' + k].innerHTML : null;
async function open(pv) { closeSheet(); drawDiff({}); RESP['/api/merge/preview'] = {replace_preview: pv || PV}; calls.length = 0; await mpLoad('docs/a.md.new', 'docs/a.md'); }
async function openPreviewed(id, cur, cand, extra) { await open(); RESP['/api/merge/finish-preview'] = {finish_preview: FPV(id || FID, cur || A, cand || B, extra)}; await mfAct('preview'); }
(async function () {
  const out = {};
  // 1. the finish entry appears only for a supported pair and never calls the server by itself
  await open(); out.intro = html('mf-box'); out.attach_calls = calls.map(c => c[0]);
  await open(Object.assign({}, PV, {eligible: false, reason: 'identical'})); out.intro_identical = html('mf-box');
  await open(Object.assign({}, PV, {eligible: false, reason: 'unsupported'})); out.intro_unsupported = html('mf-box');
  await open(Object.assign({}, PV, {schema: 'forged'})); out.intro_forged = html('mf-box');
  // 2. preview, confirm, one apply with exactly the previewed pair
  RESP['/api/merge/finish'] = {finish: {outcome: 'finished', reason: null, archive: archive(FID), pending: 'docs/a.md.new'}};
  await open(); RESP['/api/merge/finish-preview'] = {finish_preview: FPV(FID, A, B)};
  const d0 = diffs.length; await mfAct('preview');
  out.confirm_html = html('mf-box'); out.drawn = diffs.slice(d0); out.before_confirm = fcalls().length;
  const r0 = refreshed; await mfAct('confirm');
  out.finish_calls = fcalls().map(c => c[1]); out.done_html = html('mf-box'); out.done_refresh = refreshed - r0; out.tag = els['sh-tag'].innerHTML;
  await mfAct('confirm'); out.second_confirm = fcalls().length;
  // 3. replace first, then a NEW finish preview: nothing from the old pair is reused
  await open(); RESP['/api/merge/replace'] = {replace: {outcome: 'replaced', archive: {current: 'x', incoming: 'y'}}};
  RESP['/api/merge/finish-preview'] = {finish_preview: FPV(FID, A, B)};
  await mpAct('use'); out.after_replace_finish_calls = fcalls().length;
  RESP['/api/merge/finish-preview'] = {finish_preview: FPV(FID2, C, D, {identical: true})};
  RESP['/api/merge/finish'] = {finish: {outcome: 'finished', archive: archive(FID2)}};
  await mfAct('preview'); out.identical_html = html('mf-box'); await mfAct('confirm');
  out.after_replace = {paths: calls.map(c => c[0]), body: fcalls()[0][1]};
  // 4. cancel / close / keep leave nothing to confirm
  await openPreviewed(); calls.length = 0; await mfAct('back'); out.back_calls = calls.map(c => c[0]); await mfAct('confirm'); out.back_confirm = fcalls().length;
  await openPreviewed(); closeSheet(); calls.length = 0; await mfAct('confirm'); out.closed_confirm = calls.length;
  await openPreviewed(); calls.length = 0; closed = 0; await mpAct('keep'); await mfAct('confirm'); out.keep = {calls: calls.length, closed: closed};
  // 5. slow and foreign responses cannot bind a pair
  await open(); let slow; RESP['/api/merge/finish-preview'] = () => new Promise(r => { slow = r; });
  const p1 = mfAct('preview'); closeSheet(); slow({finish_preview: FPV(FID, A, B)}); await p1; out.slow_closed = {mf: MF, html: html('mf-box')};
  await open(); let slow1; RESP['/api/merge/finish-preview'] = () => new Promise(r => { slow1 = r; });
  const q1 = mfAct('preview'); RESP['/api/merge/finish-preview'] = {finish_preview: FPV(FID2, C, D)}; await mfAct('preview');
  slow1({finish_preview: FPV(FID, A, B)}); await q1; out.overtaken = MF && MF.finish_id;
  await openPreviewed(FID, A, B, {pending: 'docs/other.md.new'}); out.other_pair = {mf: MF, html: html('mf-box')};
  await openPreviewed(FID, A, B, {schema: 'forged'}); out.forged_schema = MF;
  await openPreviewed(FID, A, B, {finish_id: 'finish-xyz'}); out.bad_id = {mf: MF, html: html('mf-box')};
  await openPreviewed(null, A, B, {eligible: false, reason: 'os_unsupported', finish_id: null, archive: null}); out.ineligible = {mf: MF, html: html('mf-box')};
  await mfAct('confirm'); out.ineligible_confirm = fcalls().length;
  RESP['/api/merge/finish-preview'] = () => Promise.reject(new Error('offline')); await mfAct('preview'); out.preview_offline = {mf: MF, html: html('mf-box')};
  await openPreviewed(); let slowF; RESP['/api/merge/finish'] = () => new Promise(r => { slowF = r; });
  const c1 = mfAct('confirm'), spin = html('mf-box'); closeSheet(); const r1 = refreshed;
  slowF({finish: {outcome: 'finished', archive: archive(FID)}}); await c1; out.slow_apply = {same_html: html('mf-box') === spin, refreshed: refreshed - r1};
  // 6. one in-flight apply at most
  await openPreviewed(); let slowD; RESP['/api/merge/finish'] = () => new Promise(r => { slowD = r; });
  const d1 = mfAct('confirm'), d2 = mfAct('confirm'); out.double_during = fcalls().length;
  slowD({finish: {outcome: 'finished', archive: archive(FID)}}); await d1; await d2; out.double_after = fcalls().length;
  // 7. every outcome has its own message; refresh only when something may have been written
  const outcomes = {finished: {outcome: 'finished', archive: archive(FID)}, still: {outcome: 'still_pending', reason: 'rename_failed', archive: archive(FID)},
    partial: {outcome: 'partial', reason: 'transaction_unfinished', txid: 'TX-1', archive: archive(FID)}, unproven: {outcome: 'partial', reason: 'verify_mismatch', txid: 'TX-2', archive: archive(FID)},
    conflict: {outcome: 'conflict', reason: 'moved_bytes_differ', archive: archive(FID)}, unknown: {outcome: 'unknown', reason: 'both_absent', archive: archive(FID)},
    stale: {outcome: 'stale', reason: 'stale_preview'}, collision: {outcome: 'blocked', reason: 'archive_collision'}, pending: {outcome: 'blocked', reason: 'transaction_pending'},
    os: {outcome: 'blocked', reason: 'os_unsupported'}, error: {outcome: 'error', reason: 'write_failed'}, net: null};
  out.outcomes = {};
  for (const k of Object.keys(outcomes)) {
    await openPreviewed(); RESP['/api/merge/finish'] = outcomes[k] === null ? () => Promise.reject(new Error('offline')) : {finish: outcomes[k]};
    const r2 = refreshed; await mfAct('confirm');
    out.outcomes[k] = {html: html('mf-box'), refreshed: refreshed - r2, tag: els['sh-tag'].innerHTML};
  }
  // 8. a lost response: observe the same operation, never send the apply again
  await openPreviewed(); RESP['/api/merge/finish'] = () => Promise.reject(new Error('offline')); await mfAct('confirm');
  const n0 = fcalls().length; out.lost_html = html('mf-box');
  RESP['/api/merge/finish-observe'] = {finish_observation: {schema: FOS, state: 'finished', reason: null, finish_id: FID, archive: archive(FID)}};
  const r3 = refreshed; await mfAct('observe');
  out.observe = {bodies: calls.filter(c => c[0] === '/api/merge/finish-observe').map(c => c[1]), resent: fcalls().length - n0, html: html('mf-box'), refreshed: refreshed - r3};
  RESP['/api/merge/finish-observe'] = {finish_observation: {schema: FOS, state: 'finished', reason: null, finish_id: FID2, archive: archive(FID2)}};
  await mfAct('observe'); out.foreign_observe = html('mf-box');
  RESP['/api/merge/finish-observe'] = {finish_observation: {schema: FOS, state: 'still_pending', reason: null, finish_id: FID, archive: archive(FID)}};
  await mfAct('observe'); out.still_observe = html('mf-box');
  RESP['/api/merge/finish-observe'] = {finish_observation: {schema: FOS, state: 'partial', reason: 'archive_incomplete', finish_id: FID, txid: null, unfinished: ['TX-9'], archive: archive(FID)}};
  await mfAct('observe'); out.partial_observe = html('mf-box');
  // 9. click delegation
  const listener = sheet.listeners[sheet.listeners.length - 1]; copied.length = 0;
  listener({target: {closest: s => s === '[data-mf]' ? {dataset: {mf: 'copy', v: 'p/q'}} : null}}); listener({target: {closest: () => null}});
  out.delegated = copied.slice();
  console.log(JSON.stringify(out));
})();
"""


def ui_block_c():
    text = UI_FILE.read_text(encoding='utf-8')
    return text[text.index('// D01C-BEGIN'):text.index('// D01C-END')]


@unittest.skipUnless(shutil.which('node'), 'node is needed to run the panel script outside a browser')
class MergeFinishUiScript(unittest.TestCase):
    """Chạy đúng hai khối JS (D01B + D01C) trong node với api giả và DOM tối thiểu: chứng minh luồng bất đồng bộ, không chứng minh giao diện trực quan.
    Runs the real D01B + D01C script blocks in node with a fake api and a minimal DOM: it proves the async flow, not the visual look in a browser."""

    @classmethod
    def setUpClass(cls):
        with tempfile.TemporaryDirectory(prefix='finish-ui-') as tmp:
            script = Path(tmp) / 'harness.js'
            script.write_text(HARNESS_C.replace('%BLOCKB%', ui_block()).replace('%BLOCKC%', ui_block_c()), encoding='utf-8')
            done = subprocess.run(['node', str(script)], capture_output=True, encoding='utf-8', timeout=60)
        assert done.returncode == 0, done.stdout + done.stderr
        cls.out = json.loads(done.stdout.strip().splitlines()[-1])

    def test_the_finish_entry_is_a_separate_button_for_supported_pairs_only(self):
        o = self.out
        self.assertIn('data-mf="preview"', o['intro'])
        self.assertIn('mf_btn', o['intro'])
        self.assertNotIn('data-mp=', o['intro'])  # tách khỏi Giữ bản của tôi / Dùng bản mới / separate from Keep mine and Use the new copy
        self.assertEqual(o['attach_calls'], ['/api/merge/preview'])  # mở so sánh không tự xem trước hay lưu / opening neither previews nor stores
        self.assertIn('data-mf="preview"', o['intro_identical'])
        for key in ('intro_unsupported', 'intro_forged'):
            self.assertIsNone(o[key])

    def test_confirm_needs_a_finish_preview_then_sends_the_previewed_pair_once(self):
        o = self.out
        for key in ('mf_will_cur', 'mf_will_save', 'mf_will_move', 'mf_will_get', 'mf_will_html', 'mf_will_two', 'mf_differ'):
            self.assertIn(key, o['confirm_html'])
        self.assertIn('data-mf="confirm"', o['confirm_html'])
        self.assertIn('data-mf="back"', o['confirm_html'])
        self.assertIn('.agent/merge-backups/finish-' + '1' * 32, o['confirm_html'])
        self.assertEqual(o['drawn'], [{'added': 2, 'removed': 1, 'hunks': []}])  # diff hiển thị lấy từ chính finish preview / the diff shown is the finish preview's own
        self.assertEqual(o['before_confirm'], 0)
        self.assertEqual(o['finish_calls'], [{'path': 'docs/a.md.new', 'current_sha256': 'a' * 64, 'candidate_sha256': 'b' * 64, 'finish_id': 'finish-' + '1' * 32}])
        self.assertIn('mf_done', o['done_html'])
        self.assertIn('pending-retired.txt', o['done_html'])
        self.assertEqual(o['done_refresh'], 1)
        self.assertIn('result_success', o['tag'])
        self.assertEqual(o['second_confirm'], 1)  # preview đã dùng / the preview is spent

    def test_after_a_replace_the_finish_uses_a_new_preview_and_never_the_old_hashes(self):
        o = self.out
        self.assertEqual(o['after_replace_finish_calls'], 0)  # replace không tự finish / replace does not finish
        self.assertIn('mf_same', o['identical_html'])
        self.assertEqual(o['after_replace']['body'], {'path': 'docs/a.md.new', 'current_sha256': 'c' * 64, 'candidate_sha256': 'd' * 64, 'finish_id': 'finish-' + '2' * 32})
        paths = o['after_replace']['paths']
        self.assertEqual(paths.count('/api/merge/finish-preview'), 1)
        self.assertLess(paths.index('/api/merge/replace'), paths.index('/api/merge/finish-preview'))

    def test_back_close_and_keep_leave_nothing_to_confirm(self):
        o = self.out
        self.assertEqual(o['back_calls'], ['/api/merge/preview'])  # quay lại chỉ xem trước replace lại / going back only reloads the replace preview
        self.assertEqual(o['back_confirm'], 0)
        self.assertEqual(o['closed_confirm'], 0)
        self.assertEqual(o['keep'], {'calls': 0, 'closed': 1})  # Giữ bản của tôi vẫn chỉ đóng, 0 ghi / Keep mine still only closes, zero writes

    def test_slow_foreign_or_malformed_previews_never_bind_a_pair(self):
        o = self.out
        self.assertIsNone(o['slow_closed']['mf'])
        self.assertNotIn('data-mf="confirm"', o['slow_closed']['html'])
        self.assertEqual(o['overtaken'], 'finish-' + '2' * 32)
        for key in ('other_pair', 'bad_id', 'ineligible'):
            self.assertIsNone(o[key]['mf'], key)
            self.assertNotIn('data-mf="confirm"', o[key]['html'], key)
        self.assertIsNone(o['forged_schema'])
        self.assertIn('mf_pv_fail', o['other_pair']['html'])
        self.assertIn('mf_r_os', o['ineligible']['html'])
        self.assertEqual(o['ineligible_confirm'], 0)
        self.assertIsNone(o['preview_offline']['mf'])
        self.assertIn('mf_pv_fail', o['preview_offline']['html'])

    def test_a_response_after_the_sheet_closed_still_refreshes_but_draws_nothing(self):
        self.assertEqual(self.out['slow_apply'], {'same_html': True, 'refreshed': 1})

    def test_a_double_click_sends_one_apply(self):
        self.assertEqual((self.out['double_during'], self.out['double_after']), (1, 1))

    def test_every_outcome_has_its_own_message_and_only_possible_writes_refresh(self):
        res = self.out['outcomes']
        expect = {'finished': ('mf_done', 1), 'still': ('mf_still', 1), 'partial': ('mf_partial:TX-1', 1), 'unproven': ('mf_unproven:TX-2', 1),
                  'conflict': ('mf_conflict', 1), 'unknown': ('mf_unknown_state', 1), 'stale': ('mf_stale', 0), 'collision': ('mf_collision', 0),
                  'pending': ('mf_r_pending', 0), 'os': ('mf_r_os', 0), 'error': ('mf_err', 1), 'net': ('mf_net', 1)}
        for key, (needle, refreshed) in expect.items():
            with self.subTest(outcome=key):
                self.assertIn(needle, res[key]['html'])
                self.assertEqual(res[key]['refreshed'], refreshed)
                if key != 'finished':
                    self.assertNotIn('mf_done', res[key]['html'])  # không báo xong khi chưa xong / never says finished when it is not
        for key in ('still', 'partial', 'unproven', 'conflict', 'unknown', 'net'):
            self.assertIn('data-mf="observe"', res[key]['html'], key)
        for key in ('finished', 'stale', 'collision', 'pending', 'os', 'error'):
            self.assertNotIn('data-mf="observe"', res[key]['html'], key)
        self.assertIn('mf_no_rollback', res['still']['html'])  # không chỉ dẫn Phục hồi để hoàn tác rename / no advice to roll back a move
        self.assertNotIn('mf_no_rollback', res['partial']['html'])
        self.assertIn('mf_w_rename', res['still']['html'])
        self.assertIn('mf_w_moved', res['conflict']['html'])
        self.assertIn('result_partial', res['partial']['tag'])
        self.assertIn('result_success', res['finished']['tag'])

    def test_a_lost_response_is_checked_with_the_same_operation_and_never_resent(self):
        o = self.out
        self.assertIn('mf_net', o['lost_html'])
        self.assertEqual(o['observe']['bodies'], [{'path': 'docs/a.md.new', 'current_sha256': 'a' * 64, 'candidate_sha256': 'b' * 64, 'finish_id': 'finish-' + '1' * 32}])
        self.assertEqual(o['observe']['resent'], 0)
        self.assertIn('mf_done', o['observe']['html'])
        self.assertEqual(o['observe']['refreshed'], 1)
        self.assertNotIn('mf_done', o['foreign_observe'])  # quan sát của ID khác không được tin / an observation of another id is not trusted
        self.assertIn('data-mf="observe"', o['still_observe'])
        self.assertIn('mf_still', o['still_observe'])
        self.assertIn('mf_unproven:TX-9', o['partial_observe'])  # giao dịch dở thấy từ danh sách / the open transaction comes from the list

    def test_click_delegation_maps_only_finish_buttons(self):
        self.assertEqual(self.out['delegated'], ['p/q'])

    def test_script_text_has_no_launcher_no_state_writer_and_one_apply_path(self):
        block = ui_block_c()
        for forbidden in ('/api/open', 'openPath', 'merge_state', 'launch', 'setTimeout', 'startfile', '/api/merge/replace', "api('/api/run'"):
            self.assertNotIn(forbidden, block)
        self.assertEqual(sorted(set(re.findall(r"api\('(/api/[a-z/-]+)'", block))),
                         ['/api/merge/finish', '/api/merge/finish-observe', '/api/merge/finish-preview'])
        self.assertEqual(len(re.findall(r'mfConfirm\(', block)), 2)  # định nghĩa + mfAct('confirm') / the definition and the one dispatch
        self.assertEqual(len(re.findall(r"api\('/api/merge/finish'", block)), 1)
        old = ui_block()
        self.assertIn("if(kind==='keep'){closeSheet();toast(t('mp_kept'));return Promise.resolve('kept');}", old)  # Giữ bản của tôi không đổi nghĩa / Keep mine unchanged
        self.assertEqual(sorted(set(re.findall(r"api\('(/api/[a-z/]+)'", old))), ['/api/merge/preview', '/api/merge/replace'])


class MergeFinishUiTexts(unittest.TestCase):
    def used(self):
        return set(re.findall(r"'(mf_[a-z_]+)'", ui_block_c()))

    def test_every_key_has_both_languages_no_dead_keys_and_no_false_claims(self):
        d = ui_dictionary()
        have = {k for k in d if k.startswith('mf_')}
        self.assertEqual(self.used() - have, set())
        self.assertEqual(have - self.used(), set())
        for key in sorted(have):
            with self.subTest(key=key):
                vi, en = d[key]
                self.assertTrue(vi.strip() and en.strip())
                for bad in MergeReplaceUiTexts.FORBIDDEN:
                    for text in (vi, en):
                        self.assertIsNone(re.search(r'(?<![A-Za-z])' + re.escape(bad) + r'(?![A-Za-z])', text), (key, bad))

    def test_failures_say_what_happened_and_the_next_step(self):
        d = ui_dictionary()
        for key in ('mf_pv_fail', 'mf_still', 'mf_not_started', 'mf_partial', 'mf_unproven', 'mf_conflict', 'mf_unknown_state', 'mf_net', 'mf_stale',
                    'mf_collision', 'mf_err', 'mf_r_pending'):
            self.assertIn('Bước tiếp', d[key][0], key)
            self.assertIn('Next step', d[key][1], key)
        for key in ('mf_pv_fail', 'mf_r_missing', 'mf_r_content', 'mf_r_unsafe', 'mf_r_pending', 'mf_r_os', 'mf_r_fs', 'mf_r_other', 'mf_err'):
            self.assertIn('Không có', d[key][0], key)
            self.assertTrue('Nothing was changed' in d[key][1] or 'No file was changed' in d[key][1], key)
        for key in ('mf_stale', 'mf_collision'):
            self.assertIn('chưa ghi gì', d[key][0])
            self.assertIn('nothing was written', d[key][1])
        for key in ('mf_partial', 'mf_unproven'):
            self.assertIn('%s', d[key][0])
            self.assertIn('%s', d[key][1])

    def test_the_decision_names_what_is_kept_stored_moved_and_not_updated(self):
        d = ui_dictionary()
        for key, vi_part, en_part in (('mf_will_cur', 'giữ nguyên', 'kept exactly'), ('mf_will_save', 'được lưu', 'are stored'),
                                      ('mf_will_move', 'không xóa', 'not deleted'), ('mf_will_get', 'bằng tay', 'by hand'),
                                      ('mf_will_html', 'MANIFEST', 'MANIFEST'), ('mf_will_two', 'hai bước', 'two separate steps'),
                                      ('mf_done_note', 'pending-retired.txt', 'pending-retired.txt')):
            self.assertIn(vi_part, d[key][0], key)
            self.assertIn(en_part, d[key][1], key)
        self.assertIn('không xác nhận', d['mf_will_html'][0])
        self.assertIn('does not confirm', d['mf_will_html'][1])
        self.assertNotIn(d['mf_btn'][0], (d['mp_keep'][0], d['mp_use'][0]))  # nút riêng, không trùng nút cũ / its own button
        self.assertIn('CHƯA', d['mf_still'][0])
        self.assertIn('NOT', d['mf_still'][1])
        self.assertIn('không hoàn tác việc chuyển file', d['mf_no_rollback'][0])
        self.assertIn('does not reverse a file move', d['mf_no_rollback'][1])
        self.assertIn('không tự chuyển', d['mf_partial'][0])
        self.assertIn('does not move', d['mf_partial'][1])


@windows_only
class MergeFinishLifecycle(MergeReplaceBase):
    """Gói tối thiểu A→B→C với MANIFEST nhất quán, chạy upgrade CLI thật. Không full lifecycle .cvf/migration.
    Minimal packages A -> B -> C with a consistent MANIFEST and the real upgrade CLI. Not a full .cvf / migration lifecycle."""

    def package(self, name, version, guide_bytes):
        r = self.root / name
        write(r, 'VERSION', version + '\n')
        write(r, 'untouched.txt', 'package file\n')
        write(r, 'PROJECT.md', 'private context\n')
        for tool in ('check_package.py', 'result_schema.py'):
            write(r, 'tools/' + tool, (PKG / 'tools' / tool).read_bytes())
        write(r, self.guide, guide_bytes)
        manifest(r)
        return r

    def start(self):
        self.pkg = {n: self.package('pkg-' + n, v, b) for n, v, b in (('a', '1.0.0', b'guide A\r\nline A\r\n'), ('b', '1.1.0', b'guide B\nline B\n'),
                                                                      ('c', '1.2.0', b'guide C\nline C\n'))}
        proj = self.root / 'upgraded'
        shutil.copytree(str(self.pkg['a']), str(proj))
        write(proj, '.agent/PACKAGE_INFO.json', json.dumps({'package': 'local-agent-workspace', 'version': '1.0.0', 'language': 'vi'}))
        write(proj, self.guide, b'my own edit of the guide\r\n')
        done = run_script('upgrade_package.py', '--project', proj, '--from', self.pkg['b'], '--apply')
        self.assertEqual(done.returncode, 0, done.stdout + done.stderr)
        self.assertEqual((proj / (self.guide + '.new')).read_bytes(), b'guide B\nline B\n')  # xung đột → .new của B / conflict -> B's .new
        return proj

    def verify(self, proj):
        e = dict(os.environ, PYTHONUTF8='1', PYTHONDONTWRITEBYTECODE='1')
        return subprocess.run([sys.executable, str(proj / 'tools' / 'check_package.py'), '--verify'], cwd=str(proj), env=e,
                              capture_output=True, encoding='utf-8', timeout=120).returncode

    def finish(self, proj):
        pv = merge_center.preview_finish(str(proj), self.guide + '.new')
        self.assertTrue(pv['eligible'], pv)
        res = merge_center.apply_finish(str(proj), self.guide + '.new', pv['current_sha256'], pv['candidate_sha256'], pv['finish_id'])
        self.assertEqual(res['outcome'], 'finished', res)
        return pv

    def test_replace_with_the_candidate_then_finish_lets_the_next_upgrade_update_normally(self):
        proj = self.start()
        pv = merge_center.preview_replace(str(proj), self.guide + '.new')
        rep = merge_center.apply_replace(str(proj), self.guide + '.new', pv['current_sha256'], pv['candidate_sha256'])
        self.assertEqual(rep['outcome'], 'replaced')
        self.assertEqual(self.verify(proj), 0)  # candidate là baseline của B nên verify khớp / the candidate is B's baseline, so verify matches
        fin = self.finish(proj)
        self.assertFalse((proj / (self.guide + '.new')).exists())
        before = tree_state(proj)
        preview = run_script('upgrade_package.py', '--project', proj, '--from', self.pkg['c'])
        self.assertEqual(preview.returncode, 0, preview.stdout + preview.stderr)
        self.assertEqual(tree_state(proj), before)  # preview nâng cấp không ghi / the upgrade preview writes nothing
        done = run_script('upgrade_package.py', '--project', proj, '--from', self.pkg['c'], '--apply')
        self.assertEqual(done.returncode, 0, done.stdout + done.stderr)
        self.assertEqual((proj / self.guide).read_bytes(), b'guide C\nline C\n')
        self.assertFalse((proj / (self.guide + '.new')).exists())
        self.assertEqual((proj / '.agent' / 'merge-backups' / fin['finish_id'] / 'pending-retired.txt').read_bytes(), b'guide B\nline B\n')

    def test_keep_then_finish_removes_the_stale_pending_that_blocked_the_next_upgrade(self):
        proj = self.start()
        self.assertEqual(self.verify(proj), 1)  # giữ bản của tôi nên lệch MANIFEST B / keeping my file differs from B's MANIFEST
        before = tree_state(proj)
        blocked = run_script('upgrade_package.py', '--project', proj, '--from', self.pkg['c'], '--apply')
        self.assertNotEqual(blocked.returncode, 0)  # .new của B khác nội dung C chặn nâng cấp / B's .new differs from C and blocks it
        self.assertEqual(tree_state(proj), before)
        fin = self.finish(proj)
        self.assertEqual((proj / self.guide).read_bytes(), b'my own edit of the guide\r\n')
        archived = {n: (proj / '.agent' / 'merge-backups' / fin['finish_id'] / n).read_bytes() for n in ('current.txt', 'incoming.txt', 'pending-retired.txt', 'record.json')}
        self.assertEqual((archived['current.txt'], archived['incoming.txt'], archived['pending-retired.txt']),
                         (b'my own edit of the guide\r\n', b'guide B\nline B\n', b'guide B\nline B\n'))
        done = run_script('upgrade_package.py', '--project', proj, '--from', self.pkg['c'], '--apply')
        self.assertEqual(done.returncode, 0, done.stdout + done.stderr)
        self.assertEqual((proj / self.guide).read_bytes(), b'my own edit of the guide\r\n')  # current vẫn của người dùng / current stays the user's
        self.assertEqual((proj / (self.guide + '.new')).read_bytes(), b'guide C\nline C\n')  # candidate mới của C theo engine cũ / a fresh C candidate from the unchanged engine
        self.assertEqual({n: (proj / '.agent' / 'merge-backups' / fin['finish_id'] / n).read_bytes() for n in archived}, archived)


if __name__ == '__main__':
    unittest.main()
