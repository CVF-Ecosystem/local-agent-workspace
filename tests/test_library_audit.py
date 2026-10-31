import base64
import contextlib
import hashlib
import io
import json
import os
import pathlib
import subprocess
import sys
import tempfile
import types
import unittest
from unittest import mock

import openpyxl
import zipfile
import zlib

try:
    import pdfminer.high_level
    HAS_PDF_ENGINE = True
except ImportError:
    HAS_PDF_ENGINE = False

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / 'tools'))
import library
import library_audit as audit


# Password-protected DATAMOCK fixture generated once with reportlab 4.4.4;
# no generator dependency is needed to run the tests. Password is the literal test value secret.
ENCRYPTED_PDF = 'JVBERi0xLjMKJZOMi54gUmVwb3J0TGFiIEdlbmVyYXRlZCBQREYgZG9jdW1lbnQgaHR0cDovL3d3dy5yZXBvcnRsYWIuY29tCjEgMCBvYmoKPDwKL0YxIDIgMCBSCj4+CmVuZG9iagoyIDAgb2JqCjw8Ci9CYXNlRm9udCAvSGVsdmV0aWNhIC9FbmNvZGluZyAvV2luQW5zaUVuY29kaW5nIC9OYW1lIC9GMSAvU3VidHlwZSAvVHlwZTEgL1R5cGUgL0ZvbnQKPj4KZW5kb2JqCjMgMCBvYmoKPDwKL0NvbnRlbnRzIDggMCBSIC9NZWRpYUJveCBbIDAgMCA1OTUuMjc1NiA4NDEuODg5OCBdIC9QYXJlbnQgNyAwIFIgL1Jlc291cmNlcyA8PAovRm9udCAxIDAgUiAvUHJvY1NldCBbIC9QREYgL1RleHQgL0ltYWdlQiAvSW1hZ2VDIC9JbWFnZUkgXQo+PiAvUm90YXRlIDAgL1RyYW5zIDw8Cgo+PiAKICAvVHlwZSAvUGFnZQo+PgplbmRvYmoKNCAwIG9iago8PAovUGFnZU1vZGUgL1VzZU5vbmUgL1BhZ2VzIDcgMCBSIC9UeXBlIC9DYXRhbG9nCj4+CmVuZG9iago1IDAgb2JqCjw8Ci9BdXRob3IgKFwzNjJcMjE2XlwwMTVcMzUxXFxcMzA0SVwwMDYpIC9DcmVhdGlvbkRhdGUgKFwzMjdcMzMyXDAwM1NcMjQyXDAwN1wyMzJcMDE0Rj9cMjAzXDI2NVwyNzZcMjYxXDM1MVwyMDBcMzUxXDAyN1wyMjRKM1wzMTRcMjU1KSAvQ3JlYXRvciAoXDMwMVwyMDVBXDAxNFwzNDJFXDM0N11cMDI3L1wzNDNcMzA2XDMxMlwyNTBcMjI1XDMyMVwyNDBVXDMwMlwwMzd6XDMzNFwyNDdcMjA2XDM2MlwzMDZcMDE1TFwyNzRcMjQyXDM2N1wzMDRcMzI0XDIzNVwzMjJkdVwzMjRcMzUzXDAzMVwwMjcpIC9LZXl3b3JkcyAoKSAvTW9kRGF0ZSAoXDMyN1wzMzJcMDAzU1wyNDJcMDA3XDIzMlwwMTRGP1wyMDNcMjY1XDI3NlwyNjFcMzUxXDIwMFwzNTFcMDI3XDIyNEozXDMxNFwyNTUpIC9Qcm9kdWNlciAoXDMwMVwyMDVBXDAxNFwzNDJFXDM0N11cMDI3L1wzNDNcMzA2XDMxMlwyNTBcMjI1XDMyMVwyNDBVXDMwMlwwMzd6XDMzNFwyNDdcMjA2XDM2MlwzMDZcMDE1TFwyNzRcMjQyXDM2N1wzMDRcMzI0XDIzNVwzMjJkdVwzMjRcMzUzXDAzMVwwMjcpIAogIC9TdWJqZWN0IChcMzQ2XDIxNkJcMDIzXDM2NVJcMzAyWlwwMzRqXDMyNykgL1RpdGxlIChcMzQ2XDIxNkVcMDEyXDM0NF1cMzE2WCkgL1RyYXBwZWQgL0ZhbHNlCj4+CmVuZG9iago2IDAgb2JqCjw8Ci9GaWx0ZXIgL1N0YW5kYXJkIC9PIDxFNUE4RDI2ODdCRDlEMENGRjk0NkI3QUM1NUY1MTA4MURDRjBEMTE2NTU0QzRCRkNCMEE1RTQ0NkY2OUVBNDhBPiAvUCAtMjAgL1IgMiAvVSA8N0NFNjNFNDFFNjlFM0UwQjc4NUMxMDc4N0ZFQUY0NkI1RDdFODU4MzM5RjgwNDMyNTNDRkY4Qjk1NTg2QjE4NT4gL1YgMQo+PgplbmRvYmoKNyAwIG9iago8PAovQ291bnQgMSAvS2lkcyBbIDMgMCBSIF0gL1R5cGUgL1BhZ2VzCj4+CmVuZG9iago4IDAgb2JqCjw8Ci9GaWx0ZXIgWyAvQVNDSUk4NURlY29kZSAvRmxhdGVEZWNvZGUgXSAvTGVuZ3RoIDk1Cj4+CnN0cmVhbQqtcylK1wWoWUhYegtP1ro0zC33rmueEkjPELIcmRGxDyV4TcYt8YhhBvV3EFbQZpKq+osCrYWCXzb8bbyv7snZVIicPK/yjehC3xxRNIgMtMxFkZ0n1OO9JNn9cr+n5WVuZHN0cmVhbQplbmRvYmoKeHJlZgowIDkKMDAwMDAwMDAwMCA2NTUzNSBmIAowMDAwMDAwMDczIDAwMDAwIG4gCjAwMDAwMDAxMDQgMDAwMDAgbiAKMDAwMDAwMDIxMSAwMDAwMCBuIAowMDAwMDAwNDE0IDAwMDAwIG4gCjAwMDAwMDA0ODIgMDAwMDAgbiAKMDAwMDAwMTEzMyAwMDAwMCBuIAowMDAwMDAxMzI5IDAwMDAwIG4gCjAwMDAwMDEzODggMDAwMDAgbiAKdHJhaWxlcgo8PAovRW5jcnlwdCA2IDAgUgovSUQgCls8NTM4ZDE4MzUzMTNkNGYwMTcwNTQyMjBlNTUzYzhmZmU+PDUzOGQxODM1MzEzZDRmMDE3MDU0MjIwZTU1M2M4ZmZlPl0KJSBSZXBvcnRMYWIgZ2VuZXJhdGVkIFBERiBkb2N1bWVudCAtLSBkaWdlc3QgKGh0dHA6Ly93d3cucmVwb3J0bGFiLmNvbSkKCi9JbmZvIDUgMCBSCi9Sb290IDQgMCBSCi9TaXplIDkKPj4Kc3RhcnR4cmVmCjE1NzIKJSVFT0YK'


def pdf(pages, encrypted=False):
    if encrypted:
        return base64.b64decode(ENCRYPTED_PDF)
    objects = [b'', b'', b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
    pixels = zlib.compress(bytes([0, 255, 0, 255] * 4))
    objects.append(b'<< /Type /XObject /Subtype /Image /Width 4 /Height 4 /ColorSpace /DeviceGray /BitsPerComponent 8 /Filter /FlateDecode /Length ' + str(len(pixels)).encode() + b' >>\nstream\n' + pixels + b'\nendstream')
    kids = []
    for text in pages:
        page, stream = len(objects) + 1, len(objects) + 2
        kids.append(('%d 0 R' % page).encode())
        objects.append(('<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> /XObject << /Scan 4 0 R >> >> /Contents %d 0 R >>' % stream).encode())
        if text is None:
            content = b'q 300 0 0 300 50 400 cm /Scan Do Q'
        else:
            escaped = text.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
            content = ('BT /F1 12 Tf 50 750 Td (%s) Tj ET' % escaped).encode('ascii')
        objects.append(b'<< /Length ' + str(len(content)).encode() + b' >>\nstream\n' + content + b'\nendstream')
    objects[0] = b'<< /Type /Catalog /Pages 2 0 R >>'
    objects[1] = b'<< /Type /Pages /Count ' + str(len(kids)).encode() + b' /Kids [' + b' '.join(kids) + b'] >>'
    raw, offsets = bytearray(b'%PDF-1.4\n'), [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(raw))
        raw.extend(('%d 0 obj\n' % number).encode() + obj + b'\nendobj\n')
    xref = len(raw)
    raw.extend(('xref\n0 %d\n0000000000 65535 f \n' % len(offsets)).encode())
    for offset in offsets[1:]:
        raw.extend(('%010d 00000 n \n' % offset).encode())
    raw.extend(('trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n' % (len(offsets), xref)).encode())
    return bytes(raw)


def office(ext='.docx', wrong_relationship=False, dtd=False):
    part, namespace, root, content_type = {
        '.docx': ('word/document.xml', 'wordprocessingml', 'document', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml'),
        '.xlsx': ('xl/workbook.xml', 'spreadsheetml', 'workbook', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml'),
        '.pptx': ('ppt/presentation.xml', 'presentationml', 'presentation', 'application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml')
    }[ext]
    content = '<%s xmlns="http://schemas.openxmlformats.org/%s/2006/main"/>' % (root, namespace)
    if dtd:
        content = ('<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE %s [<!ENTITY x "fixture">]>' % root + content).encode('utf-16')
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w') as z:
        z.writestr('[Content_Types].xml', '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Override PartName="/%s" ContentType="%s"/></Types>' % (part, content_type))
        z.writestr(part, content)
        z.writestr('_rels/.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="%s"/></Relationships>' % ('wrong.xml' if wrong_relationship else part))
    return out.getvalue()


class AuditTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = pathlib.Path(self.tmp.name)
        self.ref = self.root / 'references'
        self.ref.mkdir()
        self.entries = []
        self.records = []
        self.patches = [mock.patch.object(library, 'ROOT', str(self.root)),
                        mock.patch.object(library, 'REF', str(self.ref)),
                        mock.patch.object(library, 'INDEX', str(self.ref / '_INDEX.xlsx'))]
        for p in self.patches:
            p.start()
            self.addCleanup(p.stop)

    def add(self, name, raw, ident, cached=None, verify='Cần đối chiếu', revision=1, state='PRESENT'):
        (self.ref / name).write_bytes(raw)
        sha = hashlib.sha256(raw).hexdigest()
        self.records.append({'id': ident, 'group': 'Nội bộ', 'number': ident, 'title': ident,
                             'status': 'Không rõ', 'fulltext': 'Có', 'file': name, 'verify': verify,
                             'law_revision': revision, 'law_sha': cached if cached is not None else sha,
                             'law_state': state})
        return sha

    def save(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = library.SHEET
        keys = list(library.KEYS)
        for i, key in enumerate(keys, 1):
            ws.cell(4, i, library.KEYS[key])
        for row, record in enumerate(self.records, 5):
            for col, key in enumerate(keys, 1):
                ws.cell(row, col, record.get(key))
        if self.entries:
            sheet = wb.create_sheet(library.CONTENT_SHEET)
            sheet.cell(1, 1, library.CONTENT_MARKER)
            for i, name in enumerate(library.CONTENT_HEADERS, 1):
                sheet.cell(4, i, name)
            for row, entry in enumerate(self.entries, 5):
                for col, key in enumerate(library.CONTENT_HEADERS, 1):
                    sheet.cell(row, col, entry.get(key, ''))
        wb.save(self.ref / '_INDEX.xlsx')

    def run_check(self, enabled=True):
        capture = io.StringIO()
        with contextlib.redirect_stdout(capture):
            code = library.cmd_check({'json': True, 'content_audit': enabled})
        return code, json.loads(capture.getvalue())

    def messages(self, result):
        return '\n'.join(x['msg'] for x in result['warnings'])

    @unittest.skipUnless(HAS_PDF_ENGINE, 'optional pdfminer.six absent; real PDF extraction NOT_TESTED')
    def test_real_pdfs_and_custody(self):
        rich = 'Text ' * 100
        self.add('text.pdf', pdf([rich]), 'A')
        self.add('scan.pdf', pdf([None]), 'B')
        self.add('mixed.pdf', pdf([rich, None]), 'C')
        self.add('low.pdf', pdf(['Short legitimate title']), 'D')
        self.add('encrypted.pdf', pdf(['secret'], True), 'E')
        self.add('broken.pdf', b'%PDF-1.4\nbroken', 'F')
        self.add('fake.pdf', b'<html>error</html>', 'G')
        self.save()
        before = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in self.ref.iterdir()}
        code, result = self.run_check()
        after = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in self.ref.iterdir()}
        self.assertEqual(before, after)
        self.assertEqual(code, 0)
        byrow = {row: ' '.join(x['msg'] for x in result['warnings'] if x['row'] == row) for row in range(5, 12)}
        self.assertIn('1 trang', byrow[5])
        self.assertIn('trang 1: 0', byrow[6])
        self.assertIn('trang 2: 0', byrow[7])
        self.assertIn('mật độ thấp', byrow[8])
        self.assertIn('chưa kiểm lớp chữ', byrow[9])
        self.assertIn('chưa kiểm lớp chữ', byrow[10])
        self.assertIn('chữ ký PDF', byrow[11])

    def test_hash_and_scoped_binding(self):
        raw = b'alpha'
        sha = self.add('a.txt', raw, 'A', verify='Đã đối chiếu toàn văn', revision=2, state='NEEDS_REVIEW')
        self.add('b.txt', raw, 'B', cached='1' * 64, verify='Đã đối chiếu toàn văn')
        self.add('c.txt', b'different', 'C', cached=sha)
        self.entries.append(dict(entry_id='E', source_id='A', source_revision=2, source_sha256=sha,
                                 source_path='a.txt', locator='p1', read_scope='page 1', kind='EXTRACT', text='x',
                                 recorded_by='test', recorded_at='2026-01-01', check_level='DECLARED_CHECK',
                                 check_scope='page 1', check_evidence_ref='fixture', checked_by='test', checked_at='2026-01-01',
                                 basis_status='Không rõ'))
        self.save()
        _, out = self.run_check()
        m = self.messages(out)
        self.assertIn('SHA đã ghi trùng ID A', m)
        self.assertIn('SHA file thực trùng ID A', m)
        self.assertIn('bytes chưa đồng bộ', m)
        a = ' '.join(x['msg'] for x in out['warnings'] if x['row'] == 5)
        self.assertIn('phạm vi page 1', a)
        self.assertNotIn('thiếu căn cứ gắn', a)
        self.assertIn('cần đối chiếu revision mới', a)
        b = ' '.join(x['msg'] for x in out['warnings'] if x['row'] == 6)
        self.assertIn('thiếu căn cứ gắn', b)

    def test_limits_office_and_missing_engine(self):
        self.assertEqual(audit.pdf_summary(b'x' * (audit.BYTE_LIMIT + 1))['reason'], 'byte_limit')
        with mock.patch.object(audit.subprocess, 'run', side_effect=subprocess.TimeoutExpired('x', 15)):
            self.assertEqual(audit.pdf_summary(b'x')['reason'], 'time_limit')
        with mock.patch.dict(sys.modules, {'pdfminer.high_level': None}):
            with mock.patch('sys.stdin', io.TextIOWrapper(io.BytesIO(b''), encoding='utf-8')):
                with contextlib.redirect_stdout(io.StringIO()) as out:
                    audit.pdf_worker()
            self.assertEqual(json.loads(out.getvalue())['reason'], 'engine_missing')
        fake = io.BytesIO()
        with __import__('zipfile').ZipFile(fake, 'w') as z:
            z.writestr('wrong.xml', '<x/>')
        self.assertEqual(audit.office_reason(fake.getvalue(), '.docx'), 'wrong_office_structure')
        for ext in ('.docx', '.xlsx', '.pptx'):
            self.assertEqual(audit.office_reason(office(ext), ext), '')
        self.assertEqual(audit.office_reason(office(wrong_relationship=True), '.docx'), 'wrong_office_structure')
        self.assertEqual(audit.office_reason(office(dtd=True), '.docx'), 'invalid_office')

    @unittest.skipUnless(HAS_PDF_ENGINE, 'optional pdfminer.six absent; real page limit NOT_TESTED')
    def test_pdf_page_limit(self):
        result = audit.pdf_summary(pdf([None] * 101))
        self.assertEqual(result['reason'], 'page_limit')
        self.assertEqual(len(result['pages']), 100)

    def test_partial_pdf_failure_keeps_completed_page_results(self):
        def extract(raw, maxpages):
            self.assertEqual(maxpages, 100)
            yield []
            raise ValueError('damaged later page')

        modules = {'pdfminer.high_level': types.SimpleNamespace(extract_pages=extract),
                   'pdfminer.layout': types.SimpleNamespace(LTTextContainer=object)}
        with mock.patch.dict(sys.modules, modules), mock.patch('sys.stdin', io.TextIOWrapper(io.BytesIO(b'%PDF-1.4'))):
            with contextlib.redirect_stdout(io.StringIO()) as output:
                audit.pdf_worker()
        result = json.loads(output.getvalue())
        self.assertEqual(result['pages'], [0])
        self.assertEqual(result['reason'], 'pdf_error:ValueError')

    def test_default_json_exit_and_bad_path(self):
        self.add('a.txt', b'one', 'A')
        self.records[0]['file'] = '../outside/secret.txt'
        self.records[0]['fulltext'] = 'Không'
        self.save()
        _, plain = self.run_check(False)
        with mock.patch('builtins.open', wraps=open) as opened:
            code, result = self.run_check(True)
        self.assertEqual(code, 0)
        self.assertEqual(set(plain), set(result))
        self.assertIn('out_of_scope', self.messages(result))
        self.assertFalse(any('secret.txt' in str(call.args[0]) for call in opened.call_args_list))
        self.records.append(dict(self.records[0]))
        self.save()
        self.assertEqual(self.run_check(True)[0], 1)
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(library.main(['bogus']), 2)

    def test_cli_result_and_missing_openpyxl(self):
        self.add('a.pdf', b'<html>bad</html>', 'A')
        self.save()
        script = ('import sys; sys.path.insert(0, sys.argv[1]); import library; '
                  'library.ROOT=sys.argv[2]; library.REF=sys.argv[2]+"/references"; '
                  'sys.exit(library.run_with_result(["check", "--content-audit", "--json"]))')
        run = subprocess.run([sys.executable, '-c', script, str(pathlib.Path(library.__file__).parent), str(self.root)],
                             capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(run.returncode, 0, run.stderr)
        self.assertIn('PDF', run.stdout)
        self.assertEqual(json.loads(run.stdout.split('@@RESULT_JSON@@', 1)[1])['code'], 0)
        with mock.patch.dict(sys.modules, {'openpyxl': None}):
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                with self.assertRaises(SystemExit) as error:
                    library.main(['check', '--content-audit'])
        self.assertEqual(error.exception.code, 2)

    def test_malformed_binding_and_read_error(self):
        self.add('a.txt', b'one', 'A', verify='Đã đối chiếu toàn văn')
        self.save()
        wb = openpyxl.load_workbook(self.ref / '_INDEX.xlsx')
        wb.create_sheet(library.CONTENT_SHEET)['A1'] = 'invalid'
        wb.save(self.ref / '_INDEX.xlsx')
        _, result = self.run_check()
        self.assertIn('content_sheet_unavailable_or_changed', self.messages(result))
        real_open = open
        def deny(path, *args, **kwargs):
            if str(path).endswith('a.txt'):
                raise PermissionError('denied fixture')
            return real_open(path, *args, **kwargs)
        with mock.patch('builtins.open', side_effect=deny):
            code, result = self.run_check()
        self.assertEqual(code, 0)
        self.assertIn('chưa kiểm nguồn/lớp chữ', self.messages(result))
        self.assertIn('chưa kiểm binding', self.messages(result))

    def test_source_saved_during_read_is_not_used_as_hash_evidence(self):
        self.add('a.txt', b'one', 'A', verify='Đã đối chiếu toàn văn')
        self.save()
        index_before = (self.ref / '_INDEX.xlsx').read_bytes()
        real_open = open
        target = self.ref / 'a.txt'

        @contextlib.contextmanager
        def save_after_read(path, *args, **kwargs):
            with real_open(path, *args, **kwargs) as stream:
                if str(path) == str(target) and args and args[0] == 'rb':
                    class ConcurrentSave:
                        def read(self, count):
                            raw = stream.read(count)
                            with real_open(target, 'wb') as external_writer:
                                external_writer.write(b'independently saved newer bytes')
                            return raw
                    yield ConcurrentSave()
                else:
                    yield stream

        with mock.patch('builtins.open', side_effect=save_after_read):
            code, result = self.run_check()
        self.assertEqual(code, 0)
        self.assertIn('source_changed_during_read', self.messages(result))
        self.assertNotIn('SHA file thực', self.messages(result))
        self.assertEqual(target.read_bytes(), b'independently saved newer bytes')
        self.assertEqual((self.ref / '_INDEX.xlsx').read_bytes(), index_before)

    def test_limits_and_missing_engine_through_cli_preserve_sources(self):
        self.add('a.pdf', pdf([None]), 'A')
        self.save()
        before = {p.name: p.read_bytes() for p in self.ref.iterdir()}
        with mock.patch.object(audit, 'BYTE_LIMIT', 32), mock.patch.object(audit, 'pdf_summary') as summary:
            with contextlib.redirect_stdout(io.StringIO()) as output:
                code = library.main(['check', '--content-audit', '--json'])
            self.assertEqual(code, 0)
            self.assertIn('byte_limit', output.getvalue())
            summary.assert_not_called()
        with mock.patch.object(audit, 'pdf_summary', return_value={'pages': [], 'reason': 'engine_missing'}):
            code, result = self.run_check()
            self.assertEqual(code, 0)
            self.assertIn('engine_missing', self.messages(result))
            self.assertIn('chưa kiểm lớp chữ', self.messages(result))
        self.assertEqual({p.name: p.read_bytes() for p in self.ref.iterdir()}, before)

    def test_timeout_terminates_a_real_child(self):
        real_run = subprocess.run

        def slow_child(args, **kwargs):
            return real_run([sys.executable, '-c', 'import time; time.sleep(2)'], **kwargs)

        with mock.patch.object(audit, 'TIME_LIMIT', 0.05), mock.patch.object(audit.subprocess, 'run', side_effect=slow_child):
            self.assertEqual(audit.pdf_summary(b'%PDF-1.4')['reason'], 'time_limit')

    @unittest.skipUnless(os.name == 'nt', 'Windows junction test')
    def test_junction_does_not_open_target(self):
        outside = self.root / 'outside'
        outside.mkdir()
        (outside / 'secret.txt').write_bytes(b'outside secret')
        link = self.ref / 'linked'
        made = subprocess.run(['cmd', '/c', 'mklink', '/J', str(link), str(outside)], capture_output=True)
        self.assertEqual(made.returncode, 0, made.stderr.decode(errors='replace'))
        self.add('a.txt', b'one', 'A')
        self.records[0]['file'] = 'linked/secret.txt'
        self.records[0]['fulltext'] = 'Không'
        self.save()
        with mock.patch('builtins.open', wraps=open) as opened:
            _, result = self.run_check()
        self.assertIn('link', self.messages(result))
        self.assertFalse(any('secret.txt' in str(call.args[0]) for call in opened.call_args_list))


if __name__ == '__main__':
    unittest.main()
