"""Offline intake/recovery acceptance on independent temporary projects, never real source data."""
import contextlib
import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import stat
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'tools'))
import library
import library_intake as intake
import library_sync as sync
import transaction as tx
import result_schema

try:
    import openpyxl
except ImportError:
    openpyxl = None


def snapshot(root, runtime=False):
    return {p.relative_to(root).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in root.rglob('*') if p.is_file() and (runtime or '.agent/runtime/' not in p.relative_to(root).as_posix())}


@unittest.skipUnless(openpyxl, 'openpyxl is required for the source workbook')
class Intake(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='law-intake-')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.ref = self.root / 'references'
        self.folder = self.ref / 'legal' / '_moi_tai'
        self.folder.mkdir(parents=True)
        (self.ref / 'legal' / 'team').mkdir()
        (self.root / '.agent').mkdir()
        self.index = self.ref / '_INDEX.xlsx'
        shutil.copyfile(ROOT / 'skills/local/shared/source-library/assets/source-index-template.xlsx', self.index)
        self.config = {'schema': intake.SCHEMA, 'enabled': True, 'staging': 'references/legal/_moi_tai',
                       'mappings': [{'code': 'LAW-01', 'canonical_name': 'Canonical.txt', 'destination': 'legal/team/Canonical.txt'}]}
        self.config_path = self.root / intake.CONFIG
        self.write_config()
        self.source = self.folder / 'LAW-01_source.txt'
        self.source.write_bytes(b'RAW source\r\nbytes\x00')
        self.target = self.ref / 'legal/team/Canonical.txt'
        (self.root / 'private.txt').write_bytes(b'PRIVATE sentinel')
        self.selection = intake.request()

    def write_config(self):
        self.config_path.write_text(json.dumps(self.config), encoding='utf-8')

    def plan(self, selection=None):
        return sync.intake_plan(str(self.root), selection or self.selection)

    def apply(self, result=None, selection=None):
        return sync.intake_apply(str(self.root), selection or self.selection, (result or self.plan(selection))['plan_token'])

    def edit_index(self, fn):
        wb = openpyxl.load_workbook(self.index)
        try:
            fn(wb)
            wb.save(self.index)
        finally:
            wb.close()

    def read_cells(self):
        wb = openpyxl.load_workbook(self.index, data_only=False)
        try:
            return sync.cell_values(wb)
        finally:
            wb.close()

    def fault(self, point='commit:1'):
        plan = self.plan()
        with patch.dict(os.environ, {tx.FAULT_ENV: point}):
            with self.assertRaises(sync.SyncError):
                self.apply(plan)
        pending = tx.incomplete(str(self.root), strict=True)
        return pending[0]['id'] if pending else None

    def test_preview_is_read_only_and_scanners_exclude_staging(self):
        before = snapshot(self.root, True)
        result = self.plan()
        self.assertEqual(snapshot(self.root, True), before)
        self.assertFalse((self.root / '.agent/runtime').exists())
        self.assertEqual(result['items'][0]['id'], 'SRC-000001')
        self.assertEqual(sync.sources(str(self.root))[0], {})
        with patch.object(library, 'ROOT', str(self.root)), patch.object(library, 'REF', str(self.ref)):
            self.assertNotIn('legal/_moi_tai/LAW-01_source.txt', library.library_files())
        self.assertEqual(library.locate_source(str(self.root), 'legal/_moi_tai/LAW-01_source.txt')[1], 'out_of_scope')
        # Audit consumes the same managed sources; do not invoke an extractor for staged bytes.
        import library_audit
        fake_row = {'_row': 5, 'id': 'STAGED', 'file': 'legal/_moi_tai/LAW-01_source.txt'}
        real_open = open
        def spy(path, *args, **kwargs):
            if os.path.normcase(str(path)) == os.path.normcase(str(self.source)):
                self.fail('audit attempted to read staging bytes')
            return real_open(path, *args, **kwargs)
        with patch('builtins.open', spy):
            warnings = library_audit.audit(str(self.root), [fake_row], library)
        self.assertTrue(any('out_of_scope' in text for _, text in warnings))

    def test_success_exact_bytes_metadata_blank_fulltext_repeat_refuses(self):
        raw, private = self.source.read_bytes(), (self.root / 'private.txt').read_bytes()
        self.apply()
        self.assertEqual(self.target.read_bytes(), raw)
        self.assertEqual(self.source.read_bytes(), raw)
        self.assertEqual((self.root / 'private.txt').read_bytes(), private)
        cells = self.read_cells()['Danh mục']
        self.assertEqual(cells[(5, 1)][0], 'SRC-000001')
        self.assertEqual(cells[(5, 14)][0], 'Cần đối chiếu')
        self.assertNotIn((5, 15), cells)
        self.assertEqual(cells[(5, 22)][0], 1)
        self.assertEqual(cells[(5, 23)][0], hashlib.sha256(raw).hexdigest())
        before = snapshot(self.root)
        with self.assertRaisesRegex(sync.SyncError, 'already exists'):
            self.plan()
        self.assertEqual(snapshot(self.root), before)
        journals = list((self.root / '.agent/runtime/transactions').glob('*/journal.json'))
        self.assertEqual(json.loads(journals[0].read_text())['schema'], tx.CREATE_SCHEMA)

    def test_destination_collision_same_and_different_bytes(self):
        for raw in (self.source.read_bytes(), b'foreign'):
            with self.subTest(raw=raw):
                self.target.write_bytes(raw)
                before = snapshot(self.root, True)
                with self.assertRaisesRegex(sync.SyncError, 'already exists'):
                    self.plan()
                self.assertEqual(snapshot(self.root, True), before)

    def test_exact_name_and_explicit_code_are_classification_only(self):
        self.source.rename(self.folder / 'Canonical.txt')
        self.assertIn('not verified', self.plan()['warning'])
        (self.folder / 'Canonical.txt').rename(self.folder / 'unknown.txt')
        explicit = intake.request('unknown.txt', 'LAW-01')
        self.assertEqual(self.plan(explicit)['items'][0]['destination'], 'legal/team/Canonical.txt')

    def test_missing_multiple_codes_multiple_files_and_topic_only_refuse(self):
        second = {'code': 'LAW-02', 'canonical_name': 'Other.txt', 'destination': 'legal/team/Other.txt'}
        self.config['mappings'].append(second)
        self.write_config()
        for name, reason in [('old-law-same-topic.txt', 'no unique'), ('LAW-01_LAW-02.txt', 'multiple codes')]:
            new = self.folder / name
            self.source.rename(new)
            with self.assertRaisesRegex(sync.SyncError, reason):
                self.plan()
            new.rename(self.source)
        (self.folder / 'LAW-01_other.txt').write_bytes(b'other')
        with self.assertRaisesRegex(sync.SyncError, 'multiple files'):
            self.plan()

    def test_new_row_never_inherits_prefilled_fulltext_or_user_note(self):
        self.edit_index(lambda wb: (wb['Danh mục'].cell(5, 15, 'Có'), wb['Danh mục'].cell(5, 21, 'user note')))
        result = self.apply()
        cells = self.read_cells()['Danh mục']
        self.assertEqual(result['items'][0]['row'], 6)
        self.assertEqual(cells[(5, 15)][0], 'Có')
        self.assertEqual(cells[(5, 21)][0], 'user note')
        self.assertNotIn((6, 15), cells)

    def test_missing_number_column_and_duplicate_json_keys_refuse(self):
        self.edit_index(lambda wb: setattr(wb['Danh mục'].cell(4, 5), 'value', 'unmanaged header'))
        with self.assertRaisesRegex(sync.SyncError, 'Number/code column'):
            self.plan()
        self.config_path.write_text('{"schema":"law.source-intake/1","schema":"law.source-intake/1"}', encoding='utf-8')
        with self.assertRaisesRegex(sync.SyncError, 'duplicate configuration field'):
            self.plan()

    def test_cli_known_content_error_has_no_traceback_and_partial_has_effects(self):
        def cli(*args, fault=None):
            script = ('import sys; sys.path.insert(0,sys.argv[1]); import library as l; '
                      'l.ROOT=sys.argv[2]; l.REF=l.ROOT+"/references"; l.INDEX=l.REF+"/_INDEX.xlsx"; '
                      'sys.exit(l.run_with_result(sys.argv[3:]))')
            env = dict(os.environ, PYTHONDONTWRITEBYTECODE='1', PYTHONUTF8='1')
            env.pop(tx.FAULT_ENV, None)
            if fault:
                env[tx.FAULT_ENV] = fault
            return subprocess.run([sys.executable, '-c', script, str(ROOT / 'tools'), str(self.root), 'intake', *args],
                                  env=env, capture_output=True, encoding='utf-8', timeout=30)
        planned = self.plan()
        failed = cli('--apply', '--expect-plan', planned['plan_token'], '--json', fault='commit:1')
        env = result_schema.split_output(failed.stdout)[1]
        self.assertEqual(failed.returncode, 1, failed.stderr)
        self.assertEqual(env['status'], 'partial')
        self.assertTrue(env['changes'])
        self.assertTrue(env['transaction'])
        self.assertEqual(env['files'], [{'path': 'references/legal/team/Canonical.txt', 'action': 'create'}])
        self.assertNotIn('Traceback', failed.stderr)
        tx.rollback(str(self.root), env['transaction'])
        self.edit_index(lambda wb: wb.create_sheet(library.CONTENT_SHEET).__setitem__('A1', 'invalid'))
        malformed = cli('--json')
        self.assertEqual(malformed.returncode, 1)
        self.assertNotIn('Traceback', malformed.stderr)

    def test_introduction_suffix_refuses_even_explicit(self):
        name = 'LAW-01_ban-gioi-thieu.txt'
        self.source.rename(self.folder / name)
        with self.assertRaisesRegex(sync.SyncError, 'introduction'):
            self.plan(intake.request(name, 'LAW-01'))

    def test_same_hash_different_codes_never_merge_ids(self):
        self.config['mappings'].append({'code': 'LAW-02', 'canonical_name': 'Other.txt', 'destination': 'legal/team/Other.txt'})
        self.write_config()
        (self.folder / 'LAW-02_source.txt').write_bytes(self.source.read_bytes())
        result = self.plan()
        self.assertEqual([x['id'] for x in result['items']], ['SRC-000001', 'SRC-000002'])
        self.apply(result)

    def test_bad_config_missing_disabled_schema_duplicates_and_scope(self):
        original = json.loads(json.dumps(self.config))
        variants = [dict(original, enabled=False), dict(original, enabled=1), dict(original, extra=True),
                    dict(original, staging='../outside'), dict(original, staging='references/legal/visible'),
                    dict(original, schema='unknown'), dict(original, mappings=original['mappings'] * 2)]
        for cfg in variants:
            with self.subTest(cfg=cfg):
                self.config = cfg
                self.write_config()
                before = snapshot(self.root, True)
                with self.assertRaises(sync.SyncError):
                    self.plan()
                self.assertEqual(snapshot(self.root, True), before)
        self.config_path.unlink()
        with self.assertRaisesRegex(sync.SyncError, 'configuration unavailable'):
            self.plan()

    def test_unsafe_paths_reserved_ads_and_bad_destinations(self):
        for value in ('../x.txt', 'legal/team/NUL.txt', 'legal/team/a.txt:stream', 'legal/team/x. ',
                      'legal/_hidden/a.txt', 'legal/team/../a.txt', 'legal/team/a.exe', 'policies/no-parent/a.txt'):
            with self.subTest(value=value):
                self.config['mappings'][0]['destination'] = value
                self.write_config()
                with self.assertRaises(sync.SyncError):
                    self.plan()

    def test_limits_config_entries_file_total_and_subdirectories(self):
        with patch.object(intake, 'MAX_CONFIG', 3):
            with self.assertRaisesRegex(sync.SyncError, '64 KiB'):
                self.plan()
        with patch.object(intake, 'MAX_ENTRIES', 0):
            with self.assertRaises(sync.SyncError):
                self.plan()
        with patch.object(intake, 'MAX_FILE', 1):
            with self.assertRaisesRegex(sync.SyncError, '20 MiB'):
                self.plan()
        with patch.object(intake, 'MAX_TOTAL', 1):
            with self.assertRaisesRegex(sync.SyncError, '100 MiB'):
                self.plan()
        (self.folder / 'subdir').mkdir()
        with self.assertRaisesRegex(sync.SyncError, 'plain file'):
            self.plan()

    def test_config_permission_read_failure_and_partial_selection(self):
        real_open = open
        def denied(path, *args, **kwargs):
            if str(path) == str(self.config_path):
                raise PermissionError('fixture config denied')
            return real_open(path, *args, **kwargs)
        with patch('builtins.open', denied):
            with self.assertRaisesRegex(sync.SyncError, 'configuration unavailable'):
                self.plan()
        for args in (('a.txt', None), (None, 'LAW-01'), ('../a.txt', 'LAW-01')):
            with self.assertRaises(sync.SyncError):
                intake.request(*args)

    def test_existing_path_always_replacement_even_when_file_missing(self):
        def setup(wb):
            ws = wb['Danh mục']
            ws.cell(5, 1, 'OLD-42'); ws.cell(5, 5, 'LAW-01'); ws.cell(5, 16, 'legal/team/missing.txt')
        self.edit_index(setup)
        with self.assertRaisesRegex(sync.SyncError, 'Replacement'):
            self.plan()

    def test_existing_destination_row_refuses_even_different_code(self):
        self.edit_index(lambda wb: wb['Danh mục'].cell(5, 16, 'legal/team/Canonical.txt'))
        with self.assertRaisesRegex(sync.SyncError, 'Replacement'):
            self.plan()

    def pathless(self, wb):
        ws = wb['Danh mục']
        ws.cell(5, 1, 'PRIVATE-42'); ws.cell(5, 5, 'LAW-01'); ws.cell(5, 9, 'User title')
        ws.cell(5, 21, 'User note'); ws.cell(5, 26, 'Custom heading'); ws.cell(6, 26, '=1+2')

    def test_pathless_keeps_id_title_notes_custom_formula_style(self):
        self.edit_index(self.pathless)
        before = self.read_cells()
        self.apply()
        after = self.read_cells()
        for cell in ((5, 1), (5, 5), (5, 9), (5, 21), (5, 26), (6, 26)):
            self.assertEqual(before['Danh mục'][cell], after['Danh mục'][cell])
        wb = openpyxl.load_workbook(self.index)
        try:
            self.assertTrue(wb['Danh mục'].data_validations.count)
            self.assertTrue(wb['Danh mục'].cell(4, 1).has_style)
        finally:
            wb.close()

    def test_pathless_observation_history_contentbasis_and_fulltext_refuse(self):
        self.edit_index(self.pathless)
        base = self.index.read_bytes()
        cases = [('revision', lambda wb: (wb['Danh mục'].cell(4, 27, 'LAW revision'), wb['Danh mục'].cell(5, 27, 1))),
                 ('sha', lambda wb: (wb['Danh mục'].cell(4, 27, 'LAW SHA-256'), wb['Danh mục'].cell(5, 27, 'a' * 64))),
                 ('fulltext', lambda wb: wb['Danh mục'].cell(5, 15, 'Có')),
                 ('history', lambda wb: (wb.create_sheet('LAW history'), wb['LAW history'].append(sync.HISTORY_LEGACY),
                                         wb['LAW history'].append(('PRIVATE-42', 1, 'a' * 64, 'Cần đối chiếu', '2026-01-01'))))]
        for name, fn in cases:
            with self.subTest(name=name):
                self.index.write_bytes(base)
                self.edit_index(fn)
                with self.assertRaises(sync.SyncError):
                    self.plan()
        self.index.write_bytes(base)
        fake = [{'source_id': 'PRIVATE-42'}]
        with patch.object(library, 'parse_content', return_value=fake):
            self.edit_index(lambda wb: wb.create_sheet(library.CONTENT_SHEET))
            with self.assertRaisesRegex(sync.SyncError, 'prior observation'):
                self.plan()

    def test_index_corrupt_missing_columns_duplicate_ids_refuse(self):
        base = self.index.read_bytes()
        for mode in ('corrupt', 'columns', 'duplicate'):
            with self.subTest(mode=mode):
                self.index.write_bytes(base)
                if mode == 'corrupt':
                    self.index.write_bytes(b'not a workbook')
                elif mode == 'columns':
                    self.edit_index(lambda wb: setattr(wb['Danh mục'].cell(4, 16), 'value', None))
                else:
                    self.edit_index(lambda wb: (wb['Danh mục'].cell(5, 1, 'X'), wb['Danh mục'].cell(6, 1, 'X')))
                before = snapshot(self.root, True)
                with self.assertRaises(sync.SyncError):
                    self.plan()
                self.assertEqual(snapshot(self.root, True), before)

    def test_all_preview_drifts_refuse_before_transaction(self):
        for mode in ('config', 'source', 'inventory', 'index', 'selection'):
            with self.subTest(mode=mode):
                result = self.plan()
                old_config, old_index, old_source = self.config_path.read_bytes(), self.index.read_bytes(), self.source.read_bytes()
                if mode == 'config': self.config_path.write_bytes(old_config + b' ')
                if mode == 'source': self.source.write_bytes(old_source + b'changed')
                if mode == 'inventory': (self.folder / 'extra.txt').write_bytes(b'extra')
                if mode == 'index': self.edit_index(lambda wb: wb['Danh mục'].cell(1, 10, 'user save'))
                selected = intake.request(self.source.name, 'LAW-01') if mode == 'selection' else self.selection
                before = snapshot(self.root)
                with self.assertRaises(sync.SyncError): self.apply(result, selected)
                self.assertEqual(snapshot(self.root), before)
                self.assertFalse(list((self.root / '.agent/runtime/transactions').glob('*/journal.json')))
                self.config_path.write_bytes(old_config); self.index.write_bytes(old_index); self.source.write_bytes(old_source)
                extra = self.folder / 'extra.txt'
                if extra.exists(): extra.unlink()

    def test_stage_and_backup_faults_keep_all_originals_and_allow_reuse(self):
        for point in ('stage:1', 'stage:2', 'backup:1'):
            with self.subTest(point=point):
                before = snapshot(self.root)
                self.assertIsNone(self.fault(point))
                self.assertEqual(snapshot(self.root), before)
                self.assertFalse(self.target.exists())
        self.apply()

    def test_commit_boundary_resume_completes_and_staging_is_retained(self):
        original = self.index.read_bytes()
        ident = self.fault()
        self.assertTrue(self.target.exists())
        self.assertEqual(self.index.read_bytes(), original)
        self.assertEqual(tx.resume(str(self.root), ident)['code'], 0)
        self.assertTrue(self.source.exists())
        self.assertEqual(self.read_cells()['Danh mục'][(5, 22)][0], 1)
        self.assertEqual(tx.resume(str(self.root), ident)['code'], 0)
        self.assertEqual(tx.incomplete(str(self.root), strict=True), [])

    def test_commit_boundary_rollback_restores_exact_and_can_reuse(self):
        before = snapshot(self.root)
        ident = self.fault()
        self.assertEqual(tx.rollback(str(self.root), ident)['code'], 0)
        self.assertEqual(snapshot(self.root), before)
        self.apply()

    def test_concurrent_index_save_is_not_overwritten_by_resume_or_rollback(self):
        ident = self.fault()
        self.edit_index(lambda wb: wb['Danh mục'].cell(1, 10, 'user save'))
        user = self.index.read_bytes()
        self.assertEqual(tx.resume(str(self.root), ident)['code'], 1)
        self.assertEqual(self.index.read_bytes(), user)
        self.assertEqual(tx.rollback(str(self.root), ident)['code'], 1)
        self.assertEqual(self.index.read_bytes(), user)
        self.assertFalse(self.target.exists())

    def test_newfile_user_atomic_save_samebytes_is_preserved_and_blocks_index(self):
        original = self.index.read_bytes()
        ident = self.fault()
        raw = self.target.read_bytes()
        self.target.unlink(); self.target.write_bytes(raw)
        self.assertEqual(tx.resume(str(self.root), ident)['code'], 1)
        self.assertEqual(self.index.read_bytes(), original)
        self.assertEqual(tx.rollback(str(self.root), ident)['code'], 1)
        self.assertEqual(self.target.read_bytes(), raw)
        with self.assertRaisesRegex(sync.SyncError, 'already exists'): self.plan()

    def test_newfile_user_inplace_edit_is_preserved_and_blocks_index(self):
        original = self.index.read_bytes()
        ident = self.fault()
        self.target.write_bytes(b'user inplace')
        self.assertEqual(tx.resume(str(self.root), ident)['code'], 1)
        self.assertEqual(self.index.read_bytes(), original)
        self.assertEqual(tx.rollback(str(self.root), ident)['code'], 1)
        self.assertEqual(self.target.read_bytes(), b'user inplace')

    def test_atomic_no_clobber_race_same_or_different_bytes(self):
        real_link = os.link
        for raw in (self.source.read_bytes(), b'foreign different'):
            with self.subTest(raw=raw):
                def race(src, dst, *args, **kwargs):
                    if str(dst) == str(self.target): self.target.write_bytes(raw)
                    return real_link(src, dst, *args, **kwargs)
                with patch.object(tx.os, 'link', side_effect=race):
                    with self.assertRaises(sync.SyncError): self.apply()
                self.assertEqual(self.target.read_bytes(), raw)
                self.assertNotIn((5, 1), self.read_cells()['Danh mục'])
                self.target.unlink()

    def test_hardlink_unavailable_fails_closed(self):
        before = snapshot(self.root)
        with patch.object(tx.os, 'link', side_effect=OSError('fixture filesystem unsupported')):
            with self.assertRaises(sync.SyncError): self.apply()
        self.assertEqual(snapshot(self.root), before)
        self.assertFalse(self.target.exists())

    def test_disk_permission_and_existing_open_probe_fail_before_new_file(self):
        for target, side in (('probe_writable', OSError('fixture permission denied')), ('preflight', ['fixture full disk'])):
            with self.subTest(target=target):
                before = snapshot(self.root)
                obj = tx if target == 'probe_writable' else tx.Transaction
                kwargs = {'side_effect': side} if isinstance(side, Exception) else {'return_value': side}
                with patch.object(obj, target, **kwargs):
                    with self.assertRaises(sync.SyncError): self.apply()
                self.assertEqual(snapshot(self.root), before)

    @unittest.skipUnless(os.name == 'nt', 'Windows sharing handle fixture, not actual Excel')
    def test_windows_open_index_and_readonly_are_preserved(self):
        from test_library_sync import hold_like_excel
        before = snapshot(self.root)
        release = hold_like_excel(self.index)
        try:
            with self.assertRaises(sync.SyncError): self.apply()
        finally:
            release()
        self.assertEqual(snapshot(self.root), before)
        os.chmod(self.index, stat.S_IREAD)
        try:
            with self.assertRaises(sync.SyncError): self.apply()
            self.assertFalse(self.target.exists())
        finally:
            os.chmod(self.index, stat.S_IREAD | stat.S_IWRITE)

    def test_create_schema_validation_and_old_write_compatibility(self):
        transaction = tx.Transaction(str(self.root), 'old-write')
        transaction.add('private.txt', b'updated')
        self.assertEqual(transaction.journal['schema'], tx.SCHEMA)
        transaction.run()
        self.assertEqual(tx._load(str(self.root), transaction.txid)['schema'], tx.SCHEMA)
        bad = dict(transaction.journal, ops=[dict(transaction.journal['ops'][0], action='create')])
        tx._validate_create_journal(bad)  # informational legacy create stays compatible
        transaction = tx.Transaction(str(self.root), 'create', create_only=True)
        with self.assertRaises(tx.TransactionError): transaction.add('private.txt', b'updated', action='create', expected_before=None)
        bad['schema'] = tx.CREATE_SCHEMA
        with self.assertRaises(tx.TransactionError): tx._validate_create_journal(bad)

    def test_bad_and_unfinished_journals_block_apply(self):
        folder = self.root / '.agent/runtime/transactions/bad'
        folder.mkdir(parents=True)
        (folder / 'journal.json').write_text('{', encoding='utf-8')
        with self.assertRaisesRegex(sync.SyncError, 'Cannot tell'): self.apply()
        self.assertFalse(self.target.exists())

    def test_crash_after_link_before_journal_resumes_owned_file_only(self):
        real_link = os.link
        def link_then_crash(src, dst, *args, **kwargs):
            real_link(src, dst, *args, **kwargs)
            raise KeyboardInterrupt('fixture crash immediately after hard link')
        with patch.object(tx.os, 'link', side_effect=link_then_crash):
            with self.assertRaises(KeyboardInterrupt): self.apply()
        ident = tx.incomplete(str(self.root), strict=True)[0]['id']
        self.assertEqual(tx._load(str(self.root), ident)['ops'][0]['state'], 'pending')
        self.assertEqual(tx.resume(str(self.root), ident)['code'], 0)

    def test_missing_or_corrupt_created_staging_refuses_resume_without_index_write(self):
        original = self.index.read_bytes()
        ident = self.fault()
        staged = Path(tx.txdir(str(self.root), ident)) / 'staging/references/legal/team/Canonical.txt'
        staged.unlink()
        self.assertEqual(tx.resume(str(self.root), ident)['code'], 1)
        self.assertEqual(self.index.read_bytes(), original)
        self.assertEqual(tx.rollback(str(self.root), ident)['code'], 1)
        self.assertTrue(self.target.exists())
        self.assertEqual(tx.resume(str(self.root), ident)['code'], 1)

    def test_changed_created_file_inside_commit_hook_blocks_index(self):
        original = self.index.read_bytes()
        def hook(point, path):
            if point == 'commit' and str(path) == str(self.index):
                self.target.write_bytes(b'user changed after dependency check')
        with patch.object(tx, '_HOOK', hook):
            with self.assertRaises(sync.SyncError): self.apply()
        self.assertEqual(self.index.read_bytes(), original)

    def test_v2_journal_unknown_action_and_shape_refuse(self):
        ident = self.fault()
        path = Path(tx.txdir(str(self.root), ident)) / 'journal.json'
        journal = json.loads(path.read_text())
        for change in ('action', 'status', 'shape'):
            bad = json.loads(json.dumps(journal))
            if change == 'action': bad['ops'][0]['action'] = 'unknown'
            if change == 'status': bad['status'] = 'unknown'
            if change == 'shape': del bad['ops'][0]['before']
            path.write_text(json.dumps(bad), encoding='utf-8')
            with self.assertRaises(tx.TransactionError): tx.resume(str(self.root), ident)
        path.write_text(json.dumps(journal), encoding='utf-8')

    @unittest.skipUnless(os.name == 'nt', 'real Windows junction fixture')
    def test_staging_and_destination_junctions_refuse_without_reading_target(self):
        private = self.root / 'private-staging'
        private.mkdir()
        self.source.rename(private / self.source.name)
        self.folder.rmdir()
        result = subprocess.run(['cmd', '/c', 'mklink', '/J', str(self.folder), str(private)], capture_output=True, timeout=10)
        self.assertEqual(result.returncode, 0, result.stderr)
        real_open = open
        def spy(path, *args, **kwargs):
            if str(path).startswith(str(private)) or str(path).startswith(str(self.folder)):
                self.fail('attempted to open a junction target')
            return real_open(path, *args, **kwargs)
        try:
            with patch('builtins.open', spy):
                with self.assertRaisesRegex(sync.SyncError, 'link'): self.plan()
        finally:
            os.rmdir(self.folder)
        self.folder.mkdir()
        (private / self.source.name).rename(self.source)
        dest_parent = self.target.parent
        dest_parent.rmdir()
        result = subprocess.run(['cmd', '/c', 'mklink', '/J', str(dest_parent), str(private)], capture_output=True, timeout=10)
        self.assertEqual(result.returncode, 0, result.stderr)
        try:
            with self.assertRaisesRegex(sync.SyncError, 'link'): self.plan()
        finally:
            os.rmdir(dest_parent)

    def test_real_config_file_symlink_refuses(self):
        actual = self.root / 'private-config.json'
        self.config_path.rename(actual)
        try:
            os.symlink(actual, self.config_path)
        except OSError as exc:
            self.skipTest('real config file symlink unavailable: %s' % exc)
        with self.assertRaisesRegex(sync.SyncError, 'link'): self.plan()

    def test_x14_loss_refuses_before_create(self):
        from test_library_sync import LibraryRuntimeAndBoundary
        LibraryRuntimeAndBoundary.x14_workbook(self)
        result = self.plan()
        self.assertIn('x14_data_validation', result['notices'])
        before = snapshot(self.root)
        with self.assertRaisesRegex(sync.SyncError, 'x14'): self.apply(result)
        self.assertEqual(snapshot(self.root), before)
        self.assertFalse(self.target.exists())

    def test_existing_history_and_content_binding_cells_are_preserved(self):
        old = self.ref / 'legal/team/existing.txt'
        old.write_bytes(b'existing')
        sync.apply(str(self.root))
        request = {'entry_id': 'E-OLD', 'source_id': 'SRC-000001', 'source_revision': 1,
                   'source_sha256': hashlib.sha256(old.read_bytes()).hexdigest(), 'locator': 'paragraph 1',
                   'read_scope': 'paragraph 1', 'kind': 'EXTRACT', 'text': 'fixture existing', 'recorded_by': 'tester'}
        preview = sync.content_plan(str(self.root), request)
        sync.content_apply(str(self.root), request, preview['plan_token'])
        def history(wb):
            ws = wb.create_sheet('LAW history'); ws.append(sync.HISTORY_LEGACY)
            ws.append(('OTHER', 1, 'b' * 64, 'Cần đối chiếu', '2026-01-01'))
        self.edit_index(history)
        before = self.read_cells()
        self.apply()
        after = self.read_cells()
        self.assertEqual(before['LAW history'], after['LAW history'])
        self.assertEqual(before[library.CONTENT_SHEET], after[library.CONTENT_SHEET])
        self.assertEqual(library.content_find(str(self.root))['entries'][0]['basis_state'], 'CURRENT_CATALOG_BASIS')

    def test_real_process_kill_resume_and_rollback_boundary(self):
        script = ('import sys; sys.path.insert(0,sys.argv[1]); import library_sync as s,library_intake as i; '
                  'r=s.intake_plan(sys.argv[2],i.request()); s.intake_apply(sys.argv[2],i.request(),r["plan_token"])')
        before = snapshot(self.root)
        env = dict(os.environ, PYTHONDONTWRITEBYTECODE='1', PYTHONUTF8='1', **{tx.FAULT_ENV: 'kill:1'})
        result = subprocess.run([sys.executable, '-c', script, str(ROOT / 'tools'), str(self.root)], env=env,
                                capture_output=True, timeout=30)
        self.assertEqual(result.returncode, 70, result.stderr)
        ident = tx.incomplete(str(self.root), strict=True)[0]['id']
        self.assertEqual(tx.rollback(str(self.root), ident)['code'], 0)
        self.assertEqual(snapshot(self.root), before)
        result = subprocess.run([sys.executable, '-c', script, str(ROOT / 'tools'), str(self.root)], env=env,
                                capture_output=True, timeout=30)
        self.assertEqual(result.returncode, 70, result.stderr)
        ident = tx.incomplete(str(self.root), strict=True)[0]['id']
        self.assertEqual(tx.resume(str(self.root), ident)['code'], 0)

    def test_standalone_cli_uses_callers_result_envelope(self):
        tools = self.root / 'tools'
        tools.mkdir()
        for name in ('library.py', 'library_intake.py', 'library_sync.py', 'transaction.py', 'result_schema.py', 'artifacts.py'):
            shutil.copyfile(ROOT / 'tools' / name, tools / name)
        def cli(*args, fault=None):
            env = dict(os.environ, PYTHONDONTWRITEBYTECODE='1', PYTHONUTF8='1')
            env.pop(tx.FAULT_ENV, None)
            if fault:
                env[tx.FAULT_ENV] = fault
            return subprocess.run([sys.executable, str(tools / 'library.py'), 'intake', *args, '--result-json'],
                                  cwd=self.root, env=env, capture_output=True, encoding='utf-8', timeout=30)
        preview = cli('--json')
        text, envelope = result_schema.split_output(preview.stdout)
        self.assertEqual(preview.returncode, 0, preview.stderr)
        self.assertEqual(envelope['status'], 'preview')
        token = json.loads(text)['plan_token']
        failed = cli('--apply', '--expect-plan', token, '--json', fault='commit:1')
        partial = result_schema.split_output(failed.stdout)[1]
        self.assertEqual((failed.returncode, partial['status'], partial['changes']), (1, 'partial', True))
        self.assertTrue(partial['transaction'])
        self.assertNotIn('Traceback', failed.stderr)
        tx.rollback(str(self.root), partial['transaction'])
        preview = cli('--json')
        token = json.loads(result_schema.split_output(preview.stdout)[0])['plan_token']
        applied = cli('--apply', '--expect-plan', token, '--json')
        envelope = result_schema.split_output(applied.stdout)[1]
        self.assertEqual((applied.returncode, envelope['status']), (0, 'success'), applied.stderr)
        self.assertIn({'path': 'references/_INDEX.xlsx', 'action': 'write'}, envelope['files'])

    def test_cli_default_json_apply_and_missing_token(self):
        script = ('import sys; sys.path.insert(0,sys.argv[1]); import library as l; '
                  'l.ROOT=sys.argv[2]; l.REF=l.ROOT+"/references"; l.INDEX=l.REF+"/_INDEX.xlsx"; '
                  'sys.exit(l.run_with_result(sys.argv[3:]))')
        def cli(*args):
            return subprocess.run([sys.executable, '-c', script, str(ROOT / 'tools'), str(self.root), 'intake', *args],
                                  env=dict(os.environ, PYTHONDONTWRITEBYTECODE='1', PYTHONUTF8='1'),
                                  capture_output=True, encoding='utf-8', timeout=30)
        before = snapshot(self.root, True)
        result = cli('--json')
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        rendered, envelope = result_schema.split_output(result.stdout)
        token = json.loads(rendered)['plan_token']
        self.assertEqual(envelope['status'], 'preview')
        self.assertEqual(snapshot(self.root, True), before)
        self.assertEqual(cli('--apply').returncode, 1)
        self.assertEqual(cli('--apply', '--expect-plan', token, '--json').returncode, 0)


if __name__ == '__main__':
    unittest.main()
