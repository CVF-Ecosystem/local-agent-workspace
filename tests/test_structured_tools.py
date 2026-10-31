"""Structured results (--result-json) of pack, install_skills, library and build_guides_html.

Each test works on a disposable copy of the package; the real repository is never written.
Steps inside one test share the copy because they depend on each other (install, update, remove).
"""
import datetime
import hashlib
import importlib.util
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'tools'))
import check_package as cp
import result_schema

HAS_OPENPYXL = importlib.util.find_spec('openpyxl') is not None
MARKER = result_schema.MARKER
# chạy công cụ sau một đoạn mã chèn vào / run a tool after an inserted snippet (simulated failures)
RUNNER = ('import runpy, sys\n{patch}\ntool = sys.argv[1]\nsys.argv = sys.argv[1:]\n'
          "sys.path.insert(0, __import__('os').path.dirname(__import__('os').path.abspath(tool)))\n"
          "runpy.run_path(tool, run_name='__main__')\n")
NO_OPENPYXL = "import sys\nsys.modules['openpyxl'] = None  # mô phỏng thiếu openpyxl / simulate a missing openpyxl\n"
# lần ghi file thứ hai lỗi, mô phỏng đĩa đầy giữa chừng / the second file write fails: a disk-full stand-in
FLAKY_COPY = ('import shutil\n_real_copy = shutil.copyfile\n_calls = []\n'
              'def _flaky_copy(src, dst, *a, **k):\n    _calls.append(dst)\n'
              '    if len(_calls) == 2:\n        raise OSError(28, "simulated write failure")\n'
              '    return _real_copy(src, dst, *a, **k)\n'
              'shutil.copyfile = _flaky_copy\n')
LIBRARY_PREFIX = {'id': 'Mã', 'group': 'Nhóm', 'domain': 'Lĩnh vực', 'issuer': 'Cơ quan', 'number': 'Số, ký hiệu',
                  'type': 'Loại văn bản', 'title': 'Trích yếu', 'status': 'Tình trạng', 'checked': 'Ngày kiểm',
                  'verify': 'Mức độ xác minh', 'fulltext': 'Có toàn văn', 'file': 'Đường dẫn file'}


def write(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(text.encode('utf-8'))


def read(path):
    return path.read_bytes().decode('utf-8')


def fill_library(base):
    """Năm dòng mẫu: lỗi file thiếu, trùng số, dòng trống, hai cảnh báo / five synthetic rows: missing file, duplicate, blank, warnings."""
    import openpyxl
    index = base / 'references' / '_INDEX.xlsx'
    wb = openpyxl.load_workbook(str(index))
    ws = wb['Danh mục']
    headers = {c: str(ws.cell(row=4, column=c).value or '') for c in range(1, ws.max_column + 1)}

    def column(prefix):
        return next(c for c, h in headers.items() if h.startswith(prefix))
    rows = [
        {'id': 'L-001', 'group': 'Pháp luật', 'domain': 'Đất đai', 'issuer': 'Quốc hội', 'number': '31/2024/QH15',
         'type': 'Luật', 'title': 'Luật Đất đai', 'status': 'Còn hiệu lực', 'checked': datetime.date(2026, 9, 1),
         'verify': 'Đã đối chiếu toàn văn', 'fulltext': 'Có', 'file': 'legal/dat-dai/luat-dat-dai.pdf'},
        {'id': 'L-002', 'group': 'Pháp luật', 'number': '31/2024/QH15', 'title': 'Luật trùng số',
         'status': 'Còn hiệu lực', 'checked': datetime.date(2026, 9, 1), 'fulltext': 'Không'},
        {'id': 'L-003', 'group': 'Nội bộ', 'status': 'Còn hiệu lực', 'fulltext': 'Không'},
        {'id': 'L-004', 'group': 'Pháp luật', 'number': '12/2023/NĐ-CP', 'title': 'Nghị định kiểm thử', 'fulltext': 'Không'},
        {'id': 'L-005', 'group': 'Nội bộ', 'number': 'QC-01', 'title': 'Quy chế kiểm thử', 'status': 'Còn hiệu lực',
         'fulltext': 'Có', 'file': 'policies/quy-che/quy-che.docx', 'verify': 'Đã đối chiếu toàn văn'},
    ]
    for i, rec in enumerate(rows):
        for key, value in rec.items():
            ws.cell(row=5 + i, column=column(LIBRARY_PREFIX[key]), value=value)
    wb.save(str(index))
    write(base / 'references' / 'legal' / 'dat-dai' / 'chua-lap-chi-muc.pdf', 'synthetic')
    write(base / 'references' / 'policies' / 'quy-che' / 'quy-che.docx', 'synthetic')


class StructuredBase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='structured-tools-')
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name) / 'package'
        shutil.copytree(ROOT, self.base, ignore=shutil.ignore_patterns('.git', '__pycache__', 'dist', 'roadmap', 'tests', 'feedback'))
        shutil.rmtree(str(self.base / '.agent' / 'runtime'), ignore_errors=True)  # dữ liệu vận hành của máy phát triển không vào bản sao / dev-machine runtime data never enters the copy
        [f.unlink() for f in (self.base / 'assets').rglob('*') if f.is_file() and f.name != 'README.md']  # logo, font của máy phát triển / dev-machine branding and fonts

    def run_tool(self, script, *args, patch=''):
        env = dict(os.environ, PYTHONUTF8='1', PYTHONDONTWRITEBYTECODE='1')
        tool = str(self.base / 'tools' / script)
        if patch:
            cmd = [sys.executable, '-c', RUNNER.replace('{patch}', patch), tool, *map(str, args)]
        else:
            cmd = [sys.executable, tool, *map(str, args)]
        return subprocess.run(cmd, cwd=str(self.base), env=env, capture_output=True, encoding='utf-8', timeout=120)

    def structured(self, script, *args, patch=''):
        result = self.run_tool(script, *args, '--result-json', patch=patch)
        human, data = result_schema.split_output(result.stdout)
        self.assertIsNotNone(data, result.stdout + result.stderr)
        self.assertEqual(data['schema'], result_schema.SCHEMA)
        self.assertEqual(data['code'], result.returncode, result.stdout + result.stderr)
        self.assertNotIn(MARKER, human)
        return result, human, data

    def plain(self, script, *args):
        result = self.run_tool(script, *args)
        self.assertNotIn(MARKER, result.stdout)
        return result

    def init_project(self):
        result = self.run_tool('init_project.py', '--name', 'Test')
        self.assertEqual(result.returncode, 0, result.stderr)


class PackResults(StructuredBase):
    def test_list_check_and_findings(self):
        result, _, data = self.structured('pack.py', 'list')
        self.assertEqual((result.returncode, data['status']), (0, 'success'))
        packs = {p['name']: p for p in data['packs']}
        self.assertIn('cang-logistics', packs)
        self.assertIsNone(packs['cang-logistics']['installed'])
        self.plain('pack.py', 'list')

        result, _, data = self.structured('pack.py', 'check')
        self.assertEqual((result.returncode, data['status'], data['findings']), (0, 'success', []))
        self.assertIn('cang-logistics', data['checked'])

        write(self.base / 'packs' / 'broken' / 'README.md', 'no pack.json here')
        result, _, data = self.structured('pack.py', 'check')
        self.assertEqual((result.returncode, data['status']), (1, 'issues'))
        self.assertTrue(any(f['pack'] == 'broken' and 'pack.json' in f['message'] for f in data['findings']))

        result, _, data = self.structured('pack.py', 'check', 'nope')
        self.assertEqual((result.returncode, data['status'], data['pack']), (1, 'blocked', 'nope'))

    def test_add_update_remove_new_and_partial_add(self):
        self.init_project()
        result, _, data = self.structured('pack.py', 'add', 'nope')
        self.assertEqual((result.returncode, data['status']), (1, 'blocked'))

        result, _, data = self.structured('pack.py', 'add', 'cang-logistics')
        self.assertEqual((result.returncode, data['status'], data['changes']), (0, 'success', True))
        actions = {f['path']: f['action'] for f in data['files']}
        self.assertEqual(actions['skills/local/packs/cang-logistics/pack.json'], 'add')
        self.assertEqual(actions['.agent/SKILL_INDEX.md'], 'index')

        result, _, data = self.structured('pack.py', 'add', 'cang-logistics')
        self.assertEqual((result.returncode, data['status'], data['pack']), (1, 'blocked', 'cang-logistics'))

        result, _, data = self.structured('pack.py', 'update', 'cang-logistics')
        self.assertEqual((result.returncode, data['status'], data['changes']), (0, 'success', False))

        source = self.base / 'packs' / 'cang-logistics' / 'README.md'
        installed = self.base / 'skills' / 'local' / 'packs' / 'cang-logistics' / 'README.md'
        write(source, read(source) + '\nBản đổi kiểm thử.\n')
        before = installed.read_bytes()
        result, _, data = self.structured('pack.py', 'update', 'cang-logistics')
        self.assertEqual((result.returncode, data['status'], data['merge_required']), (0, 'preview', False))
        self.assertIn({'path': 'skills/local/packs/cang-logistics/README.md', 'action': 'update'}, data['files'])
        self.assertEqual(installed.read_bytes(), before, 'preview must not write')
        result, _, data = self.structured('pack.py', 'update', 'cang-logistics', '--apply')
        self.assertEqual((result.returncode, data['status'], data['merge_required']), (0, 'success', False))

        write(installed, read(installed) + '\nGhi chú của người dùng.\n')
        write(source, read(source) + '\nBản đổi thứ hai.\n')
        result, _, data = self.structured('pack.py', 'update', 'cang-logistics')
        self.assertEqual((result.returncode, data['status']), (0, 'preview'))
        self.assertIn('merge', {f['action'] for f in data['files']})
        self.assertFalse((installed.parent / 'README.md.new').exists())
        result, _, data = self.structured('pack.py', 'update', 'cang-logistics', '--apply')
        self.assertEqual((result.returncode, data['status'], data['merge_required']), (0, 'success', True))
        self.assertTrue((installed.parent / 'README.md.new').is_file())

        # bản .new đang được hợp nhất dở không bị ghi đè khi gói đổi tiếp / a .new being merged is not overwritten by a later change
        pending = installed.parent / 'README.md.new'
        write(pending, read(pending) + '\nĐang hợp nhất dở.\n')
        merging = pending.read_bytes()
        write(source, read(source) + '\nBản đổi thứ ba.\n')
        tree = {p: p.read_bytes() for p in installed.parent.rglob('*') if p.is_file()}
        for extra in ((), ('--apply',)):
            result, _, data = self.structured('pack.py', 'update', 'cang-logistics', *extra)
            self.assertEqual((result.returncode, data['status'], data['changes'], data['merge_required']), (1, 'blocked', False, True))
            self.assertIn({'path': 'skills/local/packs/cang-logistics/README.md.new', 'action': 'blocked'}, data['files'])
            self.assertEqual(pending.read_bytes(), merging)
            self.assertEqual({p: p.read_bytes() for p in installed.parent.rglob('*') if p.is_file()}, tree, 'nothing written')
        pending.unlink()

        result, _, data = self.structured('pack.py', 'remove', 'cang-logistics')
        self.assertEqual((result.returncode, data['status']), (0, 'success'))
        self.assertFalse((self.base / 'skills' / 'local' / 'packs' / 'cang-logistics').exists())
        result, _, data = self.structured('pack.py', 'remove', 'cang-logistics')
        self.assertEqual((result.returncode, data['status']), (1, 'blocked'))

        result, _, data = self.structured('pack.py', 'new', 'demo-pack')
        self.assertEqual((result.returncode, data['status'], data['changes']), (0, 'success', True))
        self.assertIn('packs/demo-pack/pack.json', {f['path'] for f in data['files']})
        self.assertTrue(all(f['action'] == 'create' for f in data['files']))
        result, _, data = self.structured('pack.py', 'new', 'demo-pack')
        self.assertEqual((result.returncode, data['status']), (1, 'blocked'))
        result, _, data = self.structured('pack.py', 'new', 'Bad_Name')
        self.assertEqual((result.returncode, data['status']), (1, 'blocked'))

        result, _, data = self.structured('pack.py', 'add', 'hanh-chinh-nhan-su', patch=FLAKY_COPY)
        self.assertEqual((result.returncode, data['status']), (1, 'partial'))
        self.assertEqual(len(data['files']), 1)
        self.assertTrue(data['next_actions'])


class ProjectJourney(StructuredBase):
    """Hành trình của project: tạo bản sao -> thêm nguồn -> cài thử gói -> kiểm package. Luật phát hành không áp vào file của project.
    Project journey: copy -> add sources -> trial-install a pack -> check. Release rules do not apply to project files."""

    def use_clean_package(self):
        """Project đi ra từ gói sạch (projection chung với bản phát hành), không phải bản chép nguyên cây kho nguồn (có v1/v2, fixtures, evidence).
        A project comes out of the clean shipped package (shared projection), not a raw copy of the source tree."""
        self.base = Path(self.tmp.name) / 'clean-package'
        cp.materialize(str(self.base), branding=False)

    def check(self):
        result, human, data = self.structured('check_package.py')
        return result.returncode, data['status'], human

    def test_project_files_do_not_trigger_release_rules_but_package_files_still_do(self):
        self.use_clean_package()
        self.init_project()
        self.assertEqual(self.check()[:2], (0, 'success'))
        result, _, data = self.structured('pack.py', 'add', 'hanh-chinh-nhan-su')
        self.assertEqual(result.returncode, 0, data)
        self.assertIn('v0.', read(self.base / '.agent' / 'SKILL_INDEX.md'))  # dòng gói có phiên bản riêng / the pack row has its own version
        write(self.base / 'references' / 'policies' / 'quy-che.md',
              '# Quy chế thanh toán v2.1\n\nMẫu đề nghị ở `output/mau_de_nghi.md`.\n')
        write(self.base / 'working' / 'nhap.md', '# Bản nháp\n')
        write(self.base / 'output' / 'bao-cao' / 'ket-qua.md', '# Kết quả v3.0\n')
        code, status, human = self.check()
        self.assertEqual((code, status), (0, 'success'), human)
        # đối chứng: file của package vẫn bị kiểm như cũ / control: package files are still checked
        write(self.base / 'docs' / 'README.md', read(self.base / 'docs' / 'README.md') + '\nGhi chú v9.9 kiểm thử\n')
        code, status, human = self.check()
        self.assertEqual((code, status), (1, 'issues'))
        self.assertIn('docs/README.md', human)

    def test_a_readme_the_user_creates_in_a_new_folder_belongs_to_the_project(self):
        """README do người dùng tạo trong thư mục riêng là của project; README khung của package vẫn bị kiểm.
        A README the user makes in their own folder belongs to the project; the package's scaffold README is still checked."""
        self.use_clean_package()
        self.init_project()
        self.assertFalse((self.base / 'references' / 'customer-a').exists())
        for rel in (('references', 'customer-a', 'README.md'), ('working', 'khach-a', 'README.md'), ('output', 'README-khac', 'README.md')):
            write(self.base.joinpath(*rel), '# Project notes v2.1\n\nXem `output/future-report.md`.\n')
        code, status, human = self.check()
        self.assertEqual((code, status), (0, 'success'), human)
        # đối chứng: README khung có trong package vẫn bị kiểm / control: the package's scaffold README is still checked
        scaffold = self.base / 'references' / 'policies' / 'README.md'
        self.assertTrue(scaffold.is_file())
        write(scaffold, read(scaffold) + '\nGhi chú v9.9 kiểm thử\n')
        code, status, human = self.check()
        self.assertEqual((code, status), (1, 'issues'))
        self.assertIn('references/policies/README.md', human)
        self.assertNotIn('customer-a', human)

    def test_crlf_project_template_still_takes_the_answers(self):
        f = self.base / 'PROJECT.md'
        f.write_bytes(f.read_bytes().replace(b'\r\n', b'\n').replace(b'\n', b'\r\n'))
        result = self.run_tool('init_project.py', '--name', 'CRLF', '--purpose', 'Soạn quy trình')
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertNotIn('đã có nội dung', result.stdout)
        data = f.read_bytes()
        self.assertIn('- Mục đích và người đọc: Soạn quy trình\r\n'.encode('utf-8'), data)
        self.assertNotIn(b'\n', data.replace(b'\r\n', b''), 'every line ending stays CRLF')

    def test_init_on_a_source_copy_does_not_carry_the_maintainer_state(self):
        state = self.base / '.agent' / 'STATE.md'
        # Own the precondition; clean source snapshots carry no private maintainer state.
        write(state, '# PROJECT STATE\n<!-- local-agent-workspace: maintainer-state -->\n'
              '\n## Current phase\nDATAMOCK maintainer work.\n')
        maintainer = read(state)
        self.assertIn('<!-- local-agent-workspace: maintainer-state -->', maintainer)
        for lang, tpl in (('vi', 'release_templates/.agent/STATE.md'), ('en', 'release_templates/en/.agent/STATE.md')):
            if lang == 'en':
                state.write_bytes(maintainer.encode('utf-8'))
                (self.base / '.agent' / 'PACKAGE_INFO.json').unlink()
                for b in (self.base / '.agent').glob('STATE.md.bak-*'):
                    b.unlink()
            result = self.run_tool('init_project.py', '--name', 'Copy', '--lang', lang)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            now = read(state)
            self.assertNotIn('maintainer-state', now)
            clean = read(self.base / tpl)
            self.assertEqual(now.split('\n')[0], clean.split('\n')[0])
            self.assertNotIn('[Phase/', now)  # trạng thái khởi đầu đã điền / the starting phase is filled in
            baks = list((self.base / '.agent').glob('STATE.md.bak-*'))
            self.assertEqual(len(baks), 1)
            self.assertEqual(read(baks[0]), maintainer)

    def test_init_keeps_a_project_state_that_has_real_work(self):
        state = self.base / '.agent' / 'STATE.md'
        write(state, '# PROJECT STATE\n\n## Current phase\nĐang soạn quy trình mua sắm.\n')
        result = self.run_tool('init_project.py', '--name', 'Real')
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn('Đang soạn quy trình mua sắm.', read(state))
        self.assertEqual(list((self.base / '.agent').glob('STATE.md.bak-*')), [])


def tree(root, skip=('__pycache__',)):
    """Bytes của mọi file, trừ cache Python và dữ liệu vận hành / every file's bytes, minus Python caches and runtime data."""
    return {q.relative_to(root).as_posix(): q.read_bytes() for q in root.rglob('*')
            if q.is_file() and not set(q.relative_to(root).parts) & set(skip) and '.agent/runtime' not in q.relative_to(root).as_posix()}


class NewProject(StructuredBase):
    """Tạo project mới là bản sao của bộ mẫu: bộ mẫu không đổi, không ghi đè, không mang trạng thái hay file người dùng.
    A new project is a copy of the template: the template is unchanged, nothing is overwritten, no state or user files carried."""

    def release_template(self):
        """Bộ mẫu như sau khi giải nén ZIP phát hành: mẫu sạch, không có release_templates/ / as unzipped from a release ZIP."""
        clean = Path(self.tmp.name) / 'clean-package'
        cp.materialize(str(clean), branding=False)  # chính bytes và MANIFEST của bản phát hành / the very bytes and MANIFEST of a release
        self.base = clean

    def test_preview_then_create_next_to_the_template_and_never_overwrite(self):
        self.release_template()
        write(self.base / 'references' / 'policies' / 'rieng-cua-bo-mau.md', '# file người dùng đặt nhầm vào bộ mẫu\n')
        write(self.base / 'assets' / 'branding.json', '{"title": "Đơn vị Mẫu"}\n')  # nhận diện của bộ mẫu đi theo dự án / template branding carries over
        write(self.base / 'assets' / 'logo.svg', '<svg xmlns="http://www.w3.org/2000/svg"/>\n')
        write(self.base / 'assets' / 'ghi-chu.txt', 'không phải nhận diện\n')
        write(self.base / 'assets' / 'fonts' / 'DonVi-Regular.ttf', 'font giả\n')
        write(self.base / 'assets' / 'fonts' / 'ghi-chu.txt', 'không phải font\n')
        dest = Path(self.tmp.name)
        before = tree(self.base)
        result, _, data = self.structured('new_project.py', '--name', 'Biên bản họp: tổ A/B')
        self.assertEqual((result.returncode, data['status']), (0, 'preview'))
        target = dest / 'Biên bản họp tổ A B'
        self.assertEqual(Path(data['target']).resolve(), target.resolve())
        self.assertFalse(target.exists(), 'preview must write nothing')
        result, _, data = self.structured('new_project.py', '--name', 'Biên bản họp: tổ A/B', '--apply')
        self.assertEqual((result.returncode, data['status'], data['changes']), (0, 'success', True), data)
        self.assertEqual(tree(self.base), before, 'the template must not change')
        info = json.loads(read(target / '.agent' / 'PACKAGE_INFO.json'))
        self.assertEqual((info['project_name'], info['language']), ('Biên bản họp: tổ A/B', 'vi'))
        self.assertIn('- Tên dự án: Biên bản họp: tổ A/B', read(target / 'PROJECT.md'))
        self.assertFalse((target / 'references' / 'policies' / 'rieng-cua-bo-mau.md').exists(), 'user files of the template are not copied')
        self.assertTrue((target / 'references' / 'policies' / 'README.md').is_file())
        self.assertEqual(read(target / 'assets' / 'branding.json'), '{"title": "Đơn vị Mẫu"}\n')
        self.assertTrue((target / 'assets' / 'logo.svg').is_file())
        self.assertFalse((target / 'assets' / 'ghi-chu.txt').exists(), 'only branding files are carried from assets/')
        self.assertTrue((target / 'assets' / 'fonts' / 'DonVi-Regular.ttf').is_file())
        self.assertFalse((target / 'assets' / 'fonts' / 'ghi-chu.txt').exists())
        self.assertTrue((target / 'assets' / 'fonts' / 'README.md').is_file())
        self.assertEqual(list(dest.glob('.*.creating-*')), [], 'no temp folder is left behind')
        chk = subprocess.run([sys.executable, str(target / 'tools' / 'check_package.py'), '--verify'], cwd=str(target), capture_output=True)
        self.assertEqual(chk.returncode, 0, chk.stdout.decode('utf-8', 'replace'))
        # chạy lại cùng tên: chặn, không ghi đè project đã có / same name again: blocked, the existing project is kept
        write(target / 'output' / 'ket-qua.md', 'công việc thật\n')
        result, _, data = self.structured('new_project.py', '--name', 'Biên bản họp: tổ A/B', '--apply')
        self.assertEqual((result.returncode, data['status'], data['changes']), (1, 'blocked', False))
        self.assertEqual(read(target / 'output' / 'ket-qua.md'), 'công việc thật\n')
        # từ project (không phải bộ mẫu): từ chối / from a project (not the template): refused
        r = subprocess.run([sys.executable, str(target / 'tools' / 'new_project.py'), '--name', 'X', '--apply', '--result-json'],
                           cwd=str(target), capture_output=True, encoding='utf-8')
        self.assertEqual(r.returncode, 1)
        self.assertEqual(result_schema.split_output(r.stdout)[1]['role'], 'project')

    def test_source_repository_copy_uses_clean_templates(self):
        dest = Path(self.tmp.name) / 'noi-tao'
        dest.mkdir()
        result, _, data = self.structured('new_project.py', '--name', 'Từ kho nguồn', '--dest', dest, '--lang', 'en', '--apply')
        self.assertEqual((result.returncode, data['status']), (0, 'success'), data)
        state = read(dest / 'Từ kho nguồn' / '.agent' / 'STATE.md')
        self.assertNotIn('maintainer-state', state)
        self.assertIn('Just initialised', state)
        self.assertFalse((dest / 'Từ kho nguồn' / 'release_templates').exists())
        self.assertFalse((dest / 'Từ kho nguồn' / 'tests').exists())

    def test_unknown_role_and_bad_places_are_refused_without_writing(self):
        self.release_template()
        before = set(os.listdir(self.tmp.name))
        result, _, data = self.structured('new_project.py', '--name', 'Trong mẫu', '--dest', self.base / 'working', '--apply')
        self.assertEqual((result.returncode, data['status']), (1, 'blocked'))
        self.assertFalse((self.base / 'working' / 'Trong mẫu').exists())
        result, _, data = self.structured('new_project.py', '--name', 'X', '--dest', Path(self.tmp.name) / 'khong-co', '--apply')
        self.assertEqual((result.returncode, data['status']), (1, 'blocked'))
        write(self.base / 'PROJECT.md', read(self.base / 'PROJECT.md').replace('[Tên]', 'Ai đó đã sửa tay'))
        result, _, data = self.structured('new_project.py', '--name', 'Y', '--apply')
        self.assertEqual((result.returncode, data['status'], data['role']), (1, 'blocked', 'unknown'))
        self.assertEqual(set(os.listdir(self.tmp.name)), before)


class InstallSkillsResults(StructuredBase):
    def test_preview_copy_skip_overwrite_unknown_and_partial(self):
        dest = Path(self.tmp.name) / 'dest'
        self.assertNotIn(MARKER, self.run_tool('install_skills.py', '--dest', dest).stdout)
        result, _, data = self.structured('install_skills.py', '--dest', dest)  # chưa chọn: chỉ liệt kê / nothing chosen: list only
        self.assertEqual((result.returncode, data['status'], data['changes'], data['files']), (0, 'preview', False, []))
        self.assertIn('meeting-minutes', data['eligible'])
        result, _, data = self.structured('install_skills.py', '--dest', dest, '--all')
        self.assertEqual((result.returncode, data['status'], data['changes']), (0, 'preview', True))
        self.assertEqual(data['dest'], os.path.abspath(str(dest)))
        self.assertTrue(all(f['action'] == 'copy' for f in data['files']))
        self.assertFalse(dest.exists(), 'preview must not copy')

        result, _, data = self.structured('install_skills.py', '--dest', dest, '--all', '--yes')
        self.assertEqual((result.returncode, data['status'], data['changes']), (0, 'success', True))
        self.assertTrue((dest / 'meeting-minutes' / 'SKILL.md').is_file())
        result, _, data = self.structured('install_skills.py', '--dest', dest, '--all', '--yes')
        self.assertEqual((result.returncode, data['status'], data['changes']), (0, 'success', False))
        self.assertTrue(all(f['action'] == 'skip' and f['reason'] == 'exists' for f in data['files']))
        result, _, data = self.structured('install_skills.py', '--dest', dest, '--all', '--overwrite')
        self.assertEqual((result.returncode, data['status']), (0, 'preview'))
        self.assertTrue({'overwrite', 'backup'} <= {f['action'] for f in data['files']})

        result, _, data = self.structured('install_skills.py', '--dest', dest, '--skills', 'nope')
        self.assertEqual((result.returncode, data['status']), (1, 'blocked'))

        result, _, data = self.structured('install_skills.py', '--dest', Path(self.tmp.name) / 'dest2', '--all', '--yes', patch=FLAKY_COPY)
        self.assertEqual((result.returncode, data['status']), (1, 'partial'))
        self.assertEqual(len(data['files']), 1)


@unittest.skipUnless(HAS_OPENPYXL, 'openpyxl is not installed; library checks skipped')
class LibraryResults(StructuredBase):
    def test_lifecycle_findings_json_and_stats(self):
        result, _, data = self.structured('library.py', 'check', '--json')
        self.assertEqual((result.returncode, data['status']), (1, 'blocked'))  # no index before init

        result, _, data = self.structured('library.py', 'init')
        self.assertEqual((result.returncode, data['status'], data['changes']), (0, 'success', True))
        self.assertIn('references/_INDEX.xlsx', {f['path'] for f in data['files']})
        result = self.plain('library.py', 'init')  # chạy lại không ghi gì; không có dấu trong chữ / rerun writes nothing; no marker
        self.assertEqual(result.returncode, 0)
        self.assertIn('đã có đủ', result.stdout)

        fill_library(self.base)
        result = self.run_tool('library.py', 'check', '--json', '--result-json')
        human, data = result_schema.split_output(result.stdout)
        self.assertEqual((result.returncode, data['status'], data['errors']), (1, 'issues', 3))
        json.loads(human)  # phần chữ là JSON hợp lệ / the human part is valid JSON on its own
        self.assertEqual(sorted(f['row'] for f in data['findings'] if f['severity'] == 'error'), [5, 6, 7])
        self.assertTrue(any(f['row'] is None and 'chua-lap-chi-muc' in f['message'] for f in data['findings']))

        result = self.run_tool('library.py', 'find', '31/2024', '--json', '--result-json')
        human, data = result_schema.split_output(result.stdout)
        self.assertEqual((result.returncode, data['status'], data['count'], data['rows']), (0, 'success', 2, [5, 6]))
        self.assertEqual(json.loads(human)['count'], 2)

        result = self.run_tool('library.py', 'stats', '--json', '--result-json')
        human, data = result_schema.split_output(result.stdout)
        self.assertEqual((result.returncode, data['status'], data['rows']), (0, 'success', 5))
        self.assertEqual(json.loads(human)['rows'], 5)
        self.plain('library.py', 'stats')


class LibraryNotChecked(StructuredBase):
    def test_missing_openpyxl_is_not_checked_with_exit_two(self):
        result, _, data = self.structured('library.py', 'check', patch=NO_OPENPYXL)
        self.assertEqual((result.returncode, data['status']), (2, 'not_checked'))
        self.assertEqual(data['unchecked'], [{'tool': 'library.py', 'reason': 'openpyxl_missing'}])
        result, _, data = self.structured('library.py', 'stats', '--json', patch=NO_OPENPYXL)
        self.assertEqual((result.returncode, data['status']), (2, 'not_checked'))


@unittest.skipUnless(HAS_OPENPYXL, 'openpyxl is not installed; library checks skipped')
class LibraryReadFailures(StructuredBase):
    """A-07: an index that cannot be read or has no usable structure is a failure for find, stats and check, never a success with
    zero rows. A valid index with no rows, and a filter with no hits, stay successes. Synthetic workbooks in a package copy."""

    def index(self):
        return self.base / 'references' / '_INDEX.xlsx'

    def workbook(self, build):
        import openpyxl
        wb = openpyxl.Workbook()
        build(wb)
        self.index().parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(self.index()))

    def read(self, *args):
        result = self.run_tool('library.py', *args, '--result-json')
        human, data = result_schema.split_output(result.stdout)
        return result, human, data

    def test_valid_empty_index_and_filter_with_no_hit_are_successes(self):
        self.assertEqual(self.structured('library.py', 'init')[0].returncode, 0)
        result, _, data = self.read('stats', '--json')
        self.assertEqual((result.returncode, data['status'], data['rows']), (0, 'success', 0))
        result, _, data = self.read('find', 'khong-co-gi', '--json')
        self.assertEqual((result.returncode, data['status'], data['count']), (0, 'success', 0))
        fill_library(self.base)
        result, _, data = self.read('find', 'tu-khoa-khong-khop', '--json')
        self.assertEqual((result.returncode, data['status'], data['count']), (0, 'success', 0))  # lọc 0 hit khác sổ 0 dòng

    def test_unreadable_or_structureless_index_is_a_failure_with_its_own_reason(self):
        def corrupt():
            self.index().parent.mkdir(parents=True, exist_ok=True)
            self.index().write_bytes(b'not a workbook')

        def no_sheet():
            self.workbook(lambda wb: setattr(wb.active, 'title', 'Sheet khác'))

        def no_header():
            def build(wb):
                wb.active.title = 'Danh mục'
                wb.active['A1'] = 'chỉ một dòng'
            self.workbook(build)

        def no_columns():
            def build(wb):
                wb.active.title = 'Danh mục'
                wb.active.cell(4, 1, 'Mã / ID')
                wb.active.cell(4, 2, 'Trích yếu / Title')
            self.workbook(build)

        for name, build, reason in (('corrupt', corrupt, 'index_unreadable'), ('no sheet', no_sheet, 'index_missing_sheet'),
                                    ('no header row', no_header, 'index_no_header'), ('no columns', no_columns, 'index_missing_columns')):
            for command in (('find', 'x', '--json'), ('stats', '--json')):
                with self.subTest(case=name, command=command[0]):
                    build()
                    result, _, data = self.read(*command)
                    self.assertNotEqual(result.returncode, 0, 'a broken index must not look like an empty library')
                    self.assertIn(data['status'], ('blocked', 'error'))
                    self.assertEqual(data.get('reason'), reason)
                    self.assertIs(data['changes'], False)

    def test_check_still_reports_missing_columns_as_findings_and_stops_on_a_missing_header(self):
        def build(wb):
            wb.active.title = 'Danh mục'
            wb.active.cell(4, 1, 'Mã / ID')
            wb.active.cell(4, 2, 'Trích yếu / Title')
        self.workbook(build)
        result, _, data = self.read('check', '--json')
        self.assertEqual((result.returncode, data['status']), (1, 'issues'))  # hành vi cũ của check giữ nguyên / the old check behaviour is kept
        self.assertTrue(any('thiếu cột' in f['message'] for f in data['findings']))

    def test_cli_and_the_panel_reader_agree_on_identity_and_status(self):
        self.assertEqual(self.structured('library.py', 'init')[0].returncode, 0)
        fill_library(self.base)
        sys.path.insert(0, str(self.base / 'tools'))
        try:
            spec = importlib.util.spec_from_file_location('library_under_test', str(self.base / 'tools' / 'library.py'))
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
            panel = module.browse(str(self.base))
        finally:
            sys.path.remove(str(self.base / 'tools'))
        self.assertTrue(panel['ok'])
        cli = json.loads(self.run_tool('library.py', 'find', 'l-0', '--json', '--all').stdout)['hits']
        by_row = {h['row']: h for h in cli}
        for rec in panel['rows']:
            if rec['row'] in by_row:
                shown = by_row[rec['row']]
                self.assertEqual((rec['id'], rec['status'], rec['verify'], rec['title'], rec['revision'], rec['file_state']),
                                 (shown['id'], shown['status'], shown['verify'], shown['title'], shown['revision'], shown['file_state']))
        self.assertTrue(by_row)


@unittest.skipUnless(HAS_OPENPYXL, 'openpyxl is not installed; library checks skipped')
class LibraryCheckAgreesWithSync(StructuredBase):
    """M3-C1/F6: `check` must not be green where `sync` refuses (duplicate ID), shipped READMEs are told apart by ownership (never by name), and
    an unreadable index says what to do. Synthetic workbooks in a package copy."""

    def index(self):
        return self.base / 'references' / '_INDEX.xlsx'

    def with_duplicate_id(self):
        import openpyxl
        self.assertEqual(self.structured('library.py', 'init')[0].returncode, 0)
        fill_library(self.base)
        wb = openpyxl.load_workbook(str(self.index()))
        ws = wb['Danh mục']
        id_col = next(c for c in range(1, ws.max_column + 1) if str(ws.cell(row=4, column=c).value or '').startswith('Mã'))
        self.assertEqual(ws.cell(row=8, column=id_col).value, 'L-004')
        ws.cell(row=8, column=id_col, value=' L-001 ')  # the same ID as row 5 once stripped, like sync strips it
        wb.save(str(self.index()))

    def test_check_names_a_duplicate_id_with_its_rows_and_exits_nonzero(self):
        self.with_duplicate_id()
        result = self.run_tool('library.py', 'check', '--json', '--result-json')
        human, data = result_schema.split_output(result.stdout)
        self.assertEqual((result.returncode, data['status'], data['errors']), (1, 'issues', 4))  # the 3 old findings + the duplicate
        dup = [f for f in data['findings'] if 'L-001' in f['message']]
        self.assertEqual(len(dup), 1, data['findings'])
        self.assertEqual(dup[0]['severity'], 'error')
        self.assertEqual(dup[0]['row'], 8)
        self.assertIn('5, 8', dup[0]['message'])  # every row of the ID, not only the later one
        steps = ' '.join(data['next_actions'])
        self.assertIn('Excel', steps)
        self.assertIn('Đừng dựng lại sổ', steps)  # never "rebuild the index" as the way out
        plain = self.run_tool('library.py', 'check')
        self.assertEqual(plain.returncode, 1)
        self.assertIn('trùng Mã "L-001" ở các dòng 5, 8', plain.stdout)
        self.assertIn('KHÔNG ĐẠT', plain.stdout)

    def test_sync_refuses_the_same_index_for_the_same_reason(self):
        self.with_duplicate_id()
        before = self.index().read_bytes()
        result = self.run_tool('library.py', 'sync')
        self.assertNotEqual(result.returncode, 0, result.stdout + result.stderr)
        said = result.stdout + result.stderr
        self.assertIn('L-001', said)
        self.assertIn('5, 8', said)
        self.assertIn('library.py check', said)
        self.assertEqual(self.index().read_bytes(), before)  # a refusal edits nothing, IDs included

    def test_an_older_index_without_an_id_column_is_read_as_before(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Danh mục'
        for col, head in enumerate(('Nhóm', 'Số, ký hiệu', 'Trích yếu', 'Tình trạng', 'Có toàn văn', 'Đường dẫn file'), 1):
            ws.cell(row=4, column=col, value=head)
        for row, number in ((5, 'QC-01'), (6, 'QC-02')):
            for col, value in enumerate(('Nội bộ', number, 'Quy chế ' + number, 'Còn hiệu lực', 'Không', ''), 1):
                ws.cell(row=row, column=col, value=value)
        self.index().parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(self.index()))
        result = self.run_tool('library.py', 'check', '--json', '--result-json')
        human, data = result_schema.split_output(result.stdout)
        self.assertEqual((result.returncode, data['status'], data['errors']), (0, 'success', 0), result.stdout)

    def test_a_shipped_readme_is_left_out_by_ownership_and_one_you_edited_is_a_source(self):
        self.assertEqual(self.structured('library.py', 'init')[0].returncode, 0)
        ref = self.base / 'references'
        shipped = ref / 'README.md'
        self.assertTrue(shipped.is_file())
        # a manifest that really lists the READMEs the package ships, at their current bytes (the maintainer manifest of the repo is stale by design)
        lines = ['%s  %s' % (hashlib.sha256(p.read_bytes()).hexdigest(), p.relative_to(self.base).as_posix()) for p in sorted(ref.rglob('README.md'))]
        (self.base / 'MANIFEST.sha256').write_text('\n'.join(lines) + '\n', encoding='utf-8')
        group = ref / 'legal' / 'README.md'
        self.assertIn('%s  references/legal/README.md' % hashlib.sha256(group.read_bytes()).hexdigest(), lines)

        def unindexed():
            result = self.run_tool('library.py', 'check', '--json')
            return [f for f in json.loads(result.stdout)['unindexed'] if f.endswith('README.md')]

        self.assertEqual(unindexed(), [])  # the two shipped READMEs are not "missing from the index"
        keep = {p: p.read_bytes() for p in (shipped, group)}
        for path, name in ((shipped, 'README.md'), (group, 'legal/README.md')):
            path.write_bytes(path.read_bytes() + 'Ghi chú của tôi\n'.encode('utf-8'))
            self.assertEqual(unindexed(), [name], 'a README you wrote or edited is a source, whatever its file name')
            path.write_bytes(keep[path])
            self.assertEqual(unindexed(), [])
        other = ref / 'policies' / 'quy-che' / 'README.md'  # not shipped there: judged by ownership, not skipped for being named README
        write(other, 'my own notes')
        self.assertEqual(unindexed(), ['policies/quy-che/README.md'])

    @unittest.skipUnless(os.name == 'nt', 'Windows junction read-boundary regression')
    def test_readme_ownership_never_hashes_through_a_junction_outside_the_project(self):
        self.assertEqual(self.structured('library.py', 'init')[0].returncode, 0)
        outside = Path(self.tmp.name) / 'outside-project'
        outside.mkdir()
        canary = outside / 'README.md'
        canary.write_bytes(b'synthetic outside README; ownership must not read this')
        before = canary.read_bytes()
        link = self.base / 'references' / 'linked'
        made = subprocess.run(['cmd', '/c', 'mklink', '/J', str(link), str(outside)],
                              capture_output=True, encoding='utf-8', errors='replace')
        self.assertEqual(made.returncode, 0, made.stdout + made.stderr)
        self.addCleanup(os.rmdir, str(link))  # remove only the junction, before the temporary tree
        patch = ('import os, sys\n'
                 'sys.path.insert(0, os.path.dirname(os.path.abspath(sys.argv[1])))\n'
                 'import library_sync\n'
                 '_digest = library_sync.digest\n'
                 'def guarded_digest(path):\n'
                 '    if not library_sync.link_free(path, library_sync.library.ROOT):\n'
                 '        raise AssertionError("ownership attempted an outside README read")\n'
                 '    return _digest(path)\n'
                 'library_sync.digest = guarded_digest\n')
        result = self.run_tool('library.py', 'check', '--json', patch=patch)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn('linked/README.md', json.loads(result.stdout)['unindexed'])
        self.assertEqual(canary.read_bytes(), before)

    def test_an_unreadable_index_says_what_to_do_and_never_rebuilds_over_it(self):
        self.index().parent.mkdir(parents=True, exist_ok=True)
        self.index().write_bytes(b'not a workbook')
        result = self.run_tool('library.py', 'check', '--json', '--result-json')
        human, data = result_schema.split_output(result.stdout)
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(data.get('reason'), 'index_unreadable')
        self.assertTrue(data['next_actions'] and 'sao lưu' in data['next_actions'][0], data)
        self.assertIn('Đừng chép mẫu rỗng', data['next_actions'][0])
        self.assertEqual(self.index().read_bytes(), b'not a workbook')


class GuidesResults(StructuredBase):
    def test_check_rebuild_and_blocked_markdown(self):
        self.assertNotIn(MARKER, self.run_tool('build_guides_html.py', '--check').stdout)
        result, _, data = self.structured('build_guides_html.py', '--check')
        self.assertEqual((result.returncode, data['status']), (0, 'success'))

        guide = self.base / 'docs' / 'HUONG_DAN_SU_DUNG_VI.md'
        write(guide, read(guide) + '\nDòng kiểm thử.\n')
        result, _, data = self.structured('build_guides_html.py', '--check')
        self.assertEqual((result.returncode, data['status']), (1, 'issues'))
        self.assertIn('docs/HUONG_DAN_SU_DUNG_VI.html', [f['path'] for f in data['findings']])
        result, _, data = self.structured('build_guides_html.py')
        self.assertEqual((result.returncode, data['status'], data['changes']), (0, 'success', True))
        self.assertIn({'path': 'docs/HUONG_DAN_SU_DUNG_VI.html', 'action': 'rebuilt'}, data['files'])
        result, _, data = self.structured('build_guides_html.py', '--check')
        self.assertEqual((result.returncode, data['status']), (0, 'success'))

        write(guide, read(guide) + '\n### Mục con\n')
        result, _, data = self.structured('build_guides_html.py')
        self.assertEqual((result.returncode, data['status']), (2, 'blocked'))
        self.assertTrue(data['findings'][0]['message'])
        result, _, data = self.structured('build_guides_html.py', '--bogus')
        self.assertEqual((result.returncode, data['status']), (2, 'blocked'))


class UpgradeBranding(StructuredBase):
    """Nhận diện đơn vị khi nâng cấp: chỉ thêm khi project chưa có logo/branding.json/branding.js, không ghi đè; font thiếu thì thêm theo tên.
    Organization branding on upgrade: added only when the project has no logo, branding.json or branding.js of its own;
    never overwritten; missing fonts are added by name."""
    LOGO = b'\x89PNG\r\n\x1a\nlogo cua goi moi'
    FONT_A = b'font A cua goi moi'
    FONT_B = b'font B chi co trong goi moi'
    BRAND_JSON = json.dumps({'title': 'Đơn vị Mẫu', 'subtitle': 'Agent Workspace', 'logo': 'logo.png'}, ensure_ascii=False).encode('utf-8')

    def make_package(self, root, version, brand=None):
        """Gói có MANIFEST như thật. Nhận diện đặt SAU MANIFEST vì MANIFEST thật không liệt kê assets/ của đơn vị.
        A package with a real-style MANIFEST; branding is written after it, as a real MANIFEST never lists organization assets."""
        write(root / 'VERSION', version + '\n')
        write(root / 'PROJECT.md', 'private context')
        lines = [hashlib.sha256((root / rel).read_bytes()).hexdigest() + '  ' + rel for rel in ('PROJECT.md', 'VERSION')]
        write(root / 'MANIFEST.sha256', '\n'.join(lines) + '\n')
        for rel, data in (brand or {}).items():
            (root / rel).parent.mkdir(parents=True, exist_ok=True)
            (root / rel).write_bytes(data)
        return root

    def make_project(self, root, brand=None):
        self.make_package(root, '1.6.0', brand)
        write(root / '.agent' / 'PACKAGE_INFO.json', json.dumps({'package': 'local-agent-workspace', 'version': '1.6.0', 'language': 'vi'}))
        return root

    def test_upgrade_adds_template_branding_only_when_project_has_none(self):
        root = Path(self.tmp.name)
        new = self.make_package(root / 'new', '1.7.0', {
            'assets/branding.json': self.BRAND_JSON, 'assets/logo.png': self.LOGO,
            'assets/fonts/A.ttf': self.FONT_A, 'assets/fonts/B.ttf': self.FONT_B})

        # (a) dự án chưa có nhận diện: nhận logo, branding.json, font / no branding of its own: receives all of it
        bare = self.make_project(root / 'bare')
        result = self.run_tool('upgrade_package.py', '--project', bare, '--from', new, '--apply')
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        for rel in ('assets/branding.json', 'assets/logo.png', 'assets/fonts/A.ttf'):
            self.assertEqual((bare / rel).read_bytes(), (new / rel).read_bytes(), rel)

        # (b) dự án có logo và font riêng: giữ nguyên, không thêm branding.json, chỉ thêm font còn thiếu
        # own logo and font kept, no branding.json, only the missing font is added
        own = self.make_project(root / 'own', {'assets/logo.png': b'logo rieng cua don vi', 'assets/fonts/A.ttf': b'font rieng cua don vi'})
        result = self.run_tool('upgrade_package.py', '--project', own, '--from', new, '--apply')
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertEqual((own / 'assets' / 'logo.png').read_bytes(), b'logo rieng cua don vi')
        self.assertEqual((own / 'assets' / 'fonts' / 'A.ttf').read_bytes(), b'font rieng cua don vi')
        self.assertFalse((own / 'assets' / 'branding.json').exists())
        self.assertEqual((own / 'assets' / 'fonts' / 'B.ttf').read_bytes(), self.FONT_B)


if __name__ == '__main__':
    unittest.main()
