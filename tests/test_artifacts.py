"""B-02 checks: artifact versions and source–appendix–result relations (tools/artifacts.py), one writer, no scan, no AI.

Every test works in a fresh temporary project with fake data. Opening a file never launches an application (the opener is replaced).
P-CN1 revisions are written by the real owner (tools/publish_new.py); the controlled flow is Windows NTFS only and needs jsonschema.
Nothing here evaluates grants, acceptance, D4 or a checker PASS: the registry is a declaration with provenance, never a verification.
"""
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
import warnings
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'tools'))
sys.path.insert(0, str(ROOT / 'tests'))
import artifacts  # noqa: E402
import library  # noqa: E402
import library_sync  # noqa: E402
import publish_new as pn  # noqa: E402
import transaction  # noqa: E402

try:
    import openpyxl
except ImportError:
    openpyxl = None
try:
    import jsonschema  # noqa: F401
    HAVE_JSONSCHEMA = True
except ImportError:
    HAVE_JSONSCHEMA = False

TEMPLATE = ROOT / 'skills/local/shared/source-library/assets/source-index-template.xlsx'
WRITE_ERRORS = (library.ContentError, library_sync.SyncError)
NEED_PCN1 = unittest.skipUnless(os.name == 'nt' and HAVE_JSONSCHEMA, 'the P-CN1 owner needs Windows NTFS and jsonschema')


def sha(data):
    return hashlib.sha256(data).hexdigest()


def tree(root):
    """Every file under root with its bytes: the whole project, runtime included."""
    out = {}
    for p in sorted(Path(root).rglob('*')):
        out[p.relative_to(root).as_posix()] = p.read_bytes() if p.is_file() else None
    return out


def att():
    return {'source_class': 'USER_CONFIRMATION', 'channel': 'USER_MESSAGE', 'evidence_locator': 'host:msg/1',
            'recorded_by': {'actor_id': 'main', 'role': 'main'}}


@unittest.skipUnless(openpyxl, 'openpyxl is needed for the source workbook')
class Base(unittest.TestCase):
    def setUp(self):
        warnings.simplefilter('ignore')
        self.tmp = tempfile.TemporaryDirectory(prefix='law-b02-')
        self.addCleanup(self.tmp.cleanup)
        self.root = self.tmp.name
        self.project = Path(self.root)
        self.folder = self.project / 'references' / 'source-documents' / 'team'
        self.folder.mkdir(parents=True)
        for rel in ('working/w', 'output', '.agent/runtime'):
            (self.project / rel).mkdir(parents=True)
        self.index = self.project / 'references' / '_INDEX.xlsx'
        shutil.copyfile(TEMPLATE, self.index)
        self.a = self.source('a.txt', 'alpha v1')
        self.b = self.source('b.txt', 'beta v1')
        library_sync.apply(self.root)
        self.sa, self.sb = 'SRC-000001', 'SRC-000002'
        self.opened = []
        patcher = mock.patch.object(artifacts, '_launch', lambda path: self.opened.append(path))
        patcher.start()
        self.addCleanup(patcher.stop)

    # ---- fixtures
    def source(self, name, text):
        path = self.folder / name
        path.write_text(text, encoding='utf-8')
        return path

    def resync(self):
        return library_sync.apply(self.root)

    def cat(self, sid):
        rec = artifacts.load_book(self.root)['catalog'][sid][0]
        return library.revision_of(rec['law_revision']), rec['law_sha']

    def dep(self, sid):
        revision, digest = self.cat(sid)
        return {'kind': 'SOURCE', 'id': sid, 'revision': revision, 'sha256': digest}

    def adep(self, aid, revision):
        row = artifacts.load_book(self.root)['registry']['by_key'][(aid, revision)]
        return {'kind': 'ARTIFACT', 'id': aid, 'revision': revision, 'sha256': row['sha256']}

    def file(self, rel, text=None):
        path = self.project / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        if text is not None or not path.exists():
            path.write_text(text if text is not None else 'content of ' + rel, encoding='utf-8')
        return path

    def req(self, aid, revision, deps=(), path=None, **over):
        path = path or 'working/w/%s-v%d.txt' % (aid.lower(), revision)
        self.file(path)
        out = {'artifact_id': aid, 'revision': revision, 'path': path, 'title': 'Kết quả ' + aid, 'kind': 'REPORT',
               'limits': 'chỉ đọc đoạn đã khai', 'origin': 'USER_FILE', 'recorded_by': 'tester', 'dependencies': list(deps)}
        out.update(over)
        return out

    def register(self, req):
        plan = artifacts.register_plan(self.root, req)
        return artifacts.register_apply(self.root, req, plan['plan_token'])

    def new(self, aid, revision, deps=(), **over):
        res = self.register(self.req(aid, revision, deps, **over))
        self.assertEqual(res['action'], 'record', res)
        return res

    def find(self, **kw):
        res = artifacts.find(self.root, **kw)
        self.assertTrue(res['ok'], res)
        return res

    def item(self, aid, revision):
        return next(i for i in self.find(artifact_id=aid)['items'] if i['revision'] == revision)

    def refused(self, code, call, *args):
        with self.assertRaises(WRITE_ERRORS) as caught:
            call(*args)
        self.assertEqual(getattr(caught.exception, 'code', None), code, str(caught.exception))
        return caught.exception

    def sheet_names(self):
        wb = openpyxl.load_workbook(self.index)
        try:
            return wb.sheetnames
        finally:
            wb.close()

    def edit(self, fn):
        wb = openpyxl.load_workbook(self.index)
        try:
            fn(wb)
            wb.save(self.index)
        finally:
            wb.close()

    def event(self, sid, **over):
        revision, digest = self.cat(sid)
        req = {'event_id': 'V-%d' % revision, 'source_id': sid, 'expected_revision': revision, 'expected_sha256': digest,
               'mode': 'semantic_update', 'changes': {'status': 'Hết hiệu lực'}, 'reason': 'thay thế bởi văn bản mới',
               'evidence_ref': 'công văn giả 01', 'recorded_by': 'tester'}
        req.update(over)
        return library_sync.event_apply(self.root, req, library_sync.event_plan(self.root, req)['plan_token'])


class LazyAndLegacy(Base):
    def test_legacy_workbook_has_no_registry_and_reads_write_nothing(self):
        before = tree(self.root)
        res = self.find()
        self.assertEqual(res['registry'], 'absent')
        self.assertEqual((res['total'], res['items']), (0, []))
        self.assertEqual(self.find(source_id=self.sa)['matched'], 0)
        sel = {'snapshot': res['snapshot'], 'row': 5, 'id': self.sa, 'revision': 1, 'locator': 'source-documents/team/a.txt'}
        self.assertEqual(artifacts.links(self.root, sel)['results'], [])
        self.assertEqual(tree(self.root), before)
        self.assertNotIn(artifacts.ARTIFACT_SHEET, self.sheet_names())

    def test_preview_never_creates_sheets_or_writes_and_the_first_apply_creates_both(self):
        req = self.req('R1', 1, [self.dep(self.sa)])  # writes the result file the plan will read
        before = tree(self.root)
        plan = artifacts.register_plan(self.root, req)
        self.assertEqual((plan['action'], plan['sheets_created']), ('record', True))
        self.assertIsNotNone(plan['bytes'])
        rel = artifacts.relate_plan(self.root, {'source_a': self.sa, 'source_b': self.sb, 'recorded_by': 't'})
        self.assertEqual(rel['action'], 'record')
        self.assertEqual(tree(self.root), before)
        self.register(req)
        self.assertEqual([n for n in self.sheet_names() if n.startswith('LAW ')], ['LAW artifacts', 'LAW relations'])

    def test_only_one_of_the_two_sheets_is_a_structure_error_not_empty(self):
        self.new('R1', 1, [self.dep(self.sa)])
        self.edit(lambda wb: wb.remove(wb[artifacts.RELATION_SHEET]))
        res = artifacts.find(self.root)
        self.assertEqual((res['ok'], res['kind'], res['code']), (False, 'registry_malformed', 'ARTIFACT_SHEETS_INCOMPLETE'))
        self.refused('ARTIFACT_SHEETS_INCOMPLETE', artifacts.register_plan, self.root, self.req('R2', 1))
        self.edit(lambda wb: wb.remove(wb[artifacts.ARTIFACT_SHEET]))
        self.assertTrue(artifacts.find(self.root)['ok'])  # neither: valid legacy again


class Lineage(Base):
    def test_many_sources_many_outputs_and_related_does_not_propagate(self):
        self.new('R1', 1, [self.dep(self.sa), self.dep(self.sb)])
        self.new('R2', 1, [self.dep(self.sa)])
        rel = artifacts.relate_plan(self.root, {'source_a': self.sb, 'source_b': self.sa, 'recorded_by': 't'})
        artifacts.relate_apply(self.root, {'source_a': self.sb, 'source_b': self.sa, 'recorded_by': 't'}, rel['plan_token'])
        self.assertEqual({i['key'] for i in self.find(source_id=self.sa)['items']}, {'R1@1', 'R2@1'})
        self.assertEqual({i['key'] for i in self.find(source_id=self.sb)['items']}, {'R1@1'})
        self.assertEqual(self.item('R1', 1)['dependency_state'], 'CURRENT_CATALOG_BASIS')
        # a change to the related source marks nothing that only used the other one
        self.source('b.txt', 'beta v2')
        self.resync()
        self.assertEqual(self.item('R2', 1)['dependency_state'], 'CURRENT_CATALOG_BASIS')
        self.assertEqual(self.item('R1', 1)['dependency_state'], 'NEEDS_REVIEW')
        sel = {'snapshot': self.find()['snapshot'], 'row': 5, 'id': self.sa, 'revision': 1, 'locator': 'source-documents/team/a.txt'}
        out = artifacts.links(self.root, sel)
        self.assertEqual([r['id'] for r in out['related']], [self.sb])
        self.assertEqual({r['artifact_id'] for r in out['results']}, {'R1', 'R2'})

    def test_related_retry_is_a_noop_and_orders_the_pair(self):
        req = {'source_a': self.sb, 'source_b': self.sa, 'recorded_by': 't'}
        plan = artifacts.relate_plan(self.root, req)
        artifacts.relate_apply(self.root, req, plan['plan_token'])
        again = artifacts.relate_plan(self.root, dict(req, source_a=self.sa, source_b=self.sb))
        self.assertEqual(again['action'], 'already_recorded')
        self.assertIsNone(again['bytes'])
        self.refused('MALFORMED_REQUEST', artifacts.relate_plan, self.root, {'source_a': self.sa, 'source_b': self.sa, 'recorded_by': 't'})
        self.refused('UNKNOWN_SOURCE', artifacts.relate_plan, self.root, {'source_a': self.sa, 'source_b': 'SRC-999999', 'recorded_by': 't'})

    def test_a_to_b_to_c_with_a_source_event_on_the_same_sha_is_still_needs_review(self):
        self.new('B', 1, [self.dep(self.sa)])
        self.new('C', 1, [self.adep('B', 1)])
        self.assertEqual(self.item('C', 1)['dependency_state'], 'CURRENT_CATALOG_BASIS')
        _, digest = self.cat(self.sa)
        self.event(self.sa)
        self.assertEqual(self.cat(self.sa)[1], digest)  # same bytes, new revision
        for aid in ('B', 'C'):
            item = self.item(aid, 1)
            self.assertEqual(item['dependency_state'], 'NEEDS_REVIEW', aid)
        self.assertEqual([d['reason'] for d in self.item('C', 1)['diagnostics']], ['catalog_revision_newer'])

    def test_a_second_revision_never_replaces_what_the_first_was_used_for(self):
        self.new('B', 1, [self.dep(self.sa)])
        self.new('C', 1, [self.adep('B', 1)])
        self.new('B', 2, [self.dep(self.sa)])
        c1 = self.item('C', 1)
        self.assertEqual((c1['dependencies'][0]['id'], c1['dependencies'][0]['revision']), ('B', 1))
        self.assertEqual((c1['dependency_state'], c1['dependencies'][0]['reason']), ('NEEDS_REVIEW', 'upstream_superseded'))
        self.assertEqual((self.item('B', 1)['superseded_by'], self.item('B', 2)['superseded_by']), (2, ''))
        self.refused('UPSTREAM_SUPERSEDED', artifacts.prepare, self.root, {'inputs': [{'kind': 'ARTIFACT', 'id': 'B', 'revision': 1}]})
        self.refused('UPSTREAM_SUPERSEDED', artifacts.register_plan, self.root, self.req('C', 2, [self.adep('B', 1)]))

    def test_kind_limits_and_unknown_basis_travel_down_to_the_consumer(self):
        self.new('B', 1, kind='INFERENCE', limits='suy luận, giả định X chưa kiểm')  # no declared basis
        self.new('C', 1, [self.adep('B', 1)], kind='REPORT', limits='chỉ phần 2')
        c = self.item('C', 1)
        self.assertEqual(self.item('B', 1)['dependency_state'], 'NO_DECLARED_BASIS')
        self.assertEqual((c['kind'], c['inference'], c['unknown_basis'], c['dependency_state']), ('REPORT', ['B@1'], ['B@1'], 'NO_DECLARED_BASIS'))
        self.assertEqual({x['key']: (x['kind'], x['limits']) for x in c['lineage']},
                         {'C@1': ('REPORT', 'chỉ phần 2'), 'B@1': ('INFERENCE', 'suy luận, giả định X chưa kiểm')})
        text = artifacts.copy_input(self.root, self.find(artifact_id='C')['items'][0]['pick'])['text']
        for needle in ('suy luận, giả định X chưa kiểm', 'chỉ phần 2', 'INFERENCE', 'B@1', 'sha256 '):
            self.assertIn(needle, text)
        prep = artifacts.prepare(self.root, {'inputs': [{'kind': 'ARTIFACT', 'id': 'C', 'revision': 1}]})
        self.assertEqual((prep['inference'], prep['unknown_basis']), (['B@1'], ['B@1']))
        self.assertIn('suy luận, giả định X chưa kiểm', json.dumps(prep['limits'], ensure_ascii=False))
        self.assertFalse(prep['verifies_content'])

    def test_historical_basis_that_no_longer_matches_does_not_block_sync_or_source_events(self):
        self.new('R1', 1, [self.dep(self.sa)])
        self.source('a.txt', 'alpha v2')
        self.resync()  # revision 2: the registered basis is now historical
        self.assertEqual(self.item('R1', 1)['dependency_state'], 'NEEDS_REVIEW')
        self.source('b.txt', 'beta v2')
        self.assertTrue(self.resync()['changes'])  # sync still writes
        self.event(self.sa)  # a valid source-event still raises the revision
        self.assertEqual(self.cat(self.sa)[0], 3)
        self.assertEqual(self.item('R1', 1)['dependencies'][0]['revision'], 1)  # the old basis is kept, never rewritten
        (self.folder / 'a.txt').unlink()
        self.resync()  # the source file is now missing; the registry stays readable
        self.assertEqual(self.item('R1', 1)['dependency_state'], 'NEEDS_REVIEW')

    def test_find_and_links_never_hash_a_file_or_locate_a_source(self):
        self.new('B', 1, [self.dep(self.sa)])
        self.new('C', 1, [self.adep('B', 1)])
        boom = AssertionError('a query hashed or opened a file')
        sel = {'snapshot': self.find()['snapshot'], 'row': 5, 'id': self.sa, 'revision': 1, 'locator': 'source-documents/team/a.txt'}
        with mock.patch.object(library_sync, 'digest', side_effect=boom), mock.patch.object(artifacts, 'observe_file', side_effect=boom),                 mock.patch.object(artifacts, 'observe_source', side_effect=boom), mock.patch.object(library, 'locate_source', side_effect=boom),                 mock.patch.object(library_sync, 'sources', side_effect=boom), mock.patch('os.walk', side_effect=boom):
            res = self.find()
            self.assertEqual(res['bytes_observation'], 'UNOBSERVED')
            self.assertEqual([i['bytes_observation'] for i in res['items']], ['UNOBSERVED', 'UNOBSERVED'])
            self.assertEqual(len(artifacts.links(self.root, sel)['results']), 2)
            self.assertTrue(artifacts.copy_input(self.root, res['items'][1]['pick'])['ok'])

    def test_closure_limit_is_an_error_never_a_silent_cut(self):
        self.new('A1', 1, [self.dep(self.sa)])
        self.new('A2', 1, [self.adep('A1', 1)])
        self.new('A3', 1, [self.adep('A2', 1)])
        with mock.patch.object(artifacts, 'MAX_DEPTH', 1):
            res = artifacts.find(self.root)
            self.assertEqual((res['ok'], res['code']), (False, 'CLOSURE_LIMIT'))
            self.refused('CLOSURE_LIMIT', artifacts.prepare, self.root, {'inputs': [{'kind': 'ARTIFACT', 'id': 'A3', 'revision': 1}]})
        with mock.patch.object(artifacts, 'MAX_NODES', 2):
            self.refused('CLOSURE_LIMIT', artifacts.prepare, self.root, {'inputs': [{'kind': 'ARTIFACT', 'id': 'A3', 'revision': 1}]})


class Registration(Base):
    def test_retry_returns_already_without_rows_or_a_freshness_claim_and_conflict_otherwise(self):
        req = self.req('R1', 1, [self.dep(self.sa)])
        self.register(req)
        rows = len(self.find()['items'])
        before = self.index.read_bytes()
        plan = artifacts.register_plan(self.root, req)
        self.assertEqual((plan['action'], plan['current_basis_validated'], plan['bytes']), ('already_recorded', False, None))
        res = artifacts.register_apply(self.root, req, plan['plan_token'])
        self.assertEqual(res['action'], 'already_recorded')
        self.assertEqual((self.index.read_bytes(), len(self.find()['items'])), (before, rows))
        # a stale source does not turn a retry into a refusal (history, not freshness)
        self.source('a.txt', 'alpha v2')
        self.resync()
        self.assertEqual(artifacts.register_plan(self.root, req)['action'], 'already_recorded')
        self.refused('ARTIFACT_CONFLICT', artifacts.register_plan, self.root, dict(req, title='khác'))
        self.refused('ARTIFACT_CONFLICT', artifacts.register_plan, self.root, dict(req, dependencies=[]))

    def test_revision_predecessor_and_path_rules(self):
        self.new('R1', 1)
        self.refused('BAD_REVISION', artifacts.register_plan, self.root, self.req('R1', 3))
        self.refused('BAD_REVISION', artifacts.register_plan, self.root, self.req('Z', 2))
        self.refused('BAD_PREDECESSOR', artifacts.register_plan, self.root, self.req('R1', 2, predecessor=5))
        self.refused('PATH_REUSED', artifacts.register_plan, self.root, self.req('R1', 2, path='working/w/r1-v1.txt'))
        self.refused('PATH_REUSED', artifacts.register_plan, self.root, self.req('Q', 1, path='working/w/R1-V1.TXT'))
        self.new('R1', 2, predecessor=1)
        row = artifacts.load_book(self.root)['registry']['by_key'][('R1', 2)]
        self.assertEqual(row['predecessor'], 1)
        self.assertEqual((self.item('R1', 1)['superseded_by'], self.item('R1', 2)['latest_revision']), (2, 2))

    def test_paths_outside_working_or_output_and_authority_names_are_refused(self):
        for rel in ('archive/x.txt', 'references/x.txt', '../x.txt', 'working/../x.txt', 'working/AGENTS.md', 'working\\x.txt', 'x.txt'):
            with self.assertRaises(WRITE_ERRORS, msg=rel):
                artifacts.register_plan(self.root, dict(self.req('PP', 1), path=rel))
        self.assertEqual(self.find()['registry'], 'absent')

    def test_unknown_keys_missing_values_and_oversized_cells_are_refused_not_cut(self):
        base = self.req('R1', 1)
        self.refused('MALFORMED_REQUEST', artifacts.register_plan, self.root, dict(base, surprise=1))
        self.refused('MALFORMED_REQUEST', artifacts.register_plan, self.root, dict(base, kind='NOPE'))
        self.refused('MALFORMED_REQUEST', artifacts.register_plan, self.root, dict(base, origin='P_CN1'))  # rid is missing
        self.refused('MALFORMED_REQUEST', artifacts.register_plan, self.root, dict(base, rid='0' * 32))  # USER_FILE has no rid
        self.refused('MALFORMED_REQUEST', artifacts.register_plan, self.root, dict(base, limits=''))
        self.refused('MALFORMED_REQUEST', artifacts.register_plan, self.root, dict(base, title='x' * (library.CELL_LIMIT + 1)))
        self.refused('MALFORMED_REQUEST', artifacts.register_plan, self.root, dict(base, title='bad\x07char'))
        self.refused('MALFORMED_REQUEST', artifacts.register_plan, self.root, dict(base, dependencies=[{'kind': 'SOURCE', 'id': self.sa}]))
        dep = self.dep(self.sa)
        self.refused('MALFORMED_REQUEST', artifacts.register_plan, self.root, dict(base, dependencies=[dep, dep]))
        self.refused('MALFORMED_REQUEST', artifacts.register_plan, self.root, dict(base, dependencies=[self.adep_self()]))
        self.assertEqual(self.find()['registry'], 'absent')

    def adep_self(self):
        return {'kind': 'ARTIFACT', 'id': 'R1', 'revision': 1, 'sha256': '0' * 64}

    def test_strings_stay_literal_through_every_load_and_save(self):
        self.new('R1', 1, [self.dep(self.sa)], title='=1+1', limits='@SUM(A1)', recorded_by='+cmd')
        self.new('R2', 1, title='-2+3')  # a second write reloads and saves the workbook again
        wb = openpyxl.load_workbook(self.index, data_only=False)
        try:
            ws = wb[artifacts.ARTIFACT_SHEET]
            self.assertEqual([ws.cell(5, 7).value, ws.cell(5, 9).value, ws.cell(5, 13).value, ws.cell(6, 7).value],
                             ['=1+1', '@SUM(A1)', '+cmd', '-2+3'])
            self.assertTrue(all(ws.cell(r, c).data_type != 'f' for r in (5, 6) for c in range(1, 15)))
        finally:
            wb.close()
        self.assertEqual(self.item('R1', 1)['title'], '=1+1')

    def test_dependencies_must_match_the_current_sources_bytes_and_versions(self):
        dep = self.dep(self.sa)
        self.refused('STALE_SOURCE_BASIS', artifacts.register_plan, self.root, self.req('R1', 1, [dict(dep, revision=9)]))
        self.refused('STALE_SOURCE_BASIS', artifacts.register_plan, self.root, self.req('R1', 1, [dict(dep, sha256='1' * 64)]))
        self.refused('UNKNOWN_SOURCE', artifacts.register_plan, self.root, self.req('R1', 1, [dict(dep, id='SRC-424242')]))
        self.refused('UNKNOWN_ARTIFACT', artifacts.register_plan, self.root, self.req('R1', 1, [{'kind': 'ARTIFACT', 'id': 'Nope', 'revision': 1, 'sha256': '1' * 64}]))
        self.source('a.txt', 'alpha changed on disk, not synced')
        self.refused('STALE_SOURCE_BASIS', artifacts.register_plan, self.root, self.req('R1', 1, [dep]))
        self.assertEqual(self.find()['registry'], 'absent')

    def test_user_file_without_a_basis_is_no_declared_basis_and_makes_no_receipt_or_authority(self):
        self.new('U1', 1)
        item = self.item('U1', 1)
        self.assertEqual((item['dependency_state'], item['origin'], item['publish_rid'], item['publish_identity']), ('NO_DECLARED_BASIS', 'USER_FILE', '', ''))
        self.assertEqual(item['dependencies'], [])
        self.assertFalse((self.project / '.agent' / 'runtime' / 'publish').exists())
        self.assertFalse((self.project / '.agent' / 'STATE.md').exists())  # no current-draft pointer, no STATE write
        prep = artifacts.prepare(self.root, {'inputs': [{'kind': 'ARTIFACT', 'id': 'U1', 'revision': 1}]})
        self.assertEqual(prep['unknown_basis'], ['U1@1'])  # usable as a user-supplied document, unknown travels

    def test_user_file_digest_declaration_must_match_and_unsafe_files_are_refused(self):
        req = self.req('U1', 1)
        self.refused('FILE_DIGEST_MISMATCH', artifacts.register_plan, self.root, dict(req, sha256='2' * 64))
        self.assertEqual(artifacts.register_plan(self.root, dict(req, sha256=sha((self.project / req['path']).read_bytes())))['action'], 'record')
        self.refused('ARTIFACT_FILE_MISSING', artifacts.register_plan, self.root, dict(req, path='working/w/absent.txt'))
        os.makedirs(self.project / 'working' / 'dir.txt')
        self.refused('ARTIFACT_FILE_NOT_REGULAR', artifacts.register_plan, self.root, dict(req, path='working/dir.txt'))
        os.link(str(self.project / req['path']), str(self.project / 'working' / 'w' / 'twin.txt'))
        self.refused('ARTIFACT_FILE_HARDLINK', artifacts.register_plan, self.root, dict(req, path='working/w/twin.txt'))
        self.refused('ARTIFACT_FILE_HARDLINK', artifacts.register_plan, self.root, dict(req, artifact_id='U2', path=req['path']))

    def test_the_plan_token_binds_the_workbook_the_file_the_sources_and_the_request(self):
        req = self.req('R1', 1, [self.dep(self.sa)])
        token = artifacts.register_plan(self.root, req)['plan_token']
        before = self.index.read_bytes()
        (self.project / req['path']).write_text('the file changed after the preview', encoding='utf-8')
        err = self.refused_mismatch(req, token)
        self.assertEqual(err.reason, 'plan')
        self.file(req['path'], 'content of ' + req['path'])
        token = artifacts.register_plan(self.root, req)['plan_token']
        self.source('b.txt', 'beta v2')
        self.resync()  # the workbook changed after the preview, the declared source did not
        self.refused_mismatch(req, token)
        token = artifacts.register_plan(self.root, req)['plan_token']
        self.refused_mismatch(dict(req, title='another title'), token)  # a different request is a different plan
        self.source('a.txt', 'alpha v2')
        self.resync()  # the declared source itself moved: the plan cannot even be recomputed
        before = self.index.read_bytes()
        self.refused('STALE_SOURCE_BASIS', artifacts.register_apply, self.root, req, token)
        self.assertEqual(self.index.read_bytes(), before)
        self.assertEqual(self.find()['registry'], 'absent')
        self.assertFalse(transaction.incomplete(self.root))

    def refused_mismatch(self, req, token):
        with self.assertRaises(library_sync.PreviewMismatch) as caught:
            artifacts.register_apply(self.root, req, token)
        return caught.exception

    def test_apply_needs_a_wellformed_token_and_a_concurrent_registration_makes_it_stale(self):
        req = self.req('R1', 1, [self.dep(self.sa)])
        for bad in (None, '', 'zz', 'a' * 63):
            with self.assertRaises(library_sync.PreviewMismatch):
                artifacts.register_apply(self.root, req, bad)
        token = artifacts.register_plan(self.root, req)['plan_token']
        self.new('Other', 1)
        self.refused_mismatch(req, token)  # the workbook changed under the plan
        self.assertEqual({i['key'] for i in self.find()['items']}, {'Other@1'})

    def test_x14_refusal_and_staged_verification_run_before_any_transaction(self):
        req = self.req('R1', 1, [self.dep(self.sa)])
        before = self.index.read_bytes()
        with mock.patch.object(library_sync, 'lossy_notices', return_value=['x14_data_validation']):
            plan = artifacts.register_plan(self.root, req)
            self.assertEqual(plan['notices'], ['x14_data_validation'])  # preview still works and says why apply refuses
            with self.assertRaises(library_sync.SyncError):
                artifacts.register_apply(self.root, req, plan['plan_token'])
        self.assertEqual(self.index.read_bytes(), before)
        self.assertFalse(transaction.incomplete(self.root))

    def test_other_managed_sheets_and_formulas_outside_the_registry_survive(self):
        req = {'entry_id': 'E-1', 'source_id': self.sa, 'source_revision': self.cat(self.sa)[0], 'source_sha256': self.cat(self.sa)[1],
               'locator': 'đoạn 1', 'read_scope': 'đoạn 1', 'kind': 'EXTRACT', 'text': 'nội dung giả', 'recorded_by': 'tester'}
        library_sync.content_apply(self.root, req, library_sync.content_plan(self.root, req)['plan_token'])
        self.event(self.sa)
        self.edit(lambda wb: wb.create_sheet('Ghi chú riêng').__setitem__('A1', '=1+2'))
        keep = library.content_find(self.root)['entries']
        self.new('R1', 1, [self.dep(self.sa)])
        self.new('R2', 1, [self.adep('R1', 1)])
        wb = openpyxl.load_workbook(self.index, data_only=False)
        try:
            self.assertEqual(wb['Ghi chú riêng']['A1'].value, '=1+2')
            self.assertEqual(wb['Ghi chú riêng']['A1'].data_type, 'f')
            self.assertIn('LAW history', wb.sheetnames)
        finally:
            wb.close()
        self.assertEqual(library.content_find(self.root)['entries'], keep)
        # and the other writers still keep the registry
        self.source('b.txt', 'beta v2')
        self.resync()
        self.assertEqual({i['key'] for i in self.find()['items']}, {'R1@1', 'R2@1'})


class Malformed(Base):
    def setUp(self):
        super().setUp()
        self.new('B', 1, [self.dep(self.sa)])
        self.new('C', 1, [self.adep('B', 1)])

    def broken(self, mutate):
        self.edit(mutate)
        res = artifacts.find(self.root)
        self.assertFalse(res['ok'])
        self.assertEqual(res['kind'], 'registry_malformed')
        return res['code']

    def test_every_writer_refuses_a_malformed_registry_and_nothing_is_written(self):
        digest = self.cat(self.sa)[1]
        self.edit(lambda wb: wb[artifacts.ARTIFACT_SHEET].cell(5, 5).__setattr__('value', 'not-a-sha'))
        before = self.index.read_bytes()
        self.refused('MALFORMED_ARTIFACT_SHEET', artifacts.register_plan, self.root, self.req('Z', 1))
        self.refused('MALFORMED_ARTIFACT_SHEET', artifacts.relate_plan, self.root, {'source_a': self.sa, 'source_b': self.sb, 'recorded_by': 't'})
        self.refused('MALFORMED_ARTIFACT_SHEET', library_sync.event_plan, self.root, {
            'event_id': 'V-1', 'source_id': self.sa, 'expected_revision': 1, 'expected_sha256': digest, 'mode': 'semantic_update',
            'changes': {'status': 'x'}, 'reason': 'r', 'evidence_ref': 'e', 'recorded_by': 't'})
        self.refused('MALFORMED_ARTIFACT_SHEET', library_sync.content_plan, self.root, {
            'entry_id': 'E-1', 'source_id': self.sa, 'source_revision': 1, 'source_sha256': digest, 'locator': 'l',
            'read_scope': 's', 'kind': 'EXTRACT', 'text': 't', 'recorded_by': 't'})
        with self.assertRaises(library_sync.SyncError):
            library_sync.plan(self.root)
        self.assertEqual(self.index.read_bytes(), before)

    def test_structure_errors_are_named(self):
        def bad_marker(wb):
            wb[artifacts.ARTIFACT_SHEET]['A1'] = 'something else'
        self.assertEqual(self.broken(bad_marker), 'MALFORMED_ARTIFACT_SHEET')

    def test_formula_in_a_managed_cell_is_a_structure_error(self):
        code = self.broken(lambda wb: wb[artifacts.ARTIFACT_SHEET].cell(5, 7).__setattr__('value', '=1+1'))
        self.assertEqual(code, 'MALFORMED_ARTIFACT_SHEET')

    def test_header_change_and_data_outside_the_managed_columns(self):
        self.assertEqual(self.broken(lambda wb: wb[artifacts.RELATION_SHEET].cell(4, 3).__setattr__('value', 'x')), 'MALFORMED_RELATION_SHEET')

    def test_data_outside_the_managed_columns_is_refused(self):
        self.assertEqual(self.broken(lambda wb: wb[artifacts.ARTIFACT_SHEET].cell(5, 30).__setattr__('value', 'stray')), 'MALFORMED_ARTIFACT_SHEET')

    def test_duplicate_version_and_duplicate_relation(self):
        def dup_artifact(wb):
            ws = wb[artifacts.ARTIFACT_SHEET]
            for col in range(1, 15):
                ws.cell(7, col).value = ws.cell(5, col).value
        self.assertEqual(self.broken(dup_artifact), 'MALFORMED_ARTIFACT_SHEET')

    def test_duplicate_relation_row(self):
        def dup(wb):
            ws = wb[artifacts.RELATION_SHEET]
            for col in range(1, 16):
                ws.cell(7, col).value = ws.cell(5, col).value
        self.assertEqual(self.broken(dup), 'MALFORMED_RELATION_SHEET')

    def test_path_reused_by_two_versions_and_revision_gap(self):
        def same_path(wb):
            ws = wb[artifacts.ARTIFACT_SHEET]
            ws.cell(6, 4).value = ws.cell(5, 4).value
        self.assertEqual(self.broken(same_path), 'MALFORMED_ARTIFACT_SHEET')

    def test_revision_gap_without_predecessor(self):
        self.assertEqual(self.broken(lambda wb: wb[artifacts.ARTIFACT_SHEET].cell(6, 2).__setattr__('value', 3)), 'MALFORMED_ARTIFACT_SHEET')

    def test_dangling_artifact_version_is_a_structure_error_not_an_empty_dependency(self):
        self.assertEqual(self.broken(lambda wb: wb[artifacts.RELATION_SHEET].cell(6, 8).__setattr__('value', 9)), 'ARTIFACT_DANGLING')

    def test_a_relation_whose_target_hash_differs_from_the_registered_version(self):
        self.assertEqual(self.broken(lambda wb: wb[artifacts.RELATION_SHEET].cell(6, 9).__setattr__('value', '3' * 64)), 'MALFORMED_RELATION_SHEET')

    def test_a_tampered_relation_key_is_detected(self):
        self.assertEqual(self.broken(lambda wb: wb[artifacts.RELATION_SHEET].cell(5, 7).__setattr__('value', 'SRC-OTHER')), 'MALFORMED_RELATION_SHEET')

    def test_cycle_between_versions(self):
        book = artifacts.load_book(self.root)
        b_sha = book['registry']['by_key'][('B', 1)]['sha256']
        c_sha = book['registry']['by_key'][('C', 1)]['sha256']

        def cycle(wb):
            ws = wb[artifacts.RELATION_SHEET]
            key = artifacts.relation_key('DERIVED_FROM', 'ARTIFACT', 'B', 1, 'ARTIFACT', 'C')
            values = (key, 'DERIVED_FROM', 'ARTIFACT', 'B', 1, 'ARTIFACT', 'C', 1, c_sha, '', '', '', '', 'tester', '2026-01-01T00:00:00')
            for col, value in enumerate(values, 1):
                ws.cell(7, col).value = value
        self.assertEqual(self.broken(cycle), 'ARTIFACT_CYCLE')
        self.assertTrue(b_sha)

    def test_oversized_text_in_a_cell_cannot_be_registered(self):
        # the writer refuses before it could cut the text; a cell the workbook cannot hold never reaches the sheet
        with self.assertRaises(WRITE_ERRORS):
            artifacts.register_plan(self.root, self.req('Big', 1, limits='y' * (library.CELL_LIMIT + 5)))

    def test_unreadable_workbook_and_missing_index_are_reported_not_empty(self):
        self.index.write_bytes(b'not a workbook')
        res = artifacts.find(self.root)
        self.assertEqual((res['ok'], res['kind']), (False, 'unreadable'))
        self.index.unlink()
        self.assertEqual(artifacts.find(self.root)['kind'], 'no_index')


class PrepareAndGuards(Base):
    def setUp(self):
        super().setUp()
        self.new('B', 1, [self.dep(self.sa)])

    def test_prepare_returns_the_snapshot_and_exact_inputs_for_the_publish_task(self):
        before = tree(self.root)
        prep = artifacts.prepare(self.root, {'inputs': [{'kind': 'ARTIFACT', 'id': 'B', 'revision': 1}, {'kind': 'SOURCE', 'id': self.sb}]})
        paths = {i['path']: i for i in prep['inputs']}
        self.assertEqual(set(paths), {'references/_INDEX.xlsx', 'references/source-documents/team/a.txt', 'references/source-documents/team/b.txt', 'working/w/b-v1.txt'})
        self.assertEqual(paths['references/_INDEX.xlsx']['digest'], 'sha256:' + prep['snapshot'])
        self.assertEqual(prep['snapshot'], sha(self.index.read_bytes()))
        self.assertEqual(paths['references/source-documents/team/a.txt']['source_id'], self.sa)
        self.assertEqual(paths['working/w/b-v1.txt']['digest'], 'sha256:' + sha((self.project / 'working/w/b-v1.txt').read_bytes()))
        self.assertEqual(tree(self.root), before)  # prepare writes nothing: no receipt, no store

    def test_a_source_file_changed_after_selection_but_not_synced_is_refused_by_prepare(self):
        pick = self.find()['items'][0]['pick']
        self.source('a.txt', 'alpha edited after the selection, never synced')
        self.assertTrue(artifacts.copy_input(self.root, pick)['ok'])  # copying only copies; it observes nothing
        self.refused('STALE_SOURCE_BASIS', artifacts.prepare, self.root, {'inputs': [{'kind': 'ARTIFACT', 'id': 'B', 'revision': 1}]})
        self.refused('STALE_SOURCE_BASIS', artifacts.prepare, self.root, {'inputs': [{'kind': 'SOURCE', 'id': self.sa}]})

    def test_stale_known_basis_and_changed_result_files_are_refused(self):
        self.source('a.txt', 'alpha v2')
        self.resync()
        self.refused('STALE_SOURCE_BASIS', artifacts.prepare, self.root, {'inputs': [{'kind': 'ARTIFACT', 'id': 'B', 'revision': 1}]})

    def test_a_changed_result_file_is_refused(self):
        (self.project / 'working/w/b-v1.txt').write_text('edited by hand', encoding='utf-8')
        self.refused('ARTIFACT_FILE_CHANGED', artifacts.prepare, self.root, {'inputs': [{'kind': 'ARTIFACT', 'id': 'B', 'revision': 1}]})
        (self.project / 'working/w/b-v1.txt').unlink()
        self.refused('ARTIFACT_FILE_MISSING', artifacts.prepare, self.root, {'inputs': [{'kind': 'ARTIFACT', 'id': 'B', 'revision': 1}]})

    def test_prepare_request_and_selection_are_validated(self):
        self.refused('MALFORMED_REQUEST', artifacts.prepare, self.root, {'inputs': []})
        self.refused('MALFORMED_REQUEST', artifacts.prepare, self.root, {'inputs': [{'kind': 'WHAT', 'id': 'x'}]})
        self.refused('UNKNOWN_ARTIFACT', artifacts.prepare, self.root, {'inputs': [{'kind': 'ARTIFACT', 'id': 'B', 'revision': 7}]})
        self.refused('STALE_ARTIFACT_BASIS', artifacts.prepare, self.root, {'inputs': [{'kind': 'ARTIFACT', 'id': 'B', 'revision': 1, 'sha256': '4' * 64}]})
        self.refused('STALE_SELECTION', artifacts.prepare, self.root, {'inputs': [{'kind': 'SOURCE', 'id': self.sa}], 'expected_snapshot': '5' * 64})
        self.refused('UNKNOWN_SOURCE', artifacts.prepare, self.root, {'inputs': [{'kind': 'SOURCE', 'id': 'SRC-999999'}]})

    @NEED_PCN1
    def test_a_change_after_prepare_is_caught_by_the_publish_guard(self):
        prep = artifacts.prepare(self.root, {'inputs': [{'kind': 'ARTIFACT', 'id': 'B', 'revision': 1}]})
        request = PcnBase.publish_request(self, 'w-1', 'working/w-1/c1.docx', prep['inputs'])
        self.source('a.txt', 'alpha changed after prepare')
        res = pn.publish(self.root, request, b'PK c1', request_id='%032x' % 7)
        self.assertEqual((res['outcome'], res['reason']), ('REFUSED', 'INPUT_CHANGED_SINCE_READ'))
        self.assertFalse((self.project / 'working/w-1/c1.docx').exists())


class PcnBase(Base):
    """Real P-CN1 revisions for P_CN1 registrations."""

    def publish_request(self, work, target, inputs):
        return {'work_id': work, 'target': target,
                'intent_ref': {'ref_id': 'intent-1', 'locator': 'host:msg/1', 'digest': 'sha256:' + '5' * 64, 'attribution': att()},
                'main_interpretation': 'Báo cáo ngắn', 'requirements': [{'req_id': 'REQ-1', 'statement': 'Tóm tắt', 'mandatory': True}],
                'grant_refs': [{'ref_id': 'grant-1', 'locator': 'PROJECT.md#pham-vi', 'digest': 'sha256:' + '1' * 64, 'attribution': att()}],
                'subject': {'actor_id': 'main', 'role': 'main'}, 'inputs': inputs}

    def publish(self, n, target, inputs, data=None):
        res = pn.publish(self.root, self.publish_request('w-1', target, inputs), data if data is not None else ('PK b%d' % n).encode(),
                         request_id='%032x' % n)
        self.assertEqual(res['outcome'], 'PUBLISHED', res)
        return '%032x' % n

    def pcn_req(self, aid, revision, rid, target, deps, **over):
        out = {'artifact_id': aid, 'revision': revision, 'path': target, 'title': 'Kết quả ' + aid, 'kind': 'REPORT', 'limits': 'chỉ đọc đoạn đã khai',
               'origin': 'P_CN1', 'rid': rid, 'recorded_by': 'tester', 'dependencies': list(deps)}
        out.update(over)
        return out

    def receipts(self):
        folder = self.project / '.agent' / 'runtime' / 'publish'
        return sorted(p.name for p in folder.iterdir()) if folder.exists() else []


@NEED_PCN1
class PcnRegistration(PcnBase):
    def prepared(self):
        return artifacts.prepare(self.root, {'inputs': [{'kind': 'SOURCE', 'id': self.sa}]})['inputs']

    def test_register_matches_prepared_inputs_snapshot_and_the_published_evidence(self):
        rid = self.publish(1, 'working/w-1/b1.docx', self.prepared())
        req = self.pcn_req('B', 1, rid, 'working/w-1/b1.docx', [self.dep(self.sa)])
        names = self.receipts()
        plan = artifacts.register_plan(self.root, req)
        self.assertEqual(plan['action'], 'record')
        artifacts.register_apply(self.root, req, plan['plan_token'])
        item = self.item('B', 1)
        self.assertEqual((item['origin'], item['publish_rid'], item['sha256'], item['size']),
                         ('P_CN1', rid, sha(b'PK b1'), 5))
        self.assertRegex(item['publish_identity'], r'^[0-9a-f]{8}:[0-9a-f]{16}$')
        self.assertEqual(self.receipts(), names)  # no new receipt, pins and evidence untouched
        self.assertEqual(self.item('B', 1)['dependency_state'], 'CURRENT_CATALOG_BASIS')

    def test_a_stale_source_after_publish_refuses_registration_and_keeps_file_and_receipt(self):
        rid = self.publish(1, 'working/w-1/b1.docx', self.prepared())
        names, before = self.receipts(), self.index.read_bytes()
        self.source('a.txt', 'alpha changed after publish')
        self.resync()
        stale = self.index.read_bytes()
        err = self.refused('STALE_SOURCE_BASIS', artifacts.register_plan, self.root, self.pcn_req('B', 1, rid, 'working/w-1/b1.docx', [
            {'kind': 'SOURCE', 'id': self.sa, 'revision': 1, 'sha256': sha(b'alpha v1')}]))
        self.assertIn('chưa đăng ký', str(err))
        self.assertEqual((self.index.read_bytes(), self.receipts()), (stale, names))
        self.assertEqual((self.project / 'working/w-1/b1.docx').read_bytes(), b'PK b1')
        self.assertEqual(self.find()['registry'], 'absent')
        self.assertTrue(before)

    def test_the_workbook_changing_after_prepare_refuses_even_when_no_source_moved(self):
        rid = self.publish(1, 'working/w-1/b1.docx', self.prepared())
        self.new('Other', 1)  # an unrelated registration changes the raw workbook
        req = self.pcn_req('B', 1, rid, 'working/w-1/b1.docx', [self.dep(self.sa)])
        err = self.refused('REGISTRY_CHANGED_SINCE_PREPARE', artifacts.register_plan, self.root, req)
        self.assertIn('chưa đăng ký', str(err))
        self.assertEqual({i['key'] for i in self.find()['items']}, {'Other@1'})

    def test_inputs_missing_from_prepared_are_refused(self):
        workbook_only = [i for i in self.prepared() if i['path'] == 'references/_INDEX.xlsx']
        rid = self.publish(1, 'working/w-1/b1.docx', workbook_only)
        self.refused('INPUT_NOT_IN_PREPARED', artifacts.register_plan, self.root, self.pcn_req('B', 1, rid, 'working/w-1/b1.docx', [self.dep(self.sa)]))
        no_workbook = [i for i in self.prepared() if i['path'] != 'references/_INDEX.xlsx']
        rid2 = self.publish(2, 'working/w-1/b2.docx', no_workbook)
        self.refused('WORKBOOK_NOT_IN_PREPARED', artifacts.register_plan, self.root, self.pcn_req('B', 1, rid2, 'working/w-1/b2.docx', [self.dep(self.sa)]))

    def test_rid_path_digest_and_target_bytes_must_agree_with_the_evidence(self):
        rid = self.publish(1, 'working/w-1/b1.docx', self.prepared())
        good = self.pcn_req('B', 1, rid, 'working/w-1/b1.docx', [self.dep(self.sa)])
        self.refused('EVIDENCE_UNUSABLE', artifacts.register_plan, self.root, dict(good, rid='%032x' % 99))
        self.refused('RID_PATH_MISMATCH', artifacts.register_plan, self.root, dict(good, path='working/w-1/other.docx'))
        self.refused('RID_DIGEST_MISMATCH', artifacts.register_plan, self.root, dict(good, sha256='6' * 64))
        (self.project / 'working/w-1/b1.docx').write_bytes(b'PK edited afterwards')
        self.refused('TARGET_NOT_MATCHING', artifacts.register_plan, self.root, good)

    def test_a_p_cn1_registration_without_dependencies_needs_no_workbook_input(self):
        rid = self.publish(1, 'working/w-1/n1.docx', [])
        self.assertEqual(artifacts.register_plan(self.root, self.pcn_req('N', 1, rid, 'working/w-1/n1.docx', []))['action'], 'record')

    def test_open_uses_the_owners_evidence_and_identity(self):
        rid = self.publish(1, 'working/w-1/b1.docx', self.prepared())
        req = self.pcn_req('B', 1, rid, 'working/w-1/b1.docx', [self.dep(self.sa)])
        artifacts.register_apply(self.root, req, artifacts.register_plan(self.root, req)['plan_token'])
        pick = self.find()['items'][0]['pick']
        out = artifacts.open_artifact(self.root, pick)
        self.assertEqual((out['ok'], out['state'], out['bytes_observation']), (True, 'OPENED', 'OBSERVED_BEFORE_OPEN'))
        self.assertEqual(self.opened, [str(self.project / 'working' / 'w-1' / 'b1.docx')])
        # a historical version whose source went stale still opens when the file is right
        self.source('a.txt', 'alpha v2')
        self.resync()
        pick = self.find()['items'][0]['pick']
        out = artifacts.open_artifact(self.root, pick)
        self.assertEqual((out['ok'], out['dependency_state']), (True, 'NEEDS_REVIEW'))
        # replacing the file (same name) is refused
        (self.project / 'working/w-1/b1.docx').write_bytes(b'PK other bytes')
        self.assertEqual(artifacts.open_artifact(self.root, pick), {'ok': False, 'reason': 'file_changed'})
        (self.project / 'working/w-1/b1.docx').unlink()
        self.assertEqual(artifacts.open_artifact(self.root, pick), {'ok': False, 'reason': 'missing'})
        self.assertEqual(len(self.opened), 2)

    def test_a_registry_identity_that_differs_from_the_evidence_is_an_identity_conflict(self):
        rid = self.publish(1, 'working/w-1/b1.docx', [])
        req = self.pcn_req('N', 1, rid, 'working/w-1/b1.docx', [])
        artifacts.register_apply(self.root, req, artifacts.register_plan(self.root, req)['plan_token'])
        self.edit(lambda wb: wb[artifacts.ARTIFACT_SHEET].cell(5, 12).__setattr__('value', '00000001:0000000000000002'))
        pick = self.find()['items'][0]['pick']
        self.assertEqual(artifacts.open_artifact(self.root, pick), {'ok': False, 'reason': 'identity_conflict'})
        self.assertEqual(self.opened, [])

    def test_owner_unavailable_refuses_without_falling_back_to_the_bytes(self):
        rid = self.publish(1, 'working/w-1/b1.docx', [])
        req = self.pcn_req('N', 1, rid, 'working/w-1/b1.docx', [])
        import current_draft as cd
        with mock.patch.object(cd.Env, '__init__', side_effect=cd.Unavailable('OWNER_UNAVAILABLE')):
            err = self.refused('OWNER_UNAVAILABLE', artifacts.register_plan, self.root, req)
        self.assertIn('chưa đăng ký', str(err))


@NEED_PCN1
class ReviewPcnRegressions(PcnBase):
    def test_same_bytes_replacement_before_commit_is_not_registered_with_old_identity(self):
        inputs = artifacts.prepare(self.root, {'inputs': [{'kind': 'SOURCE', 'id': self.sa}]})['inputs']
        target = 'working/w-1/review.docx'
        rid = self.publish(91, target, inputs)
        req = self.pcn_req('Review', 1, rid, target, [self.dep(self.sa)])
        plan = artifacts.register_plan(self.root, req)
        before = self.index.read_bytes()
        def replace_target(_notices):
            path = self.project / target
            other = path.with_suffix('.replacement')
            other.write_bytes(path.read_bytes())
            os.replace(other, path)
        with mock.patch.object(library_sync, 'refuse_if_lossy', side_effect=replace_target):
            self.refused('TARGET_NOT_MATCHING', artifacts.register_apply, self.root, req, plan['plan_token'])
        self.assertEqual(self.index.read_bytes(), before)
        self.assertEqual(artifacts.find(self.root)['registry'], 'absent')


class OpenAndCopy(Base):
    def setUp(self):
        super().setUp()
        self.new('B', 1, [self.dep(self.sa)])
        self.pick = self.find()['items'][0]['pick']

    def test_open_observes_the_file_right_before_the_opener_and_reports_what_it_did_not_do(self):
        before = tree(self.root)
        out = artifacts.open_artifact(self.root, self.pick)
        self.assertEqual((out['ok'], out['state'], out['path']), (True, 'OPENED', 'working/w/b-v1.txt'))
        self.assertEqual(self.opened, [str(self.project / 'working' / 'w' / 'b-v1.txt')])
        self.assertEqual(tree(self.root), before)  # no ops log, no STATE, no registry write

    def test_a_stale_changed_or_malformed_selection_is_refused_and_never_opens(self):
        self.assertEqual(artifacts.open_artifact(self.root, dict(self.pick, snapshot='7' * 64)), {'ok': False, 'reason': 'stale'})
        self.assertEqual(artifacts.open_artifact(self.root, dict(self.pick, revision=2)), {'ok': False, 'reason': 'stale'})
        self.assertEqual(artifacts.open_artifact(self.root, dict(self.pick, sha256='8' * 64)), {'ok': False, 'reason': 'changed'})
        self.assertEqual(artifacts.open_artifact(self.root, dict(self.pick, path='working/w/x.txt')), {'ok': False, 'reason': 'changed'})
        for bad in (None, [], {}, dict(self.pick, extra=1), dict(self.pick, revision=True), dict(self.pick, revision='1'), dict(self.pick, id=5)):
            self.assertEqual(artifacts.open_artifact(self.root, bad), {'ok': False, 'reason': 'bad_selection'}, bad)
        self.assertEqual(self.opened, [])
        self.new('Other', 1)  # a registration after the list was shown makes the selection stale
        self.assertEqual(artifacts.open_artifact(self.root, self.pick), {'ok': False, 'reason': 'stale'})
        self.assertEqual(artifacts.copy_input(self.root, self.pick), {'ok': False, 'reason': 'stale'})

    def test_changed_missing_hardlinked_and_linked_files_are_refused(self):
        path = self.project / 'working/w/b-v1.txt'
        path.write_text('changed', encoding='utf-8')
        self.assertEqual(artifacts.open_artifact(self.root, self.pick), {'ok': False, 'reason': 'file_changed'})
        path.write_text('content of working/w/b-v1.txt', encoding='utf-8')
        os.link(str(path), str(self.project / 'working' / 'twin.txt'))
        self.assertEqual(artifacts.open_artifact(self.root, self.pick), {'ok': False, 'reason': 'hardlink'})
        os.unlink(str(self.project / 'working' / 'twin.txt'))
        path.unlink()
        self.assertEqual(artifacts.open_artifact(self.root, self.pick), {'ok': False, 'reason': 'missing'})
        self.assertEqual(self.opened, [])

    @unittest.skipUnless(os.name == 'nt', 'junctions need Windows')
    def test_a_parent_folder_that_became_a_junction_is_refused(self):
        self.new('J', 1, path='output/real/j.txt')
        pick = [i for i in self.find()['items'] if i['artifact_id'] == 'J'][0]['pick']
        outside = self.project.parent / (self.project.name + '-outside')
        os.makedirs(outside)
        self.addCleanup(shutil.rmtree, str(outside), True)
        shutil.move(str(self.project / 'output' / 'real'), str(outside / 'real'))
        r = subprocess.run(['cmd', '/c', 'mklink', '/J', str(self.project / 'output' / 'real'), str(outside / 'real')], capture_output=True, text=True)
        if r.returncode != 0:
            self.skipTest('cannot create a junction here')
        self.addCleanup(lambda: os.rmdir(str(self.project / 'output' / 'real')))
        self.assertEqual(artifacts.open_artifact(self.root, pick), {'ok': False, 'reason': 'link'})
        self.assertEqual(self.opened, [])

    def test_html_can_be_registered_but_is_not_opened_and_no_opener_framework_exists(self):
        self.new('H', 1, path='working/w/report.html')
        item = [i for i in self.find()['items'] if i['artifact_id'] == 'H'][0]
        self.assertFalse(item['openable'])
        self.assertEqual(artifacts.open_artifact(self.root, item['pick']), {'ok': False, 'reason': 'unsupported_type'})
        import current_draft
        self.assertEqual(set(current_draft.ALLOWED_EXT), {'.docx', '.xlsx', '.pptx', '.pdf', '.txt', '.md', '.csv'})

    def test_copy_carries_the_exact_key_path_hash_snapshot_basis_and_limits_and_writes_nothing(self):
        before = tree(self.root)
        out = artifacts.copy_input(self.root, self.pick)
        self.assertTrue(out['ok'])
        text = out['text']
        for needle in ('B@1', 'working/w/b-v1.txt', self.pick['sha256'], self.pick['snapshot'][:12], 'CURRENT_CATALOG_BASIS', 'chỉ đọc đoạn đã khai', 'UNOBSERVED'):
            self.assertIn(needle, text)
        self.assertEqual(artifacts.copy_input(self.root, self.pick, 'en')['text'].splitlines()[0][:16], 'Result as input ')
        self.assertEqual(tree(self.root), before)
        self.assertEqual(self.opened, [])


class ReviewRegressions(Base):
    def test_closure_error_is_reported_before_the_opener_is_called(self):
        self.new('B', 1, [self.dep(self.sa)])
        pick = self.item('B', 1)['pick']
        failure = artifacts.err('CLOSURE_LIMIT', 'closure quá lớn', 'closure too large')
        with mock.patch.object(artifacts, 'describe', side_effect=failure):
            result = artifacts.open_artifact(self.root, pick)
        self.assertFalse(result['ok'])
        self.assertEqual(result['reason'], 'CLOSURE_LIMIT')
        self.assertEqual(self.opened, [])

    def test_mixed_known_and_unknown_basis_never_becomes_current(self):
        self.new('B', 1, kind='INFERENCE', limits='unknown input limits')
        self.new('C', 1, [self.adep('B', 1), self.dep(self.sa)])
        item = self.item('C', 1)
        self.assertEqual((item['dependency_state'], item['unknown_basis']), ('NO_DECLARED_BASIS', ['B@1']))
        prep = artifacts.prepare(self.root, {'inputs': [{'kind': 'ARTIFACT', 'id': 'C', 'revision': 1}]})
        self.assertEqual(prep['selected'][0]['dependency_state'], 'NO_DECLARED_BASIS')
        self.event(self.sa)
        self.assertEqual(self.item('C', 1)['dependency_state'], 'NEEDS_REVIEW')
        self.refused('STALE_SOURCE_BASIS', artifacts.prepare, self.root,
                     {'inputs': [{'kind': 'ARTIFACT', 'id': 'C', 'revision': 1}]})

    def test_copy_retains_every_declared_edge_and_semantic_snapshot(self):
        self.event(self.sa)
        self.new('B', 1, [self.dep(self.sa)], kind='INFERENCE')
        self.new('C', 1, [self.adep('B', 1), self.dep(self.sb)])
        item = self.item('C', 1)
        for lang in ('vi', 'en'):
            text = artifacts.copy_input(self.root, item['pick'], lang)['text']
            for dep in (self.dep(self.sa), self.dep(self.sb), self.adep('B', 1)):
                self.assertIn('%s@%d' % (dep['id'], dep['revision']), text)
                self.assertIn(dep['sha256'], text)
            self.assertIn('Hết hiệu lực', text)
            self.assertIn('basis_status', text)
            self.assertIn(item['pick']['snapshot'], text)

    def test_linked_index_is_refused_before_reading_any_workbook(self):
        with mock.patch.object(transaction, 'link_free', return_value=False), \
                mock.patch('builtins.open', side_effect=AssertionError('must not read an unsafe index')):
            result = artifacts.load_book(self.root)
        self.assertFalse(result['ok'])
        self.assertEqual(result['kind'], 'unsafe_index')

    @unittest.skipUnless(os.name == 'nt', 'junction probe uses Windows')
    def test_real_junction_cannot_supply_an_external_index(self):
        refs = self.project / 'references'
        saved = self.project / 'references-safe'
        with tempfile.TemporaryDirectory(prefix='law-b02-external-') as external:
            shutil.copyfile(self.index, Path(external) / '_INDEX.xlsx')
            refs.rename(saved)
            try:
                linked = subprocess.run(['cmd', '/c', 'mklink', '/J', str(refs), external], capture_output=True)
                self.assertEqual(linked.returncode, 0, linked.stderr)
                result = artifacts.load_book(self.root)
                self.assertFalse(result['ok'])
                self.assertEqual(result['kind'], 'unsafe_index')
            finally:
                if refs.exists():
                    os.rmdir(refs)  # remove only the fixture junction, never its target
                saved.rename(refs)


class Panel(Base):
    """The panel helpers (`links`, open, copy) and their escaping/read-only contract, without a server."""

    def test_links_rebinds_the_source_row_and_a_changed_workbook_is_stale(self):
        self.new('B', 1, [self.dep(self.sa)], title='<img src=x onerror=alert(1)>')
        res = self.find()
        sel = {'snapshot': res['snapshot'], 'row': 5, 'id': self.sa, 'revision': 1, 'locator': 'source-documents/team/a.txt'}
        out = artifacts.links(self.root, sel)
        self.assertTrue(out['ok'], out)
        self.assertEqual(out['results'][0]['title'], '<img src=x onerror=alert(1)>')  # raw data; the page escapes every dynamic value
        self.assertEqual(out['results'][0]['pick']['snapshot'], res['snapshot'])
        self.assertEqual(artifacts.links(self.root, dict(sel, snapshot='9' * 64)), {'ok': False, 'reason': 'stale'})
        self.assertEqual(artifacts.links(self.root, dict(sel, id='SRC-000002'))['reason'], 'changed')
        self.assertEqual(artifacts.links(self.root, None)['reason'], 'bad_selection')

    def test_the_page_escapes_the_dynamic_values_of_the_links_block(self):
        html = (ROOT / 'tools' / 'app_ui.html').read_text(encoding='utf-8')
        block = html[html.index('function lbLinksBox'):html.index('function lbLinksLoad')]
        for expression in ('esc(o.title)', 'esc(o.path)', "esc(o.artifact_id+'@'+o.revision)", 'esc(o.id)', 'esc(o.kind)', 'esc(o.origin)'):
            self.assertEqual(block.count(expression), 1, expression)
        self.assertEqual(block.count('<span class="mono">'), 2)
        self.assertNotIn('+o.title', block)
        self.assertNotIn('+o.path', block)


class Cli(Base):
    """The command line: a new process reads the saved workbook, so the versions and relations survive the session."""

    def run_cli(self, *args, env=None):
        e = dict(os.environ, PYTHONUTF8='1', PYTHONDONTWRITEBYTECODE='1')
        e.update(env or {})
        proc = subprocess.run([sys.executable, '-B', str(ROOT / 'tools' / 'artifacts.py'), '--project', self.root] + [str(a) for a in args],
                              capture_output=True, encoding='utf-8', env=e, timeout=300)
        return proc.returncode, proc.stdout, proc.stderr

    def request_file(self, name, data):
        path = self.project / 'working' / name
        path.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
        return str(path)

    def test_register_preview_apply_find_copy_open_and_prepare_through_new_processes(self):
        req = self.request_file('req1.json', self.req('R1', 1, [self.dep(self.sa)]))
        code, out, err = self.run_cli('register', 'preview', '--request', req, '--json')
        self.assertEqual(code, 0, err)
        plan = json.loads(out)
        self.assertEqual((plan['action'], plan['index_changed'], plan['applied']), ('record', True, False))
        self.assertEqual(self.find()['registry'], 'absent')  # preview wrote nothing
        code, out, err = self.run_cli('register', 'apply', '--request', req, '--expect-plan', plan['plan_token'], '--json')
        self.assertEqual(code, 0, err)
        self.assertTrue(json.loads(out)['applied'])
        code, out, err = self.run_cli('find', '--source-id', self.sa, '--json')  # a fresh process: the saved workbook is the only memory
        found = json.loads(out)
        self.assertEqual((code, found['matched'], found['items'][0]['key'], found['bytes_observation']), (0, 1, 'R1@1', 'UNOBSERVED'))
        code, out, err = self.run_cli('find', 'ket qua')
        self.assertEqual(code, 0, err)
        self.assertIn('R1@1', out)
        log = str(self.project / 'open.log')
        code, out, err = self.run_cli('open', '--artifact-id', 'R1', '--revision', 1, env={'LOCAL_AGENT_WORKSPACE_TEST_OPEN_LOG': log})
        self.assertEqual(code, 0, err)
        self.assertEqual(Path(log).read_text(encoding='utf-8').split(), [str(self.project / 'working' / 'w' / 'r1-v1.txt')])
        code, out, err = self.run_cli('copy', '--artifact-id', 'R1', '--revision', 1, '--lang', 'en')
        self.assertEqual(code, 0, err)
        self.assertIn('Result as input', out)
        prep = self.request_file('prep.json', {'inputs': [{'kind': 'ARTIFACT', 'id': 'R1', 'revision': 1}]})
        code, out, err = self.run_cli('prepare', '--request', prep, '--json')
        self.assertEqual(code, 0, err)
        self.assertEqual(json.loads(out)['bytes_observation'], 'OBSERVED_AT_PREPARE')

    def test_usage_errors_and_refusals_have_distinct_exit_codes(self):
        req = self.request_file('req2.json', self.req('R1', 1))
        self.assertEqual(self.run_cli('register', 'apply', '--request', req)[0], 2)  # apply needs --expect-plan
        self.assertEqual(self.run_cli('nonsense')[0], 2)
        code, out, err = self.run_cli('register', 'apply', '--request', req, '--expect-plan', 'a' * 64)
        self.assertEqual(code, 1)  # a token that matches no plan is refused
        self.assertEqual(self.find()['registry'], 'absent')
        code, out, err = self.run_cli('open', '--artifact-id', 'Nope', '--revision', 1)
        self.assertEqual(code, 1)
        rel = self.request_file('rel.json', {'source_a': self.sa, 'source_b': self.sb, 'recorded_by': 'cli'})
        code, out, err = self.run_cli('relate', 'preview', '--request', rel, '--json')
        self.assertEqual(code, 0, err)
        token = json.loads(out)['plan_token']
        self.assertEqual(self.run_cli('relate', 'apply', '--request', rel, '--expect-plan', token)[0], 0)
        self.assertEqual(self.find(source_id=self.sa)['related'][0]['b'], self.sb)


@unittest.skipUnless(openpyxl, 'openpyxl is needed for the source workbook')
class Routes(unittest.TestCase):
    """The panel's dedicated routes over HTTP: token and Host checks, no operations log, no generic run path."""

    @classmethod
    def setUpClass(cls):
        import test_ui_browser as ub
        cls.ub = ub
        cls.tmp = tempfile.TemporaryDirectory(prefix='law-b02-routes-')
        cls.pkg = ub.copy_package(Path(cls.tmp.name) / 'package')
        ub.init_project(cls.pkg)
        assert ub.run_tool(cls.pkg, 'library.py', 'init').returncode == 0
        path = cls.pkg / 'references' / 'source-documents' / 'a.txt'
        path.write_text('alpha', encoding='utf-8')
        assert ub.run_tool(cls.pkg, 'library.py', 'sync', '--apply').returncode == 0
        root = str(cls.pkg)
        (cls.pkg / 'working').mkdir(exist_ok=True)
        (cls.pkg / 'working' / 'r1.txt').write_text('report', encoding='utf-8')
        rec = artifacts.load_book(root)['catalog']['SRC-000001'][0]
        req = {'artifact_id': 'R1', 'revision': 1, 'path': 'working/r1.txt', 'title': 'Báo cáo', 'kind': 'REPORT', 'limits': 'chỉ đoạn 1',
               'origin': 'USER_FILE', 'recorded_by': 'tester',
               'dependencies': [{'kind': 'SOURCE', 'id': 'SRC-000001', 'revision': 1, 'sha256': rec['law_sha']}]}
        artifacts.register_apply(root, req, artifacts.register_plan(root, req)['plan_token'])
        cls.open_log = Path(cls.tmp.name) / 'open.log'
        cls.panel = ub.Panel(cls.pkg, cls.tmp.name)
        cls.panel.env = {'LOCAL_AGENT_WORKSPACE_TEST_OPEN_LOG': str(cls.open_log)}
        cls.panel.start()

    @classmethod
    def tearDownClass(cls):
        cls.panel.stop()
        cls.tmp.cleanup()

    def post(self, path, body=None, token=True, host=None):
        import urllib.error
        import urllib.request
        headers = {'Content-Type': 'application/json'}
        if token:
            headers['X-Panel-Token'] = self.panel.token
        if host:
            headers['Host'] = host
        req = urllib.request.Request(self.panel.base + path, data=json.dumps(body or {}).encode('utf-8'), method='POST', headers=headers)
        try:
            with urllib.request.urlopen(req, timeout=60) as r:
                return r.status, json.loads(r.read().decode('utf-8'))
        except urllib.error.HTTPError as e:
            return e.code, json.loads(e.read().decode('utf-8') or '{}')

    def test_routes_need_the_token_and_a_local_host(self):
        for path in ('/api/library/links', '/api/library/artifact-open', '/api/library/artifact-copy'):
            self.assertEqual(self.post(path, {}, token=False)[0], 403, path)
            self.assertEqual(self.post(path, {}, host='evil.example')[0], 403, path)

    def test_links_open_and_copy_work_over_http_and_the_operations_log_stays_untouched(self):
        log = self.pkg / '.agent' / 'runtime' / 'operations.jsonl'
        before = log.read_bytes() if log.exists() else None
        book = artifacts.load_book(str(self.pkg))
        sel = {'snapshot': book['snapshot'], 'row': 5, 'id': 'SRC-000001', 'revision': 1, 'locator': 'source-documents/a.txt'}
        code, res = self.post('/api/library/links', {'selection': sel})
        self.assertEqual((code, res['ok'], res['results'][0]['artifact_id']), (200, True, 'R1'))
        pick = res['results'][0]['pick']
        code, out = self.post('/api/library/artifact-open', {'selection': pick})
        self.assertEqual((code, out['ok'], out['state']), (200, True, 'OPENED'))
        self.assertEqual(self.open_log.read_text(encoding='utf-8').split(), [str(self.pkg / 'working' / 'r1.txt')])
        code, out = self.post('/api/library/artifact-copy', {'selection': pick})
        self.assertEqual((code, out['ok']), (200, True))
        self.assertEqual(self.post('/api/library/artifact-open', {'selection': dict(pick, sha256='1' * 64)})[1], {'ok': False, 'reason': 'changed'})
        self.assertEqual(log.read_bytes() if log.exists() else None, before)  # the dedicated routes are outside run_api/_record

    def test_the_generic_run_route_has_no_artifact_action(self):
        code, r = self.post('/api/run', {'action': 'artifact_register', 'params': {}})
        self.assertEqual(r['code'], 2)
        self.assertTrue(r.get('input_error'))


if __name__ == '__main__':
    unittest.main()
