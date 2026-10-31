"""APP intake: synthetic projects only; real CLI and headless loopback browser."""
import hashlib
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import threading
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'tools'))
import app
try:
    import openpyxl
except ImportError:
    openpyxl = None


def snap(root):
    return {p.relative_to(root).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in root.rglob('*') if p.is_file()}


@unittest.skipUnless(openpyxl, "openpyxl unavailable; no installation attempted")
class PanelIntake(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='law-panel-intake-')
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        tools = self.root / 'tools'
        tools.mkdir()
        # Explicit public runtime files only, never copytree of a workspace.
        for f in (ROOT / 'tools').iterdir():
            if f.is_file() and f.suffix in ('.py', '.html'):
                shutil.copyfile(f, tools / f.name)
        (self.root / '.agent').mkdir()
        # Browser navigation requires a configured project, not an unknown folder.
        (self.root / '.agent/PACKAGE_INFO.json').write_text(
            json.dumps({'package': 'local-agent-workspace', 'version': '1.9.0', 'lang': 'vi'}), encoding='utf-8')
        (self.root / 'PROJECT.md').write_text('- Project name: Synthetic\n', encoding='utf-8')
        (self.root / 'VERSION').write_text('test', encoding='utf-8')
        self.staging = self.root / 'references/legal/_waiting'
        self.staging.mkdir(parents=True)
        (self.root / 'references/legal/team').mkdir()
        self.index = self.root / 'references/_INDEX.xlsx'
        shutil.copyfile(ROOT / 'skills/local/shared/source-library/assets/source-index-template.xlsx', self.index)
        self.cfg = self.root / '.agent/SOURCE_INTAKE.json'
        self.config = {'schema': 'law.source-intake/1', 'enabled': True,
                       'staging': 'references/legal/_waiting', 'mappings': [
                           {'code': 'LAW-01', 'canonical_name': 'Canonical.txt', 'destination': 'legal/team/Canonical.txt'}]}
        self.cfg.write_text(json.dumps(self.config), encoding='utf-8')
        self.source = self.staging / 'LAW-01_<img onerror=alert(1)>.txt'
        # Windows disallows < >; ampersand still exercises HTML escaping.
        self.source = self.staging / 'LAW-01_&source.txt'
        self.source.write_bytes(b'raw source\r\n')
        self.target = self.root / 'references/legal/team/Canonical.txt'
        (self.root / 'private.txt').write_bytes(b'PRIVATE')
        self.bind = patch.multiple(app, ROOT=str(self.root), TOOLS=str(tools), UI_FILE=str(tools / 'app_ui.html'))
        self.bind.start()
        self.addCleanup(self.bind.stop)

    def call(self, **p):
        return app.run_api({'action': 'library_intake', 'params': p})[1]

    def test_strict_input_and_fixed_argv(self):
        with patch.object(app, 'run_tool') as child:
            for p in ({'apply': 1}, {'apply': 'false'}, {'apply': None}, {'file': '../x'},
                      {'path': 'x'}, {'code': 'LAW-01'}, {'args': []}, {'plan_token': 'a'*64},
                      {'apply': True}, {'apply': True, 'plan_token': 'A'*64}, {'lang': []}):
                self.assertEqual(self.call(**p)['result']['status'], 'blocked', p)
            self.assertEqual(app.run_api({'action': 'library_intake', 'params': []})[1]['code'], 2)
            child.assert_not_called()
            app.action('library_intake', {})
            child.assert_called_once_with('library.py', ['intake'])
            child.reset_mock()
            app.action('library_intake', {'apply': True, 'plan_token': 'a'*64})
            child.assert_called_once_with('library.py', ['intake', '--apply', '--expect-plan', 'a'*64])

    def test_real_cli_preservation_stale_collision_config_partial(self):
        before = snap(self.root)
        with patch.dict(os.environ):
            os.environ.pop('PYTHONDONTWRITEBYTECODE', None)
            preview = self.call()
        self.assertEqual(preview['result']['status'], 'preview', preview)
        self.assertEqual(snap(self.root), before)
        token = preview['result']['plan_token']
        self.source.write_bytes(b'changed')
        index = self.index.read_bytes()
        self.assertEqual(self.call(apply=True, plan_token=token)['result']['status'], 'blocked')
        self.assertFalse(self.target.exists())
        self.assertEqual(self.index.read_bytes(), index)
        preview = self.call()
        self.cfg.write_text(json.dumps(dict(self.config, enabled=False)), encoding='utf-8')
        self.assertEqual(self.call(apply=True, plan_token=preview['result']['plan_token'])['result']['status'], 'blocked')
        self.assertEqual(self.index.read_bytes(), index)
        self.assertFalse(self.target.exists())
        self.assertEqual(self.call()['result']['status'], 'blocked')
        self.cfg.unlink()
        self.assertEqual(self.call()['result']['status'], 'blocked')
        self.cfg.write_text(json.dumps(self.config), encoding='utf-8')
        for raw in (self.source.read_bytes(), b'foreign'):
            self.target.write_bytes(raw)
            self.assertEqual(self.call()['result']['status'], 'blocked')
            self.assertEqual(self.target.read_bytes(), raw)
            self.assertEqual(self.index.read_bytes(), index)
        self.target.unlink()
        preview = self.call()
        with patch.dict(os.environ, {'LOCAL_AGENT_WORKSPACE_TEST_FAULT': 'commit:1'}):
            result = self.call(apply=True, plan_token=preview['result']['plan_token'])['result']
        self.assertEqual(result['status'], 'partial', result)
        self.assertTrue(result['transaction'])
        self.assertTrue(result['next_actions'])
        self.assertEqual(self.source.read_bytes(), b'changed')

    def test_browser_success_stale_missing_and_token_closure(self):
        try:
            from playwright.sync_api import sync_playwright, expect
        except ImportError:
            self.skipTest('playwright missing; no installation attempted')
        server = app.Server(('127.0.0.1', 0), app.Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            with sync_playwright() as pw:
                try:
                    browser = pw.chromium.launch(headless=True)
                except Exception as exc:
                    self.skipTest('Chromium unavailable: ' + type(exc).__name__)
                try:
                    page = browser.new_page()
                    page.set_default_timeout(8000)
                    errors = []
                    page.on('pageerror', lambda e: errors.append(str(e)))
                    requests = []
                    page.on('request', lambda req: requests.append(json.loads(req.post_data))
                            if req.url.endswith('/api/run') and req.post_data else None)
                    page.goto('http://127.0.0.1:%d/?t=%s' % (server.server_address[1], app.TOKEN))
                    page.locator('[data-page="library"]').click()
                    page.locator('.lang button[data-lang="en"]').click()
                    expect(page.locator('#lb-intake')).to_have_text('Import waiting files')
                    before = snap(self.root)
                    page.locator('#lb-intake').click()
                    expect(page.locator('#sh-apply')).to_be_visible()
                    self.assertEqual(snap(self.root), before)
                    self.assertIn('LAW-01_&source.txt', page.locator('#sh-out').inner_text())
                    self.assertEqual(page.locator('#sh-out img').count(), 0)
                    token = page.locator('#sh-out').inner_text().split('plan_token: ')[1].strip()
                    page.locator('#sh-apply').click()
                    expect(page.locator('#sh-apply')).to_be_hidden()
                    expect(page.locator('#sh-tag')).to_contain_text('Succeeded')
                    self.assertEqual(requests[-1]['params']['plan_token'], token)
                    self.assertEqual(len(requests), 2)
                    self.assertEqual(self.target.read_bytes(), self.source.read_bytes())
                    self.assertEqual(self.cfg.read_text(), json.dumps(self.config))
                    self.assertEqual((self.root / 'private.txt').read_bytes(), b'PRIVATE')
                    wb = openpyxl.load_workbook(self.index)
                    try:
                        self.assertEqual(wb['Danh mục'].cell(5, 14).value, 'Cần đối chiếu')
                        headers = {cell.value: cell.column for cell in wb['Danh mục'][4]}
                        self.assertEqual(wb['Danh mục'].cell(5, headers['LAW revision']).value, 1)
                    finally:
                        wb.close()
                    # Fresh synthetic batch after removing only our own committed file/index.
                    self.target.unlink()
                    shutil.copyfile(ROOT / 'skills/local/shared/source-library/assets/source-index-template.xlsx', self.index)
                    page.locator('#sh-close').click()
                    page.locator('#lb-intake').click()
                    expect(page.locator('#sh-apply')).to_be_visible()
                    self.source.write_bytes(b'stale source')
                    index = self.index.read_bytes()
                    page.locator('#sh-apply').click()
                    expect(page.locator('#sh-tag')).to_contain_text('Not carried out')
                    expect(page.locator('#sh-apply')).to_be_hidden()
                    self.assertEqual(len(requests), 4)
                    self.assertFalse(self.target.exists())
                    self.assertEqual(self.index.read_bytes(), index)
                    page.locator('#sh-close').click()
                    self.cfg.unlink()
                    page.locator('#lb-intake').click()
                    expect(page.locator('#sh-tag')).to_contain_text('Not carried out')
                    expect(page.locator('#sh-apply')).to_be_hidden()
                    self.assertEqual(len(requests), 5)
                    page.locator('#sh-close').click()
                    page.locator('.lang button[data-lang="vi"]').click()
                    expect(page.locator('#lb-intake')).to_have_text('Nhập file chờ')
                    self.assertEqual(errors, [])
                finally:
                    browser.close()
        finally:
            server.shutdown()
            server.server_close()
            thread.join(5)
            self.assertFalse(thread.is_alive())


if __name__ == '__main__':
    unittest.main()
